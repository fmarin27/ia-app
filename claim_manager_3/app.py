from __future__ import annotations

import base64
import json
import os
import re
import shutil
import smtplib
import ssl
import struct
import subprocess
import sys
import tempfile
import threading
import time
import textwrap
import webbrowser
from copy import copy
from datetime import datetime, timedelta
from email.message import EmailMessage
from io import BytesIO
from pathlib import Path
from urllib.error import URLError
from urllib.parse import quote
from urllib.request import Request, urlopen

def _runtime_app_dir() -> Path:
    portable_home = (os.environ.get("CLAIM_MANAGER_HOME") or "").strip()
    if portable_home:
        return Path(portable_home).expanduser().resolve()
    if getattr(sys, "frozen", False):
        exe_dir = Path(sys.executable).resolve().parent
        for candidate in (exe_dir, exe_dir.parent):
            if (candidate / "claims_data.json").exists() or (candidate / "PENDING CLAIMS").exists():
                return candidate
        return exe_dir
    return Path(__file__).resolve().parents[1]


APP_DIR = _runtime_app_dir()
CLAIM_MANAGER_DIR = APP_DIR / "claim_manager_3" if (APP_DIR / "claim_manager_3").exists() else Path(__file__).resolve().parent
VENDOR_DIR = CLAIM_MANAGER_DIR / "_vendor"
if VENDOR_DIR.exists():
    vendor_path = str(VENDOR_DIR)
    if vendor_path not in sys.path:
        sys.path.insert(0, vendor_path)

from PySide6.QtCore import QDate, QObject, QPoint, Qt, QTimer, QUrl, Signal
from PySide6.QtGui import QAction, QColor, QDesktopServices, QFontMetrics
from PySide6.QtWidgets import (
    QApplication,
    QCheckBox,
    QComboBox,
    QDialog,
    QDialogButtonBox,
    QDateEdit,
    QFileDialog,
    QFormLayout,
    QFrame,
    QGridLayout,
    QHeaderView,
    QHBoxLayout,
    QInputDialog,
    QLabel,
    QLineEdit,
    QListWidget,
    QListWidgetItem,
    QMainWindow,
    QMenu,
    QMessageBox,
    QPushButton,
    QSplitter,
    QPlainTextEdit,
    QTableWidget,
    QTableWidgetItem,
    QStackedWidget,
    QToolButton,
    QTextEdit,
    QVBoxLayout,
    QWidget,
    QSizePolicy,
    QAbstractItemView,
)

from data_access import (
    BodyShopEntry,
    ClaimToolContact,
    ClaimToolFile,
    ClaimView,
    ClaimsRepository,
    InsuranceCompanyEntry,
    render_note_history,
)

try:
    from docx import Document
except Exception:
    Document = None

try:
    from docx.enum.section import WD_ORIENT
except Exception:
    WD_ORIENT = None

try:
    from docx.shared import Inches
except Exception:
    Inches = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    from pypdf import PdfReader, PdfWriter
except Exception:
    PdfReader = None
    PdfWriter = None

try:
    from reportlab.pdfgen import canvas as reportlab_canvas
except Exception:
    reportlab_canvas = None


APP_NAME = "Claim Manager 3.0"
APP_VERSION = "3.0.1"
DESKTOP_UPDATE_BRANCH = "desktop-updates"
DESKTOP_UPDATE_MANIFEST_URL = f"https://api.github.com/repos/fmarin27/ia-app/contents/claim_manager_3/desktop_update/latest.json?ref={DESKTOP_UPDATE_BRANCH}"
DESKTOP_UPDATE_CONFIG_FILE = APP_DIR / "desktop_update_config.json"
SECRETS_FILE = APP_DIR / "claims_secrets.json"
APPTRAK_AUTOMATION_DIR = Path(r"C:\AMobile\automation")
APPTRAK_IMPORT_BAT = APPTRAK_AUTOMATION_DIR / "Run-AppTrakImport.bat"
APPTRAK_EXE = Path(r"C:\AMobile\app\apptrak.exe")
APPTRAK_DBF_DIR = Path(r"C:\AMobile\dbf")
APPTRAK_DOCS_DIR = Path(r"C:\AMobile\docs")
APPTRAK_PICS_DIR = Path(r"C:\AMobile\pics")
REFRESH_SCAN_HELPER = CLAIM_MANAGER_DIR / "refresh_scan_helper.py"
AUTOSOURCE_PDA_TEMPLATE = Path(r"C:\Users\ferna\Downloads\Autosource_PDA_Sub_Level.pdf")
OFFICE_UPDATE_EMAIL_FROM = "fernandomarin27@gmail.com"
OFFICE_UPDATE_EMAIL_TO = "ldellacorte@duhamels.com"
OFFICE_UPDATE_EMAIL_CC = "joe@lasalallc.com, dnoone@duhamels.com"
OFFICE_UPDATE_EMAIL_SUBJECT = "Open sheet"
OFFICE_UPDATE_EMAIL_BODY = "thank you!"
OFFICE_REVIEW_EMAIL_TO = "joe@lasalallc.com"
OFFICE_RMC_EMAIL_TO = "gferreira@duhamels.com"
OFFICE_SUPPLEMENT_EMAIL_TO = "ldellacorte@duhamels.com"
OFFICE_SUPPLEMENT_EMAIL_SUBJECT = "Supplement request"
ROUTE_HOME_ADDRESS = "5 Richlee Rd, Norwalk, CT 06851"
PAYROLL_OUTPUT_DIR = APP_DIR / "Payroll"
CLAIM_TABLE_HEADERS = ["Claim ID", "Customer", "Insurance", "Vehicle", "Shop", "Town", "Date Of Loss", "Status"]
OPEN_CLAIM_TABLE_HEADERS = CLAIM_TABLE_HEADERS + ["Progress"]
OFFICE_TABLE_HEADERS = ["Claim ID", "Customer", "Shop", "Inspection When", "Progress", "Additional Notes"]
REPORT_TABLE_HEADERS = ["Closed Date", "Claim ID", "Customer", "Insurance", "Type"]
REPORT_NULL_QDATE = QDate(2000, 1, 1)


def _version_key(version: str) -> tuple[int, ...]:
    parts = [int(part) for part in re.findall(r"\d+", version)]
    return tuple(parts) if parts else (0,)


def _load_desktop_update_config() -> dict[str, object]:
    default_config: dict[str, object] = {
        "enabled": True,
        "channel": "stable",
        "manifest_url": DESKTOP_UPDATE_MANIFEST_URL,
    }
    if not DESKTOP_UPDATE_CONFIG_FILE.exists():
        return default_config
    try:
        raw = json.loads(DESKTOP_UPDATE_CONFIG_FILE.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return default_config
    config = {**default_config, **raw}
    config["enabled"] = bool(config.get("enabled", True))
    config["channel"] = str(config.get("channel", "") or "stable").strip() or "stable"
    config["manifest_url"] = str(config.get("manifest_url", "") or DESKTOP_UPDATE_MANIFEST_URL).strip() or DESKTOP_UPDATE_MANIFEST_URL
    return config


class RecordEditorDialog(QDialog):
    _last_position: QPoint | None = None

    def __init__(
        self,
        parent: QWidget | None,
        title: str,
        fields: list[tuple],
        values: dict[str, str] | None = None,
        browse_keys: set[str] | None = None,
        stacked_labels: bool = False,
    ) -> None:
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setModal(True)
        self.resize(860 if stacked_labels else 520, 420 if stacked_labels else 320)
        if self.__class__._last_position is not None:
            self.move(self.__class__._last_position)
        self._widgets: dict[str, QWidget] = {}
        self._browse_keys = browse_keys or set()
        self._stacked_labels = stacked_labels

        layout = QVBoxLayout(self)
        form = QFormLayout()
        form.setHorizontalSpacing(12)
        form.setVerticalSpacing(10)
        stacked_form = QVBoxLayout()
        stacked_form.setSpacing(10)
        values = values or {}

        for field in fields:
            key = field[0]
            label = field[1]
            multiline = bool(field[2])
            choices = field[3] if len(field) > 3 else None
            if key in self._browse_keys:
                row = QWidget()
                row_layout = QHBoxLayout(row)
                row_layout.setContentsMargins(0, 0, 0, 0)
                row_layout.setSpacing(8)
                entry = QLineEdit(values.get(key, ""))
                browse = QPushButton("Browse")
                browse.clicked.connect(lambda _=False, target=entry: self._pick_file(target))
                row_layout.addWidget(entry, 1)
                row_layout.addWidget(browse, 0)
                self._widgets[key] = entry
                if self._stacked_labels:
                    label_widget = QLabel(label)
                    label_widget.setWordWrap(True)
                    stacked_form.addWidget(label_widget)
                    stacked_form.addWidget(row)
                else:
                    form.addRow(label, row)
                continue

            if choices:
                widget = QComboBox()
                widget.addItems([str(choice) for choice in choices])
                current_value = values.get(key, "")
                match_index = widget.findText(current_value)
                if match_index >= 0:
                    widget.setCurrentIndex(match_index)
                self._widgets[key] = widget
                if self._stacked_labels:
                    label_widget = QLabel(label)
                    label_widget.setWordWrap(True)
                    stacked_form.addWidget(label_widget)
                    stacked_form.addWidget(widget)
                else:
                    form.addRow(label, widget)
                continue

            if multiline:
                widget = QTextEdit()
                widget.setPlainText(values.get(key, ""))
                widget.setMinimumHeight(88)
            else:
                widget = QLineEdit(values.get(key, ""))
            self._widgets[key] = widget
            if self._stacked_labels:
                label_widget = QLabel(label)
                label_widget.setWordWrap(True)
                stacked_form.addWidget(label_widget)
                stacked_form.addWidget(widget)
            else:
                form.addRow(label, widget)

        if self._stacked_labels:
            layout.addLayout(stacked_form)
        else:
            layout.addLayout(form)
        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def _pick_file(self, target: QLineEdit) -> None:
        current = target.text().strip()
        start_dir = current if current else str(Path.home())
        file_path, _ = QFileDialog.getOpenFileName(self, "Choose File", start_dir)
        if file_path:
            target.setText(file_path)

    def values(self) -> dict[str, str]:
        result: dict[str, str] = {}
        for key, widget in self._widgets.items():
            if isinstance(widget, QTextEdit):
                result[key] = widget.toPlainText().strip()
            elif isinstance(widget, QComboBox):
                result[key] = widget.currentText().strip()
            elif isinstance(widget, QLineEdit):
                result[key] = widget.text().strip()
        return result

    def moveEvent(self, event) -> None:  # type: ignore[override]
        self.__class__._last_position = self.pos()
        super().moveEvent(event)


class AttachmentPickerDialog(QDialog):
    def __init__(self, parent: QWidget | None, title: str, paths: list[Path]) -> None:
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setModal(True)
        self.resize(560, 360)
        self._paths = paths

        layout = QVBoxLayout(self)
        self.list_widget = QListWidget()
        self.list_widget.setSelectionMode(QAbstractItemView.MultiSelection)
        for path in paths:
            item = QListWidgetItem(path.name)
            item.setToolTip(str(path))
            self.list_widget.addItem(item)
        layout.addWidget(self.list_widget, 1)

        buttons = QDialogButtonBox(QDialogButtonBox.Ok | QDialogButtonBox.Cancel)
        buttons.accepted.connect(self.accept)
        buttons.rejected.connect(self.reject)
        layout.addWidget(buttons)

    def selected_paths(self) -> list[Path]:
        selected_rows = {self.list_widget.row(item) for item in self.list_widget.selectedItems()}
        return [path for index, path in enumerate(self._paths) if index in selected_rows]


class ClaimNotesDialog(QDialog):
    _last_position: QPoint | None = None

    def __init__(self, parent: QWidget | None, claim: ClaimView) -> None:
        super().__init__(parent)
        self.setWindowTitle(f"Claim Notes - {claim.claim_id or claim.display_customer}")
        self.setModal(True)
        self.resize(760, 520)
        self.setMinimumSize(620, 420)
        if self.__class__._last_position is not None:
            self.move(self.__class__._last_position)

        layout = QVBoxLayout(self)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(10)

        saved_label = QLabel("Saved Notes")
        saved_label.setObjectName("sectionTitle")
        layout.addWidget(saved_label)

        self.saved_notes_view = QPlainTextEdit()
        self.saved_notes_view.setReadOnly(True)
        self.saved_notes_view.setMinimumHeight(220)
        self.saved_notes_view.setPlainText(render_note_history(claim.note_history).strip() or claim.notes or "")
        layout.addWidget(self.saved_notes_view, 1)

        new_label = QLabel("New Note")
        new_label.setObjectName("sectionTitle")
        layout.addWidget(new_label)

        self.new_note_edit = QTextEdit()
        self.new_note_edit.setMinimumHeight(120)
        layout.addWidget(self.new_note_edit)

        buttons = QDialogButtonBox()
        self.add_button = buttons.addButton("Add Note", QDialogButtonBox.AcceptRole)
        self.close_button = buttons.addButton("Close", QDialogButtonBox.RejectRole)
        self.add_button.clicked.connect(self.accept)
        self.close_button.clicked.connect(self.reject)
        layout.addWidget(buttons)

    def note_text(self) -> str:
        return self.new_note_edit.toPlainText().strip()

    def moveEvent(self, event) -> None:  # type: ignore[override]
        self.__class__._last_position = self.pos()
        super().moveEvent(event)


class WizardDialog(QDialog):
    _last_position: QPoint | None = None

    def __init__(self, parent: QWidget | None, title: str, fields: list[tuple], values: dict[str, str] | None = None) -> None:
        super().__init__(parent)
        self.setWindowTitle(title)
        self.setModal(True)
        self.resize(660, 270)
        self.setMinimumSize(560, 240)
        if self.__class__._last_position is not None:
            self.move(self.__class__._last_position)
        self._fields = fields
        self._answers = dict(values or {})
        self._index = 0
        self._current_key = ""
        self._current_widget: QWidget | None = None

        layout = QVBoxLayout(self)
        layout.setContentsMargins(18, 18, 18, 18)
        layout.setSpacing(12)

        self.step_label = QLabel("")
        self.step_label.setObjectName("workspaceStripTitle")
        layout.addWidget(self.step_label)

        self.question_label = QLabel("")
        self.question_label.setWordWrap(True)
        self.question_label.setObjectName("sectionTitle")
        layout.addWidget(self.question_label)

        self.input_host = QFrame()
        self.input_host_layout = QVBoxLayout(self.input_host)
        self.input_host_layout.setContentsMargins(0, 0, 0, 0)
        self.input_host_layout.setSpacing(0)
        layout.addWidget(self.input_host, 1)

        buttons = QHBoxLayout()
        buttons.addStretch(1)
        self.back_button = QPushButton("Back")
        self.next_button = QPushButton("Next")
        self.finish_button = QPushButton("Finish")
        self.cancel_button = QPushButton("Cancel")
        self.back_button.setAutoDefault(False)
        self.back_button.setDefault(False)
        self.back_button.clicked.connect(self._go_back)
        self.next_button.setAutoDefault(True)
        self.next_button.clicked.connect(self._go_next)
        self.finish_button.setAutoDefault(True)
        self.finish_button.clicked.connect(self._finish)
        self.cancel_button.setAutoDefault(False)
        self.cancel_button.setDefault(False)
        self.cancel_button.clicked.connect(self.reject)
        buttons.addWidget(self.back_button)
        buttons.addWidget(self.next_button)
        buttons.addWidget(self.finish_button)
        buttons.addWidget(self.cancel_button)
        layout.addLayout(buttons)

        self._render_step()

    def _clear_input(self) -> None:
        if self._current_widget is not None:
            self._current_widget.setParent(None)
            self._current_widget.deleteLater()
            self._current_widget = None

    def _build_widget(self, field: tuple) -> QWidget:
        key = field[0]
        multiline = bool(field[2])
        choices = field[3] if len(field) > 3 else None
        current_value = self._answers.get(key, "")

        if choices:
            widget = QComboBox()
            widget.addItems([str(choice) for choice in choices])
            match_index = widget.findText(current_value)
            if match_index >= 0:
                widget.setCurrentIndex(match_index)
            widget.setMinimumHeight(40)
            widget.installEventFilter(self)
            return widget

        if multiline:
            widget = QTextEdit()
            widget.setPlainText(current_value)
            widget.setMinimumHeight(120)
            widget.installEventFilter(self)
            return widget

        widget = QLineEdit(current_value)
        widget.setMinimumHeight(40)
        widget.returnPressed.connect(self._advance_from_enter)
        return widget

    def _advance_from_enter(self) -> None:
        if self._index >= len(self._fields) - 1:
            self._finish()
        else:
            self._go_next()

    def _read_current_value(self) -> None:
        if not self._current_widget:
            return
        if isinstance(self._current_widget, QTextEdit):
            self._answers[self._current_key] = self._current_widget.toPlainText().strip()
        elif isinstance(self._current_widget, QComboBox):
            self._answers[self._current_key] = self._current_widget.currentText().strip()
        elif isinstance(self._current_widget, QLineEdit):
            self._answers[self._current_key] = self._current_widget.text().strip()

    def _render_step(self) -> None:
        self._clear_input()
        field = self._fields[self._index]
        self._current_key = field[0]
        label = field[1]
        self.step_label.setText(f"Step {self._index + 1} of {len(self._fields)}")
        self.question_label.setText(str(label))
        self._current_widget = self._build_widget(field)
        self.input_host_layout.addWidget(self._current_widget)
        self.back_button.setEnabled(self._index > 0)
        self.next_button.setVisible(self._index < len(self._fields) - 1)
        self.finish_button.setVisible(self._index == len(self._fields) - 1)
        self.next_button.setDefault(self._index < len(self._fields) - 1)
        self.finish_button.setDefault(self._index == len(self._fields) - 1)
        self._current_widget.setFocus()

    def _go_back(self) -> None:
        self._read_current_value()
        if self._index > 0:
            self._index -= 1
            self._render_step()

    def _go_next(self) -> None:
        self._read_current_value()
        if self._index < len(self._fields) - 1:
            self._index += 1
            self._render_step()

    def _finish(self) -> None:
        self._read_current_value()
        self.accept()

    def values(self) -> dict[str, str]:
        return dict(self._answers)

    def eventFilter(self, watched, event) -> bool:  # type: ignore[override]
        if event.type() == event.Type.KeyPress and event.key() in (Qt.Key_Return, Qt.Key_Enter):
            if isinstance(watched, QTextEdit):
                if event.modifiers() & Qt.ShiftModifier:
                    return super().eventFilter(watched, event)
                self._advance_from_enter()
                return True
            if isinstance(watched, QComboBox):
                self._advance_from_enter()
                return True
        return super().eventFilter(watched, event)

    def moveEvent(self, event) -> None:  # type: ignore[override]
        self.__class__._last_position = self.pos()
        super().moveEvent(event)


class AppTrakWorkerSignals(QObject):
    finished = Signal(dict)


class RefreshScanWorkerSignals(QObject):
    finished = Signal(dict)


class ClaimsDashboard(QMainWindow):
    update_manifest_ready = Signal(object)
    update_download_ready = Signal(object)

    def __init__(self) -> None:
        super().__init__()
        self.repo = ClaimsRepository()
        self.claims: list[ClaimView] = []
        self.filtered_all: list[ClaimView] = []
        self.filtered_open: list[ClaimView] = []
        self.filtered_closed: list[ClaimView] = []
        self.settings: dict[str, object] = {}
        self.claim_tools_contacts: list[ClaimToolContact] = []
        self.claim_tools_files: list[ClaimToolFile] = []
        self.body_shop_entries: list[BodyShopEntry] = []
        self.insurance_company_entries: list[InsuranceCompanyEntry] = []
        self.route_plan_keys: list[str] = []
        self.route_selected_stop_key: str | None = None
        self.current_claim: ClaimView | None = None
        self.report_selected_date = ""
        self.processed_apptrak_pdfs: list[str] = []
        self.apptrak_import_running = False
        self.refresh_scan_running = False
        self.desktop_update_config = _load_desktop_update_config()
        self._update_check_in_progress = False
        self._update_download_in_progress = False
        self.apptrak_signals = AppTrakWorkerSignals()
        self.apptrak_signals.finished.connect(self._finish_apptrak_import_cycle)
        self.refresh_scan_signals = RefreshScanWorkerSignals()
        self.refresh_scan_signals.finished.connect(self._finish_refresh_scan)
        self.update_manifest_ready.connect(self._handle_update_manifest_result)
        self.update_download_ready.connect(self._handle_update_download_result)

        self.setWindowTitle(APP_NAME)
        self.resize(1440, 900)
        self.setMinimumSize(1180, 760)
        self._build_ui()
        self._load_claims()
        if getattr(sys, "frozen", False) and bool(self.desktop_update_config.get("enabled", True)):
            QTimer.singleShot(1500, self._check_for_desktop_updates_on_launch)

    def _build_ui(self) -> None:
        central = QWidget()
        central.setObjectName("central")
        self.setCentralWidget(central)

        outer = QVBoxLayout(central)
        outer.setContentsMargins(24, 0, 24, 24)
        outer.setSpacing(10)

        top_band = QHBoxLayout()
        top_band.setContentsMargins(0, 0, 0, 0)
        top_band.setSpacing(12)
        outer.addLayout(top_band)

        title = QLabel(APP_NAME)
        title.setObjectName("title")
        top_band.addWidget(title, 0, Qt.AlignLeft | Qt.AlignTop)
        top_band.addStretch(1)

        workspace_card = QFrame()
        workspace_card.setObjectName("workspaceStrip")
        workspace_card.setMinimumWidth(480)
        workspace_card.setMaximumWidth(680)
        workspace_card.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Fixed)
        workspace_layout = QVBoxLayout(workspace_card)
        workspace_layout.setContentsMargins(10, 4, 10, 4)
        workspace_layout.setSpacing(2)

        workspace_title_row = QHBoxLayout()
        workspace_title_row.setSpacing(10)
        watched_title = QLabel("Watched Folders")
        watched_title.setObjectName("workspaceStripTitle")
        tools_title = QLabel("Claim Tools Folder")
        tools_title.setObjectName("workspaceStripTitle")
        workspace_title_row.addWidget(watched_title, 1)
        workspace_title_row.addWidget(tools_title, 1)
        workspace_layout.addLayout(workspace_title_row)

        workspace_value_row = QHBoxLayout()
        workspace_value_row.setSpacing(10)

        self.watched_folders_display = QLabel("-")
        self.watched_folders_display.setObjectName("workspaceStripValueBox")
        self.watched_folders_display.setWordWrap(False)
        self.watched_folders_display.setTextInteractionFlags(Qt.TextSelectableByMouse)
        workspace_value_row.addWidget(self.watched_folders_display, 1)

        self.claim_tools_folder_compact = QLabel("-")
        self.claim_tools_folder_compact.setObjectName("workspaceStripValueBox")
        self.claim_tools_folder_compact.setWordWrap(False)
        self.claim_tools_folder_compact.setTextInteractionFlags(Qt.TextSelectableByMouse)
        workspace_value_row.addWidget(self.claim_tools_folder_compact, 1)
        workspace_layout.addLayout(workspace_value_row)

        top_band.addWidget(workspace_card, 0, Qt.AlignRight | Qt.AlignTop)

        settings_row = QHBoxLayout()
        settings_row.setContentsMargins(0, 0, 0, 0)
        settings_row.setSpacing(0)
        outer.addLayout(settings_row)
        settings_row.addStretch(1)

        self.settings_button = QPushButton("Settings")
        self.settings_button.setObjectName("primaryButton")
        self.settings_button.setMinimumWidth(150)
        self.settings_button.clicked.connect(self._open_settings_dialog)
        settings_row.addWidget(self.settings_button, 0, Qt.AlignCenter)

        settings_row.addStretch(1)

        controls_row = QHBoxLayout()
        controls_row.setSpacing(12)
        outer.addLayout(controls_row)
        controls_row.addStretch(1)

        search_wrap = QFrame()
        search_wrap.setObjectName("toolbarCard")
        search_layout = QHBoxLayout(search_wrap)
        search_layout.setContentsMargins(12, 10, 12, 10)
        search_layout.setSpacing(12)

        self.search_input = QLineEdit()
        self.search_input.setPlaceholderText("Search claims, insurance, town, shop...")
        self.search_input.setClearButtonEnabled(True)
        self.search_input.setMinimumWidth(300)
        self.search_input.textChanged.connect(self._apply_filters)
        search_layout.addWidget(self.search_input, 1)

        refresh_button = QPushButton("Refresh")
        refresh_button.setObjectName("primaryButton")
        refresh_button.clicked.connect(self._load_claims)
        search_layout.addWidget(refresh_button, 0)

        self.refresh_scan_button = QPushButton("Refresh Scan")
        self.refresh_scan_button.clicked.connect(self._run_refresh_scan)
        search_layout.addWidget(self.refresh_scan_button, 0)

        self.check_apptrak_button = QPushButton("Check AppTrak")
        self.check_apptrak_button.clicked.connect(self.run_apptrak_import_now)
        search_layout.addWidget(self.check_apptrak_button, 0)

        controls_row.addWidget(search_wrap, 0, Qt.AlignCenter)
        controls_row.addStretch(1)

        self.apptrak_status_label = QLabel("AppTrak import ready. Manual check only.")
        self.apptrak_status_label.setObjectName("workspaceStripHint")
        self.apptrak_status_label.setAlignment(Qt.AlignRight | Qt.AlignVCenter)
        outer.addWidget(self.apptrak_status_label, 0, Qt.AlignRight)

        splitter = QSplitter(Qt.Horizontal)
        splitter.setChildrenCollapsible(False)
        splitter.setObjectName("mainSplitter")
        self.main_splitter = splitter
        outer.addWidget(splitter, 1)

        left_panel = QWidget()
        left_layout = QVBoxLayout(left_panel)
        left_layout.setContentsMargins(0, 0, 12, 0)
        left_layout.setSpacing(16)
        splitter.addWidget(left_panel)

        stats_row = QHBoxLayout()
        stats_row.setSpacing(12)
        left_layout.addLayout(stats_row)

        self.all_card = self._build_stat_card("All Claims")
        self.open_card = self._build_stat_card("Open Claims")
        self.closed_card = self._build_stat_card("Closed Claims")
        stats_row.addWidget(self.all_card)
        stats_row.addWidget(self.open_card)
        stats_row.addWidget(self.closed_card)

        self.tab_names = [
            "All Claims",
            "Open Claims",
            "Closed Claims",
            "Reports",
            "Office Update",
            "Route Planner",
            "Claim Tools",
            "Body Shops",
            "Insurance Companies",
        ]
        self.primary_tab_indices = [0, 1, 2, 3, 4, 5]
        self.library_tab_indices = [6, 7, 8]
        self.tab_buttons: list[QPushButton] = []
        self.primary_tab_buttons: dict[int, QPushButton] = {}

        tabs_wrap = QWidget()
        tabs_wrap_layout = QVBoxLayout(tabs_wrap)
        tabs_wrap_layout.setContentsMargins(0, 0, 0, 0)
        tabs_wrap_layout.setSpacing(8)
        left_layout.addWidget(tabs_wrap, 1)

        tabs_nav = QWidget()
        tabs_nav_layout = QHBoxLayout(tabs_nav)
        tabs_nav_layout.setContentsMargins(0, 0, 0, 0)
        tabs_nav_layout.setSpacing(8)

        for index in self.primary_tab_indices:
            tab_name = self.tab_names[index]
            button = QPushButton(tab_name)
            button.setCheckable(True)
            button.setObjectName("navTabButton")
            button.clicked.connect(lambda _checked=False, idx=index: self._set_current_tab(idx))
            tabs_nav_layout.addWidget(button)
            self.tab_buttons.append(button)
            self.primary_tab_buttons[index] = button

        tabs_nav_layout.addStretch(1)

        self.library_button = QToolButton()
        self.library_button.setObjectName("navLibraryButton")
        self.library_button.setPopupMode(QToolButton.InstantPopup)
        self.library_button.setToolButtonStyle(Qt.ToolButtonTextOnly)
        self.library_button.setCheckable(True)
        self.library_button.setText("Library")
        self.library_menu = QMenu(self.library_button)
        self.library_actions: dict[int, QAction] = {}
        for index in self.library_tab_indices:
            action = self.library_menu.addAction(self.tab_names[index])
            action.triggered.connect(lambda _checked=False, idx=index: self._set_current_tab(idx))
            self.library_actions[index] = action
        self.library_button.setMenu(self.library_menu)
        tabs_nav_layout.addWidget(self.library_button)

        self.tab_stack = QStackedWidget()
        self.tab_stack.currentChanged.connect(self._handle_tab_changed)

        tabs_wrap_layout.addWidget(tabs_nav)
        tabs_wrap_layout.addWidget(self.tab_stack, 1)

        self.all_table = self._build_claim_table()
        self.open_table = self._build_claim_table(headers=OPEN_CLAIM_TABLE_HEADERS)
        self.closed_table = self._build_claim_table()
        self.office_table = self._build_claim_table(headers=OFFICE_TABLE_HEADERS)
        self.reports_table = self._build_reports_table()
        self.office_sort_column = "Claim ID"
        self.office_sort_order = Qt.AscendingOrder

        self._set_custom_widths(self.office_table, [125, 180, 170, 140, 240, 280])

        self.tab_stack.addWidget(self._wrap_table(self.all_table))
        self.tab_stack.addWidget(self._wrap_table(self.open_table))
        self.tab_stack.addWidget(self._wrap_table(self.closed_table))
        self.tab_stack.addWidget(self._build_reports_tab())
        self.tab_stack.addWidget(self._build_office_tab())
        self.tab_stack.addWidget(self._build_route_tab())
        self.tab_stack.addWidget(self._build_claim_tools_tab())
        self.tab_stack.addWidget(self._build_body_shops_tab())
        self.tab_stack.addWidget(self._build_insurance_companies_tab())

        right_panel = QWidget()
        self.right_panel = right_panel
        right_panel.setMinimumWidth(520)
        right_panel.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Expanding)
        right_layout = QVBoxLayout(right_panel)
        right_layout.setContentsMargins(12, 0, 0, 0)
        right_layout.setSpacing(16)
        splitter.addWidget(right_panel)
        splitter.setSizes([920, 440])
        splitter.setStretchFactor(0, 3)
        splitter.setStretchFactor(1, 2)

        detail_card = QFrame()
        detail_card.setObjectName("card")
        detail_layout = QVBoxLayout(detail_card)
        detail_layout.setContentsMargins(18, 18, 18, 18)
        detail_layout.setSpacing(8)
        right_layout.addWidget(detail_card, 1)

        detail_title = QLabel("Claim Details")
        detail_title.setObjectName("sectionTitle")
        detail_layout.addWidget(detail_title)

        claim_id_row = QHBoxLayout()
        claim_id_row.setSpacing(10)
        claim_id_label = QLabel("Claim ID")
        claim_id_label.setObjectName("detailLabel")
        self.claim_id_value = QLabel("-")
        self.claim_id_value.setWordWrap(False)
        self.claim_id_value.setTextInteractionFlags(Qt.TextSelectableByMouse)
        self.claim_id_value.setObjectName("detailValueChip")
        claim_id_row.addWidget(claim_id_label, 0)
        claim_id_row.addWidget(self.claim_id_value, 0)
        claim_id_row.addStretch(1)
        detail_layout.addLayout(claim_id_row)

        form = QGridLayout()
        form.setContentsMargins(0, 2, 0, 0)
        form.setHorizontalSpacing(16)
        form.setVerticalSpacing(6)
        form.setColumnMinimumWidth(0, 140)
        form.setColumnStretch(0, 0)
        form.setColumnStretch(1, 1)
        self.detail_inputs: dict[str, QLineEdit] = {}
        detail_fields = [
            ("customer_name", "Customer"),
            ("insurance_company", "Insurance"),
            ("claim_number", "Claim #"),
            ("date_of_loss", "Date Of Loss"),
            ("town", "Town"),
            ("shop_name", "Shop"),
            ("shop_email", "Shop Email"),
            ("contact_phone", "Contact Phone"),
            ("vehicle", "Vehicle"),
            ("vin", "VIN"),
            ("owner_address", "Address"),
            ("location_of_vehicle", "Inspection Location"),
        ]

        for row, (key, label) in enumerate(detail_fields):
            label_widget = QLabel(label)
            label_widget.setObjectName("detailFormLabel")
            label_widget.setAlignment(Qt.AlignLeft | Qt.AlignVCenter)
            label_widget.setMinimumHeight(30)
            entry = QLineEdit()
            entry.setObjectName("detailInput")
            entry.setFrame(False)
            entry.setMinimumHeight(26)
            entry.setMaximumHeight(26)
            entry.setAlignment(Qt.AlignLeft | Qt.AlignBottom)
            entry.setSizePolicy(QSizePolicy.Expanding, QSizePolicy.Fixed)
            self.detail_inputs[key] = entry
            form.addWidget(label_widget, row, 0, Qt.AlignLeft | Qt.AlignVCenter)
            form.addWidget(entry, row, 1)
        detail_layout.addLayout(form)

        self.save_details_button = QPushButton("Save Details")
        self.save_details_button.setObjectName("detailPrimaryButton")
        self.save_details_button.setFixedHeight(30)
        self.save_details_button.clicked.connect(self._save_claim_details)
        detail_layout.addWidget(self.save_details_button)

        actions_row = QHBoxLayout()
        actions_row.setContentsMargins(0, 0, 0, 0)
        actions_row.setSpacing(6)
        self.open_assign_button = QPushButton("Open Assignment Sheet")
        self.open_assign_button.setObjectName("detailActionButton")
        self.open_assign_button.setFixedHeight(24)
        self.open_assign_button.setMaximumWidth(210)
        self.open_assign_button.clicked.connect(self._open_assignment_sheet)
        self.open_folder_button = QPushButton("Open Claim Folder")
        self.open_folder_button.setObjectName("detailActionButton")
        self.open_folder_button.setFixedHeight(24)
        self.open_folder_button.setMaximumWidth(190)
        self.open_folder_button.clicked.connect(self._open_claim_folder)
        actions_row.addWidget(self.open_assign_button)
        actions_row.addWidget(self.open_folder_button)
        detail_layout.addLayout(actions_row)

        status_row = QHBoxLayout()
        status_row.setContentsMargins(0, 0, 0, 0)
        status_row.setSpacing(6)
        self.reopen_button = QPushButton("Reopen")
        self.reopen_button.setObjectName("detailActionButton")
        self.reopen_button.setFixedHeight(24)
        self.reopen_button.setMaximumWidth(140)
        self.reopen_button.clicked.connect(self._reopen_current_claim)
        self.mark_closed_button = QPushButton("Mark Closed")
        self.mark_closed_button.setObjectName("detailActionButton")
        self.mark_closed_button.setFixedHeight(24)
        self.mark_closed_button.setMaximumWidth(160)
        self.mark_closed_button.clicked.connect(self._mark_current_claim_closed)
        status_row.addWidget(self.reopen_button)
        status_row.addWidget(self.mark_closed_button)
        detail_layout.addLayout(status_row)

        notes_card = QFrame()
        notes_card.setObjectName("card")
        notes_layout = QVBoxLayout(notes_card)
        notes_layout.setContentsMargins(16, 14, 16, 16)
        notes_layout.setSpacing(8)
        right_layout.addWidget(notes_card, 1)

        notes_title_row = QHBoxLayout()
        notes_title_row.setSpacing(8)
        notes_title = QLabel("Assignment Notes")
        notes_title.setObjectName("sectionTitle")
        notes_title_row.addWidget(notes_title)
        notes_title_row.addStretch(1)
        self.claim_notes_button = QPushButton("Claim Notes")
        self.claim_notes_button.setObjectName("compactActionButton")
        self.claim_notes_button.clicked.connect(self._open_claim_notes_dialog)
        notes_title_row.addWidget(self.claim_notes_button)
        self.save_notes_button = QPushButton("Save Notes")
        self.save_notes_button.setObjectName("compactActionButton")
        self.save_notes_button.clicked.connect(self._save_assignment_notes)
        notes_title_row.addWidget(self.save_notes_button)
        notes_layout.addLayout(notes_title_row)

        self.notes_view = QTextEdit()
        self.notes_view.setReadOnly(False)
        notes_layout.addWidget(self.notes_view, 1)

        quit_action = QAction("Quit", self)
        quit_action.triggered.connect(self.close)
        self.addAction(quit_action)
        quit_action.setShortcut("Ctrl+Q")

        self._apply_styles()
        self._set_current_tab(0)

    def _build_stat_card(self, label: str) -> QFrame:
        card = QFrame()
        card.setObjectName("statCard")
        layout = QVBoxLayout(card)
        layout.setContentsMargins(14, 8, 14, 8)
        layout.setSpacing(2)

        title = QLabel(label)
        title.setObjectName("statTitle")
        value = QLabel("0")
        value.setObjectName("statValue")
        layout.addWidget(title)
        layout.addWidget(value)
        card.value_label = value  # type: ignore[attr-defined]
        return card

    def _set_current_tab(self, index: int) -> None:
        if not hasattr(self, "tab_stack"):
            return
        self.tab_stack.setCurrentIndex(index)
        for button_index, button in getattr(self, "primary_tab_buttons", {}).items():
            button.blockSignals(True)
            button.setChecked(button_index == index)
            button.blockSignals(False)
        if hasattr(self, "library_button"):
            is_library = index in getattr(self, "library_tab_indices", [])
            self.library_button.blockSignals(True)
            self.library_button.setChecked(is_library)
            self.library_button.setText(self.tab_names[index] if is_library else "Library")
            self.library_button.blockSignals(False)

    def _handle_tab_changed(self, index: int) -> None:
        if not hasattr(self, "tab_stack") or not hasattr(self, "right_panel") or not hasattr(self, "main_splitter"):
            return
        tab_label = self.tab_names[index] if 0 <= index < len(self.tab_names) else ""
        for button_index, button in getattr(self, "primary_tab_buttons", {}).items():
            button.blockSignals(True)
            button.setChecked(button_index == index)
            button.blockSignals(False)
        if hasattr(self, "library_button"):
            is_library = index in getattr(self, "library_tab_indices", [])
            self.library_button.blockSignals(True)
            self.library_button.setChecked(is_library)
            self.library_button.setText(tab_label if is_library else "Library")
            self.library_button.blockSignals(False)
        full_width_tabs = {"Reports", "Route Planner", "Claim Tools", "Body Shops", "Insurance Companies"}
        show_right = tab_label not in full_width_tabs
        self.right_panel.setVisible(show_right)
        if show_right:
            self.main_splitter.setSizes([920, 440])
        else:
            self.main_splitter.setSizes([1380, 0])

    def _build_claim_table(self, include_office: bool = False, headers: list[str] | None = None) -> QTableWidget:
        if headers is None:
            headers = list(CLAIM_TABLE_HEADERS)
            if include_office:
                headers.extend(["Inspection", "Office Notes"])
        table = QTableWidget(0, len(headers))
        table.setHorizontalHeaderLabels(headers)
        table.setSortingEnabled(True)
        table.setSelectionBehavior(QTableWidget.SelectRows)
        table.setSelectionMode(QTableWidget.SingleSelection)
        table.setEditTriggers(QTableWidget.NoEditTriggers)
        table.verticalHeader().setVisible(False)
        table.setAlternatingRowColors(True)
        table.setWordWrap(False)
        table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        table.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        table.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
        table.itemSelectionChanged.connect(lambda current_table=table: self._handle_table_selection(current_table))
        table.cellDoubleClicked.connect(lambda row, column, current_table=table: self._handle_claim_table_double_click(current_table, row, column))
        table.setContextMenuPolicy(Qt.CustomContextMenu)
        table.customContextMenuRequested.connect(lambda pos, current_table=table: self._show_claim_context_menu(current_table, pos))
        header = table.horizontalHeader()
        header.setStretchLastSection(False)
        header.setMinimumSectionSize(90)
        header.setSectionsMovable(True)
        header.setDefaultAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        for index in range(len(headers)):
            header.setSectionResizeMode(index, QHeaderView.Interactive)
        if headers in (CLAIM_TABLE_HEADERS, OPEN_CLAIM_TABLE_HEADERS) or (include_office and len(headers) == len(CLAIM_TABLE_HEADERS) + 2):
            self._set_table_column_widths(table, include_office)
        return table

    def _build_reports_table(self) -> QTableWidget:
        table = QTableWidget(0, len(REPORT_TABLE_HEADERS))
        table.setHorizontalHeaderLabels(REPORT_TABLE_HEADERS)
        table.setSelectionBehavior(QTableWidget.SelectRows)
        table.setSelectionMode(QTableWidget.SingleSelection)
        table.setEditTriggers(QTableWidget.NoEditTriggers)
        table.verticalHeader().setVisible(False)
        table.setAlternatingRowColors(True)
        table.setWordWrap(False)
        table.setSortingEnabled(True)
        table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        table.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
        table.itemSelectionChanged.connect(lambda current_table=table: self._handle_table_selection(current_table))
        table.cellDoubleClicked.connect(lambda row, column, current_table=table: self._handle_claim_table_double_click(current_table, row, column))
        table.setContextMenuPolicy(Qt.CustomContextMenu)
        table.customContextMenuRequested.connect(lambda pos, current_table=table: self._show_claim_context_menu(current_table, pos))
        header = table.horizontalHeader()
        header.setStretchLastSection(False)
        header.setSectionsMovable(True)
        header.setDefaultAlignment(Qt.AlignLeft | Qt.AlignVCenter)
        for index in range(len(REPORT_TABLE_HEADERS)):
            header.setSectionResizeMode(index, QHeaderView.Interactive)
        self._set_custom_widths(table, [120, 120, 190, 220, 110])
        return table

    def _set_table_column_widths(self, table: QTableWidget, include_office: bool = False) -> None:
        widths = [150, 190, 210, 210, 150, 120, 120, 110]
        if table.columnCount() == len(OPEN_CLAIM_TABLE_HEADERS):
            widths.append(180)
        if include_office:
            widths.extend([135, 260])
        for index, width in enumerate(widths):
            table.setColumnWidth(index, width)

    def _set_custom_widths(self, table: QTableWidget, widths: list[int]) -> None:
        for index, width in enumerate(widths):
            table.setColumnWidth(index, width)

    def _auto_size_table_columns(self, table: QTableWidget, include_office: bool = False) -> None:
        minimums = [95, 130, 150, 125, 115, 85, 90, 80]
        maximums = [170, 230, 280, 240, 200, 135, 120, 120]
        if table.columnCount() == len(OPEN_CLAIM_TABLE_HEADERS):
            minimums.append(130)
            maximums.append(220)
        if include_office:
            minimums.extend([110, 150])
            maximums.extend([155, 280])

        header = table.horizontalHeader()
        header_metrics = QFontMetrics(header.font())
        cell_metrics = QFontMetrics(table.font())
        for index in range(table.columnCount()):
            header_text = table.horizontalHeaderItem(index).text() if table.horizontalHeaderItem(index) else ""
            header_width = header_metrics.horizontalAdvance(header_text) + 24

            sample_widths: list[int] = []
            row_count = table.rowCount()
            for row in range(row_count):
                item = table.item(row, index)
                if not item:
                    continue
                text = item.text().strip()
                if not text:
                    continue
                sample_widths.append(cell_metrics.horizontalAdvance(text) + 22)

            if sample_widths:
                sample_widths.sort()
                median_width = sample_widths[len(sample_widths) // 2]
                if len(sample_widths) % 2 == 0 and len(sample_widths) > 1:
                    median_width = int((sample_widths[(len(sample_widths) // 2) - 1] + sample_widths[len(sample_widths) // 2]) / 2)
                target = max(header_width, median_width)
            else:
                target = header_width

            target = max(minimums[index], min(target, maximums[index]))
            table.setColumnWidth(index, target)

    def _auto_size_columns_with_bounds(self, table: QTableWidget, minimums: list[int], maximums: list[int]) -> None:
        header = table.horizontalHeader()
        header_metrics = QFontMetrics(header.font())
        cell_metrics = QFontMetrics(table.font())
        for index in range(table.columnCount()):
            header_text = table.horizontalHeaderItem(index).text() if table.horizontalHeaderItem(index) else ""
            header_width = header_metrics.horizontalAdvance(header_text) + 18
            sample_widths: list[int] = []
            for row in range(table.rowCount()):
                item = table.item(row, index)
                if not item:
                    continue
                text = item.text().strip()
                if text and text != "-":
                    sample_widths.append(cell_metrics.horizontalAdvance(text) + 20)
            if sample_widths:
                sample_widths.sort()
                median_width = sample_widths[len(sample_widths) // 2]
                if len(sample_widths) % 2 == 0 and len(sample_widths) > 1:
                    median_width = int((sample_widths[(len(sample_widths) // 2) - 1] + sample_widths[len(sample_widths) // 2]) / 2)
                target = max(header_width, median_width)
            else:
                target = header_width
            target = max(minimums[index], min(target, maximums[index]))
            table.setColumnWidth(index, target)

    def _apply_office_sort_indicator(self) -> None:
        header = self.office_table.horizontalHeader()
        if self.office_sort_column in OFFICE_TABLE_HEADERS:
            header.setSortIndicatorShown(True)
            header.setSortIndicator(OFFICE_TABLE_HEADERS.index(self.office_sort_column), self.office_sort_order)
        else:
            header.setSortIndicatorShown(False)

    def _handle_office_header_click(self, column_name: str) -> None:
        if self.office_sort_column == column_name:
            self.office_sort_order = Qt.DescendingOrder if self.office_sort_order == Qt.AscendingOrder else Qt.AscendingOrder
        else:
            self.office_sort_column = column_name
            self.office_sort_order = Qt.AscendingOrder
        self._populate_office_tables(self.filtered_open)

    def _office_sort_value(self, claim: ClaimView, column_name: str) -> str:
        mapping = {
            "Claim ID": claim.claim_id,
            "Customer": claim.display_customer,
            "Shop": claim.shop_name,
            "Inspection When": claim.office_appt_when,
            "Progress": claim.office_progress_status or "No update",
            "Additional Notes": claim.office_additional_notes,
        }
        return (mapping.get(column_name, "") or "").lower()

    def _populate_office_tables(self, claims: list[ClaimView]) -> None:
        sorted_claims = sorted(
            claims,
            key=lambda claim: self._office_sort_value(claim, self.office_sort_column),
            reverse=self.office_sort_order == Qt.DescendingOrder,
        )

        self.office_table.blockSignals(True)
        self.office_table.setSortingEnabled(False)
        self.office_table.setRowCount(len(sorted_claims))

        for row, claim in enumerate(sorted_claims):
            values = [
                claim.claim_id,
                claim.display_customer,
                claim.shop_name,
                claim.office_appt_when,
                claim.office_progress_status or "No update",
                claim.office_additional_notes or "",
            ]

            for col, value in enumerate(values):
                item = QTableWidgetItem(value or "-")
                item.setData(Qt.UserRole, claim.key)
                item.setToolTip(value or "-")
                if col >= 4:
                    font = item.font()
                    font.setBold(True)
                    item.setFont(font)
                self.office_table.setItem(row, col, item)

        self._auto_size_columns_with_bounds(
            self.office_table,
            [100, 150, 150, 130, 220, 240],
            [170, 260, 230, 200, 340, 420],
        )
        self._apply_office_sort_indicator()
        self.office_table.setSortingEnabled(True)
        self.office_table.blockSignals(False)

    def _wrap_table(self, table: QTableWidget) -> QWidget:
        wrapper = QWidget()
        layout = QVBoxLayout(wrapper)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.addWidget(table)
        return wrapper

    def _build_office_tab(self) -> QWidget:
        wrapper = QWidget()
        layout = QVBoxLayout(wrapper)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(12)

        office_splitter = QSplitter(Qt.Vertical)
        office_splitter.setChildrenCollapsible(False)
        office_splitter.setObjectName("officeSplitter")
        layout.addWidget(office_splitter, 1)

        table_card = QFrame()
        table_card.setObjectName("card")
        table_layout = QVBoxLayout(table_card)
        table_layout.setContentsMargins(18, 12, 18, 14)
        table_layout.setSpacing(6)
        office_splitter.addWidget(table_card)

        table_title = QLabel("Open Claims")
        table_title.setObjectName("sectionTitle")
        table_layout.addWidget(table_title)

        self.office_table.setVerticalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.office_table.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        table_layout.addWidget(self.office_table, 1)

        lower_row = QHBoxLayout()
        lower_row.setSpacing(0)
        lower_wrapper = QWidget()
        lower_wrapper.setMinimumHeight(220)
        lower_wrapper.setSizePolicy(QSizePolicy.Preferred, QSizePolicy.Minimum)
        lower_wrapper.setLayout(lower_row)
        office_splitter.addWidget(lower_wrapper)
        office_splitter.setSizes([820, 40])

        office_card = QFrame()
        office_card.setObjectName("card")
        office_card.setMinimumHeight(220)
        office_layout = QVBoxLayout(office_card)
        office_layout.setContentsMargins(10, 8, 10, 16)
        office_layout.setSpacing(6)
        lower_row.addWidget(office_card, 1)

        office_title_row = QHBoxLayout()
        office_title_row.setContentsMargins(0, 0, 0, 6)
        office_title_row.setSpacing(10)
        office_title = QLabel("Office Update")
        office_title.setObjectName("sectionTitle")
        office_title_row.addWidget(office_title)
        office_title_row.addStretch(1)

        self.office_preview_button = QPushButton("Preview PDF")
        self.office_preview_button.setObjectName("compactActionButton")
        self.office_preview_button.clicked.connect(self._preview_office_update_pdf)
        office_title_row.addWidget(self.office_preview_button, 0, Qt.AlignRight)

        self.office_email_button = QPushButton("Email Office Update")
        self.office_email_button.setObjectName("compactActionButton")
        self.office_email_button.clicked.connect(self._email_office_update_pdf)
        office_title_row.addWidget(self.office_email_button, 0, Qt.AlignRight)

        self.office_uninspected_button = QPushButton("Email Uninspected Only")
        self.office_uninspected_button.setObjectName("compactActionButton")
        self.office_uninspected_button.clicked.connect(self._email_uninspected_only_to_self)
        office_title_row.addWidget(self.office_uninspected_button, 0, Qt.AlignRight)

        office_form = QFormLayout()
        office_form.setHorizontalSpacing(8)
        office_form.setVerticalSpacing(6)
        self.office_progress_input = QComboBox()
        self.office_progress_input.addItems(
            [
                "No Contact Yet",
                "Contacted",
                "Appointment Scheduled",
                "Seen - Need To Write",
                "Written - Under Review",
                "Supplement - Waiting For Paperwork",
                "Waiting For Paperwork",
            ]
        )
        self.office_appt_input = QLineEdit()
        self.office_progress_input.setObjectName("officeCompactInput")
        self.office_appt_input.setObjectName("officeCompactInput")
        self.office_progress_input.setMaximumHeight(28)
        self.office_appt_input.setMaximumHeight(28)
        self.office_notes_input = QTextEdit()
        self.office_notes_input.setMinimumHeight(48)
        self.office_notes_input.setMaximumHeight(64)
        office_form.addRow("Progress", self.office_progress_input)
        office_form.addRow("Date Of Inspection", self.office_appt_input)
        office_form.addRow("Additional Notes", self.office_notes_input)

        self.save_office_button = QPushButton("Save Office Update")
        self.save_office_button.setObjectName("officeSaveButton")
        self.save_office_button.setFixedHeight(24)
        self.save_office_button.setFixedWidth(230)
        self.save_office_button.clicked.connect(self._save_office_update)
        office_title_row.addWidget(self.save_office_button, 0, Qt.AlignRight)
        office_layout.addLayout(office_title_row)
        office_layout.addLayout(office_form)
        office_layout.addSpacing(4)

        return wrapper

    def _build_email_destinations_text(self) -> str:
        lines = [
            "Office Update",
            f"  To: {self._office_update_email_to()}",
            f"  Cc: {self._office_update_email_cc() or '-'}",
            "",
            "Review Email",
            f"  To: {self._office_review_email_to()}",
            "  Greeting: Hi Joe,",
            "",
            "RMC Email",
            f"  To: {self._office_rmc_email_to()}",
            "  Greeting: Hi Gina,",
            "",
            "Supplement Request",
            f"  To: {self._office_supplement_email_to()}",
            "  Greeting: Hi Lisa,",
            "",
            "Route Email",
            "  To: your configured sender email",
            "",
            "Office Update Self-Copy",
            "  To: your configured sender email",
            "",
            "Assignment Sheet",
            "  To: your configured sender email",
            "",
            "Prelim Email",
            "  To: selected claim shop email",
            "",
            "Shop Draft",
            "  To: shop email you enter in Outlook",
        ]
        return "\n".join(lines)

    def _copy_email_destinations(self) -> None:
        QApplication.clipboard().setText(self._build_email_destinations_text())
        QMessageBox.information(self, APP_NAME, "Email destination list copied to clipboard.")

    def _get_email_setting(self, key: str, default: str) -> str:
        return str(self.settings.get(key, "") or "").strip() or default

    def _office_update_email_to(self) -> str:
        return self._get_email_setting("office_update_email_to", OFFICE_UPDATE_EMAIL_TO)

    def _office_update_email_cc(self) -> str:
        return self._get_email_setting("office_update_email_cc", OFFICE_UPDATE_EMAIL_CC)

    def _office_review_email_to(self) -> str:
        return self._get_email_setting("office_review_email_to", OFFICE_REVIEW_EMAIL_TO)

    def _office_rmc_email_to(self) -> str:
        return self._get_email_setting("office_rmc_email_to", OFFICE_RMC_EMAIL_TO)

    def _office_supplement_email_to(self) -> str:
        return self._get_email_setting("office_supplement_email_to", OFFICE_SUPPLEMENT_EMAIL_TO)

    def _build_workspace_settings_text(self) -> str:
        watched_folders = [str(folder) for folder in self.settings.get("watched_folders", [])]
        claim_tools_folder = str(self.settings.get("claim_tools_folder", "") or "").strip() or "-"
        lines = ["Watched Folders"]
        if watched_folders:
            for index, folder in enumerate(watched_folders, start=1):
                lines.append(f"  {index}. {folder}")
        else:
            lines.append("  No watched folders configured.")
        lines.extend(["", "Claim Tools Folder", f"  {claim_tools_folder}"])
        return "\n".join(lines)

    def _build_future_settings_text(self) -> str:
        return "\n".join(
            [
                f"Desktop Version: {APP_VERSION}",
                f"Desktop Update Channel: {str(self.desktop_update_config.get('channel', 'stable') or 'stable')}",
                f"Desktop Update Source: {str(self.desktop_update_config.get('manifest_url', '') or '-')}",
                "",
                "Planned sections",
                "  User Profile",
                "  Email Preferences",
                "  Notification Preferences",
                "  Appearance / Layout",
                "",
                "These are placeholders for future settings so we can keep expanding from one home base.",
            ]
        )

    def _open_settings_dialog(self) -> None:
        dialog = QDialog(self)
        dialog.setWindowTitle(f"{APP_NAME} Settings")
        dialog.resize(860, 760)

        layout = QVBoxLayout(dialog)
        layout.setContentsMargins(16, 16, 16, 16)
        layout.setSpacing(12)

        title = QLabel("Settings")
        title.setObjectName("sectionTitle")
        layout.addWidget(title)

        workspace_card = QFrame()
        workspace_card.setObjectName("card")
        workspace_layout = QVBoxLayout(workspace_card)
        workspace_layout.setContentsMargins(14, 12, 14, 14)
        workspace_layout.setSpacing(8)
        layout.addWidget(workspace_card, 1)

        workspace_header = QHBoxLayout()
        workspace_header.setSpacing(10)
        workspace_title = QLabel("Workspace")
        workspace_title.setObjectName("sectionTitle")
        workspace_header.addWidget(workspace_title)
        workspace_header.addStretch(1)

        edit_workspace_button = QPushButton("Edit Folder")
        edit_workspace_button.setObjectName("compactActionButton")
        remove_folder_button = QPushButton("Remove Watched Folder")
        remove_folder_button.setObjectName("compactActionButton")
        workspace_header.addWidget(edit_workspace_button)
        workspace_header.addWidget(remove_folder_button)
        workspace_layout.addLayout(workspace_header)

        workspace_help = QLabel("Manage watched folders and the Claim Tools folder from here.")
        workspace_help.setObjectName("workspaceStripHint")
        workspace_layout.addWidget(workspace_help)

        workspace_view = QPlainTextEdit()
        workspace_view.setReadOnly(True)
        workspace_view.setPlainText(self._build_workspace_settings_text())
        workspace_layout.addWidget(workspace_view, 1)

        email_card = QFrame()
        email_card.setObjectName("card")
        email_layout = QVBoxLayout(email_card)
        email_layout.setContentsMargins(14, 12, 14, 14)
        email_layout.setSpacing(8)
        layout.addWidget(email_card, 1)

        email_header = QHBoxLayout()
        email_header.setSpacing(10)
        email_title = QLabel("Email Destinations")
        email_title.setObjectName("sectionTitle")
        email_header.addWidget(email_title)
        email_header.addStretch(1)

        save_email_button = QPushButton("Save Email Settings")
        save_email_button.setObjectName("compactActionButton")
        reset_email_button = QPushButton("Reset Defaults")
        reset_email_button.setObjectName("compactActionButton")
        copy_email_button = QPushButton("Copy List")
        copy_email_button.setObjectName("compactActionButton")
        email_header.addWidget(save_email_button)
        email_header.addWidget(reset_email_button)
        email_header.addWidget(copy_email_button)
        email_layout.addLayout(email_header)

        email_help = QLabel("Edit the addresses used by each outgoing email flow. Leave a field blank to fall back to the built-in default.")
        email_help.setObjectName("workspaceStripHint")
        email_help.setWordWrap(True)
        email_layout.addWidget(email_help)

        email_form = QFormLayout()
        email_form.setHorizontalSpacing(12)
        email_form.setVerticalSpacing(8)

        office_update_to_input = QLineEdit(self._office_update_email_to())
        office_update_cc_input = QLineEdit(self._office_update_email_cc())
        review_to_input = QLineEdit(self._office_review_email_to())
        rmc_to_input = QLineEdit(self._office_rmc_email_to())
        supplement_to_input = QLineEdit(self._office_supplement_email_to())

        email_form.addRow("Office Update To", office_update_to_input)
        email_form.addRow("Office Update Cc", office_update_cc_input)
        email_form.addRow("Review Email To", review_to_input)
        email_form.addRow("RMC Email To", rmc_to_input)
        email_form.addRow("Supplement Request To", supplement_to_input)
        email_layout.addLayout(email_form)

        future_card = QFrame()
        future_card.setObjectName("card")
        future_layout = QVBoxLayout(future_card)
        future_layout.setContentsMargins(14, 12, 14, 14)
        future_layout.setSpacing(8)
        layout.addWidget(future_card, 0)

        future_title = QLabel("Future Settings")
        future_title.setObjectName("sectionTitle")
        future_layout.addWidget(future_title)

        desktop_update_row = QHBoxLayout()
        desktop_update_row.setSpacing(10)
        desktop_update_label = QLabel(f"Desktop Version {APP_VERSION}")
        desktop_update_label.setObjectName("workspaceStripHint")
        desktop_update_row.addWidget(desktop_update_label)
        desktop_update_row.addStretch(1)
        check_update_button = QPushButton("Check for Desktop Update")
        check_update_button.setObjectName("compactActionButton")
        desktop_update_row.addWidget(check_update_button)
        future_layout.addLayout(desktop_update_row)

        future_view = QPlainTextEdit()
        future_view.setReadOnly(True)
        future_view.setMaximumHeight(150)
        future_view.setPlainText(self._build_future_settings_text())
        future_layout.addWidget(future_view)

        button_box = QDialogButtonBox(QDialogButtonBox.Close)
        button_box.rejected.connect(dialog.reject)
        button_box.accepted.connect(dialog.accept)
        layout.addWidget(button_box)

        def collect_email_settings() -> dict[str, str]:
            return {
                "office_update_email_to": office_update_to_input.text().strip(),
                "office_update_email_cc": office_update_cc_input.text().strip(),
                "office_review_email_to": review_to_input.text().strip(),
                "office_rmc_email_to": rmc_to_input.text().strip(),
                "office_supplement_email_to": supplement_to_input.text().strip(),
            }

        def refresh_dialog_views() -> None:
            workspace_view.setPlainText(self._build_workspace_settings_text())

        def edit_workspace() -> None:
            self._edit_workspace_settings()
            refresh_dialog_views()

        def remove_folder() -> None:
            self._remove_selected_watched_folder()
            refresh_dialog_views()

        def save_email_settings() -> None:
            new_settings = collect_email_settings()
            self.repo.save_email_settings(new_settings)
            self.settings.update(new_settings)
            QMessageBox.information(dialog, APP_NAME, "Email settings saved.")

        def reset_email_defaults() -> None:
            office_update_to_input.setText(OFFICE_UPDATE_EMAIL_TO)
            office_update_cc_input.setText(OFFICE_UPDATE_EMAIL_CC)
            review_to_input.setText(OFFICE_REVIEW_EMAIL_TO)
            rmc_to_input.setText(OFFICE_RMC_EMAIL_TO)
            supplement_to_input.setText(OFFICE_SUPPLEMENT_EMAIL_TO)

        def copy_email_settings() -> None:
            snapshot = collect_email_settings()
            original_settings = self.settings
            try:
                self.settings = {**self.settings, **snapshot}
                QApplication.clipboard().setText(self._build_email_destinations_text())
            finally:
                self.settings = original_settings
            QMessageBox.information(dialog, APP_NAME, "Email destination list copied to clipboard.")

        edit_workspace_button.clicked.connect(edit_workspace)
        remove_folder_button.clicked.connect(remove_folder)
        save_email_button.clicked.connect(save_email_settings)
        reset_email_button.clicked.connect(reset_email_defaults)
        copy_email_button.clicked.connect(copy_email_settings)
        check_update_button.clicked.connect(lambda: self._check_for_desktop_update(manual=True))

        dialog.exec()

    def _check_for_desktop_updates_on_launch(self) -> None:
        self._check_for_desktop_update(manual=False)

    def _check_for_desktop_update(self, manual: bool) -> None:
        self.desktop_update_config = _load_desktop_update_config()
        if not getattr(sys, "frozen", False):
            if manual:
                QMessageBox.information(self, APP_NAME, "Desktop updates are only available from the installed desktop app build.")
            return
        if not bool(self.desktop_update_config.get("enabled", True)):
            if manual:
                QMessageBox.information(self, APP_NAME, "Desktop updates are disabled for this install.")
            return
        manifest_url = str(self.desktop_update_config.get("manifest_url", "") or "").strip()
        if not manifest_url:
            if manual:
                QMessageBox.warning(self, APP_NAME, "No desktop update manifest URL is configured for this install.")
            return
        if self._update_check_in_progress:
            if manual:
                QMessageBox.information(self, APP_NAME, "A desktop update check is already running.")
            return
        self._update_check_in_progress = True
        threading.Thread(
            target=self._desktop_update_manifest_worker,
            args=(manifest_url, manual),
            daemon=True,
        ).start()

    def _desktop_update_manifest_worker(self, manifest_url: str, manual: bool) -> None:
        result: dict[str, object] = {"manual": manual}
        try:
            request = Request(
                manifest_url,
                headers={"User-Agent": f"ClaimManager3Desktop/{APP_VERSION}"},
            )
            with urlopen(request, timeout=12) as response:
                payload = json.loads(response.read().decode("utf-8"))
            if "content" in payload and str(payload.get("encoding", "") or "").lower() == "base64":
                decoded = base64.b64decode(str(payload.get("content", "") or ""))
                payload = json.loads(decoded.decode("utf-8"))
            version = str(payload.get("version", "") or "").strip()
            installer_url = str(payload.get("installer_url", "") or "").strip()
            if not version or not installer_url:
                raise ValueError("Update manifest is missing version or installer_url.")
            result["manifest"] = {
                "version": version,
                "installer_url": installer_url,
                "notes": str(payload.get("notes", "") or "").strip(),
                "published_at": str(payload.get("published_at", "") or "").strip(),
            }
        except Exception as exc:
            result["error"] = str(exc)
        self.update_manifest_ready.emit(result)

    def _handle_update_manifest_result(self, result: object) -> None:
        payload = result if isinstance(result, dict) else {}
        self._update_check_in_progress = False
        manual = bool(payload.get("manual", False))
        error = str(payload.get("error", "") or "").strip()
        if error:
            if manual:
                QMessageBox.warning(self, APP_NAME, f"Could not check for a desktop update.\n\n{error}")
            return
        manifest = payload.get("manifest")
        if not isinstance(manifest, dict):
            if manual:
                QMessageBox.warning(self, APP_NAME, "Desktop update check returned an invalid manifest.")
            return
        remote_version = str(manifest.get("version", "") or "").strip()
        if _version_key(remote_version) <= _version_key(APP_VERSION):
            if manual:
                QMessageBox.information(self, APP_NAME, f"You are already on the latest desktop version.\n\nCurrent version: {APP_VERSION}")
            return
        message = [
            "A desktop update is available.",
            "",
            f"Current version: {APP_VERSION}",
            f"New version: {remote_version}",
        ]
        notes = str(manifest.get("notes", "") or "").strip()
        if notes:
            message.extend(["", notes])
        message.extend(["", "Download and install it now?"])
        if QMessageBox.question(
            self,
            APP_NAME,
            "\n".join(message),
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.Yes,
        ) != QMessageBox.Yes:
            return
        self._download_desktop_update(manifest)

    def _download_desktop_update(self, manifest: dict[str, object]) -> None:
        if self._update_download_in_progress:
            QMessageBox.information(self, APP_NAME, "A desktop update download is already running.")
            return
        self._update_download_in_progress = True
        threading.Thread(
            target=self._desktop_update_download_worker,
            args=(dict(manifest),),
            daemon=True,
        ).start()

    def _desktop_update_download_worker(self, manifest: dict[str, object]) -> None:
        result: dict[str, object] = {"manifest": manifest}
        try:
            version = str(manifest.get("version", "") or "").strip() or "update"
            installer_url = str(manifest.get("installer_url", "") or "").strip()
            if not installer_url:
                raise ValueError("The update manifest did not include an installer URL.")
            download_dir = Path(tempfile.gettempdir()) / "claim_manager_3_updates"
            download_dir.mkdir(parents=True, exist_ok=True)
            installer_path = download_dir / f"Claim-Manager-3-Setup-{version}.exe"
            request = Request(
                installer_url,
                headers={"User-Agent": f"ClaimManager3Desktop/{APP_VERSION}"},
            )
            with urlopen(request, timeout=60) as response, installer_path.open("wb") as output_file:
                shutil.copyfileobj(response, output_file)
            result["installer_path"] = str(installer_path)
        except Exception as exc:
            result["error"] = str(exc)
        self.update_download_ready.emit(result)

    def _handle_update_download_result(self, result: object) -> None:
        payload = result if isinstance(result, dict) else {}
        self._update_download_in_progress = False
        error = str(payload.get("error", "") or "").strip()
        if error:
            QMessageBox.warning(self, APP_NAME, f"Could not download the desktop update.\n\n{error}")
            return
        installer_path = Path(str(payload.get("installer_path", "") or "").strip())
        if not installer_path.exists():
            QMessageBox.warning(self, APP_NAME, "The desktop update finished downloading, but the installer could not be found.")
            return
        manifest = payload.get("manifest")
        version = ""
        if isinstance(manifest, dict):
            version = str(manifest.get("version", "") or "").strip()
        message = "The new desktop installer is ready."
        if version:
            message += f"\n\nVersion: {version}"
        message += "\n\nClaim Manager will close and start the installer."
        if QMessageBox.question(
            self,
            APP_NAME,
            message,
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.Yes,
        ) != QMessageBox.Yes:
            return
        self._launch_desktop_update_installer(installer_path)

    def _launch_desktop_update_installer(self, installer_path: Path) -> None:
        runner_path = Path(tempfile.gettempdir()) / f"claim-manager-update-{int(time.time())}.cmd"
        runner_path.write_text(
            "\n".join(
                [
                    "@echo off",
                    "setlocal",
                    "timeout /t 2 /nobreak >nul",
                    f"start \"\" /wait \"{installer_path}\" /SP- /VERYSILENT /SUPPRESSMSGBOXES /NORESTART /CLOSEAPPLICATIONS /FORCECLOSEAPPLICATIONS",
                    "endlocal",
                ]
            ),
            encoding="ascii",
        )
        creation_flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
        subprocess.Popen(["cmd.exe", "/c", str(runner_path)], creationflags=creation_flags)
        QApplication.instance().quit()

    def _build_reports_tab(self) -> QWidget:
        wrapper = QWidget()
        layout = QVBoxLayout(wrapper)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(12)

        filters_card = QFrame()
        filters_card.setObjectName("card")
        filters_layout = QGridLayout(filters_card)
        filters_layout.setContentsMargins(16, 12, 16, 12)
        filters_layout.setHorizontalSpacing(12)
        filters_layout.setVerticalSpacing(6)
        filters_layout.setColumnStretch(1, 1)
        filters_layout.setColumnStretch(3, 1)
        layout.addWidget(filters_card, 0)

        filters_title = QLabel("Filters")
        filters_title.setObjectName("sectionTitle")
        filters_layout.addWidget(filters_title, 0, 0, 1, 5)

        self.report_specific_date_input = self._build_report_date_picker("Any")
        self.report_specific_date_input.dateChanged.connect(self._on_reports_specific_date_changed)
        self.report_start_input = self._build_report_date_picker("Start")
        self.report_end_input = self._build_report_date_picker("End")

        filters_layout.addWidget(QLabel("Specific Date"), 1, 0)
        filters_layout.addWidget(self.report_specific_date_input, 1, 1)
        filters_layout.addWidget(QLabel("Start Date"), 1, 2)
        filters_layout.addWidget(self.report_start_input, 1, 3)
        self.report_apply_button = QPushButton("Apply Date Range")
        self.report_apply_button.clicked.connect(self._apply_reports_range_filter)
        filters_layout.addWidget(self.report_apply_button, 1, 4)

        filters_layout.addWidget(QLabel("End Date"), 2, 2)
        filters_layout.addWidget(self.report_end_input, 2, 3)
        self.report_clear_button = QPushButton("Clear Filters")
        self.report_clear_button.clicked.connect(self._clear_reports_filters)
        filters_layout.addWidget(self.report_clear_button, 2, 4)

        body_frame = QFrame()
        body_frame.setObjectName("reportsBody")
        body_layout = QGridLayout(body_frame)
        body_layout.setContentsMargins(0, 0, 0, 0)
        body_layout.setHorizontalSpacing(12)
        body_layout.setVerticalSpacing(0)
        body_layout.setColumnStretch(0, 0)
        body_layout.setColumnStretch(1, 1)
        layout.addWidget(body_frame, 1)

        stats_card = QFrame()
        stats_card.setObjectName("card")
        stats_card.setMinimumWidth(300)
        stats_layout = QGridLayout(stats_card)
        stats_layout.setContentsMargins(16, 14, 16, 14)
        stats_layout.setHorizontalSpacing(12)
        stats_layout.setVerticalSpacing(6)
        body_layout.addWidget(stats_card, 0, 0)

        stats_title = QLabel("Closed Volume")
        stats_title.setObjectName("sectionTitle")
        stats_layout.addWidget(stats_title, 0, 0, 1, 2)

        stat_pairs = [
            ("Closed Today", "today"),
            ("Closed This Week", "week"),
            ("Closed This Month", "month"),
            ("Selected Date", "selected_date"),
            ("Count For Date", "selected_count"),
        ]
        self.report_stats_labels: dict[str, QLabel] = {}
        for row, (label, key) in enumerate(stat_pairs, start=1):
            stats_layout.addWidget(QLabel(label), row, 0)
            value_label = QLabel("0" if key != "selected_date" else "-")
            value_label.setObjectName("detailValue")
            self.report_stats_labels[key] = value_label
            stats_layout.addWidget(value_label, row, 1)

        self.report_payroll_button = QPushButton("Calculate Payroll")
        self.report_payroll_button.setObjectName("compactPrimaryButton")
        self.report_payroll_button.clicked.connect(self._calculate_payroll)
        stats_layout.addWidget(self.report_payroll_button, len(stat_pairs) + 1, 0, 1, 2)

        list_card = QFrame()
        list_card.setObjectName("card")
        list_layout = QVBoxLayout(list_card)
        list_layout.setContentsMargins(16, 14, 16, 14)
        list_layout.setSpacing(8)
        body_layout.addWidget(list_card, 0, 1)

        list_title = QLabel("Claims In Filter")
        list_title.setObjectName("sectionTitle")
        list_layout.addWidget(list_title)
        list_layout.addWidget(self.reports_table, 1)

        return wrapper

    def _build_report_date_picker(self, empty_text: str) -> QDateEdit:
        picker = QDateEdit()
        picker.setCalendarPopup(True)
        picker.setDisplayFormat("yyyy-MM-dd")
        picker.setMinimumDate(REPORT_NULL_QDATE)
        picker.setSpecialValueText(empty_text)
        picker.setDate(REPORT_NULL_QDATE)
        picker.setKeyboardTracking(False)
        return picker

    def _report_picker_has_value(self, picker: QDateEdit) -> bool:
        return picker.date() != picker.minimumDate()

    def _report_picker_to_datetime(self, picker: QDateEdit) -> datetime | None:
        if not self._report_picker_has_value(picker):
            return None
        value = picker.date()
        return datetime(value.year(), value.month(), value.day())

    def _build_route_tab(self) -> QWidget:
        wrapper = QWidget()
        layout = QVBoxLayout(wrapper)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(14)

        settings_card = QFrame()
        settings_card.setObjectName("card")
        settings_layout = QGridLayout(settings_card)
        settings_layout.setContentsMargins(16, 12, 16, 14)
        settings_layout.setHorizontalSpacing(12)
        settings_layout.setVerticalSpacing(8)
        layout.addWidget(settings_card, 0)

        settings_title = QLabel("Route Settings")
        settings_title.setObjectName("sectionTitle")
        settings_layout.addWidget(settings_title, 0, 0, 1, 4)

        settings_layout.addWidget(QLabel("Start From"), 1, 0)
        self.route_home_input = QLineEdit()
        settings_layout.addWidget(self.route_home_input, 1, 1)

        settings_layout.addWidget(QLabel("Show"), 1, 2)
        self.route_status_filter = QComboBox()
        self.route_status_filter.addItems(["Open", "Closed", "All"])
        self.route_status_filter.currentTextChanged.connect(lambda _text: self._populate_route_planner())
        settings_layout.addWidget(self.route_status_filter, 1, 3)

        self.route_appt_only_checkbox = QCheckBox("Appointment Scheduled Only")
        self.route_appt_only_checkbox.setChecked(True)
        self.route_appt_only_checkbox.toggled.connect(lambda _checked: self._populate_route_planner())
        settings_layout.addWidget(self.route_appt_only_checkbox, 2, 0, 1, 2)

        self.route_status_label = QLabel("Build tomorrow's route from your chosen start address.")
        self.route_status_label.setObjectName("workspaceStripTitle")
        settings_layout.addWidget(self.route_status_label, 2, 2, 1, 2)
        settings_layout.setColumnStretch(1, 1)
        settings_layout.setColumnStretch(3, 1)

        tables_row = QHBoxLayout()
        tables_row.setSpacing(14)
        layout.addLayout(tables_row, 1)

        available_card = QFrame()
        available_card.setObjectName("card")
        available_layout = QVBoxLayout(available_card)
        available_layout.setContentsMargins(16, 12, 16, 14)
        available_layout.setSpacing(8)
        tables_row.addWidget(available_card, 1)

        available_title = QLabel("Available Claims")
        available_title.setObjectName("sectionTitle")
        available_layout.addWidget(available_title)

        self.route_available_table = QTableWidget(0, 3)
        self.route_available_table.setHorizontalHeaderLabels(["Claim ID", "Customer", "Address"])
        self.route_available_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.route_available_table.setSelectionMode(QTableWidget.MultiSelection)
        self.route_available_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.route_available_table.verticalHeader().setVisible(False)
        self.route_available_table.setAlternatingRowColors(True)
        self.route_available_table.setWordWrap(False)
        self.route_available_table.setSortingEnabled(True)
        self.route_available_table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.route_available_table.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.route_available_table.horizontalHeader().setStretchLastSection(False)
        self.route_available_table.horizontalHeader().setSectionsMovable(True)
        for index in range(3):
            self.route_available_table.horizontalHeader().setSectionResizeMode(index, QHeaderView.Interactive)
        self.route_available_table.itemSelectionChanged.connect(
            lambda current_table=self.route_available_table: self._handle_table_selection(current_table)
        )
        self.route_available_table.cellDoubleClicked.connect(
            lambda row, column, current_table=self.route_available_table: self._handle_claim_table_double_click(
                current_table, row, column
            )
        )
        self.route_available_table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.route_available_table.customContextMenuRequested.connect(
            lambda pos, current_table=self.route_available_table: self._show_claim_context_menu(current_table, pos)
        )
        available_layout.addWidget(self.route_available_table, 1)

        self.route_add_button = QPushButton("Add Selected To Route")
        self.route_add_button.clicked.connect(self._add_selected_claims_to_route)
        available_layout.addWidget(self.route_add_button)

        selected_card = QFrame()
        selected_card.setObjectName("card")
        selected_layout = QVBoxLayout(selected_card)
        selected_layout.setContentsMargins(16, 12, 16, 14)
        selected_layout.setSpacing(8)
        tables_row.addWidget(selected_card, 1)

        selected_title = QLabel("Tomorrow's Stops")
        selected_title.setObjectName("sectionTitle")
        selected_layout.addWidget(selected_title)

        self.route_selected_table = QTableWidget(0, 4)
        self.route_selected_table.setHorizontalHeaderLabels(["Order", "Claim ID", "Customer", "Address"])
        self.route_selected_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.route_selected_table.setSelectionMode(QTableWidget.SingleSelection)
        self.route_selected_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.route_selected_table.verticalHeader().setVisible(False)
        self.route_selected_table.setAlternatingRowColors(True)
        self.route_selected_table.setWordWrap(False)
        self.route_selected_table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.route_selected_table.setVerticalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.route_selected_table.horizontalHeader().setStretchLastSection(False)
        self.route_selected_table.horizontalHeader().setSectionsMovable(True)
        for index in range(4):
            self.route_selected_table.horizontalHeader().setSectionResizeMode(index, QHeaderView.Interactive)
        self.route_selected_table.itemSelectionChanged.connect(self._on_route_stop_selected)
        self.route_selected_table.itemSelectionChanged.connect(
            lambda current_table=self.route_selected_table: self._handle_table_selection(current_table)
        )
        self.route_selected_table.cellDoubleClicked.connect(
            lambda row, column, current_table=self.route_selected_table: self._handle_claim_table_double_click(
                current_table, row, column
            )
        )
        self.route_selected_table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.route_selected_table.customContextMenuRequested.connect(
            lambda pos, current_table=self.route_selected_table: self._show_claim_context_menu(current_table, pos)
        )
        selected_layout.addWidget(self.route_selected_table, 1)

        address_row = QHBoxLayout()
        address_row.setSpacing(8)
        self.route_address_input = QLineEdit()
        self.route_address_input.setPlaceholderText("Stop address")
        address_row.addWidget(self.route_address_input, 1)
        self.route_move_up_button = QPushButton("Move Up")
        self.route_move_up_button.clicked.connect(lambda: self._move_route_stop(-1))
        address_row.addWidget(self.route_move_up_button)
        self.route_move_down_button = QPushButton("Move Down")
        self.route_move_down_button.clicked.connect(lambda: self._move_route_stop(1))
        address_row.addWidget(self.route_move_down_button)
        self.route_save_address_button = QPushButton("Save Address")
        self.route_save_address_button.clicked.connect(self._save_route_stop_address)
        address_row.addWidget(self.route_save_address_button)
        self.route_remove_stop_button = QPushButton("Remove Stop")
        self.route_remove_stop_button.clicked.connect(self._remove_route_stop)
        address_row.addWidget(self.route_remove_stop_button)
        selected_layout.addLayout(address_row)

        actions_row = QHBoxLayout()
        actions_row.setSpacing(8)
        self.route_open_maps_button = QPushButton("Open Route")
        self.route_open_maps_button.clicked.connect(self._open_route_in_google_maps)
        actions_row.addWidget(self.route_open_maps_button)
        self.route_clear_button = QPushButton("Clear Route")
        self.route_clear_button.clicked.connect(self._clear_route_plan)
        actions_row.addWidget(self.route_clear_button)
        selected_layout.addLayout(actions_row)

        return wrapper

    def _build_claim_tools_tab(self) -> QWidget:
        wrapper = QWidget()
        layout = QVBoxLayout(wrapper)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(0)

        tools_splitter = QSplitter(Qt.Vertical)
        tools_splitter.setChildrenCollapsible(False)
        tools_splitter.setObjectName("officeSplitter")
        layout.addWidget(tools_splitter, 1)

        contacts_card = QFrame()
        contacts_card.setObjectName("card")
        contacts_card.setMinimumHeight(220)
        contacts_layout = QVBoxLayout(contacts_card)
        contacts_layout.setContentsMargins(16, 12, 16, 14)
        contacts_layout.setSpacing(8)
        tools_splitter.addWidget(contacts_card)

        contacts_header = QHBoxLayout()
        contacts_header.setSpacing(10)
        contacts_title = QLabel("Claim Tools Contacts")
        contacts_title.setObjectName("sectionTitle")
        contacts_header.addWidget(contacts_title)
        contacts_header.addStretch(1)
        self.tools_add_contact_button = QPushButton("Add")
        self.tools_edit_contact_button = QPushButton("Edit")
        self.tools_remove_contact_button = QPushButton("Remove")
        for button in (self.tools_add_contact_button, self.tools_edit_contact_button, self.tools_remove_contact_button):
            button.setObjectName("compactActionButton")
        self.tools_add_contact_button.clicked.connect(self._add_claim_tool_contact)
        self.tools_edit_contact_button.clicked.connect(self._edit_claim_tool_contact)
        self.tools_remove_contact_button.clicked.connect(self._remove_claim_tool_contact)
        contacts_header.addWidget(self.tools_add_contact_button)
        contacts_header.addWidget(self.tools_edit_contact_button)
        contacts_header.addWidget(self.tools_remove_contact_button)
        contacts_layout.addLayout(contacts_header)

        self.claim_tools_contacts_table = QTableWidget(0, 4)
        self.claim_tools_contacts_table.setHorizontalHeaderLabels(["Name", "Number", "Prompt Guide", "Notes"])
        self.claim_tools_contacts_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.claim_tools_contacts_table.setSelectionMode(QTableWidget.SingleSelection)
        self.claim_tools_contacts_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.claim_tools_contacts_table.verticalHeader().setVisible(False)
        self.claim_tools_contacts_table.setAlternatingRowColors(True)
        self.claim_tools_contacts_table.setWordWrap(False)
        self.claim_tools_contacts_table.setSortingEnabled(True)
        self.claim_tools_contacts_table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.claim_tools_contacts_table.horizontalHeader().setStretchLastSection(False)
        self.claim_tools_contacts_table.horizontalHeader().setSectionsMovable(True)
        for index in range(4):
            self.claim_tools_contacts_table.horizontalHeader().setSectionResizeMode(index, QHeaderView.Interactive)
        self.claim_tools_contacts_table.cellDoubleClicked.connect(lambda *_: self._call_selected_claim_tool_contact())
        contacts_layout.addWidget(self.claim_tools_contacts_table, 1)

        files_card = QFrame()
        files_card.setObjectName("card")
        files_card.setMinimumHeight(220)
        files_layout = QVBoxLayout(files_card)
        files_layout.setContentsMargins(16, 12, 16, 14)
        files_layout.setSpacing(8)
        tools_splitter.addWidget(files_card)

        files_header = QHBoxLayout()
        files_header.setSpacing(10)
        files_title = QLabel("Claim Tools Files")
        files_title.setObjectName("sectionTitle")
        files_header.addWidget(files_title)
        files_header.addStretch(1)
        self.tools_add_file_button = QPushButton("Add")
        self.tools_edit_file_button = QPushButton("Edit")
        self.tools_open_file_button = QPushButton("Open")
        self.tools_remove_file_button = QPushButton("Remove")
        for button in (
            self.tools_add_file_button,
            self.tools_edit_file_button,
            self.tools_open_file_button,
            self.tools_remove_file_button,
        ):
            button.setObjectName("compactActionButton")
        self.tools_add_file_button.clicked.connect(self._add_claim_tool_file)
        self.tools_edit_file_button.clicked.connect(self._edit_claim_tool_file)
        self.tools_open_file_button.clicked.connect(self._open_selected_claim_tool_file)
        self.tools_remove_file_button.clicked.connect(self._remove_claim_tool_file)
        files_header.addWidget(self.tools_add_file_button)
        files_header.addWidget(self.tools_edit_file_button)
        files_header.addWidget(self.tools_open_file_button)
        files_header.addWidget(self.tools_remove_file_button)
        files_layout.addLayout(files_header)

        self.claim_tools_files_table = QTableWidget(0, 4)
        self.claim_tools_files_table.setHorizontalHeaderLabels(["Section", "Label", "File", "Notes"])
        self.claim_tools_files_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.claim_tools_files_table.setSelectionMode(QTableWidget.SingleSelection)
        self.claim_tools_files_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.claim_tools_files_table.verticalHeader().setVisible(False)
        self.claim_tools_files_table.setAlternatingRowColors(True)
        self.claim_tools_files_table.setWordWrap(False)
        self.claim_tools_files_table.setSortingEnabled(True)
        self.claim_tools_files_table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.claim_tools_files_table.horizontalHeader().setStretchLastSection(False)
        self.claim_tools_files_table.horizontalHeader().setSectionsMovable(True)
        for index in range(4):
            self.claim_tools_files_table.horizontalHeader().setSectionResizeMode(index, QHeaderView.Interactive)
        self.claim_tools_files_table.cellDoubleClicked.connect(lambda *_: self._open_selected_claim_tool_file())
        files_layout.addWidget(self.claim_tools_files_table, 1)

        tools_splitter.setSizes([460, 320])

        return wrapper

    def _build_body_shops_tab(self) -> QWidget:
        wrapper = QWidget()
        layout = QVBoxLayout(wrapper)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(14)

        shops_card = QFrame()
        shops_card.setObjectName("card")
        shops_layout = QVBoxLayout(shops_card)
        shops_layout.setContentsMargins(18, 14, 18, 18)
        shops_layout.setSpacing(10)
        layout.addWidget(shops_card, 1)

        shops_header = QHBoxLayout()
        shops_header.setSpacing(10)
        shops_title = QLabel("Body Shops")
        shops_title.setObjectName("sectionTitle")
        shops_header.addWidget(shops_title)
        shops_header.addStretch(1)

        self.body_shops_add_button = QPushButton("Add")
        self.body_shops_edit_button = QPushButton("Edit")
        self.body_shops_remove_button = QPushButton("Remove")
        self.body_shops_add_button.clicked.connect(self._add_body_shop_entry)
        self.body_shops_edit_button.clicked.connect(self._edit_body_shop_entry)
        self.body_shops_remove_button.clicked.connect(self._remove_body_shop_entry)
        shops_header.addWidget(self.body_shops_add_button)
        shops_header.addWidget(self.body_shops_edit_button)
        shops_header.addWidget(self.body_shops_remove_button)
        shops_layout.addLayout(shops_header)

        self.body_shops_table = QTableWidget(0, 13)
        self.body_shops_table.setHorizontalHeaderLabels(
            [
                "Shop Name",
                "Contact",
                "Phone",
                "Email",
                "Tax ID",
                "Address",
                "Body Rate",
                "Paint Rate",
                "Frame Rate",
                "Mechanical Rate",
                "Certifications",
                "Negotiation Notes",
                "Notes",
            ]
        )
        self.body_shops_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.body_shops_table.setSelectionMode(QTableWidget.SingleSelection)
        self.body_shops_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.body_shops_table.verticalHeader().setVisible(False)
        self.body_shops_table.setAlternatingRowColors(True)
        self.body_shops_table.setWordWrap(False)
        self.body_shops_table.setSortingEnabled(True)
        self.body_shops_table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.body_shops_table.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.body_shops_table.horizontalHeader().setStretchLastSection(False)
        self.body_shops_table.horizontalHeader().setSectionsMovable(True)
        self.body_shops_table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.body_shops_table.cellDoubleClicked.connect(self._handle_body_shop_double_click)
        self.body_shops_table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.body_shops_table.customContextMenuRequested.connect(self._show_body_shop_context_menu)
        shops_layout.addWidget(self.body_shops_table, 1)

        return wrapper

    def _build_insurance_companies_tab(self) -> QWidget:
        wrapper = QWidget()
        layout = QVBoxLayout(wrapper)
        layout.setContentsMargins(0, 0, 0, 0)
        layout.setSpacing(14)

        insurance_card = QFrame()
        insurance_card.setObjectName("card")
        insurance_layout = QVBoxLayout(insurance_card)
        insurance_layout.setContentsMargins(18, 14, 18, 18)
        insurance_layout.setSpacing(10)
        layout.addWidget(insurance_card, 1)

        insurance_header = QHBoxLayout()
        insurance_header.setSpacing(10)
        insurance_title = QLabel("Insurance Companies")
        insurance_title.setObjectName("sectionTitle")
        insurance_header.addWidget(insurance_title)
        insurance_header.addStretch(1)

        self.insurance_add_button = QPushButton("Add")
        self.insurance_edit_button = QPushButton("Edit")
        self.insurance_remove_button = QPushButton("Remove")
        self.insurance_add_button.clicked.connect(self._add_insurance_company_entry)
        self.insurance_edit_button.clicked.connect(self._edit_insurance_company_entry)
        self.insurance_remove_button.clicked.connect(self._remove_insurance_company_entry)
        insurance_header.addWidget(self.insurance_add_button)
        insurance_header.addWidget(self.insurance_edit_button)
        insurance_header.addWidget(self.insurance_remove_button)
        insurance_layout.addLayout(insurance_header)

        self.insurance_companies_table = QTableWidget(0, 17)
        self.insurance_companies_table.setHorizontalHeaderLabels(
            [
                "Company Name",
                "Quick Summary",
                "Fatal Errors",
                "Photo Rules",
                "Estimate / Supp Rules",
                "Parts Rules",
                "Total Loss Rules",
                "Tow Rules",
                "Supplement Rules",
                "Betterment / Depreciation",
                "Documentation",
                "Rates & Tax",
                "Misc Rules",
                "Contact Info",
                "Labor Rates",
                "Total Loss Threshold",
                "Notes",
            ]
        )
        self.insurance_companies_table.setSelectionBehavior(QTableWidget.SelectRows)
        self.insurance_companies_table.setSelectionMode(QTableWidget.SingleSelection)
        self.insurance_companies_table.setEditTriggers(QTableWidget.NoEditTriggers)
        self.insurance_companies_table.verticalHeader().setVisible(False)
        self.insurance_companies_table.setAlternatingRowColors(True)
        self.insurance_companies_table.setWordWrap(False)
        self.insurance_companies_table.setSortingEnabled(True)
        self.insurance_companies_table.setHorizontalScrollMode(QAbstractItemView.ScrollPerPixel)
        self.insurance_companies_table.setHorizontalScrollBarPolicy(Qt.ScrollBarAsNeeded)
        self.insurance_companies_table.horizontalHeader().setStretchLastSection(False)
        self.insurance_companies_table.horizontalHeader().setSectionsMovable(True)
        self.insurance_companies_table.horizontalHeader().setSectionResizeMode(QHeaderView.Interactive)
        self.insurance_companies_table.cellDoubleClicked.connect(self._handle_insurance_company_double_click)
        self.insurance_companies_table.setContextMenuPolicy(Qt.CustomContextMenu)
        self.insurance_companies_table.customContextMenuRequested.connect(self._show_insurance_company_context_menu)
        insurance_layout.addWidget(self.insurance_companies_table, 1)

        return wrapper

    def _apply_styles(self) -> None:
        self.setStyleSheet(
            """
            QWidget#central {
                background: #eef3f9;
            }
            QFrame#card, QFrame#statCard, QFrame#toolbarCard, QFrame#workspaceStrip {
                background: #fbfdff;
                border: 1px solid #d5dfeb;
                border-radius: 22px;
            }
            QFrame#workspaceStrip {
                background: rgba(251, 253, 255, 0.58);
                border-radius: 12px;
            }
            QLabel#title {
                color: #142334;
                font-size: 20px;
                font-weight: 700;
                letter-spacing: 0.01em;
            }
            QLabel#subtitle {
                color: #6d7f92;
                font-size: 11px;
            }
            QLabel#sectionTitle {
                color: #213346;
                font-size: 14px;
                font-weight: 700;
            }
            QLabel#workspaceStripTitle {
                color: #5e7388;
                font-size: 9px;
                font-weight: 700;
                letter-spacing: 0.04em;
                text-transform: uppercase;
            }
            QLabel#workspaceStripValue {
                color: #62758a;
                font-size: 9px;
            }
            QLabel#workspaceStripValueBox {
                background: rgba(255, 255, 255, 0.88);
                border: 1px solid #d7e0ea;
                border-radius: 9px;
                padding: 3px 8px;
                color: #62758a;
                font-size: 9px;
            }
            QLabel#statTitle {
                color: #728399;
                font-size: 9px;
                font-weight: 600;
                letter-spacing: 0.04em;
                text-transform: uppercase;
            }
            QLabel#statValue {
                color: #112233;
                font-size: 24px;
                font-weight: 700;
            }
            QLabel#detailLabel {
                color: #6a7b8e;
                font-size: 12px;
                min-width: 126px;
            }
            QLabel#detailFormLabel {
                color: #55687d;
                font-size: 11px;
                font-weight: 600;
                padding: 0px 0px 1px 0px;
            }
            QLabel#detailValue {
                color: #18293a;
                font-size: 12px;
            }
            QLabel#detailValueChip {
                background: #edf3f9;
                color: #1f3347;
                border: 1px solid #d5e0ec;
                border-radius: 10px;
                padding: 4px 10px;
                font-size: 12px;
                font-weight: 600;
            }
            QLineEdit, QTextEdit, QPlainTextEdit, QTableWidget, QListWidget {
                background: #ffffff;
                border: 1px solid #d7e0ea;
                border-radius: 16px;
            }
            QStackedWidget {
                background: transparent;
                border: none;
            }
            QLineEdit, QTextEdit, QPlainTextEdit {
                padding: 8px 11px;
                color: #182738;
                selection-background-color: #d7e8fb;
                selection-color: #122131;
            }
            QLineEdit#detailInput {
                background: transparent;
                border: none;
                border-bottom: 1px solid #d5dfeb;
                border-radius: 0px;
                padding: 2px 2px 0px 2px;
                color: #1a2c3d;
            }
            QLineEdit#detailInput:focus {
                border-bottom: 2px solid #7ea2c8;
            }
            QComboBox {
                background: #ffffff;
                border: 1px solid #d7e0ea;
                border-radius: 16px;
                padding: 10px 12px;
                color: #182738;
            }
            QComboBox::drop-down {
                border: none;
                width: 26px;
            }
            QPushButton {
                background: #e8eef5;
                border: 1px solid #d6dfeb;
                border-radius: 16px;
                color: #1f2f40;
                padding: 11px 15px;
                font-weight: 600;
            }
            QPushButton:hover {
                background: #dde7f2;
            }
            QPushButton:pressed {
                background: #d2ddea;
            }
            QPushButton#primaryButton {
                background: #20384f;
                color: #f6fbff;
                border: 1px solid #20384f;
            }
            QPushButton#detailPrimaryButton {
                background: #20384f;
                color: #f6fbff;
                border: 1px solid #20384f;
                border-radius: 15px;
                padding: 4px 12px;
                min-height: 30px;
                font-size: 11px;
                font-weight: 700;
            }
            QPushButton#detailPrimaryButton:hover {
                background: #27445e;
            }
            QPushButton#detailPrimaryButton:pressed {
                background: #1b3146;
            }
            QPushButton#compactPrimaryButton {
                background: #20384f;
                color: #f6fbff;
                border: 1px solid #20384f;
                border-radius: 14px;
                padding: 5px 12px;
                min-height: 28px;
                font-size: 11px;
                font-weight: 700;
            }
            QPushButton#compactPrimaryButton:hover {
                background: #27445e;
            }
            QPushButton#compactPrimaryButton:pressed {
                background: #1b3146;
            }
            QPushButton#primaryButton:hover {
                background: #27445e;
            }
            QPushButton#primaryButton:pressed {
                background: #1b3146;
            }
            QPushButton#officeSaveButton {
                background: #20384f;
                color: #f6fbff;
                border: 1px solid #20384f;
                border-radius: 14px;
                padding: 4px 14px;
                min-height: 28px;
                font-size: 11px;
                font-weight: 700;
            }
            QPushButton#officeSaveButton:hover {
                background: #27445e;
            }
            QPushButton#officeSaveButton:pressed {
                background: #1b3146;
            }
            QPushButton#compactActionButton {
                border-radius: 12px;
                padding: 5px 12px;
                min-height: 24px;
                font-size: 11px;
                font-weight: 600;
            }
            QPushButton#detailActionButton {
                border-radius: 12px;
                padding: 2px 10px;
                min-height: 24px;
                font-size: 10px;
                font-weight: 600;
            }
            QFrame#workspaceStrip QPushButton {
                border-radius: 9px;
                padding: 2px 8px;
                min-height: 16px;
                max-height: 18px;
                font-size: 9px;
            }
            QComboBox#officeCompactInput, QLineEdit#officeCompactInput {
                padding: 3px 9px;
                border-radius: 11px;
                font-size: 10px;
            }
            QPushButton#navTabButton {
                background: #e5ebf3;
                color: #596b80;
                border: 1px solid #dce4ee;
                border-radius: 14px;
                padding: 9px 12px;
                min-height: 22px;
                font-size: 11px;
                font-weight: 600;
            }
            QPushButton#navTabButton:hover {
                background: #dde7f2;
            }
            QPushButton#navTabButton:checked {
                background: #ffffff;
                color: #162739;
                border: 1px solid #d5dfeb;
            }
            QToolButton#navLibraryButton {
                background: #e5ebf3;
                color: #596b80;
                border: 1px solid #dce4ee;
                border-radius: 14px;
                padding: 9px 16px;
                min-height: 22px;
                font-size: 11px;
                font-weight: 600;
            }
            QToolButton#navLibraryButton:hover {
                background: #dde7f2;
            }
            QToolButton#navLibraryButton:checked {
                background: #ffffff;
                color: #162739;
                border: 1px solid #d5dfeb;
            }
            QHeaderView::section {
                background: #edf2f8;
                color: #4f6378;
                border: none;
                border-bottom: 1px solid #dbe3ed;
                padding: 0px 11px;
                font-weight: 600;
            }
            QTableWidget {
                gridline-color: #edf2f7;
                alternate-background-color: #f7fafe;
                selection-background-color: #d7e7f8;
                selection-color: #142230;
            }
            QListWidget {
                padding: 8px;
                color: #182738;
                outline: none;
            }
            QListWidget::item:selected {
                background: #d7e7f8;
                color: #142230;
                border-radius: 10px;
            }
            QTableWidget::item {
                padding: 8px;
            }
            QScrollBar:vertical {
                background: transparent;
                width: 13px;
                margin: 8px 0 8px 0;
            }
            QScrollBar::handle:vertical {
                background: #c6d2e1;
                min-height: 28px;
                border-radius: 6px;
            }
            QScrollBar::add-line:vertical, QScrollBar::sub-line:vertical,
            QScrollBar::add-page:vertical, QScrollBar::sub-page:vertical {
                background: transparent;
                border: none;
                height: 0;
            }
            QSplitter#mainSplitter::handle {
                background: transparent;
                width: 12px;
            }
            """
        )

    def _load_claims(self) -> None:
        self.settings = self.repo.load_settings()
        self.claims = self.repo.load_claims()
        self.claim_tools_contacts = self.repo.load_claim_tools_contacts()
        self.claim_tools_files = self.repo.load_claim_tools_files()
        self.body_shop_entries = self.repo.load_body_shop_database()
        self.insurance_company_entries = self.repo.load_insurance_company_database()
        if self.repo.sync_reference_data_from_claims(include_claim_folders=False):
            self.claims = self.repo.load_claims()
            self.claim_tools_files = self.repo.load_claim_tools_files()
            self.body_shop_entries = self.repo.load_body_shop_database()
            self.insurance_company_entries = self.repo.load_insurance_company_database()
        self.route_plan_keys = list(self.settings.get("route_plan_keys", []))
        self.processed_apptrak_pdfs = self.repo.load_processed_apptrak_pdfs()
        if hasattr(self, "route_home_input"):
            self.route_home_input.setText(str(self.settings.get("route_home_address", "") or ""))
        self._apply_filters()

    def _run_refresh_scan(self) -> None:
        if self.refresh_scan_running:
            QMessageBox.information(self, APP_NAME, "Refresh Scan is already running.")
            return

        self.refresh_scan_running = True
        self.refresh_scan_button.setEnabled(False)
        self.apptrak_status_label.setText("Refreshing watched folders...")

        def worker() -> None:
            result: dict[str, object] = {"error": None}
            try:
                if not REFRESH_SCAN_HELPER.exists():
                    raise RuntimeError(f"Refresh scan helper was not found:\n{REFRESH_SCAN_HELPER}")
                subprocess.run(
                    [sys.executable, str(REFRESH_SCAN_HELPER)],
                    cwd=str(APP_DIR),
                    check=True,
                    capture_output=True,
                    text=True,
                    timeout=300,
                )
            except Exception as exc:
                result["error"] = str(exc)
            self.refresh_scan_signals.finished.emit(result)

        threading.Thread(target=worker, daemon=True).start()

    def _finish_refresh_scan(self, result: dict[str, object]) -> None:
        self.refresh_scan_running = False
        self.refresh_scan_button.setEnabled(True)

        error = str(result.get("error") or "").strip()
        if error:
            self.apptrak_status_label.setText("Refresh Scan failed.")
            QMessageBox.warning(self, APP_NAME, f"Could not complete Refresh Scan.\n\n{error}")
            return

        self._load_claims()
        self.apptrak_status_label.setText("Refresh Scan completed.")

    def run_apptrak_import_now(self) -> None:
        self._run_apptrak_import_cycle(manual=True)

    def _run_apptrak_import_cycle(self, manual: bool) -> None:
        if self.apptrak_import_running:
            if manual:
                QMessageBox.information(self, APP_NAME, "AppTrak import is already running.")
            return

        self.apptrak_import_running = True
        self.check_apptrak_button.setEnabled(False)
        self.apptrak_status_label.setText("Running AppTrak import...")
        processed_snapshot = set(self.processed_apptrak_pdfs)

        def worker() -> None:
            result: dict[str, object] = {
                "manual": manual,
                "error": None,
                "staged": [],
                "import_ran": False,
            }
            run_started_at = datetime.now()
            try:
                if not APPTRAK_IMPORT_BAT.exists():
                    raise RuntimeError(f"AppTrak import script was not found:\n{APPTRAK_IMPORT_BAT}")
                self._ensure_apptrak_running()
                subprocess.run(
                    ["cmd", "/c", str(APPTRAK_IMPORT_BAT)],
                    cwd=str(APPTRAK_AUTOMATION_DIR),
                    check=True,
                    capture_output=True,
                    text=True,
                    timeout=300,
                )
                result["import_ran"] = True
                result["staged"] = self._stage_new_apptrak_imports(
                    processed_snapshot,
                    run_started_at,
                    newest_fallback=True,
                )
            except Exception as exc:
                result["error"] = str(exc)
            self.apptrak_signals.finished.emit(result)

        threading.Thread(target=worker, daemon=True).start()

    def _ensure_apptrak_running(self) -> None:
        if not APPTRAK_EXE.exists():
            raise RuntimeError(f"AppTrak executable was not found:\n{APPTRAK_EXE}")
        if self._is_apptrak_running(require_window=True):
            return
        if self._is_apptrak_running():
            try:
                subprocess.run(
                    ["taskkill", "/F", "/IM", "apptrak.exe"],
                    check=False,
                    capture_output=True,
                    text=True,
                    timeout=15,
                )
                time.sleep(2)
            except Exception:
                pass
        try:
            subprocess.Popen([str(APPTRAK_EXE)], cwd=str(APPTRAK_EXE.parent))
        except Exception as exc:
            raise RuntimeError(f"Could not launch AppTrak:\n{exc}") from exc

        deadline = time.time() + 30
        while time.time() < deadline:
            if self._is_apptrak_running(require_window=True):
                time.sleep(5)
                return
            time.sleep(1)
        raise RuntimeError("AppTrak did not finish opening in time.")

    def _is_apptrak_running(self, require_window: bool = False) -> bool:
        try:
            result = subprocess.run(
                ["tasklist", "/FI", "IMAGENAME eq apptrak.exe"],
                check=False,
                capture_output=True,
                text=True,
                timeout=15,
            )
        except Exception:
            return False

        running = "apptrak.exe" in (result.stdout or "").lower()
        if not running or not require_window:
            return running

        try:
            window_result = subprocess.run(
                [
                    "powershell",
                    "-NoProfile",
                    "-Command",
                    "(Get-Process apptrak -ErrorAction SilentlyContinue | Where-Object { $_.MainWindowHandle -ne 0 -or $_.MainWindowTitle -ne '' } | Measure-Object).Count",
                ],
                check=False,
                capture_output=True,
                text=True,
                timeout=15,
            )
        except Exception:
            return False
        return (window_result.stdout or "").strip() not in {"", "0"}

    def _finish_apptrak_import_cycle(self, result: dict[str, object]) -> None:
        self.apptrak_import_running = False
        self.check_apptrak_button.setEnabled(True)

        manual = bool(result.get("manual"))
        error = str(result.get("error") or "").strip()
        staged = list(result.get("staged") or [])

        if error:
            self.apptrak_status_label.setText("AppTrak import failed.")
            if manual:
                QMessageBox.warning(self, APP_NAME, f"Could not complete the AppTrak import.\n\n{error}")
            return

        if staged:
            processed = set(self.processed_apptrak_pdfs)
            for item in staged:
                processed_keys = item.get("processed_keys") or []
                processed.update(str(value) for value in processed_keys if str(value).strip())
            self.processed_apptrak_pdfs = sorted(processed)
            self.repo.save_processed_apptrak_pdfs(self.processed_apptrak_pdfs)
            self._load_claims()
            staged_names = ", ".join(str(item.get("folder_name") or "") for item in staged if str(item.get("folder_name") or "").strip())
            self.apptrak_status_label.setText(
                f"AppTrak staged {len(staged)} new import(s){': ' + staged_names if staged_names else '.'}"
            )
            return

        self.apptrak_status_label.setText("AppTrak checked. No new imports found.")

    def _stage_new_apptrak_imports(
        self,
        processed_snapshot: set[str],
        imported_since: datetime | None = None,
        newest_fallback: bool = False,
    ) -> list[dict[str, object]]:
        pending_root = self._pending_claims_root()
        if not pending_root:
            raise RuntimeError("Pending Claims folder is not configured in Watched Folders.")
        if not APPTRAK_DBF_DIR.exists():
            raise RuntimeError(f"AppTrak dbf folder was not found:\n{APPTRAK_DBF_DIR}")

        pending_root.mkdir(parents=True, exist_ok=True)
        existing_ids = {str(claim.claim_id or "").strip() for claim in self.claims if str(claim.claim_id or "").strip()}
        staged: list[dict[str, object]] = []

        for item in self._recent_apptrak_import_records(imported_since, newest_fallback):
            claim_id = str(item.get("claim_id") or "").strip()
            if not claim_id or claim_id in processed_snapshot or claim_id in existing_ids:
                continue

            title = str(item.get("title") or claim_id).strip() or claim_id
            target_folder = self._unique_destination(pending_root / title)
            target_folder.mkdir(parents=True, exist_ok=False)

            assignment_source = item.get("assignment_pdf")
            if not isinstance(assignment_source, Path) or not assignment_source.exists():
                continue

            assign_target = target_folder / "assign.pdf"
            shutil.copy2(assignment_source, assign_target)

            claim_key = f"fs::{target_folder}"
            claim_record = self._build_apptrak_claim_record(item, target_folder, assign_target, claim_key)
            self.repo.upsert_claim_record(claim_key, claim_record)

            staged.append(
                {
                    "claim_id": claim_id,
                    "folder_name": target_folder.name,
                    "folder_path": str(target_folder),
                    "processed_keys": [claim_id],
                }
            )
        return staged

    def _build_apptrak_claim_record(
        self,
        item: dict[str, object],
        target_folder: Path,
        assign_target: Path,
        claim_key: str,
    ) -> dict[str, object]:
        timestamp = datetime.now().isoformat(timespec="seconds")
        claim_id = str(item.get("claim_id") or "").strip()
        customer_name = str(item.get("customer_name") or "").strip()
        title = str(item.get("title") or claim_id).strip() or claim_id
        vehicle = str(item.get("vehicle") or "").strip()

        return {
            "key": claim_key,
            "claim_id": claim_id,
            "title": title,
            "status": "Open",
            "source_path": str(target_folder),
            "claim_type": "Original",
            "total_loss": False,
            "updated_at": timestamp,
            "closed_date": "",
            "customer_name": customer_name,
            "insurance_company": "",
            "claim_number": "",
            "date_of_loss": "",
            "town": "",
            "owner_address": "",
            "location_of_vehicle": "",
            "vehicle": vehicle,
            "vin": "",
            "shop_name": "No Shop Chosen",
            "shop_phone": "",
            "shop_email": "",
            "contact_phone": "",
            "contact_email": "",
            "assignment_claim_notes": "",
            "office_progress_status": "",
            "office_appt_when": "",
            "office_waiting_for_paperwork": "",
            "office_additional_notes": "",
            "route_address_override": "",
            "assign_pdf_path": str(assign_target),
            "manual_overrides": {},
            "note_history": [],
            "notes": "",
        }

    def _recent_apptrak_import_records(
        self,
        imported_since: datetime | None = None,
        newest_fallback: bool = False,
    ) -> list[dict[str, object]]:
        apprroot_rows = self._read_dbf_rows(APPTRAK_DBF_DIR / "apprroot.dbf")
        apprname_rows = self._read_dbf_rows(APPTRAK_DBF_DIR / "apprname.dbf")
        apprveh_rows = self._read_dbf_rows(APPTRAK_DBF_DIR / "apprveh.dbf")
        if not apprroot_rows:
            return []

        names_by_tag: dict[str, list[dict[str, str]]] = {}
        for row in apprname_rows:
            tag = str(row.get("NAMTAG") or "").strip()
            if tag:
                names_by_tag.setdefault(tag, []).append(row)

        vehicles_by_tag: dict[str, dict[str, str]] = {}
        for row in apprveh_rows:
            tag = str(row.get("VEHTAG") or "").strip()
            if tag:
                vehicles_by_tag[tag] = row

        recent_cutoff = imported_since - timedelta(seconds=20) if imported_since else None
        candidates: list[dict[str, object]] = []

        for row in apprroot_rows:
            claim_id = str(row.get("APRNO") or "").strip()
            if not re.fullmatch(r"26\d+", claim_id):
                continue

            tag = str(row.get("APRTAG") or "").strip() or f"{claim_id}{str(row.get('APRSEQ') or '00').zfill(2)}"
            assignment_pdf = APPTRAK_DOCS_DIR / f"{claim_id[3:]}.pdf"
            if not assignment_pdf.exists() or not assignment_pdf.is_file():
                continue

            related_times = [datetime.fromtimestamp(assignment_pdf.stat().st_mtime)]
            if APPTRAK_PICS_DIR.exists():
                for photo in APPTRAK_PICS_DIR.glob(f"{tag}-*.jpg"):
                    related_times.append(datetime.fromtimestamp(photo.stat().st_mtime))
            last_activity = max(related_times)
            if recent_cutoff and last_activity < recent_cutoff:
                continue

            owner_rows = names_by_tag.get(claim_id, [])
            owner = next((entry for entry in owner_rows if str(entry.get("NAMTYPE") or "").strip().upper() == "O"), None)
            insured = next((entry for entry in owner_rows if str(entry.get("NAMTYPE") or "").strip().upper() == "I"), None)
            chosen = owner or insured or {}
            first = str(chosen.get("NAMFIRST") or "").strip()
            last = str(chosen.get("NAMLAST") or "").strip()
            customer_name = " ".join(part for part in (first, last) if part).strip() or last or first or claim_id
            customer_name = self._sanitize_folder_part(customer_name)

            vehicle_row = vehicles_by_tag.get(tag, {})
            vehicle_year = str(vehicle_row.get("VEHYEAR") or vehicle_row.get("YEAR") or "").strip()
            vehicle_make = str(vehicle_row.get("VEHMAKE") or vehicle_row.get("MAKE") or "").strip()
            vehicle_model = str(vehicle_row.get("VEHMODEL") or vehicle_row.get("MODEL") or "").strip()
            vehicle = " ".join(part for part in (vehicle_year, vehicle_make, vehicle_model) if part).strip()

            candidates.append(
                {
                    "claim_id": claim_id,
                    "customer_name": customer_name,
                    "title": f"{claim_id} - {customer_name}",
                    "assignment_pdf": assignment_pdf,
                    "vehicle": vehicle,
                    "last_activity": last_activity,
                }
            )

        candidates.sort(key=lambda item: item["last_activity"])  # oldest to newest, matching import order
        if candidates:
            return candidates

        if newest_fallback:
            fallback_candidates: list[dict[str, object]] = []
            for row in apprroot_rows:
                claim_id = str(row.get("APRNO") or "").strip()
                if not re.fullmatch(r"26\d+", claim_id):
                    continue

                tag = str(row.get("APRTAG") or "").strip() or f"{claim_id}{str(row.get('APRSEQ') or '00').zfill(2)}"
                assignment_pdf = APPTRAK_DOCS_DIR / f"{claim_id[3:]}.pdf"
                if not assignment_pdf.exists() or not assignment_pdf.is_file():
                    continue

                related_times = [datetime.fromtimestamp(assignment_pdf.stat().st_mtime)]
                if APPTRAK_PICS_DIR.exists():
                    for photo in APPTRAK_PICS_DIR.glob(f"{tag}-*.jpg"):
                        related_times.append(datetime.fromtimestamp(photo.stat().st_mtime))
                last_activity = max(related_times)

                owner_rows = names_by_tag.get(claim_id, [])
                owner = next((entry for entry in owner_rows if str(entry.get("NAMTYPE") or "").strip().upper() == "O"), None)
                insured = next((entry for entry in owner_rows if str(entry.get("NAMTYPE") or "").strip().upper() == "I"), None)
                chosen = owner or insured or {}
                first = str(chosen.get("NAMFIRST") or "").strip()
                last = str(chosen.get("NAMLAST") or "").strip()
                customer_name = " ".join(part for part in (first, last) if part).strip() or last or first or claim_id
                customer_name = self._sanitize_folder_part(customer_name)

                vehicle_row = vehicles_by_tag.get(tag, {})
                vehicle_year = str(vehicle_row.get("VEHYEAR") or vehicle_row.get("YEAR") or "").strip()
                vehicle_make = str(vehicle_row.get("VEHMAKE") or vehicle_row.get("MAKE") or "").strip()
                vehicle_model = str(vehicle_row.get("VEHMODEL") or vehicle_row.get("MODEL") or "").strip()
                vehicle = " ".join(part for part in (vehicle_year, vehicle_make, vehicle_model) if part).strip()

                fallback_candidates.append(
                    {
                        "claim_id": claim_id,
                        "customer_name": customer_name,
                        "title": f"{claim_id} - {customer_name}",
                        "assignment_pdf": assignment_pdf,
                        "vehicle": vehicle,
                        "last_activity": last_activity,
                    }
                )
            if fallback_candidates:
                return [max(fallback_candidates, key=lambda item: item["last_activity"])]
        return []

    def _read_dbf_rows(self, path: Path) -> list[dict[str, str]]:
        if not path.exists() or not path.is_file():
            return []
        try:
            with path.open("rb") as handle:
                header = handle.read(32)
                if len(header) < 32:
                    return []
                num_records = struct.unpack("<I", header[4:8])[0]
                header_len = struct.unpack("<H", header[8:10])[0]
                record_len = struct.unpack("<H", header[10:12])[0]
                fields: list[tuple[str, int]] = []
                while True:
                    desc = handle.read(32)
                    if not desc or desc[0] == 0x0D:
                        break
                    name = desc[:11].split(b"\x00", 1)[0].decode("ascii", "ignore").strip()
                    flen = desc[16]
                    fields.append((name, flen))
                handle.seek(header_len)
                rows: list[dict[str, str]] = []
                for _ in range(num_records):
                    record = handle.read(record_len)
                    if not record or record[0] == 0x2A:
                        continue
                    pos = 1
                    row: dict[str, str] = {}
                    for name, flen in fields:
                        raw = record[pos:pos + flen]
                        pos += flen
                        row[name] = raw.decode("latin1", "ignore").strip().strip("\x00")
                    rows.append(row)
                return rows
        except Exception:
            return []

    def _sanitize_folder_part(self, value: str) -> str:
        cleaned = re.sub(r'[\\/:*?"<>|]+', " ", value or "")
        cleaned = re.sub(r"\s+", " ", cleaned).strip(" .-_")
        return cleaned or "UNKNOWN"

    def _pending_claims_root(self) -> Path | None:
        for folder in self.settings.get("watched_folders", []):
            text = str(folder).strip()
            if text and "pending" in text.lower():
                return Path(text)
        return None

    def _apply_filters(self) -> None:
        search = self.search_input.text().strip().lower()
        filtered = [claim for claim in self.claims if self._matches_search(claim, search)]
        self.filtered_all = filtered
        self.filtered_open = [claim for claim in filtered if claim.status.lower() == "open"]
        self.filtered_closed = [claim for claim in filtered if claim.status.lower() == "closed"]

        self._set_stat_value(self.all_card, len(self.filtered_all))
        self._set_stat_value(self.open_card, len(self.filtered_open))
        self._set_stat_value(self.closed_card, len(self.filtered_closed))

        self._populate_table(self.all_table, self.filtered_all)
        self._populate_table(self.open_table, self.filtered_open)
        self._populate_table(self.closed_table, self.filtered_closed)
        self._populate_reports()
        self._populate_office_tables(self.filtered_open)
        self._populate_workspace()
        self._populate_claim_tools()
        self._populate_body_shops()
        self._populate_insurance_companies()
        self._populate_route_planner()

    def _matches_search(self, claim: ClaimView, search: str) -> bool:
        if not search:
            return True
        haystack = " ".join(
            [
                claim.claim_id,
                claim.title,
                claim.customer_name,
                claim.insurance_company,
                claim.claim_number,
                claim.shop_name,
                claim.town,
                claim.vehicle,
            ]
        ).lower()
        return search in haystack

    def _set_stat_value(self, card: QFrame, value: int) -> None:
        card.value_label.setText(str(value))  # type: ignore[attr-defined]

    def _closed_date_for_claim(self, claim: ClaimView) -> datetime | None:
        if claim.status.strip().lower() != "closed":
            return None
        source_value = (claim.closed_date or claim.updated_at or "").strip()
        if not source_value:
            return None
        normalized = source_value.replace("Z", "+00:00")
        try:
            return datetime.fromisoformat(normalized)
        except ValueError:
            pass
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d"):
            try:
                return datetime.strptime(normalized, fmt)
            except ValueError:
                continue
        return None

    def _parse_report_date(self, value: str) -> datetime | None:
        if not value:
            return None
        normalized = value.strip()
        for fmt in ("%Y-%m-%d", "%m/%d/%y", "%m/%d/%Y", "%m-%d-%y", "%m-%d-%Y"):
            try:
                return datetime.strptime(normalized, fmt)
            except ValueError:
                continue
        return None

    def _report_records(self) -> list[ClaimView]:
        return [claim for claim in self.filtered_all if self._closed_date_for_claim(claim)]

    def _filtered_report_claims(self) -> list[ClaimView]:
        records = self._report_records()
        start = self._report_picker_to_datetime(self.report_start_input) if hasattr(self, "report_start_input") else None
        end = self._report_picker_to_datetime(self.report_end_input) if hasattr(self, "report_end_input") else None

        if start or end:
            if start:
                records = [
                    claim for claim in records
                    if (closed_date := self._closed_date_for_claim(claim)) and closed_date.date() >= start.date()
                ]
            if end:
                records = [
                    claim for claim in records
                    if (closed_date := self._closed_date_for_claim(claim)) and closed_date.date() <= end.date()
                ]
        elif self.report_selected_date:
            records = [
                claim for claim in records
                if (closed_date := self._closed_date_for_claim(claim))
                and closed_date.strftime("%Y-%m-%d") == self.report_selected_date
            ]

        return sorted(records, key=lambda claim: ((self._closed_date_for_claim(claim) or datetime.min), claim.claim_id))

    def _populate_reports(self) -> None:
        if not hasattr(self, "reports_table"):
            return

        records = self._report_records()
        closed_dates = [value for value in (self._closed_date_for_claim(claim) for claim in records) if value]
        now = datetime.now()
        today = now.date()
        iso_year, iso_week, _ = now.isocalendar()

        today_count = sum(1 for value in closed_dates if value.date() == today)
        week_count = sum(1 for value in closed_dates if value.isocalendar()[:2] == (iso_year, iso_week))
        month_count = sum(1 for value in closed_dates if value.year == now.year and value.month == now.month)

        if hasattr(self, "report_stats_labels"):
            self.report_stats_labels["today"].setText(str(today_count))
            self.report_stats_labels["week"].setText(str(week_count))
            self.report_stats_labels["month"].setText(str(month_count))

        date_counts: dict[str, int] = {}
        for value in closed_dates:
            label = value.strftime("%Y-%m-%d")
            date_counts[label] = date_counts.get(label, 0) + 1

        options = sorted(date_counts.keys(), reverse=True)
        if hasattr(self, "report_specific_date_input"):
            self.report_specific_date_input.blockSignals(True)
            if self.report_selected_date:
                selected_dt = self._parse_report_date(self.report_selected_date)
                self.report_specific_date_input.setDate(
                    QDate(selected_dt.year, selected_dt.month, selected_dt.day) if selected_dt else REPORT_NULL_QDATE
                )
            else:
                self.report_specific_date_input.setDate(REPORT_NULL_QDATE)
            self.report_specific_date_input.blockSignals(False)

        selected_label = self.report_selected_date if self.report_selected_date else "-"
        selected_count = date_counts.get(self.report_selected_date, 0) if self.report_selected_date else 0
        if hasattr(self, "report_stats_labels"):
            self.report_stats_labels["selected_date"].setText(selected_label)
            self.report_stats_labels["selected_count"].setText(str(selected_count))

        filtered = self._filtered_report_claims()
        self.reports_table.setSortingEnabled(False)
        self.reports_table.setRowCount(len(filtered))
        for row, claim in enumerate(filtered):
            closed_date = self._closed_date_for_claim(claim)
            values = [
                closed_date.strftime("%Y-%m-%d") if closed_date else "",
                claim.claim_id,
                claim.display_customer,
                claim.insurance_company,
                claim.claim_type or "-",
            ]
            for col, value in enumerate(values):
                item = QTableWidgetItem(value or "-")
                item.setData(Qt.UserRole, claim.key)
                self.reports_table.setItem(row, col, item)
        self._auto_size_columns_with_bounds(self.reports_table, [105, 105, 155, 160, 90], [135, 140, 220, 250, 120])
        self.reports_table.setSortingEnabled(True)

    def _on_reports_specific_date_changed(self, _value: QDate) -> None:
        self.report_selected_date = ""
        if self._report_picker_has_value(self.report_specific_date_input):
            self.report_selected_date = self.report_specific_date_input.date().toString("yyyy-MM-dd")
        if self.report_selected_date:
            self.report_start_input.setDate(REPORT_NULL_QDATE)
            self.report_end_input.setDate(REPORT_NULL_QDATE)
        self._populate_reports()

    def _apply_reports_range_filter(self) -> None:
        start = self._report_picker_to_datetime(self.report_start_input)
        end = self._report_picker_to_datetime(self.report_end_input)
        if start and end and start.date() > end.date():
            QMessageBox.warning(self, "Invalid Date Range", "Start Date must be on or before End Date.")
            return
        self.report_selected_date = ""
        self.report_specific_date_input.blockSignals(True)
        self.report_specific_date_input.setDate(REPORT_NULL_QDATE)
        self.report_specific_date_input.blockSignals(False)
        self._populate_reports()

    def _clear_reports_filters(self) -> None:
        self.report_selected_date = ""
        self.report_start_input.setDate(REPORT_NULL_QDATE)
        self.report_end_input.setDate(REPORT_NULL_QDATE)
        self.report_specific_date_input.blockSignals(True)
        self.report_specific_date_input.setDate(REPORT_NULL_QDATE)
        self.report_specific_date_input.blockSignals(False)
        self._populate_reports()

    def _payroll_template_path(self) -> Path:
        claim_tools_folder = str(self.settings.get("claim_tools_folder", "") or "").strip()
        if claim_tools_folder:
            return Path(claim_tools_folder) / "APPRAISER PAYROLL SPREADSHEET.xlsx"
        return APP_DIR / "Claim Tools" / "APPRAISER PAYROLL SPREADSHEET.xlsx"

    def _calculate_payroll(self) -> None:
        if load_workbook is None:
            QMessageBox.warning(self, "Excel Support Missing", "openpyxl is not available for payroll export.")
            return

        start = self._report_picker_to_datetime(self.report_start_input)
        end = self._report_picker_to_datetime(self.report_end_input)
        if not start or not end:
            QMessageBox.warning(self, "Missing Date Range", "Enter both Start Date and End Date in Reports before calculating payroll.")
            return
        if start.date() > end.date():
            QMessageBox.warning(self, "Invalid Date Range", "Start Date must be on or before End Date.")
            return

        filtered = self._filtered_report_claims()
        if not filtered:
            QMessageBox.information(self, "No Claims", "No closed claims match the current report filters.")
            return

        template_path = self._payroll_template_path()
        if not template_path.exists():
            QMessageBox.warning(self, "Template Missing", f"Payroll template not found:\n{template_path}")
            return

        wb = load_workbook(template_path)
        ws = wb[wb.sheetnames[0]]
        ws["A2"] = "Bill Date:"
        ws["B2"] = f"{start.month}/{start.day}/{start.year} - {end.month}/{end.day}/{end.year}"
        ws["B2"].number_format = "@"

        start_row = 5
        template_end_row = 46
        exported = filtered
        total_row = start_row + len(exported)
        clear_through_row = max(template_end_row + 1, total_row)

        for row in range(start_row, clear_through_row + 1):
            for col in ("A", "B", "C", "D", "E", "F"):
                ws[f"{col}{row}"] = None

        self._clear_payroll_trailing_cells(ws, clear_through_row)

        for idx, claim in enumerate(exported, start=start_row):
            owner_name = (claim.customer_name or claim.title or "").strip()
            if owner_name and owner_name.isupper():
                owner_name = owner_name.title()
            ws[f"A{idx}"] = (claim.claim_id or claim.claim_number or "").strip()
            ws[f"B{idx}"] = (claim.insurance_company or "").strip()
            ws[f"C{idx}"] = owner_name
            ws[f"D{idx}"] = 25.0 if claim.claim_type == "Supplement" else 57.0
            ws[f"E{idx}"] = 5.0 if claim.total_loss else None
            ws[f"F{idx}"] = f"=D{idx}+E{idx}"
            self._normalize_payroll_row_style(ws, idx)
            ws[f"D{idx}"].number_format = "0.00"
            ws[f"E{idx}"].number_format = "0.00"
            ws[f"F{idx}"].number_format = "0.00"

        self._normalize_payroll_total_row_style(ws, total_row)
        self._finalize_payroll_total_row(ws, total_row)
        ws[f"F{total_row}"] = f"=SUM(F{start_row}:F{total_row - 1})"
        ws[f"F{total_row}"].number_format = "0.00"
        self._configure_payroll_sheet_for_export(ws, total_row)

        PAYROLL_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        base_name = f"Fernando pay {end.month}-{end.day}-{end.strftime('%y')}"
        output_path = PAYROLL_OUTPUT_DIR / f"{base_name}.xlsx"
        pdf_path = PAYROLL_OUTPUT_DIR / f"{base_name}.pdf"
        wb.save(output_path)

        try:
            self._export_workbook_to_pdf(output_path, pdf_path, total_row)
            os.startfile(str(pdf_path))  # type: ignore[attr-defined]
            QMessageBox.information(self, "Payroll Created", f"Saved payroll files:\n{output_path}\n{pdf_path}")
        except Exception as exc:
            QMessageBox.warning(self, "Payroll Saved", f"Saved workbook:\n{output_path}\n\nCould not create payroll PDF preview:\n{exc}")

    def _normalize_payroll_row_style(self, ws, row: int) -> None:
        source_row = 5
        for col in ("A", "B", "C", "D", "E", "F"):
            target = ws[f"{col}{row}"]
            source = ws[f"{col}{source_row}"]
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.protection = copy(source.protection)

    def _normalize_payroll_total_row_style(self, ws, row: int) -> None:
        source_row = 47
        for col in ("A", "B", "C", "D", "E", "F"):
            target = ws[f"{col}{row}"]
            source = ws[f"{col}{source_row}"]
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.protection = copy(source.protection)

    def _finalize_payroll_total_row(self, ws, row: int) -> None:
        label_cell = ws[f"E{row}"]
        value_cell = ws[f"F{row}"]
        label_cell.value = "TOTAL"
        label_cell.border = copy(ws["E5"].border)
        value_cell.border = copy(ws["F5"].border)
        label_cell.alignment = copy(ws["F4"].alignment)
        value_cell.alignment = copy(ws["F5"].alignment)
        label_font = copy(ws["E5"].font)
        label_font.bold = True
        label_cell.font = label_font
        value_font = copy(ws["F5"].font)
        value_font.bold = True
        value_cell.font = value_font

    def _clear_payroll_trailing_cells(self, ws, last_row: int) -> None:
        max_column = ws.max_column
        if max_column <= 6:
            return
        for row in range(1, last_row + 1):
            for col in range(7, max_column + 1):
                ws.cell(row, col).value = None

    def _configure_payroll_sheet_for_export(self, ws, last_row: int) -> None:
        ws.print_area = f"A1:F{last_row}"
        ws.print_title_rows = "$1:$4"
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 999
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.page_margins.header = 0.15
        ws.page_margins.footer = 0.15
        ws.print_options.horizontalCentered = True

    def _export_workbook_to_pdf(self, workbook_path: Path, pdf_path: Path, last_row: int) -> None:
        workbook_path_text = str(workbook_path).replace("'", "''")
        pdf_path_text = str(pdf_path).replace("'", "''")
        script = textwrap.dedent(
            f"""
            $excel = New-Object -ComObject Excel.Application
            $excel.Visible = $false
            $excel.DisplayAlerts = $false
            $workbook = $null
            $worksheet = $null
            $pageSetup = $null
            try {{
                $workbook = $excel.Workbooks.Open('{workbook_path_text}')
                $worksheet = $workbook.Worksheets.Item(1)
                $pageSetup = $worksheet.PageSetup
                $pageSetup.Orientation = 2
                $pageSetup.Zoom = $false
                $pageSetup.FitToPagesWide = 1
                $pageSetup.FitToPagesTall = 999
                $pageSetup.LeftMargin = $excel.InchesToPoints(0.25)
                $pageSetup.RightMargin = $excel.InchesToPoints(0.25)
                $pageSetup.TopMargin = $excel.InchesToPoints(0.35)
                $pageSetup.BottomMargin = $excel.InchesToPoints(0.35)
                $pageSetup.HeaderMargin = $excel.InchesToPoints(0.15)
                $pageSetup.FooterMargin = $excel.InchesToPoints(0.15)
                $pageSetup.CenterHorizontally = $true
                $pageSetup.PrintTitleRows = '$1:$4'
                $pageSetup.PrintArea = 'A1:F{last_row}'
                $workbook.ExportAsFixedFormat(0, '{pdf_path_text}')
                $workbook.Close($false)
            }} finally {{
                if ($pageSetup -ne $null) {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($pageSetup) }}
                if ($worksheet -ne $null) {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($worksheet) }}
                if ($workbook -ne $null) {{ [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($workbook) }}
                if ($excel -ne $null) {{
                    $excel.Quit()
                    [void][System.Runtime.InteropServices.Marshal]::ReleaseComObject($excel)
                }}
                [GC]::Collect()
                [GC]::WaitForPendingFinalizers()
            }}
            """
        ).strip()
        subprocess.run(["powershell", "-NoProfile", "-Command", script], check=True, capture_output=True, text=True)

    def _populate_table(self, table: QTableWidget, claims: list[ClaimView], include_office: bool = False) -> None:
        table.setSortingEnabled(False)
        table.setRowCount(len(claims))
        include_progress = table.columnCount() == len(OPEN_CLAIM_TABLE_HEADERS)
        for row, claim in enumerate(claims):
            values = [
                claim.claim_id,
                claim.display_customer,
                claim.insurance_company,
                claim.vehicle,
                claim.shop_name,
                claim.town,
                claim.date_of_loss,
                claim.status,
            ]
            if include_progress:
                values.append(claim.office_progress_status or "-")
            if include_office:
                values.extend([claim.office_appt_when, claim.office_additional_notes])
            for col, value in enumerate(values):
                item = QTableWidgetItem(value or "-")
                item.setData(Qt.UserRole, claim.key)
                table.setItem(row, col, item)
        self._auto_size_table_columns(table, include_office)
        table.setSortingEnabled(True)
        if table is self.open_table:
            table.sortItems(0, Qt.AscendingOrder)
            table.horizontalHeader().setSortIndicator(0, Qt.AscendingOrder)

    def _populate_workspace(self) -> None:
        watched_folders = [str(folder) for folder in self.settings.get("watched_folders", [])]
        watched_summary = "  |  ".join(watched_folders) if watched_folders else "-"
        self.watched_folders_display.setText(watched_summary)
        self.watched_folders_display.setToolTip("\n".join(watched_folders) if watched_folders else "-")

        claim_tools_folder = str(self.settings.get("claim_tools_folder", "")) or "-"
        self.claim_tools_folder_compact.setText(claim_tools_folder)
        self.claim_tools_folder_compact.setToolTip(claim_tools_folder)

    def _populate_claim_tools(self) -> None:
        self.claim_tools_contacts_table.setSortingEnabled(False)
        self.claim_tools_contacts_table.setRowCount(len(self.claim_tools_contacts))
        for row, entry in enumerate(self.claim_tools_contacts):
            values = [entry.name, entry.number, entry.prompt_guide, entry.notes]
            for col, value in enumerate(values):
                item = QTableWidgetItem(value or "-")
                item.setData(Qt.UserRole, row)
                self.claim_tools_contacts_table.setItem(row, col, item)
        self._auto_size_aux_table(
            self.claim_tools_contacts_table,
            minimums=[150, 120, 100, 220],
            maximums=[260, 170, 150, 900],
            stretch_column=3,
        )
        self.claim_tools_contacts_table.setSortingEnabled(True)

        self.claim_tools_files_table.setSortingEnabled(False)
        self.claim_tools_files_table.setRowCount(len(self.claim_tools_files))
        for row, entry in enumerate(self.claim_tools_files):
            values = [entry.section, entry.label, entry.file_name, entry.notes]
            for col, value in enumerate(values):
                item = QTableWidgetItem(value or "-")
                item.setData(Qt.UserRole, row)
                self.claim_tools_files_table.setItem(row, col, item)
        self._auto_size_aux_table(
            self.claim_tools_files_table,
            minimums=[150, 150, 220, 220],
            maximums=[230, 220, 340, 900],
            stretch_column=3,
        )
        self.claim_tools_files_table.setSortingEnabled(True)

    def _populate_body_shops(self) -> None:
        self.body_shops_table.setSortingEnabled(False)
        self.body_shops_table.setRowCount(len(self.body_shop_entries))
        for row, entry in enumerate(self.body_shop_entries):
            values = [
                entry.shop_name,
                entry.contact_name,
                entry.phone,
                entry.email,
                entry.tax_id,
                entry.address,
                entry.body_rate,
                entry.paint_rate,
                entry.frame_rate,
                entry.mechanical_rate,
                entry.certifications,
                entry.negotiation_notes,
                entry.notes,
            ]
            for col, value in enumerate(values):
                item = QTableWidgetItem(value or "-")
                item.setData(Qt.UserRole, row)
                self.body_shops_table.setItem(row, col, item)
        self._auto_size_aux_table(
            self.body_shops_table,
            minimums=[170, 120, 120, 170, 120, 180, 90, 90, 90, 110, 140, 180, 220],
            maximums=[280, 180, 150, 240, 180, 300, 120, 120, 120, 150, 220, 320, 420],
        )
        self.body_shops_table.setSortingEnabled(True)

    def _populate_insurance_companies(self) -> None:
        self.insurance_companies_table.setSortingEnabled(False)
        self.insurance_companies_table.setRowCount(len(self.insurance_company_entries))
        for row, entry in enumerate(self.insurance_company_entries):
            values = [
                entry.company_name,
                entry.quick_summary,
                entry.fatal_errors,
                entry.photo_rules,
                entry.estimate_supp_rules,
                entry.parts_rules,
                entry.total_loss_rules,
                entry.tow_rules,
                entry.supplement_rules,
                entry.betterment_depreciation_rules,
                entry.documentation_requirements,
                entry.rates_and_sales_tax_rules,
                entry.miscellaneous_rules,
                entry.contact_information,
                entry.labor_rates,
                entry.total_loss_threshold,
                entry.notes,
            ]
            for col, value in enumerate(values):
                item = QTableWidgetItem(value or "-")
                item.setData(Qt.UserRole, row)
                self.insurance_companies_table.setItem(row, col, item)
        self._auto_size_aux_table(
            self.insurance_companies_table,
            minimums=[180, 180, 180, 220, 220, 180, 220, 150, 180, 180, 220, 180, 180, 180, 180, 160, 220],
            maximums=[280, 260, 260, 340, 340, 260, 340, 260, 300, 280, 340, 280, 280, 260, 260, 220, 360],
        )
        self.insurance_companies_table.setSortingEnabled(True)

    def _meaningful_route_value(self, value: str) -> str:
        cleaned = (value or "").strip()
        if not cleaned:
            return ""
        if cleaned.lower() in {"no shop chosen", "-", "n/a", "unknown"}:
            return ""
        return cleaned

    def _route_default_address(self, claim: ClaimView) -> str:
        override = self._meaningful_route_value(claim.route_address_override or "")
        if override:
            return override
        inspection_location = self._meaningful_route_value(claim.location_of_vehicle or "")
        if inspection_location:
            return inspection_location
        shop_name = self._meaningful_route_value(claim.shop_name or "")
        if shop_name:
            return shop_name
        return self._meaningful_route_value(claim.owner_address or "")

    def _populate_route_planner(self) -> None:
        if not hasattr(self, "route_available_table") or not hasattr(self, "route_selected_table"):
            return

        route_records = list(self.filtered_all)
        route_filter = (self.route_status_filter.currentText().strip() or "Open").lower()
        if route_filter == "open":
            route_records = [claim for claim in route_records if claim.status.lower() == "open"]
        elif route_filter == "closed":
            route_records = [claim for claim in route_records if claim.status.lower() == "closed"]

        if self.route_appt_only_checkbox.isChecked():
            route_records = [
                claim
                for claim in route_records
                if (claim.office_progress_status or "").strip().lower() == "appointment scheduled"
            ]

        claims_by_key = {claim.key: claim for claim in self.claims}
        self.route_plan_keys = [key for key in self.route_plan_keys if key in claims_by_key]
        available_claims = [claim for claim in route_records if claim.key not in self.route_plan_keys]
        selected_claims = [claims_by_key[key] for key in self.route_plan_keys if key in claims_by_key]
        self.route_status_label.setText(
            f"{len(available_claims)} available claims, {len(selected_claims)} stop(s) in tomorrow's route."
        )

        self.route_available_table.setSortingEnabled(False)
        self.route_available_table.setRowCount(len(available_claims))
        for row, claim in enumerate(available_claims):
            values = [claim.claim_id, claim.display_customer, self._route_default_address(claim) or "-"]
            for col, value in enumerate(values):
                item = QTableWidgetItem(value)
                item.setData(Qt.UserRole, claim.key)
                self.route_available_table.setItem(row, col, item)
        self._auto_size_aux_table(
            self.route_available_table,
            minimums=[100, 170, 260],
            maximums=[150, 240, 900],
            stretch_column=2,
        )
        self.route_available_table.setSortingEnabled(True)

        self.route_selected_table.setSortingEnabled(False)
        self.route_selected_table.setRowCount(len(selected_claims))
        for row, claim in enumerate(selected_claims):
            values = [str(row + 1), claim.claim_id, claim.display_customer, self._route_default_address(claim) or "-"]
            for col, value in enumerate(values):
                item = QTableWidgetItem(value)
                item.setData(Qt.UserRole, claim.key)
                self.route_selected_table.setItem(row, col, item)
        self._auto_size_aux_table(
            self.route_selected_table,
            minimums=[60, 100, 160, 260],
            maximums=[70, 140, 220, 900],
            stretch_column=3,
        )

        self.route_selected_stop_key = None
        self.route_address_input.setText("")
        self.route_open_maps_button.setEnabled(bool(selected_claims))
        self.route_remove_stop_button.setEnabled(False)
        self.route_save_address_button.setEnabled(False)
        self.route_move_up_button.setEnabled(False)
        self.route_move_down_button.setEnabled(False)

    def _add_selected_claims_to_route(self) -> None:
        selected_rows = sorted({item.row() for item in self.route_available_table.selectedItems()})
        if not selected_rows:
            return
        for row in selected_rows:
            item = self.route_available_table.item(row, 0)
            if not item:
                continue
            key = item.data(Qt.UserRole)
            if key and key not in self.route_plan_keys:
                self.route_plan_keys.append(key)
        self.repo.save_route_plan_keys(self.route_plan_keys)
        self._load_claims()

    def _on_route_stop_selected(self) -> None:
        row = self.route_selected_table.currentRow()
        if row < 0:
            self.route_selected_stop_key = None
            self.route_address_input.setText("")
            self.route_remove_stop_button.setEnabled(False)
            self.route_save_address_button.setEnabled(False)
            self.route_move_up_button.setEnabled(False)
            self.route_move_down_button.setEnabled(False)
            return
        item = self.route_selected_table.item(row, 0)
        if not item:
            return
        key = item.data(Qt.UserRole)
        claim = next((record for record in self.claims if record.key == key), None)
        self.route_selected_stop_key = key
        self.route_address_input.setText(self._route_default_address(claim) if claim else "")
        self.route_remove_stop_button.setEnabled(True)
        self.route_save_address_button.setEnabled(True)
        self.route_move_up_button.setEnabled(row > 0)
        self.route_move_down_button.setEnabled(row < len(self.route_plan_keys) - 1)

    def _save_route_stop_address(self) -> None:
        if not self.route_selected_stop_key:
            return
        self.repo.update_route_address(self.route_selected_stop_key, self.route_address_input.text().strip())
        self._load_claims()

    def _remove_route_stop(self) -> None:
        if not self.route_selected_stop_key:
            return
        self.route_plan_keys = [key for key in self.route_plan_keys if key != self.route_selected_stop_key]
        self.repo.save_route_plan_keys(self.route_plan_keys)
        self._load_claims()

    def _move_route_stop(self, direction: int) -> None:
        if not self.route_selected_stop_key or direction not in (-1, 1):
            return

        try:
            current_index = self.route_plan_keys.index(self.route_selected_stop_key)
        except ValueError:
            return

        new_index = current_index + direction
        if new_index < 0 or new_index >= len(self.route_plan_keys):
            return

        self.route_plan_keys[current_index], self.route_plan_keys[new_index] = (
            self.route_plan_keys[new_index],
            self.route_plan_keys[current_index],
        )
        self.repo.save_route_plan_keys(self.route_plan_keys)
        self._populate_route_planner()
        self._reselect_route_stop(self.route_selected_stop_key)

    def _clear_route_plan(self) -> None:
        self.route_plan_keys = []
        self.repo.save_route_plan_keys([])
        self._load_claims()

    def _reselect_route_stop(self, claim_key: str | None) -> None:
        if not claim_key:
            return
        for row in range(self.route_selected_table.rowCount()):
            item = self.route_selected_table.item(row, 0)
            if item and item.data(Qt.UserRole) == claim_key:
                self.route_selected_table.setCurrentCell(row, 0)
                self.route_selected_table.selectRow(row)
                return

    def _open_route_in_google_maps(self) -> None:
        claims_by_key = {claim.key: claim for claim in self.claims}
        selected_claims = [claims_by_key[key] for key in self.route_plan_keys if key in claims_by_key]
        stops = [self._route_default_address(claim) for claim in selected_claims if self._route_default_address(claim)]
        if not stops:
            QMessageBox.information(self, APP_NAME, "Add stops to the route first.")
            return

        start_from = self.route_home_input.text().strip()
        self.repo.save_route_home_address(start_from)
        origin = start_from or stops[0]
        destination = stops[-1]
        waypoints = stops[:-1] if not start_from else stops[:-1]
        query_parts = [
            "https://www.google.com/maps/dir/?api=1",
            f"origin={quote(origin)}",
            f"destination={quote(destination)}",
            "travelmode=driving",
        ]
        if waypoints:
            encoded_waypoints = "|".join(quote(stop) for stop in waypoints)
            query_parts.append(f"waypoints={encoded_waypoints}")
        route_url = " & ".join(query_parts).replace(" & ", "&")
        webbrowser.open(route_url)
        try:
            self._send_route_email(route_url)
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Opened the route, but could not email it:\n{exc}")

    def _send_route_email(self, route_url: str) -> None:
        sender_email, app_password = self._get_email_credentials()
        if not sender_email or not app_password:
            raise RuntimeError("A Gmail address and app password are required for route email.")

        message = EmailMessage()
        message["From"] = sender_email
        message["To"] = sender_email
        message["Subject"] = "Tomorrow Route"
        message.set_content(f"Google Maps route link:\n\n{route_url}")
        self._send_email_message(message)

    def _auto_size_aux_table(
        self,
        table: QTableWidget,
        minimums: list[int],
        maximums: list[int],
        stretch_column: int | None = None,
    ) -> None:
        header = table.horizontalHeader()
        header_metrics = QFontMetrics(header.font())
        cell_metrics = QFontMetrics(table.font())
        for index in range(table.columnCount()):
            if stretch_column is not None and index == stretch_column:
                header.setSectionResizeMode(index, QHeaderView.Stretch)
                continue
            header.setSectionResizeMode(index, QHeaderView.Interactive)
            header_text = table.horizontalHeaderItem(index).text() if table.horizontalHeaderItem(index) else ""
            header_width = header_metrics.horizontalAdvance(header_text) + 24
            sample_widths: list[int] = []
            for row in range(table.rowCount()):
                item = table.item(row, index)
                if not item:
                    continue
                text = item.text().strip()
                if text and text != "-":
                    sample_widths.append(cell_metrics.horizontalAdvance(text) + 22)
            if sample_widths:
                sample_widths.sort()
                median_width = sample_widths[len(sample_widths) // 2]
                target = max(header_width, median_width)
            else:
                target = header_width
            target = max(minimums[index], min(target, maximums[index]))
            table.setColumnWidth(index, target)

    def _remove_selected_watched_folder(self) -> None:
        watched_folders = [str(folder) for folder in self.settings.get("watched_folders", [])]
        if not watched_folders:
            QMessageBox.information(self, APP_NAME, "There are no watched folders to remove.")
            return

        target, ok = QInputDialog.getItem(
            self,
            APP_NAME,
            "Which watched folder do you want to remove?",
            watched_folders,
            0,
            False,
        )
        if not ok or not target:
            return

        updated = [folder for folder in watched_folders if folder != target]
        self.repo.update_settings(updated, str(self.settings.get("claim_tools_folder", "")))
        self._load_claims()

    def _edit_workspace_settings(self) -> None:
        watched_folders = [str(folder) for folder in self.settings.get("watched_folders", [])]
        current_claim_tools = str(self.settings.get("claim_tools_folder", "") or "")

        options: list[tuple[str, str, int | None]] = []
        if len(watched_folders) > 0:
            options.append(("Open Claims Folder", watched_folders[0], 0))
        if len(watched_folders) > 1:
            options.append(("Closed Claims Folder", watched_folders[1], 1))
        for index, folder in enumerate(watched_folders[2:], start=2):
            options.append((f"Watched Folder {index + 1}", folder, index))
        options.append(("Claim Tools Folder", current_claim_tools, None))

        choice_labels = [label for label, _, _ in options]
        selected_label, ok = QInputDialog.getItem(
            self,
            APP_NAME,
            "Which folder do you want to edit?",
            choice_labels,
            0,
            False,
        )
        if not ok or not selected_label:
            return

        selected = next((option for option in options if option[0] == selected_label), None)
        if not selected:
            return

        _, current_path, watched_index = selected
        new_path = QFileDialog.getExistingDirectory(
            self,
            f"Select {selected_label}",
            current_path or os.path.expanduser("~"),
        )
        if not new_path:
            return

        if watched_index is None:
            self.repo.update_settings(watched_folders, new_path)
        else:
            updated = list(watched_folders)
            while len(updated) <= watched_index:
                updated.append("")
            updated[watched_index] = new_path
            self.repo.update_settings(updated, current_claim_tools)
        self._load_claims()

    def _handle_table_selection(self, table: QTableWidget) -> None:
        row = table.currentRow()
        if row < 0:
            return
        claim = self._claim_from_table_row(table, row)
        if not claim:
            return
        self.current_claim = claim
        self._render_claim(claim)

    def _claim_from_table_row(self, table: QTableWidget, row: int) -> ClaimView | None:
        if row < 0:
            return None
        item = table.item(row, 0)
        if not item:
            return None
        key = item.data(Qt.UserRole)
        if not key:
            return None
        return next((record for record in self.claims if record.key == key), None)

    def _show_claim_context_menu(self, table: QTableWidget, pos) -> None:
        row = table.rowAt(pos.y())
        if row < 0:
            return
        table.selectRow(row)
        claim = self._claim_from_table_row(table, row)
        if not claim:
            return
        self.current_claim = claim
        self._render_claim(claim)

        menu = QMenu(self)
        open_assign_action = menu.addAction("Open Working Sheet")
        email_assign_action = menu.addAction("Email Assignment Sheet")
        open_folder_action = menu.addAction("Open Claim Folder")
        menu.addSeparator()
        review_action = menu.addAction("Review Email")
        rmc_action = menu.addAction("RMC")
        supplement_action = menu.addAction("Request Supplement")
        prelim_action = menu.addAction("Send Prelim")
        email_shop_copy_action = menu.addAction("Email Copy To Shop")
        nada_action = menu.addAction("NADA Value PDF")
        shop_info_action = menu.addAction("Edit Shop Info")
        claim_notes_action = menu.addAction("Claim Notes")
        menu.addSeparator()
        reopen_action = menu.addAction("Reopen")
        mark_closed_action = menu.addAction("Mark Closed")
        claim_tools_menu = menu.addMenu("Claim Tools")
        self._populate_claim_tools_menu(claim_tools_menu)
        menu.addSeparator()
        copy_claim_id_action = menu.addAction("Copy Claim ID")
        copy_customer_action = menu.addAction("Copy Customer")

        chosen = menu.exec(table.viewport().mapToGlobal(pos))
        if chosen == open_assign_action:
            self._open_assignment_sheet()
        elif chosen == email_assign_action:
            self._email_assignment_sheet()
        elif chosen == open_folder_action:
            self._open_claim_folder()
        elif chosen == review_action:
            self._send_review_email()
        elif chosen == rmc_action:
            self._send_rmc_email()
        elif chosen == supplement_action:
            self._request_supplement_email()
        elif chosen == prelim_action:
            self._send_prelim_email()
        elif chosen == email_shop_copy_action:
            self._open_shop_email_draft()
        elif chosen == nada_action:
            self._open_nada_value()
        elif chosen == shop_info_action:
            self._edit_shop_info_for_claim(claim)
        elif chosen == claim_notes_action:
            self._open_claim_notes_dialog()
        elif chosen == reopen_action:
            self._set_current_claim_status("Open")
        elif chosen == mark_closed_action:
            self._set_current_claim_status("Closed")
        elif chosen == copy_claim_id_action:
            QApplication.clipboard().setText(claim.claim_id or "")
        elif chosen == copy_customer_action:
            QApplication.clipboard().setText(claim.display_customer or "")

    def _populate_claim_tools_menu(self, menu: QMenu) -> None:
        menu.addAction("AIG Claim Summary", self._generate_aig_claim_summary)
        menu.addAction("Default Claim Notes", self._generate_default_claim_notes)
        menu.addAction("Mitchell Total Loss", self._generate_mitchell_total_loss)
        menu.addAction("Autosource PDA Sub Level", self._generate_autosource_pda_sub_level)
        menu.addSeparator()
        tools_root = Path(str(self.settings.get("claim_tools_folder", "") or ""))
        if not tools_root.exists():
            action = menu.addAction("Claim Tools Folder Missing")
            action.setEnabled(False)
            return
        files = [path for path in sorted(tools_root.iterdir(), key=lambda path: path.name.lower()) if path.is_file()]
        if not files:
            action = menu.addAction("No Files Found")
            action.setEnabled(False)
            return
        for path in files:
            menu.addAction(path.name, lambda file_path=path: self._open_claim_tool_path(file_path))

    def _render_claim(self, claim: ClaimView) -> None:
        values = {
            "customer_name": claim.display_customer or "",
            "insurance_company": claim.insurance_company or "",
            "claim_number": claim.claim_number or "",
            "date_of_loss": claim.date_of_loss or "",
            "town": claim.town or "",
            "shop_name": claim.shop_name or "",
            "shop_email": claim.shop_email or "",
            "contact_phone": claim.contact_phone or "",
            "vehicle": claim.vehicle or "",
            "vin": claim.vin or "",
            "owner_address": claim.owner_address or "",
            "location_of_vehicle": claim.location_of_vehicle or claim.shop_name or claim.owner_address or "",
        }
        self.claim_id_value.setText(claim.claim_id or "-")
        for key, widget in self.detail_inputs.items():
            widget.setText(values.get(key, ""))
        self._set_combo_value(self.office_progress_input, claim.office_progress_status or "No Contact Yet")
        self.office_appt_input.setText(claim.office_appt_when or "")
        self.office_notes_input.setPlainText(claim.office_additional_notes or "")
        if hasattr(self, "office_summary_view"):
            self.office_summary_view.setPlainText(self._build_office_summary(claim))
        self.notes_view.setPlainText(claim.assignment_claim_notes or "")
        self.reopen_button.setEnabled(claim.status.strip().lower() == "closed")
        self.mark_closed_button.setEnabled(claim.status.strip().lower() != "closed")

    def _handle_claim_table_double_click(self, table: QTableWidget, row: int, column: int) -> None:
        claim = self._claim_from_table_row(table, row)
        if not claim:
            return
        header_item = table.horizontalHeaderItem(column)
        column_name = header_item.text() if header_item else ""
        if column_name == "Shop":
            self._edit_shop_info_for_claim(claim)

    def _normalize_shop_name(self, value: str) -> str:
        return re.sub(r"[^a-z0-9]+", "", (value or "").lower())

    def _match_body_shop_index(self, claim: ClaimView) -> int | None:
        claim_shop = (claim.shop_name or "").strip()
        if not claim_shop or claim_shop == "No Shop Chosen":
            return None

        normalized_claim = self._normalize_shop_name(claim_shop)
        if not normalized_claim:
            return None

        exact_match = next(
            (
                index
                for index, entry in enumerate(self.body_shop_entries)
                if self._normalize_shop_name(entry.shop_name) == normalized_claim
            ),
            None,
        )
        if exact_match:
            return exact_match

        contains_match = next(
            (
                index
                for index, entry in enumerate(self.body_shop_entries)
                if normalized_claim in self._normalize_shop_name(entry.shop_name)
                or self._normalize_shop_name(entry.shop_name) in normalized_claim
            ),
            None,
        )
        return contains_match

    def _body_shop_dialog_fields(self) -> list[tuple[str, str, bool]]:
        return [
            ("shop_name", "Shop Name", False),
            ("contact_name", "Contact Name", False),
            ("phone", "Phone", False),
            ("email", "Email", False),
            ("tax_id", "Tax ID", False),
            ("address", "Address", False),
            ("body_rate", "Body Rate", False),
            ("paint_rate", "Paint Rate", False),
            ("frame_rate", "Frame Rate", False),
            ("mechanical_rate", "Mechanical Rate", False),
            ("certifications", "Certifications", True),
            ("negotiation_notes", "Negotiation Notes", True),
            ("notes", "Notes", True),
        ]

    def _body_shop_dialog_values(self, entry: BodyShopEntry) -> dict[str, str]:
        return {
            "shop_name": entry.shop_name,
            "contact_name": entry.contact_name,
            "phone": entry.phone,
            "email": entry.email,
            "tax_id": entry.tax_id,
            "address": entry.address,
            "body_rate": entry.body_rate,
            "paint_rate": entry.paint_rate,
            "frame_rate": entry.frame_rate,
            "mechanical_rate": entry.mechanical_rate,
            "certifications": entry.certifications,
            "negotiation_notes": entry.negotiation_notes,
            "notes": entry.notes,
        }

    def _body_shop_entry_from_values(self, values: dict[str, str]) -> BodyShopEntry:
        return BodyShopEntry(
            shop_name=values["shop_name"],
            contact_name=values["contact_name"],
            phone=values["phone"],
            email=values["email"],
            tax_id=values["tax_id"],
            address=values["address"],
            body_rate=values["body_rate"],
            paint_rate=values["paint_rate"],
            frame_rate=values["frame_rate"],
            mechanical_rate=values["mechanical_rate"],
            certifications=values["certifications"],
            negotiation_notes=values["negotiation_notes"],
            notes=values["notes"],
        )

    def _shop_info_text(self, entry: BodyShopEntry) -> str:
        lines = [
            f"Shop: {entry.shop_name or '-'}",
            f"Contact: {entry.contact_name or '-'}",
            f"Phone: {entry.phone or '-'}",
            f"Email: {entry.email or '-'}",
            f"Tax ID: {entry.tax_id or '-'}",
            f"Address: {entry.address or '-'}",
            f"Body Rate: {entry.body_rate or '-'}",
            f"Paint Rate: {entry.paint_rate or '-'}",
            f"Frame Rate: {entry.frame_rate or '-'}",
            f"Mechanical Rate: {entry.mechanical_rate or '-'}",
            f"Certifications: {entry.certifications or '-'}",
            f"Negotiation Notes: {entry.negotiation_notes or '-'}",
            f"Notes: {entry.notes or '-'}",
        ]
        return "\n".join(lines)

    def _edit_shop_info_for_claim(self, claim: ClaimView) -> None:
        index = self._match_body_shop_index(claim)
        if index is None:
            claim_shop = (claim.shop_name or "").strip()
            if not claim_shop or claim_shop == "No Shop Chosen":
                QMessageBox.information(self, APP_NAME, "There is no shop on this claim yet.")
                return
            entry = BodyShopEntry(
                shop_name=claim_shop,
                contact_name="",
                phone=claim.shop_phone or "",
                email=claim.shop_email or "",
                tax_id="",
                address="",
                body_rate="",
                paint_rate="",
                frame_rate="",
                mechanical_rate="",
                certifications="",
                negotiation_notes="",
                notes="",
            )
            dialog_title = f"Add Shop Info - {entry.shop_name}"
        else:
            entry = self.body_shop_entries[index]
            dialog_title = f"Edit Shop Info - {entry.shop_name}"

        dialog = RecordEditorDialog(
            self,
            dialog_title,
            self._body_shop_dialog_fields(),
            self._body_shop_dialog_values(entry),
        )
        if dialog.exec() != QDialog.Accepted:
            return
        values = dialog.values()
        updated_entry = self._body_shop_entry_from_values(values)
        if index is None:
            self.body_shop_entries.append(updated_entry)
        else:
            self.body_shop_entries[index] = updated_entry
        self.repo.save_body_shop_database(self.body_shop_entries)
        self._populate_body_shops()
        if self.current_claim:
            self._reselect_current_claim()

    def _selected_claim_tool_contact_index(self) -> int | None:
        row = self.claim_tools_contacts_table.currentRow()
        if row < 0:
            return None
        item = self.claim_tools_contacts_table.item(row, 0)
        if not item:
            return None
        value = item.data(Qt.UserRole)
        return int(value) if value is not None else None

    def _selected_claim_tool_file_index(self) -> int | None:
        row = self.claim_tools_files_table.currentRow()
        if row < 0:
            return None
        item = self.claim_tools_files_table.item(row, 0)
        if not item:
            return None
        value = item.data(Qt.UserRole)
        return int(value) if value is not None else None

    def _selected_body_shop_index(self) -> int | None:
        row = self.body_shops_table.currentRow()
        if row < 0:
            return None
        item = self.body_shops_table.item(row, 0)
        if not item:
            return None
        value = item.data(Qt.UserRole)
        return int(value) if value is not None else None

    def _selected_insurance_company_index(self) -> int | None:
        row = self.insurance_companies_table.currentRow()
        if row < 0:
            return None
        item = self.insurance_companies_table.item(row, 0)
        if not item:
            return None
        value = item.data(Qt.UserRole)
        return int(value) if value is not None else None

    def _add_claim_tool_contact(self) -> None:
        dialog = RecordEditorDialog(
            self,
            "Add Claim Tools Contact",
            [
                ("name", "Name", False),
                ("number", "Number", False),
                ("prompt_guide", "Prompt Guide", False),
                ("notes", "Notes", True),
            ],
        )
        if dialog.exec() != QDialog.Accepted:
            return
        values = dialog.values()
        self.claim_tools_contacts.append(
            ClaimToolContact(
                name=values["name"],
                number=values["number"],
                prompt_guide=values["prompt_guide"],
                notes=values["notes"],
            )
        )
        self.repo.save_claim_tools_contacts(self.claim_tools_contacts)
        self._populate_claim_tools()

    def _edit_claim_tool_contact(self) -> None:
        index = self._selected_claim_tool_contact_index()
        if index is None:
            QMessageBox.information(self, APP_NAME, "Choose a claim tools contact first.")
            return
        entry = self.claim_tools_contacts[index]
        dialog = RecordEditorDialog(
            self,
            "Edit Claim Tools Contact",
            [
                ("name", "Name", False),
                ("number", "Number", False),
                ("prompt_guide", "Prompt Guide", False),
                ("notes", "Notes", True),
            ],
            {
                "name": entry.name,
                "number": entry.number,
                "prompt_guide": entry.prompt_guide,
                "notes": entry.notes,
            },
        )
        if dialog.exec() != QDialog.Accepted:
            return
        values = dialog.values()
        self.claim_tools_contacts[index] = ClaimToolContact(
            name=values["name"],
            number=values["number"],
            prompt_guide=values["prompt_guide"],
            notes=values["notes"],
        )
        self.repo.save_claim_tools_contacts(self.claim_tools_contacts)
        self._populate_claim_tools()

    def _remove_claim_tool_contact(self) -> None:
        index = self._selected_claim_tool_contact_index()
        if index is None:
            QMessageBox.information(self, APP_NAME, "Choose a claim tools contact first.")
            return
        del self.claim_tools_contacts[index]
        self.repo.save_claim_tools_contacts(self.claim_tools_contacts)
        self._populate_claim_tools()

    def _call_selected_claim_tool_contact(self) -> None:
        index = self._selected_claim_tool_contact_index()
        if index is None:
            return
        number = self.claim_tools_contacts[index].number.strip()
        if not number:
            QMessageBox.information(self, APP_NAME, "This contact does not have a number yet.")
            return
        tel_value = "".join(ch for ch in number if ch.isalnum() or ch == "+")
        if not tel_value:
            QMessageBox.information(self, APP_NAME, "This number cannot be used for a call action.")
            return
        QDesktopServices.openUrl(QUrl(f"tel:{tel_value}"))

    def _add_claim_tool_file(self) -> None:
        dialog = RecordEditorDialog(
            self,
            "Add Claim Tools File",
            [
                ("section", "Section", False),
                ("label", "Label", False),
                ("file_path", "File", False),
                ("notes", "Notes", True),
            ],
            browse_keys={"file_path"},
        )
        if dialog.exec() != QDialog.Accepted:
            return
        values = dialog.values()
        self.claim_tools_files.append(
            ClaimToolFile(
                section=values["section"],
                label=values["label"],
                file_path=values["file_path"],
                notes=values["notes"],
            )
        )
        self.repo.save_claim_tools_files(self.claim_tools_files)
        self._populate_claim_tools()

    def _edit_claim_tool_file(self) -> None:
        index = self._selected_claim_tool_file_index()
        if index is None:
            QMessageBox.information(self, APP_NAME, "Choose a claim tools file first.")
            return
        entry = self.claim_tools_files[index]
        dialog = RecordEditorDialog(
            self,
            "Edit Claim Tools File",
            [
                ("section", "Section", False),
                ("label", "Label", False),
                ("file_path", "File", False),
                ("notes", "Notes", True),
            ],
            {
                "section": entry.section,
                "label": entry.label,
                "file_path": entry.file_path,
                "notes": entry.notes,
            },
            browse_keys={"file_path"},
        )
        if dialog.exec() != QDialog.Accepted:
            return
        values = dialog.values()
        self.claim_tools_files[index] = ClaimToolFile(
            section=values["section"],
            label=values["label"],
            file_path=values["file_path"],
            notes=values["notes"],
        )
        self.repo.save_claim_tools_files(self.claim_tools_files)
        self._populate_claim_tools()

    def _remove_claim_tool_file(self) -> None:
        index = self._selected_claim_tool_file_index()
        if index is None:
            QMessageBox.information(self, APP_NAME, "Choose a claim tools file first.")
            return
        del self.claim_tools_files[index]
        self.repo.save_claim_tools_files(self.claim_tools_files)
        self._populate_claim_tools()

    def _open_selected_claim_tool_file(self) -> None:
        index = self._selected_claim_tool_file_index()
        if index is None:
            return
        file_path = self.claim_tools_files[index].file_path.strip()
        if not file_path:
            QMessageBox.information(self, APP_NAME, "This entry does not have a file path yet.")
            return
        path = Path(file_path)
        if not path.exists():
            QMessageBox.warning(self, APP_NAME, f"File not found:\n{file_path}")
            return
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(path)))

    def _add_body_shop_entry(self) -> None:
        dialog = RecordEditorDialog(
            self,
            "Add Body Shop",
            [
                ("shop_name", "Shop Name", False),
                ("contact_name", "Contact Name", False),
                ("phone", "Phone", False),
                ("email", "Email", False),
                ("tax_id", "Tax ID", False),
                ("address", "Address", False),
                ("body_rate", "Body Rate", False),
                ("paint_rate", "Paint Rate", False),
                ("frame_rate", "Frame Rate", False),
                ("mechanical_rate", "Mechanical Rate", False),
                ("certifications", "Certifications", True),
                ("negotiation_notes", "Negotiation Notes", True),
                ("notes", "Notes", True),
            ],
        )
        if dialog.exec() != QDialog.Accepted:
            return
        values = dialog.values()
        self.body_shop_entries.append(
            BodyShopEntry(
                shop_name=values["shop_name"],
                contact_name=values["contact_name"],
                phone=values["phone"],
                email=values["email"],
                tax_id=values["tax_id"],
                address=values["address"],
                body_rate=values["body_rate"],
                paint_rate=values["paint_rate"],
                frame_rate=values["frame_rate"],
                mechanical_rate=values["mechanical_rate"],
                certifications=values["certifications"],
                negotiation_notes=values["negotiation_notes"],
                notes=values["notes"],
            )
        )
        self.repo.save_body_shop_database(self.body_shop_entries)
        self._populate_body_shops()

    def _edit_body_shop_entry(self) -> None:
        index = self._selected_body_shop_index()
        if index is None:
            QMessageBox.information(self, APP_NAME, "Choose a body shop first.")
            return
        entry = self.body_shop_entries[index]
        dialog = RecordEditorDialog(
            self,
            "Edit Body Shop",
            [
                ("shop_name", "Shop Name", False),
                ("contact_name", "Contact Name", False),
                ("phone", "Phone", False),
                ("email", "Email", False),
                ("tax_id", "Tax ID", False),
                ("address", "Address", False),
                ("body_rate", "Body Rate", False),
                ("paint_rate", "Paint Rate", False),
                ("frame_rate", "Frame Rate", False),
                ("mechanical_rate", "Mechanical Rate", False),
                ("certifications", "Certifications", True),
                ("negotiation_notes", "Negotiation Notes", True),
                ("notes", "Notes", True),
            ],
            {
                "shop_name": entry.shop_name,
                "contact_name": entry.contact_name,
                "phone": entry.phone,
                "email": entry.email,
                "tax_id": entry.tax_id,
                "address": entry.address,
                "body_rate": entry.body_rate,
                "paint_rate": entry.paint_rate,
                "frame_rate": entry.frame_rate,
                "mechanical_rate": entry.mechanical_rate,
                "certifications": entry.certifications,
                "negotiation_notes": entry.negotiation_notes,
                "notes": entry.notes,
            },
        )
        if dialog.exec() != QDialog.Accepted:
            return
        values = dialog.values()
        self.body_shop_entries[index] = BodyShopEntry(
            shop_name=values["shop_name"],
            contact_name=values["contact_name"],
            phone=values["phone"],
            email=values["email"],
            tax_id=values["tax_id"],
            address=values["address"],
            body_rate=values["body_rate"],
            paint_rate=values["paint_rate"],
            frame_rate=values["frame_rate"],
            mechanical_rate=values["mechanical_rate"],
            certifications=values["certifications"],
            negotiation_notes=values["negotiation_notes"],
            notes=values["notes"],
        )
        self.repo.save_body_shop_database(self.body_shop_entries)
        self._populate_body_shops()

    def _remove_body_shop_entry(self) -> None:
        index = self._selected_body_shop_index()
        if index is None:
            QMessageBox.information(self, APP_NAME, "Choose a body shop first.")
            return
        del self.body_shop_entries[index]
        self.repo.save_body_shop_database(self.body_shop_entries)
        self._populate_body_shops()

    def _handle_body_shop_double_click(self, row: int, column: int) -> None:
        if row < 0:
            return
        self.body_shops_table.selectRow(row)
        if column == 2:
            index = self._selected_body_shop_index()
            if index is None:
                return
            number = self.body_shop_entries[index].phone.strip()
            if not number:
                QMessageBox.information(self, APP_NAME, "This body shop does not have a phone number yet.")
                return
            tel_value = "".join(ch for ch in number if ch.isalnum() or ch == "+")
            if not tel_value:
                QMessageBox.information(self, APP_NAME, "This number cannot be used for a call action.")
                return
            QDesktopServices.openUrl(QUrl(f"tel:{tel_value}"))
            return
        self._edit_body_shop_entry()

    def _show_body_shop_context_menu(self, pos) -> None:
        item = self.body_shops_table.itemAt(pos)
        if item is None:
            return

        self.body_shops_table.selectRow(item.row())
        cell_text = (item.text() or "").strip()
        header_item = self.body_shops_table.horizontalHeaderItem(item.column())
        header_text = header_item.text() if header_item else "Value"

        menu = QMenu(self)
        copy_value_action = menu.addAction(f"Copy {header_text}")
        edit_action = menu.addAction("Edit Body Shop")

        chosen = menu.exec(self.body_shops_table.viewport().mapToGlobal(pos))
        if chosen == copy_value_action:
            QApplication.clipboard().setText(cell_text)
        elif chosen == edit_action:
            self._edit_body_shop_entry()

    def _insurance_company_dialog_fields(self) -> list[tuple[str, str, bool]]:
        return [
            ("company_name", "Company Name", False),
            ("quick_summary", "Quick Summary", True),
            ("fatal_errors", "Fatal Errors", True),
            ("photo_rules", "Photo Rules", True),
            ("estimate_supp_rules", "Estimate / Supplement Rules", True),
            ("parts_rules", "Parts Rules", True),
            ("total_loss_rules", "Total Loss Rules", True),
            ("tow_rules", "Tow Rules", True),
            ("supplement_rules", "Supplement Rules", True),
            ("betterment_depreciation_rules", "Betterment / Depreciation Rules", True),
            ("documentation_requirements", "Documentation Requirements", True),
            ("rates_and_sales_tax_rules", "Rates & Sales Tax Rules", True),
            ("miscellaneous_rules", "Miscellaneous Rules", True),
            ("contact_information", "Contact Information", True),
            ("labor_rates", "Labor Rates", True),
            ("total_loss_threshold", "Total Loss Threshold", True),
            ("notes", "Notes", True),
        ]

    def _insurance_company_dialog_values(self, entry: InsuranceCompanyEntry) -> dict[str, str]:
        return {
            "company_name": entry.company_name,
            "quick_summary": entry.quick_summary,
            "fatal_errors": entry.fatal_errors,
            "photo_rules": entry.photo_rules,
            "estimate_supp_rules": entry.estimate_supp_rules,
            "parts_rules": entry.parts_rules,
            "total_loss_rules": entry.total_loss_rules,
            "tow_rules": entry.tow_rules,
            "supplement_rules": entry.supplement_rules,
            "betterment_depreciation_rules": entry.betterment_depreciation_rules,
            "documentation_requirements": entry.documentation_requirements,
            "rates_and_sales_tax_rules": entry.rates_and_sales_tax_rules,
            "miscellaneous_rules": entry.miscellaneous_rules,
            "contact_information": entry.contact_information,
            "labor_rates": entry.labor_rates,
            "total_loss_threshold": entry.total_loss_threshold,
            "notes": entry.notes,
        }

    def _insurance_company_entry_from_values(self, values: dict[str, str]) -> InsuranceCompanyEntry:
        return InsuranceCompanyEntry(
            company_name=values["company_name"],
            quick_summary=values["quick_summary"],
            fatal_errors=values["fatal_errors"],
            photo_rules=values["photo_rules"],
            estimate_supp_rules=values["estimate_supp_rules"],
            parts_rules=values["parts_rules"],
            total_loss_rules=values["total_loss_rules"],
            tow_rules=values["tow_rules"],
            supplement_rules=values["supplement_rules"],
            betterment_depreciation_rules=values["betterment_depreciation_rules"],
            documentation_requirements=values["documentation_requirements"],
            rates_and_sales_tax_rules=values["rates_and_sales_tax_rules"],
            miscellaneous_rules=values["miscellaneous_rules"],
            contact_information=values["contact_information"],
            labor_rates=values["labor_rates"],
            total_loss_threshold=values["total_loss_threshold"],
            notes=values["notes"],
        )

    def _add_insurance_company_entry(self) -> None:
        dialog = RecordEditorDialog(
            self,
            "Add Insurance Company",
            self._insurance_company_dialog_fields(),
        )
        if dialog.exec() != QDialog.Accepted:
            return
        self.insurance_company_entries.append(self._insurance_company_entry_from_values(dialog.values()))
        self.repo.save_insurance_company_database(self.insurance_company_entries)
        self._populate_insurance_companies()

    def _edit_insurance_company_entry(self) -> None:
        index = self._selected_insurance_company_index()
        if index is None:
            QMessageBox.information(self, APP_NAME, "Choose an insurance company first.")
            return
        entry = self.insurance_company_entries[index]
        dialog = RecordEditorDialog(
            self,
            "Edit Insurance Company",
            self._insurance_company_dialog_fields(),
            self._insurance_company_dialog_values(entry),
        )
        if dialog.exec() != QDialog.Accepted:
            return
        self.insurance_company_entries[index] = self._insurance_company_entry_from_values(dialog.values())
        self.repo.save_insurance_company_database(self.insurance_company_entries)
        self._populate_insurance_companies()

    def _remove_insurance_company_entry(self) -> None:
        index = self._selected_insurance_company_index()
        if index is None:
            QMessageBox.information(self, APP_NAME, "Choose an insurance company first.")
            return
        del self.insurance_company_entries[index]
        self.repo.save_insurance_company_database(self.insurance_company_entries)
        self._populate_insurance_companies()

    def _handle_insurance_company_double_click(self, row: int, _column: int) -> None:
        if row < 0:
            return
        self.insurance_companies_table.selectRow(row)
        self._edit_insurance_company_entry()

    def _show_insurance_company_context_menu(self, pos) -> None:
        item = self.insurance_companies_table.itemAt(pos)
        if item is None:
            return

        self.insurance_companies_table.selectRow(item.row())
        cell_text = (item.text() or "").strip()
        header_item = self.insurance_companies_table.horizontalHeaderItem(item.column())
        header_text = header_item.text() if header_item else "Value"

        menu = QMenu(self)
        copy_value_action = menu.addAction(f"Copy {header_text}")
        edit_action = menu.addAction("Edit Insurance Company")
        chosen = menu.exec(self.insurance_companies_table.viewport().mapToGlobal(pos))
        if chosen == copy_value_action:
            QApplication.clipboard().setText(cell_text)
        elif chosen == edit_action:
            self._edit_insurance_company_entry()

    def _build_office_summary(self, claim: ClaimView) -> str:
        lines = [
            f"{claim.claim_id or '-'} - {claim.display_customer or '-'}",
            f"Insurance: {claim.insurance_company or '-'}",
            f"Progress: {claim.office_progress_status or 'No Contact Yet'}",
            f"Inspection When: {claim.office_appt_when or '-'}",
            f"Claim #: {claim.claim_number or '-'}",
            f"Shop: {claim.shop_name or '-'}",
            f"Town: {claim.town or '-'}",
            "",
            f"Additional Notes: {claim.office_additional_notes or '-'}",
        ]
        return "\n".join(lines)

    def _get_open_claims_for_office_update(self) -> list[ClaimView]:
        return sorted(
            self.filtered_open,
            key=lambda claim: self._office_sort_value(claim, self.office_sort_column),
            reverse=self.office_sort_order == Qt.DescendingOrder,
        )

    def _get_uninspected_claims_for_self_copy(self) -> list[ClaimView]:
        records: list[ClaimView] = []
        for claim in self._get_open_claims_for_office_update():
            progress = (claim.office_progress_status or "").strip().lower()
            waiting = (claim.office_waiting_for_paperwork or "").strip().lower()
            if waiting == "yes":
                continue
            if progress in {"seen - need to write", "written - under review", "supplement - waiting for paperwork", "waiting for paperwork"}:
                continue
            records.append(claim)
        return records

    def _customer_last_name(self, claim: ClaimView) -> str:
        customer = (claim.customer_name or claim.title or "").strip()
        if not customer:
            return ""
        if "," in customer:
            return re.sub(r"[^A-Za-z0-9-]", "", customer.split(",", 1)[0].strip())
        tokens = [token for token in re.split(r"\s+", customer) if token]
        if not tokens:
            return ""
        return re.sub(r"[^A-Za-z0-9-]", "", tokens[-1])

    def _find_assignment_sheet_path(self, claim: ClaimView) -> Path | None:
        candidates: list[Path] = []
        if claim.assignment_path:
            candidates.append(Path(claim.assignment_path))

        if claim.source_path:
            folder = Path(claim.source_path)
            if folder.exists() and folder.is_dir():
                candidates.append(folder / "assign.pdf")
                for child in folder.iterdir():
                    if not child.is_file() or child.suffix.lower() != ".pdf":
                        continue
                    lower_name = child.name.lower()
                    if "assign" in lower_name or re.fullmatch(r"\d+\.pdf", lower_name):
                        candidates.append(child)

        for candidate in candidates:
            try:
                if not candidate.exists() or not candidate.is_file() or candidate.suffix.lower() != ".pdf":
                    continue
                lower_name = candidate.name.lower()
                if "assign" in lower_name or re.fullmatch(r"\d+\.pdf", lower_name):
                    return candidate
            except OSError:
                continue
        return None

    def _generate_office_update_pdf(self, open_claims: list[ClaimView]) -> Path:
        if Document is None:
            raise RuntimeError("python-docx is not available.")

        output_dir = APP_DIR / "Office Updates"
        output_dir.mkdir(parents=True, exist_ok=True)
        base_name = f"Office Update {datetime.now().strftime('%Y-%m-%d %H-%M')}"
        docx_path = self._unique_destination(output_dir / f"{base_name}.docx")
        pdf_path = docx_path.with_suffix(".pdf")

        document = Document()
        document.add_heading("Office Update", level=0)
        document.add_paragraph(f"Generated: {datetime.now().strftime('%m/%d/%Y %I:%M %p')}")

        if WD_ORIENT is not None:
            section = document.sections[0]
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
            if Inches is not None:
                section.left_margin = Inches(0.4)
                section.right_margin = Inches(0.4)
                section.top_margin = Inches(0.4)
                section.bottom_margin = Inches(0.4)

        columns = ["Job #", "Customer", "Insurance", "Inspection", "Shop", "Progress", "Additional Notes"]
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for idx, label in enumerate(columns):
            table.rows[0].cells[idx].text = label

        for claim in open_claims:
            row = table.add_row().cells
            row[0].text = claim.claim_id or "-"
            row[1].text = claim.display_customer or "-"
            row[2].text = claim.insurance_company or "-"
            row[3].text = claim.office_appt_when or "-"
            row[4].text = claim.shop_name or "-"
            row[5].text = claim.office_progress_status or "-"
            row[6].text = claim.office_additional_notes or ""

        document.save(docx_path)
        self._export_doc_to_pdf(docx_path, pdf_path)
        return pdf_path

    def _generate_detailed_office_update_pdf(
        self,
        open_claims: list[ClaimView],
        title: str = "Office Update - Detailed",
        base_name_prefix: str = "Office Update Detailed",
    ) -> Path:
        if Document is None:
            raise RuntimeError("python-docx is not available.")

        output_dir = APP_DIR / "Office Updates"
        output_dir.mkdir(parents=True, exist_ok=True)
        base_name = f"{base_name_prefix} {datetime.now().strftime('%Y-%m-%d %H-%M')}"
        docx_path = self._unique_destination(output_dir / f"{base_name}.docx")
        pdf_path = docx_path.with_suffix(".pdf")

        document = Document()
        document.add_heading(title, level=0)
        document.add_paragraph(f"Generated: {datetime.now().strftime('%m/%d/%Y %I:%M %p')}")

        if WD_ORIENT is not None:
            section = document.sections[0]
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
            if Inches is not None:
                section.left_margin = Inches(0.3)
                section.right_margin = Inches(0.3)
                section.top_margin = Inches(0.3)
                section.bottom_margin = Inches(0.3)

        columns = [
            "Job #",
            "Claim #",
            "Customer",
            "Vehicle",
            "Insurance",
            "Address",
            "Phone",
            "Progress",
            "Appraisal Notes",
            "Claim Notes",
            "Additional Notes",
        ]
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        for idx, label in enumerate(columns):
            table.rows[0].cells[idx].text = label

        for claim in open_claims:
            row = table.add_row().cells
            row[0].text = claim.claim_id or "-"
            row[1].text = claim.claim_number or "-"
            row[2].text = claim.display_customer or "-"
            row[3].text = claim.vehicle or "-"
            row[4].text = claim.insurance_company or "-"
            row[5].text = self._route_default_address(claim) or "-"
            row[6].text = claim.contact_phone or "-"
            row[7].text = claim.office_progress_status or "-"
            row[8].text = claim.assignment_claim_notes or "-"
            row[9].text = claim.notes or "-"
            row[10].text = claim.office_additional_notes or ""

        document.save(docx_path)
        self._export_doc_to_pdf(docx_path, pdf_path)
        return pdf_path

    def _send_office_update_email(self, attachment_path: Path) -> None:
        message = EmailMessage()
        message["From"] = OFFICE_UPDATE_EMAIL_FROM
        message["To"] = self._office_update_email_to()
        cc_value = self._office_update_email_cc()
        if cc_value:
            message["Cc"] = cc_value
        message["Subject"] = OFFICE_UPDATE_EMAIL_SUBJECT
        message.set_content(OFFICE_UPDATE_EMAIL_BODY)
        message.add_attachment(
            attachment_path.read_bytes(),
            maintype="application",
            subtype="pdf",
            filename=attachment_path.name,
        )
        self._send_email_message(message)

    def _send_office_update_email_to_self(
        self,
        attachment_path: Path,
        subject: str | None = None,
        extra_attachments: list[tuple[Path, str | None]] | None = None,
    ) -> None:
        sender_email, _ = self._get_email_credentials()
        message = EmailMessage()
        message["From"] = sender_email
        message["To"] = sender_email
        message["Subject"] = subject or f"{OFFICE_UPDATE_EMAIL_SUBJECT} - Copy"
        message.set_content(OFFICE_UPDATE_EMAIL_BODY)
        message.add_attachment(
            attachment_path.read_bytes(),
            maintype="application",
            subtype="pdf",
            filename=attachment_path.name,
        )
        for extra_path, attachment_name in extra_attachments or []:
            try:
                suffix = extra_path.suffix.lower().lstrip(".") or "pdf"
                message.add_attachment(
                    extra_path.read_bytes(),
                    maintype="application",
                    subtype=suffix,
                    filename=attachment_name or extra_path.name,
                )
            except OSError:
                continue
        self._send_email_message(message)

    def _preview_office_update_pdf(self) -> None:
        open_claims = self._get_open_claims_for_office_update()
        if not open_claims:
            QMessageBox.information(self, APP_NAME, "There is no office summary to export.")
            return
        try:
            pdf_path = self._generate_office_update_pdf(open_claims)
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(pdf_path)))
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not create office update PDF:\n{exc}")

    def _email_office_update_pdf(self) -> None:
        open_claims = self._get_open_claims_for_office_update()
        if not open_claims:
            QMessageBox.information(self, APP_NAME, "There is no office summary to export.")
            return
        try:
            pdf_path = self._generate_office_update_pdf(open_claims)
            self._send_office_update_email(pdf_path)
            QMessageBox.information(self, APP_NAME, f"Exported and emailed:\n{pdf_path}")
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not create office update PDF:\n{exc}")

    def _email_uninspected_only_to_self(self) -> None:
        records = self._get_uninspected_claims_for_self_copy()
        if not records:
            QMessageBox.information(self, APP_NAME, "There are no uninspected open claims to email.")
            return
        try:
            pdf_path = self._generate_detailed_office_update_pdf(
                records,
                title="Uninspected Open Claims",
                base_name_prefix="Uninspected Open Claims",
            )
            extra_attachments: list[tuple[Path, str | None]] = []
            for claim in records:
                assign_path = self._find_assignment_sheet_path(claim)
                if not assign_path:
                    continue
                suffix = assign_path.suffix.lower() or ".pdf"
                last_name = self._customer_last_name(claim)
                attachment_name = f"{last_name}-assign{suffix}" if last_name else assign_path.name
                extra_attachments.append((assign_path, attachment_name))
            self._send_office_update_email_to_self(
                pdf_path,
                subject="Uninspected Open Claims",
                extra_attachments=extra_attachments,
            )
            QMessageBox.information(self, APP_NAME, f"Uninspected-only copy emailed to yourself:\n{pdf_path}")
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not create office update PDF:\n{exc}")

    def _set_combo_value(self, combo: QComboBox, value: str) -> None:
        index = combo.findText(value)
        if index >= 0:
            combo.setCurrentIndex(index)
        elif combo.count():
            combo.setCurrentIndex(0)

    def _save_claim_details(self) -> None:
        if not self.current_claim:
            return
        updates = {key: widget.text().strip() for key, widget in self.detail_inputs.items()}
        self.repo.update_claim_fields(self.current_claim.key, updates)
        self._load_claims()
        self._reselect_current_claim()

    def _save_office_update(self) -> None:
        if not self.current_claim:
            return
        previous_appt_when = (self.current_claim.office_appt_when or "").strip()
        progress_value = self.office_progress_input.currentText().strip()
        updates = {
            "office_progress_status": progress_value,
            "office_appt_when": self.office_appt_input.text().strip(),
            "office_waiting_for_paperwork": "Yes" if "waiting for paperwork" in progress_value.lower() else "No",
            "office_additional_notes": self.office_notes_input.toPlainText().strip(),
        }
        self.repo.update_office_fields(self.current_claim.key, updates)
        self._load_claims()
        self._reselect_current_claim()
        refreshed_claim = self.current_claim
        if (
            refreshed_claim
            and refreshed_claim.office_appt_when.strip()
            and refreshed_claim.office_appt_when.strip() != previous_appt_when
        ):
            try:
                self._send_office_appointment_calendar_invite(refreshed_claim)
            except Exception as exc:
                QMessageBox.warning(
                    self,
                    APP_NAME,
                    f"Saved the office update, but could not send the calendar reminder:\n{exc}",
                )

    def _save_assignment_notes(self) -> None:
        if not self.current_claim:
            return
        self.repo.update_claim_fields(
            self.current_claim.key,
            {"assignment_claim_notes": self.notes_view.toPlainText().strip()},
        )
        self._load_claims()
        self._reselect_current_claim()

    def _open_claim_notes_dialog(self) -> None:
        if not self.current_claim:
            return
        dialog = ClaimNotesDialog(self, self.current_claim)
        if dialog.exec() != QDialog.Accepted:
            return
        note_text = dialog.note_text()
        if not note_text:
            QMessageBox.information(self, APP_NAME, "Enter a note before saving.")
            return
        self.repo.append_claim_note(self.current_claim.key, note_text)
        self._load_claims()
        self._reselect_current_claim()

    def _set_current_claim_status(self, status: str) -> None:
        if not self.current_claim:
            return
        try:
            new_key = self.repo.move_claim_for_status(self.current_claim.key, status)
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not update claim status:\n{exc}")
            return
        self._load_claims()
        self._reselect_current_claim(new_key)

    def _mark_current_claim_closed(self) -> None:
        self._set_current_claim_status("Closed")

    def _reopen_current_claim(self) -> None:
        self._set_current_claim_status("Open")

    def _parse_office_appointment_datetime(self, value: str) -> datetime | None:
        normalized = " ".join((value or "").split()).strip()
        if not normalized:
            return None
        for fmt in (
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d %I:%M %p",
            "%m/%d/%y %H:%M",
            "%m/%d/%Y %H:%M",
            "%m/%d/%y %I:%M %p",
            "%m/%d/%Y %I:%M %p",
            "%m-%d-%y %H:%M",
            "%m-%d-%Y %H:%M",
            "%m-%d-%y %I:%M %p",
            "%m-%d-%Y %I:%M %p",
        ):
            try:
                return datetime.strptime(normalized, fmt)
            except ValueError:
                continue
        date_only = self._parse_report_date(normalized)
        if date_only:
            return date_only.replace(hour=10, minute=0)
        return None

    def _build_office_appointment_ics(self, claim: ClaimView, start_at: datetime) -> str:
        end_at = start_at + timedelta(hours=1)
        uid = f"{claim.claim_id}-{start_at.strftime('%Y%m%dT%H%M%S')}@claimmanager3"
        summary = f"Inspection - {claim.claim_id} {claim.display_customer}"
        description_lines = [
            f"Claim ID: {claim.claim_id or '-'}",
            f"Customer: {claim.display_customer or '-'}",
            f"Insurance: {claim.insurance_company or '-'}",
            f"Shop: {claim.shop_name or '-'}",
            f"Notes: {claim.office_additional_notes or '-'}",
        ]
        location = claim.owner_address or claim.shop_name or claim.town or ""
        summary = summary.replace(",", "\\,").replace(";", "\\;")
        description = "\\n".join(line.replace(",", "\\,").replace(";", "\\;") for line in description_lines)
        location = location.replace(",", "\\,").replace(";", "\\;")
        return "\r\n".join(
            [
                "BEGIN:VCALENDAR",
                "VERSION:2.0",
                "PRODID:-//Claim Manager 3.0//Office Update//EN",
                "CALSCALE:GREGORIAN",
                "BEGIN:VEVENT",
                f"UID:{uid}",
                f"DTSTAMP:{datetime.utcnow().strftime('%Y%m%dT%H%M%SZ')}",
                f"DTSTART;TZID=America/New_York:{start_at.strftime('%Y%m%dT%H%M%S')}",
                f"DTEND;TZID=America/New_York:{end_at.strftime('%Y%m%dT%H%M%S')}",
                f"SUMMARY:{summary}",
                f"DESCRIPTION:{description}",
                f"LOCATION:{location}",
                "BEGIN:VALARM",
                "ACTION:DISPLAY",
                "DESCRIPTION:Inspection tomorrow night reminder",
                "TRIGGER:-PT15H",
                "END:VALARM",
                "BEGIN:VALARM",
                "ACTION:DISPLAY",
                "DESCRIPTION:Inspection 9 AM reminder",
                "TRIGGER:-PT1H",
                "END:VALARM",
                "END:VEVENT",
                "END:VCALENDAR",
                "",
            ]
        )

    def _send_office_appointment_calendar_invite(self, claim: ClaimView) -> None:
        sender_email, app_password = self._get_email_credentials()
        if not sender_email or not app_password:
            raise RuntimeError("A Gmail address and app password are required for calendar reminders.")
        start_at = self._parse_office_appointment_datetime(claim.office_appt_when)
        if not start_at:
            raise RuntimeError("Use a date like YYYY-MM-DD or MM/DD/YY for the inspection date.")

        ics_content = self._build_office_appointment_ics(claim, start_at)
        message = EmailMessage()
        message["From"] = sender_email
        message["To"] = sender_email
        message["Subject"] = f"Calendar Reminder - {claim.claim_id or 'Inspection'}"
        message.set_content("Calendar invite attached for your inspection reminder.")
        message.add_attachment(
            ics_content.encode("utf-8"),
            maintype="text",
            subtype="calendar",
            filename=f"{claim.claim_id or 'inspection'}-appointment.ics",
            params={"method": "REQUEST", "charset": "utf-8"},
        )
        self._send_email_message(message)

    def _get_email_credentials(self) -> tuple[str, str]:
        app_password = ""
        if SECRETS_FILE.exists():
            try:
                payload = json.loads(SECRETS_FILE.read_text(encoding="utf-8"))
                app_password = str(payload.get("office_update_app_password") or "").strip()
            except (OSError, json.JSONDecodeError):
                app_password = ""
        return OFFICE_UPDATE_EMAIL_FROM, app_password

    def _send_email_message(self, message: EmailMessage) -> None:
        sender_email, app_password = self._get_email_credentials()
        if not sender_email or not app_password:
            raise RuntimeError("A Gmail address and app password are required.")
        context = ssl.create_default_context()
        with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as smtp:
            smtp.ehlo()
            smtp.starttls(context=context)
            smtp.ehlo()
            smtp.login(sender_email, app_password)
            smtp.send_message(message)

    def _claim_folder_file_candidates(self, claim: ClaimView) -> list[Path]:
        if not claim.source_path:
            return []
        folder = Path(claim.source_path)
        if not folder.exists() or not folder.is_dir():
            return []
        ignored = {"desktop.ini", "thumbs.db"}
        candidates = [path for path in sorted(folder.iterdir()) if path.is_file() and path.name.lower() not in ignored]
        return candidates

    def _claim_folder_prelim_candidates(self, claim: ClaimView) -> list[Path]:
        candidates = [path for path in self._claim_folder_file_candidates(claim) if path.suffix.lower() == ".pdf"]
        prelims = [path for path in candidates if "prelim" in path.name.lower()]
        return prelims or candidates

    def _pick_attachments(self, title: str, paths: list[Path]) -> list[Path] | None:
        dialog = AttachmentPickerDialog(self, title, paths)
        if dialog.exec() != QDialog.Accepted:
            return None
        return dialog.selected_paths()

    def _update_current_claim_shop_info(self, shop_name: str, shop_email: str) -> None:
        if not self.current_claim:
            return
        updates = {
            "shop_name": (shop_name or self.current_claim.shop_name or "").strip(),
            "shop_email": (shop_email or self.current_claim.shop_email or "").strip(),
        }
        self.repo.update_claim_fields(self.current_claim.key, updates)
        self._load_claims()
        self._reselect_current_claim()

    def _meaningful_shop_name(self, value: str) -> str:
        cleaned = (value or "").strip()
        if cleaned.lower() in {"", "-", "n/a", "unknown", "no shop chosen"}:
            return ""
        return cleaned

    def _ensure_shop_name_for_supplement_request(self, claim: ClaimView) -> str | None:
        existing = self._meaningful_shop_name(claim.shop_name or "")
        if existing:
            return existing

        shop_name, ok = QInputDialog.getText(
            self,
            APP_NAME,
            "Enter the shop name for this supplement request:",
            text="",
        )
        if not ok:
            return None

        shop_name = self._meaningful_shop_name(shop_name)
        if not shop_name:
            QMessageBox.information(self, APP_NAME, "Enter a shop name before sending the supplement request.")
            return None

        self.repo.update_claim_fields(claim.key, {"shop_name": shop_name})
        self._load_claims()
        self._reselect_current_claim()
        refreshed_claim = self.current_claim if self.current_claim and self.current_claim.key == claim.key else claim
        return self._meaningful_shop_name(refreshed_claim.shop_name or "") or shop_name

    def _send_review_email(self) -> None:
        if not self.current_claim:
            return
        question_dialog = RecordEditorDialog(
            self,
            "Review Questions",
            [("questions", "Questions (one per line)", True)],
            {"questions": ""},
        )
        if question_dialog.exec() != QDialog.Accepted:
            return
        questions_text = question_dialog.values().get("questions", "")
        custom_questions: list[str] = []
        for line in questions_text.splitlines():
            cleaned = line.strip()
            if not cleaned:
                continue
            cleaned = re.sub(r"^\d+\.\s*", "", cleaned)
            cleaned = cleaned.lstrip("-").strip()
            if cleaned:
                custom_questions.append(cleaned)

        fields = [
            ("insurance_company", "Insurance Company", False),
            ("shop_name", "Shop Name", False),
        ]
        values = {
            "insurance_company": self.current_claim.insurance_company or "",
            "shop_name": self.current_claim.shop_name or "",
        }
        for index, question in enumerate(custom_questions, start=1):
            fields.append((f"custom_{index}", question, False))
            values[f"custom_{index}"] = ""

        detail_dialog = RecordEditorDialog(self, "Review Email", fields, values, stacked_labels=True)
        if detail_dialog.exec() != QDialog.Accepted:
            return
        answers = detail_dialog.values()

        attachments = self._claim_folder_file_candidates(self.current_claim)
        selected_attachments: list[Path] = []
        if attachments:
            picked = self._pick_attachments("Review Email Attachments", attachments)
            if picked is None:
                return
            selected_attachments = picked

        body_lines = [
            "Hi Joe,",
            "",
            "Here is the review information:",
            "",
            f"Claim ID: {self.current_claim.claim_id or '-'}",
            f"Customer: {self.current_claim.display_customer or '-'}",
            f"Claim #: {self.current_claim.claim_number or '-'}",
            f"Insurance Company: {(answers.get('insurance_company') or '').strip() or '-'}",
            f"Shop Name: {(answers.get('shop_name') or '').strip() or '-'}",
            "",
        ]
        if custom_questions:
            body_lines.append("Questions:")
            body_lines.append("")
            for index, question in enumerate(custom_questions, start=1):
                answer = (answers.get(f"custom_{index}") or "").strip() or "-"
                body_lines.append(f"{index}. {question}")
                body_lines.append(f"   {answer}")
                body_lines.append("")
        else:
            body_lines.extend(["No additional questions were entered.", ""])
        body_lines.extend(["Thank you,", OFFICE_UPDATE_EMAIL_FROM])

        message = EmailMessage()
        message["From"] = OFFICE_UPDATE_EMAIL_FROM
        message["To"] = self._office_review_email_to()
        subject_parts = [part for part in (self.current_claim.claim_id, self.current_claim.display_customer) if part]
        message["Subject"] = f"Review - {' - '.join(subject_parts)}" if subject_parts else "Review"
        message.set_content("\n".join(body_lines).strip())
        for attachment_path in selected_attachments:
            try:
                message.add_attachment(
                    attachment_path.read_bytes(),
                    maintype="application",
                    subtype="octet-stream",
                    filename=attachment_path.name,
                )
            except OSError:
                continue

        try:
            self._send_email_message(message)
            QMessageBox.information(self, APP_NAME, f"The review email was sent to {message['To']}.")
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not send the review email:\n{exc}")

    def _send_rmc_email(self) -> None:
        if not self.current_claim:
            return
        dialog = RecordEditorDialog(
            self,
            "RMC Email",
            [
                ("customer_name", "Name of Customer", False),
                ("vehicle", "Year Make and Model", False),
                ("paint_code", "Paint Code", False),
                ("paint_hours", "Paint Hrs", False),
            ],
            {
                "customer_name": self.current_claim.display_customer or "",
                "vehicle": self.current_claim.vehicle or "",
                "paint_code": "",
                "paint_hours": "",
            },
        )
        if dialog.exec() != QDialog.Accepted:
            return
        answers = dialog.values()

        message = EmailMessage()
        message["From"] = OFFICE_UPDATE_EMAIL_FROM
        message["To"] = self._office_rmc_email_to()
        message["Subject"] = "Need RMC please"
        message.set_content(
            "\n".join(
                [
                    "Hi Gina,",
                    "",
                    f"Name of Customer: {(answers.get('customer_name') or '').strip() or '-'}",
                    f"Year Make and Model: {(answers.get('vehicle') or '').strip() or '-'}",
                    f"Paint Code: {(answers.get('paint_code') or '').strip() or '-'}",
                    f"Paint Hrs: {(answers.get('paint_hours') or '').strip() or '-'}",
                ]
            )
        )
        try:
            self._send_email_message(message)
            QMessageBox.information(self, APP_NAME, f"The RMC email was sent to {message['To']}.")
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not send the RMC email:\n{exc}")

    def _request_supplement_email(self) -> None:
        if not self.current_claim:
            return
        shop_name = self._ensure_shop_name_for_supplement_request(self.current_claim)
        if not shop_name or not self.current_claim:
            return
        message = EmailMessage()
        message["From"] = OFFICE_UPDATE_EMAIL_FROM
        message["To"] = self._office_supplement_email_to()
        message["Subject"] = OFFICE_SUPPLEMENT_EMAIL_SUBJECT
        message.set_content(self._supplement_request_email_body(self.current_claim, shop_name=shop_name))
        try:
            self._send_email_message(message)
            QMessageBox.information(self, APP_NAME, f"The supplement request email was sent to {message['To']}.")
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not send the supplement request email:\n{exc}")

    def _supplement_request_email_body(
        self,
        claim: ClaimView,
        greeting: str = "Hi Lisa,",
        *,
        shop_name: str | None = None,
    ) -> str:
        resolved_shop_name = self._meaningful_shop_name(shop_name or claim.shop_name or "") or "-"
        return "\n".join(
            [
                greeting,
                "",
                f"Job #: {claim.claim_id or '-'}",
                f"Customer: {claim.display_customer or '-'}",
                f"Vehicle: {claim.vehicle or '-'}",
                f"Insurance Company: {claim.insurance_company or '-'}",
                f"Shop: {resolved_shop_name}",
                "",
                "Thank you!",
            ]
        )

    def _generate_supplement_request_preview_pdf(self, claim: ClaimView) -> Path:
        if reportlab_canvas is None:
            raise RuntimeError("PDF preview tools are not available.")

        target_folder = self._record_folder_path(claim)
        target_folder.mkdir(parents=True, exist_ok=True)
        claim_suffix = claim.claim_id or "claim"
        pdf_path = self._unique_destination(target_folder / f"{claim_suffix}-SUPPLEMENT-REQUEST-PREVIEW.pdf")

        pdf_canvas = reportlab_canvas.Canvas(str(pdf_path), pagesize=(612, 792))
        width, height = 612, 792
        margin_x = 54
        y = height - 64

        pdf_canvas.setTitle("Supplement Request Preview")
        pdf_canvas.setFillColorRGB(0.07, 0.14, 0.22)
        pdf_canvas.setFont("Helvetica-Bold", 18)
        pdf_canvas.drawString(margin_x, y, "Supplement Request Preview")
        y -= 24

        pdf_canvas.setFont("Helvetica", 10)
        pdf_canvas.drawString(margin_x, y, f"Generated: {datetime.now().strftime('%m/%d/%Y %I:%M %p')}")
        y -= 28
        pdf_canvas.drawString(margin_x, y, f"To: {self._office_supplement_email_to()}")
        y -= 18
        pdf_canvas.drawString(margin_x, y, f"Subject: {OFFICE_SUPPLEMENT_EMAIL_SUBJECT}")
        y -= 28

        pdf_canvas.setFont("Helvetica-Bold", 11)
        pdf_canvas.drawString(margin_x, y, "Email Body")
        y -= 22

        pdf_canvas.setFont("Helvetica", 11)
        max_width_chars = 72
        for paragraph in self._supplement_request_email_body(claim).splitlines():
            wrapped_lines = textwrap.wrap(paragraph, width=max_width_chars) or [""]
            for line in wrapped_lines:
                if y < 72:
                    pdf_canvas.showPage()
                    y = height - 64
                    pdf_canvas.setFillColorRGB(0.07, 0.14, 0.22)
                    pdf_canvas.setFont("Helvetica", 11)
                pdf_canvas.drawString(margin_x, y, line)
                y -= 16
            y -= 4

        pdf_canvas.save()
        return pdf_path

    def _send_prelim_email(self) -> None:
        if not self.current_claim:
            return
        pdf_candidates = self._claim_folder_prelim_candidates(self.current_claim)
        if not pdf_candidates:
            QMessageBox.information(self, APP_NAME, "Could not find any PDF files in this claim folder.")
            return
        picked = self._pick_attachments("Select Prelim PDF", pdf_candidates)
        if picked is None:
            return
        if not picked:
            QMessageBox.information(self, APP_NAME, "Select at least one prelim PDF first.")
            return

        leave_message = QMessageBox.question(
            self,
            APP_NAME,
            "Do you want to leave a message as well?",
            QMessageBox.Yes | QMessageBox.No,
            QMessageBox.No,
        ) == QMessageBox.Yes

        dialog = RecordEditorDialog(
            self,
            "Send Prelim",
            [
                ("shop_name", f"Shop Name ({self.current_claim.shop_name or 'current value'})", False),
                ("shop_email", f"Shop Email ({self.current_claim.shop_name or 'shop'})", False),
                ("message_text", "Message", True),
            ],
            {
                "shop_name": self.current_claim.shop_name or "",
                "shop_email": self.current_claim.shop_email or "",
                "message_text": "",
            },
        )
        if dialog.exec() != QDialog.Accepted:
            return
        answers = dialog.values()
        shop_name = (answers.get("shop_name") or "").strip() or self.current_claim.shop_name or ""
        shop_email = (answers.get("shop_email") or "").strip()
        message_text = (answers.get("message_text") or "").strip()
        if not shop_email:
            QMessageBox.information(self, APP_NAME, "Enter the shop email address first.")
            return

        self._update_current_claim_shop_info(shop_name, shop_email)

        body = f"{message_text}\n\nthank you!\nfernando" if leave_message and message_text else "thank you!\nfernando"
        message = EmailMessage()
        message["From"] = OFFICE_UPDATE_EMAIL_FROM
        message["To"] = shop_email
        message["Subject"] = "please review and advise"
        message.set_content(body)
        for attachment_path in picked:
            try:
                message.add_attachment(
                    attachment_path.read_bytes(),
                    maintype="application",
                    subtype="pdf",
                    filename=attachment_path.name,
                )
            except OSError:
                continue
        try:
            self._send_email_message(message)
            QMessageBox.information(self, APP_NAME, "The prelim email was sent to the shop.")
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not send the prelim email:\n{exc}")

    def _email_assignment_sheet(self) -> None:
        if not self.current_claim:
            return
        assignment_path = self._find_assignment_sheet_path(self.current_claim)
        if not assignment_path:
            QMessageBox.information(self, APP_NAME, "Could not find the original assignment sheet for this claim.")
            return
        sender_email, app_password = self._get_email_credentials()
        if not sender_email or not app_password:
            QMessageBox.warning(self, APP_NAME, "A Gmail address and app password are required to email the assignment sheet.")
            return

        assign_file = assignment_path
        subject_parts = [part for part in (self.current_claim.claim_id, self.current_claim.display_customer) if part]
        subject = f"Assignment Sheet - {' - '.join(subject_parts)}" if subject_parts else "Assignment Sheet"
        message = EmailMessage()
        message["From"] = sender_email
        message["To"] = sender_email
        message["Subject"] = subject
        message.set_content("Attached is the original assignment sheet.")
        try:
            message.add_attachment(
                assign_file.read_bytes(),
                maintype="application",
                subtype="pdf",
                filename=assign_file.name,
            )
            self._send_email_message(message)
            QMessageBox.information(self, APP_NAME, f"Emailed to {sender_email}.")
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not email the assignment sheet:\n{exc}")

    def _open_outlook_draft(self, to_address: str, subject: str, body: str, attachments: list[Path]) -> None:
        escaped_to = to_address.replace("'", "''")
        escaped_subject = subject.replace("'", "''")
        escaped_body = body.replace("'", "''")
        escaped_attachments = ["'" + str(path).replace("'", "''") + "'" for path in attachments]
        attachment_values = ", ".join(escaped_attachments)
        script = f"""
$outlook = New-Object -ComObject Outlook.Application
$mail = $outlook.CreateItem(0)
$mail.To = '{escaped_to}'
$mail.Subject = '{escaped_subject}'
$mail.Body = '{escaped_body}'
foreach ($attachment in @({attachment_values})) {{
    if ($attachment) {{
        $mail.Attachments.Add($attachment) | Out-Null
    }}
}}
$mail.Display()
"""
        subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
            check=True,
            capture_output=True,
            text=True,
            timeout=30,
        )

    def _open_shop_email_draft(self) -> None:
        if not self.current_claim:
            return
        shop_email, ok = QInputDialog.getText(
            self,
            APP_NAME,
            "Enter the shop email address for this draft:",
            text=self.current_claim.shop_email or "",
        )
        if not ok:
            return
        shop_email = shop_email.strip()
        if not shop_email:
            QMessageBox.information(self, APP_NAME, "Enter a shop email address first.")
            return

        attachments = self._claim_folder_file_candidates(self.current_claim)
        selected_attachments: list[Path] = []
        if attachments:
            picked = self._pick_attachments("Shop Email Attachments", attachments)
            if picked is None:
                return
            selected_attachments = picked

        subject_parts = [part for part in (self.current_claim.claim_id, self.current_claim.display_customer) if part]
        subject = f"Claim Copy - {' - '.join(subject_parts)}" if subject_parts else "Claim Copy"
        body_lines = [
            "Hi,",
            "",
            "Attached are the selected claim files.",
            "",
            f"Claim ID: {self.current_claim.claim_id or '-'}",
            f"Customer: {self.current_claim.display_customer or '-'}",
            f"Claim #: {self.current_claim.claim_number or '-'}",
            f"Insurance Company: {self.current_claim.insurance_company or '-'}",
            f"Shop Name: {self.current_claim.shop_name or '-'}",
            "",
            "Thank you,",
            OFFICE_UPDATE_EMAIL_FROM,
        ]
        try:
            self._open_outlook_draft(shop_email, subject, "\n".join(body_lines).strip(), selected_attachments)
            QMessageBox.information(self, APP_NAME, "A new email draft was opened for the shop.")
        except Exception as exc:
            QMessageBox.warning(
                self,
                APP_NAME,
                "Could not open the shop email draft.\n\nThis feature depends on Outlook desktop being available.\n\n"
                f"{exc}",
            )

    def _open_nada_value(self) -> None:
        webbrowser.open("https://www.jdpower.com/cars/manufacturers")

    def _open_claim_tool_path(self, file_path: Path) -> None:
        if not file_path.exists():
            QMessageBox.warning(self, APP_NAME, f"Could not find:\n{file_path}")
            return
        QDesktopServices.openUrl(QUrl.fromLocalFile(str(file_path)))

    def _record_folder_path(self, claim: ClaimView) -> Path:
        if claim.source_path:
            source = Path(claim.source_path)
            if source.exists() and source.is_dir():
                return source
        fallback = APP_DIR / "working_sheets"
        fallback.mkdir(parents=True, exist_ok=True)
        return fallback

    def _unique_destination(self, destination: Path) -> Path:
        if not destination.exists():
            return destination
        counter = 2
        while True:
            candidate = destination.with_name(f"{destination.stem} ({counter}){destination.suffix}")
            if not candidate.exists():
                return candidate
            counter += 1

    def _apply_choice_to_checkbox_indices(
        self,
        checkbox_values: list[bool],
        indices: list[int],
        labels: list[str],
        value: str,
    ) -> None:
        normalized = (value or "").strip().lower()
        selected_idx = None
        for idx, label in enumerate(labels):
            if normalized == label.lower() or normalized in label.lower():
                selected_idx = idx
                break
        if selected_idx is None:
            return
        for idx, field_number in enumerate(indices):
            checkbox_values[field_number - 1] = idx == selected_idx

    def _fill_mitchell_total_loss_template(
        self,
        output_path: Path,
        top_answers: dict[str, str],
        condition_answers: list[tuple[str, str, str]],
        aftermarket: str,
        final_comments: str,
    ) -> None:
        template_root = Path(str(self.settings.get("claim_tools_folder", "") or ""))
        template_path = template_root / "MITCHELL TOTAL LOSS.doc"
        if not template_path.exists():
            raise RuntimeError("Mitchell total loss template is missing.")

        text_field_map = {
            "Text1": top_answers.get("Claim-Suffix ID", ""),
            "Text2": top_answers.get("Claimant Name", ""),
            "Text3": top_answers.get("Claimant Phone", ""),
            "Text4": top_answers.get("Loss Date", ""),
            "Text5": top_answers.get("License Plate", ""),
            "Text6": top_answers.get("Insured Name", ""),
            "Text7": top_answers.get("Insured Phone", ""),
            "Text8": top_answers.get("Loss Type", ""),
            "Text9": top_answers.get("VIN", ""),
            "Text10": top_answers.get("Year", ""),
            "Text11": top_answers.get("Make", ""),
            "Text12": top_answers.get("Model", ""),
            "Text13": top_answers.get("Sub-model", ""),
            "Text14": top_answers.get("Mileage", ""),
            "Text15": top_answers.get("Body Style", ""),
            "Text16": top_answers.get("Ext. Color", ""),
            "Text17": top_answers.get("Engine", ""),
            "Text18": top_answers.get("Location of Vehicle", ""),
            "Text19": top_answers.get("Zip Code", ""),
            "Text20": top_answers.get("Inspected By", ""),
            "Text21": top_answers.get("Date", ""),
            "Text22": condition_answers[0][2] if len(condition_answers) > 0 else "",
            "Text23": condition_answers[1][2] if len(condition_answers) > 1 else "",
            "Text24": condition_answers[2][2] if len(condition_answers) > 2 else "",
            "Text25": condition_answers[3][2] if len(condition_answers) > 3 else "",
            "Text26": condition_answers[4][2] if len(condition_answers) > 4 else "",
            "Text27": condition_answers[5][2] if len(condition_answers) > 5 else "",
            "Text28": condition_answers[6][2] if len(condition_answers) > 6 else "",
            "Text29": condition_answers[7][2] if len(condition_answers) > 7 else "",
            "Text30": condition_answers[8][2] if len(condition_answers) > 8 else "",
            "Text31": condition_answers[9][2] if len(condition_answers) > 9 else "",
            "Text32": condition_answers[10][2] if len(condition_answers) > 10 else "",
            "Text33": condition_answers[11][2] if len(condition_answers) > 11 else "",
            "Text34": condition_answers[12][2] if len(condition_answers) > 12 else "",
            "Text35": aftermarket,
            "Text36": final_comments,
        }

        checkbox_values: list[bool] = [False] * 216
        for field_number in [31, 58, 63, 66, 79, 80, 105, 108, 109]:
            checkbox_values[field_number - 1] = True

        self._apply_choice_to_checkbox_indices(
            checkbox_values,
            [18, 19, 20, 21],
            ["Adaptive", "Automatic", "Interactive", "Manual"],
            top_answers.get("Transmission", ""),
        )
        self._apply_choice_to_checkbox_indices(
            checkbox_values,
            [22, 23, 24],
            ["2WD", "4WD", "AWD"],
            top_answers.get("Drive Train", ""),
        )

        rating_groups = [
            [124, 125, 126, 127, 128, 129],
            [131, 132, 133, 134, 135, 136],
            [138, 139, 140, 141, 142, 143],
            [145, 146, 147, 148, 149, 150],
            [152, 153, 154, 155, 156, 157],
            [159, 160, 161, 162, 163, 164],
            [166, 167, 168, 169, 170, 171],
            [173, 174, 175, 176, 177, 178],
            [180, 181, 182, 183, 184, 185],
            [187, 188, 189, 190, 191, 192],
            [194, 195, 196, 197, 198, 199],
            [201, 202, 203, 204, 205, 206],
            [208, 209, 210, 211, 212, 213],
        ]
        rating_labels = ["5 - Excellent", "4 - Very good", "3 - Good", "2 - Fair", "1 - Poor", "U - Unknown"]
        for idx, (_, rating, _) in enumerate(condition_answers):
            if idx < len(rating_groups):
                self._apply_choice_to_checkbox_indices(checkbox_values, rating_groups[idx], rating_labels, rating)

        payload = {
            "template_path": str(template_path),
            "output_path": str(output_path),
            "text_fields": text_field_map,
            "checkbox_values": checkbox_values,
        }
        payload_path = APP_DIR / "mitchell_fill_payload_3.json"
        payload_path.write_text(json.dumps(payload), encoding="utf-8")

        script = r"""
$ErrorActionPreference = 'Stop'
$ProgressPreference = 'SilentlyContinue'
$word = $null
$doc = $null
function Set-MitchellFieldResult($field, [string]$value) {
    if ($null -eq $value) {
        $candidate = ""
    }
    else {
        $candidate = [string]$value
    }
    try {
        $field.Result = $candidate
        return
    }
    catch {
        if ($_.Exception.Message -notlike '*String too long*') {
            throw
        }
    }

    $low = 0
    $high = [Math]::Min(255, $candidate.Length)
    while ($low -lt $high) {
        $mid = [int][Math]::Floor(($low + $high + 1) / 2)
        if ($mid -lt $candidate.Length -and $mid -gt 3) {
            $trial = $candidate.Substring(0, $mid - 3) + '...'
        }
        else {
            $trial = $candidate.Substring(0, $mid)
        }
        try {
            $field.Result = $trial
            $low = $mid
        }
        catch {
            if ($_.Exception.Message -like '*String too long*') {
                $high = $mid - 1
            }
            else {
                throw
            }
        }
    }

    if ($low -lt $candidate.Length -and $low -gt 3) {
        $field.Result = $candidate.Substring(0, $low - 3) + '...'
    }
    else {
        $field.Result = $candidate.Substring(0, $low)
    }
}
$payload = Get-Content -Raw '__PAYLOAD__' | ConvertFrom-Json
try {
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $word.ScreenUpdating = $false
    $word.Options.SaveNormalPrompt = $false
    $word.Options.ConfirmConversions = $false
    $word.Options.WarnBeforeSavingPrintingSendingMarkup = $false
    $word.AutomationSecurity = 3
    $doc = $word.Documents.Open($payload.template_path)
    foreach ($field in $doc.FormFields) {
        $name = [string]$field.Name
        if ($payload.text_fields.PSObject.Properties.Name -contains $name) {
            Set-MitchellFieldResult $field ([string]$payload.text_fields.$name)
        }
    }
    for ($i = 0; $i -lt $doc.FormFields.Count; $i++) {
        $field = $doc.FormFields.Item($i + 1)
        if ($field.Type -eq 71) {
            $field.CheckBox.Value = [bool]$payload.checkbox_values[$i]
        }
    }
    $doc.SaveAs2($payload.output_path, 0)
}
finally {
    if ($null -ne $doc) { $doc.Close($false) }
    if ($null -ne $word) { $word.Quit() }
}
"""
        script = script.replace("__PAYLOAD__", str(payload_path))
        self._run_word_automation_script(script, "filling the Mitchell template")

    def _run_word_automation_script(self, script: str, action: str, timeout: int = 90) -> None:
        encoded_script = base64.b64encode(script.encode("utf-16le")).decode("ascii")
        try:
            subprocess.run(
                ["powershell.exe", "-NoProfile", "-NonInteractive", "-Sta", "-EncodedCommand", encoded_script],
                check=True,
                capture_output=True,
                text=True,
                timeout=timeout,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError(f"Word timed out while {action}. Close any open Word windows and try again.") from exc
        except subprocess.CalledProcessError as exc:
            details = (exc.stderr or exc.stdout or "").strip().replace("\r", "")
            if "0x80070520" in details or "specified logon session does not exist" in details.lower():
                raise RuntimeError(
                    "Word automation could not start. Open Microsoft Word on the desktop, then try again.\n\n"
                    f"{details}"
                ) from exc
            if details:
                raise RuntimeError(f"Word failed while {action}:\n{details}") from exc
            raise RuntimeError(f"Word failed while {action} with exit code {exc.returncode}.") from exc

    def _export_doc_to_pdf(self, doc_path: Path, pdf_path: Path) -> None:
        doc_literal = str(doc_path).replace("'", "''")
        pdf_literal = str(pdf_path).replace("'", "''")
        script = rf"""
$ErrorActionPreference = 'Stop'
$ProgressPreference = 'SilentlyContinue'
$word = $null
$doc = $null
try {{
    $word = New-Object -ComObject Word.Application
    $word.Visible = $false
    $word.DisplayAlerts = 0
    $word.ScreenUpdating = $false
    $word.Options.SaveNormalPrompt = $false
    $word.Options.ConfirmConversions = $false
    $word.Options.WarnBeforeSavingPrintingSendingMarkup = $false
    $word.AutomationSecurity = 3
    $doc = $word.Documents.Open('{doc_literal}')
    $doc.ExportAsFixedFormat('{pdf_literal}', 17)
}}
finally {{
    if ($null -ne $doc) {{ $doc.Close($false) }}
    if ($null -ne $word) {{ $word.Quit() }}
}}
"""
        self._run_word_automation_script(script, "exporting the Mitchell PDF")

    def _load_existing_autosource_answers(self, claim: ClaimView) -> dict[str, str]:
        target_folder = self._record_folder_path(claim)
        candidates = sorted(target_folder.glob("*AUTOSOURCE-PDA-SUB-LEVEL.answers.json"), key=lambda path: path.stat().st_mtime)
        if not candidates:
            return {}
        try:
            return json.loads(candidates[-1].read_text(encoding="utf-8"))
        except Exception:
            return {}

    def _autosource_party_defaults(self, claim: ClaimView) -> tuple[str, str]:
        customer = (claim.display_customer or claim.customer_name or "").strip()
        return customer, customer

    def _autosource_company_branch_default(self, claim: ClaimView) -> str:
        pieces = [part.strip() for part in [claim.insurance_company, claim.town] if part and part.strip()]
        return " - ".join(pieces[:2])

    def _autosource_market_area_default(self, claim: ClaimView) -> str:
        town = (claim.town or "").strip()
        zip_code = self._zip_from_address(claim.owner_address or "")
        if town and zip_code:
            return f"{town} {zip_code}"
        return town or zip_code

    def _autosource_inspection_default(self, claim: ClaimView) -> str:
        if claim.location_of_vehicle:
            return claim.location_of_vehicle
        if claim.owner_address:
            return claim.owner_address
        return claim.town or ""

    def _autosource_drive_default(self, claim: ClaimView) -> str:
        vehicle_text = (claim.vehicle or "").upper()
        if any(token in vehicle_text for token in ("AWD", "4WD", "4X4")):
            return "4W"
        return ""

    def _autosource_answers_path(self, pdf_path: Path) -> Path:
        return pdf_path.with_suffix(".answers.json")

    def _autosource_px_rect(self, x1: int, y1: int, x2: int, y2: int) -> tuple[float, float, float, float]:
        scale_x = 612.0 / 1275.0
        scale_y = 792.0 / 1650.0
        return (
            x1 * scale_x,
            792.0 - (y2 * scale_y),
            (x2 - x1) * scale_x,
            (y2 - y1) * scale_y,
        )

    def _draw_text_in_pdf_rect(
        self,
        pdf_canvas,
        text: str,
        rect: tuple[float, float, float, float],
        *,
        font_name: str = "Helvetica",
        font_size: float = 8.0,
        min_font_size: float = 6.0,
        max_lines: int = 1,
        line_gap: float = 1.4,
        bottom_pad: float = 2.0,
    ) -> None:
        text = (text or "").strip()
        if not text:
            return
        x, y, width, height = rect
        available_width = max(8.0, width - 4.0)
        size = font_size
        words = text.split()
        lines: list[str] = []

        def wrap_for_size(current_size: float) -> list[str]:
            wrapped: list[str] = []
            current_line = ""
            for word in words if words else [text]:
                candidate = f"{current_line} {word}".strip()
                if pdf_canvas.stringWidth(candidate, font_name, current_size) <= available_width or not current_line:
                    current_line = candidate
                else:
                    wrapped.append(current_line)
                    current_line = word
            if current_line:
                wrapped.append(current_line)
            return wrapped or [text]

        lines = wrap_for_size(size)
        while (len(lines) > max_lines or any(pdf_canvas.stringWidth(line, font_name, size) > available_width for line in lines)) and size > min_font_size:
            size -= 0.4
            lines = wrap_for_size(size)

        if len(lines) > max_lines:
            lines = lines[:max_lines]
            last = lines[-1]
            while last and pdf_canvas.stringWidth(f"{last}...", font_name, size) > available_width:
                last = last[:-1]
            lines[-1] = f"{last}..." if last else "..."

        pdf_canvas.setFont(font_name, size)
        line_height = size + line_gap
        start_y = y + height - line_height
        minimum_y = y + bottom_pad
        if start_y < minimum_y:
            start_y = minimum_y
        for index, line in enumerate(lines):
            draw_y = start_y - (index * line_height)
            if draw_y < minimum_y - 1:
                break
            pdf_canvas.drawString(x + 2.0, draw_y, line)

    def _fill_autosource_pda_template(self, output_path: Path, answers: dict[str, str]) -> None:
        if PdfReader is None or PdfWriter is None or reportlab_canvas is None:
            raise RuntimeError("PDF form libraries are not available for the Autosource form.")
        if not AUTOSOURCE_PDA_TEMPLATE.exists():
            raise RuntimeError(f"Autosource template is missing:\n{AUTOSOURCE_PDA_TEMPLATE}")

        reader = PdfReader(str(AUTOSOURCE_PDA_TEMPLATE))
        writer = PdfWriter()
        first_page = reader.pages[0]
        page_width = float(first_page.mediabox.width)
        page_height = float(first_page.mediabox.height)

        overlay_buffer = BytesIO()
        pdf_canvas = reportlab_canvas.Canvas(overlay_buffer, pagesize=(page_width, page_height))
        pdf_canvas.setFillColorRGB(0.07, 0.14, 0.22)

        field_rects = {
            "product_type": self._autosource_px_rect(808, 57, 1210, 79),
            "email_address": self._autosource_px_rect(188, 81, 730, 102),
            "fax_number": self._autosource_px_rect(744, 81, 892, 102),
            "point_of_impact": self._autosource_px_rect(1040, 81, 1210, 102),
            "gross_estimate": self._autosource_px_rect(925, 102, 1210, 126),
            "audatex_id": self._autosource_px_rect(55, 103, 228, 127),
            "company_branch": self._autosource_px_rect(228, 103, 521, 127),
            "claim_rep_name": self._autosource_px_rect(521, 103, 798, 127),
            "claim_rep_phone": self._autosource_px_rect(798, 103, 1008, 127),
            "loss_date": self._autosource_px_rect(1008, 103, 1210, 127),
            "claim_number": self._autosource_px_rect(55, 128, 228, 152),
            "policy_number": self._autosource_px_rect(228, 128, 521, 152),
            "insured_name": self._autosource_px_rect(521, 128, 798, 152),
            "insured_phone": self._autosource_px_rect(798, 128, 1008, 152),
            "loss_type": self._autosource_px_rect(1008, 128, 1210, 152),
            "market_area": self._autosource_px_rect(55, 153, 228, 177),
            "inspection_city_state": self._autosource_px_rect(228, 153, 521, 177),
            "claimant_name": self._autosource_px_rect(521, 153, 798, 177),
            "claimant_phone": self._autosource_px_rect(798, 153, 1008, 177),
            "license_plate": self._autosource_px_rect(1008, 153, 1210, 177),
            "odometer": self._autosource_px_rect(55, 178, 210, 201),
            "ext_color": self._autosource_px_rect(210, 178, 375, 201),
            "vin": self._autosource_px_rect(375, 178, 795, 201),
            "year": self._autosource_px_rect(55, 202, 150, 225),
            "make": self._autosource_px_rect(150, 202, 285, 225),
            "model": self._autosource_px_rect(285, 202, 468, 225),
            "doors": self._autosource_px_rect(468, 202, 554, 225),
            "bodystyle": self._autosource_px_rect(554, 202, 688, 225),
            "drive": self._autosource_px_rect(882, 202, 958, 225),
            "edition": self._autosource_px_rect(958, 202, 1210, 225),
            "engine_type": self._autosource_px_rect(55, 226, 175, 250),
            "engine_other": self._autosource_px_rect(175, 226, 322, 250),
            "engine_size": self._autosource_px_rect(470, 226, 610, 250),
            "cylinders": self._autosource_px_rect(610, 226, 720, 250),
            "trans_type": self._autosource_px_rect(842, 226, 950, 250),
            "trans_gears": self._autosource_px_rect(950, 226, 1210, 250),
            "capacity_passenger": self._autosource_px_rect(55, 252, 168, 274),
            "bed_length": self._autosource_px_rect(168, 252, 338, 274),
            "capacity_tonnage": self._autosource_px_rect(338, 252, 545, 274),
            "van_type": self._autosource_px_rect(545, 252, 930, 274),
            "conversion_name": self._autosource_px_rect(930, 252, 1210, 274),
            "inspected_by": self._autosource_px_rect(55, 276, 235, 298),
            "inspected_date": self._autosource_px_rect(235, 276, 344, 298),
            "location_of_vehicle": self._autosource_px_rect(344, 276, 1022, 298),
            "pool_number": self._autosource_px_rect(1022, 276, 1210, 298),
            "equipment_notes": self._autosource_px_rect(184, 732, 1210, 758),
            "general_comments": self._autosource_px_rect(48, 1514, 1214, 1594),
        }

        for key, rect in field_rects.items():
            max_lines = 4 if key == "general_comments" else 2 if key in {"equipment_notes", "location_of_vehicle"} else 1
            font_size = 7.5 if key == "general_comments" else 8.0
            self._draw_text_in_pdf_rect(pdf_canvas, answers.get(key, ""), rect, font_size=font_size, max_lines=max_lines)

        pdf_canvas.save()
        overlay_reader = PdfReader(BytesIO(overlay_buffer.getvalue()))

        for index, page in enumerate(reader.pages):
            page_copy = page
            if index == 0:
                page_copy.merge_page(overlay_reader.pages[0])
            writer.add_page(page_copy)

        with output_path.open("wb") as handle:
            writer.write(handle)

    def _generate_autosource_pda_sub_level(self) -> None:
        if not self.current_claim:
            return
        if PdfReader is None or PdfWriter is None or reportlab_canvas is None:
            QMessageBox.warning(self, APP_NAME, "PDF overlay tools are not available for the Autosource form.")
            return

        existing = self._load_existing_autosource_answers(self.current_claim)
        claimant_default, insured_default = self._autosource_party_defaults(self.current_claim)
        today_default = datetime.now().strftime("%#m/%#d/%y") if os.name == "nt" else datetime.now().strftime("%-m/%-d/%y")

        claim_defaults = {
            "product_type": existing.get("product_type", ""),
            "email_address": existing.get("email_address", ""),
            "fax_number": existing.get("fax_number", ""),
            "point_of_impact": existing.get("point_of_impact", ""),
            "gross_estimate": existing.get("gross_estimate", ""),
            "audatex_id": existing.get("audatex_id", ""),
            "company_branch": existing.get("company_branch", self._autosource_company_branch_default(self.current_claim)),
            "claim_rep_name": existing.get("claim_rep_name", ""),
            "claim_rep_phone": existing.get("claim_rep_phone", ""),
            "loss_date": existing.get("loss_date", self.current_claim.date_of_loss or ""),
            "claim_number": existing.get("claim_number", self.current_claim.claim_number or self.current_claim.claim_id or ""),
            "policy_number": existing.get("policy_number", ""),
            "insured_name": existing.get("insured_name", insured_default),
            "insured_phone": existing.get("insured_phone", self.current_claim.contact_phone or ""),
            "loss_type": existing.get("loss_type", "Collision"),
            "market_area": existing.get("market_area", self._autosource_market_area_default(self.current_claim)),
            "inspection_city_state": existing.get("inspection_city_state", self._autosource_inspection_default(self.current_claim)),
            "claimant_name": existing.get("claimant_name", claimant_default),
            "claimant_phone": existing.get("claimant_phone", self.current_claim.contact_phone or ""),
            "license_plate": existing.get("license_plate", ""),
        }

        vehicle_defaults = {
            "odometer": existing.get("odometer", ""),
            "ext_color": existing.get("ext_color", ""),
            "vin": existing.get("vin", self.current_claim.vin or ""),
            "year": existing.get("year", self._mitchell_year_from_vehicle(self.current_claim.vehicle)),
            "make": existing.get("make", self._mitchell_make_from_vehicle(self.current_claim.vehicle)),
            "model": existing.get("model", self._mitchell_model_from_vehicle(self.current_claim.vehicle)),
            "doors": existing.get("doors", ""),
            "bodystyle": existing.get("bodystyle", ""),
            "drive": existing.get("drive", self._autosource_drive_default(self.current_claim)),
            "edition": existing.get("edition", ""),
            "engine_type": existing.get("engine_type", "Gas"),
            "engine_other": existing.get("engine_other", ""),
            "engine_size": existing.get("engine_size", ""),
            "cylinders": existing.get("cylinders", ""),
            "trans_type": existing.get("trans_type", "Auto"),
            "trans_gears": existing.get("trans_gears", ""),
            "capacity_passenger": existing.get("capacity_passenger", ""),
            "bed_length": existing.get("bed_length", ""),
            "capacity_tonnage": existing.get("capacity_tonnage", ""),
            "van_type": existing.get("van_type", ""),
            "conversion_name": existing.get("conversion_name", ""),
            "inspected_by": existing.get("inspected_by", "Fernando Marin"),
            "inspected_date": existing.get("inspected_date", today_default),
            "location_of_vehicle": existing.get("location_of_vehicle", self.current_claim.location_of_vehicle or self.current_claim.shop_name or self.current_claim.owner_address or ""),
            "pool_number": existing.get("pool_number", ""),
        }

        notes_defaults = {
            "equipment_notes": existing.get("equipment_notes", ""),
            "general_comments": existing.get("general_comments", self.current_claim.assignment_claim_notes or self.current_claim.office_additional_notes or ""),
        }

        claim_questions = [
            ("product_type", "Product Type", False, ["", "Valuation", "Market Search", "Total Loss"]),
            ("email_address", "Report Retrieval E-mail Address", False),
            ("fax_number", "FaxBack Fax #", False),
            ("point_of_impact", "Point of Impact", False),
            ("gross_estimate", "Gross Estimate Amount", False),
            ("audatex_id", "Audatex ID #", False),
            ("company_branch", "Co. Name & Branch", False),
            ("claim_rep_name", "Claim Rep Name", False),
            ("claim_rep_phone", "Claim Rep Phone #", False),
            ("loss_date", "Loss Date", False),
            ("claim_number", "Claim #", False),
            ("policy_number", "Policy/Member #", False),
            ("insured_name", "Insured Name", False),
            ("insured_phone", "Insured Phone #", False),
            ("loss_type", "Loss Type", False),
            ("market_area", "Market Area (City/Zip/Postal Code)", False),
            ("inspection_city_state", "Inspection (City & State)", False),
            ("claimant_name", "Claimant Name", False),
            ("claimant_phone", "Claimant Phone #", False),
            ("license_plate", "License #", False),
        ]
        claim_dialog = WizardDialog(self, "Autosource PDA - Claim Information", claim_questions, claim_defaults)
        if claim_dialog.exec() != QDialog.Accepted:
            return
        claim_answers = claim_dialog.values()

        vehicle_questions = [
            ("odometer", "Odometer", False),
            ("ext_color", "Ext. Color", False),
            ("vin", "VIN", False),
            ("year", "Year", False),
            ("make", "Make", False),
            ("model", "Model", False),
            ("doors", "Doors", False),
            ("bodystyle", "Bodystyle", False),
            ("drive", "Drive", False, ["", "2W", "4W"]),
            ("edition", "Edition", False),
            ("engine_type", "Fuel Type", False, ["Gas", "Diesel", "Other"]),
            ("engine_other", "Fuel Type - Other", False),
            ("engine_size", "Engine Size", False),
            ("cylinders", "# Cylinders", False),
            ("trans_type", "Transmission Type", False, ["Auto", "Manual", "CVT", "Other"]),
            ("trans_gears", "Transmission Gears", False, ["", "2 sp", "3 sp", "4 sp", "5 sp", "6 sp"]),
            ("capacity_passenger", "# Passenger Capacity", False),
            ("bed_length", "Bed Length", False, ["", "XShrt", "Short", "Long"]),
            ("capacity_tonnage", "Capacity/Tonnage", False),
            ("van_type", "Van Type", False, ["", "Passenger", "Cargo", "Regular", "Extended"]),
            ("conversion_name", "Conversion Name", False),
            ("inspected_by", "Inspected By", False),
            ("inspected_date", "Inspection Date", False),
            ("location_of_vehicle", "Location of Vehicle", False),
            ("pool_number", "Pool #", False),
        ]
        vehicle_dialog = WizardDialog(self, "Autosource PDA - Vehicle Information", vehicle_questions, vehicle_defaults)
        if vehicle_dialog.exec() != QDialog.Accepted:
            return
        vehicle_answers = vehicle_dialog.values()

        notes_questions = [
            ("equipment_notes", "Equipment Notes", True),
            ("general_comments", "General Comments", True),
        ]
        notes_dialog = WizardDialog(self, "Autosource PDA - Comments", notes_questions, notes_defaults)
        if notes_dialog.exec() != QDialog.Accepted:
            return
        notes_answers = notes_dialog.values()

        answers = {}
        answers.update(claim_answers)
        answers.update(vehicle_answers)
        answers.update(notes_answers)

        target_folder = self._record_folder_path(self.current_claim)
        claim_suffix = (self.current_claim.claim_id or "claim")[-4:]
        pdf_path = self._unique_destination(target_folder / f"{claim_suffix}-AUTOSOURCE-PDA-SUB-LEVEL.pdf")
        answers_path = self._autosource_answers_path(pdf_path)

        try:
            self._fill_autosource_pda_template(pdf_path, answers)
            answers_path.write_text(json.dumps(answers, indent=2), encoding="utf-8")
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(pdf_path)))
            QMessageBox.information(self, APP_NAME, f"Created:\n{pdf_path}")
        except Exception as exc:
            QMessageBox.warning(self, APP_NAME, f"Could not fill the Autosource template:\n{exc}")

    def _generate_aig_claim_summary(self) -> None:
        if not self.current_claim:
            return
        is_total_loss = "total" in (self.current_claim.claim_type or "").lower() or "total" in (self.current_claim.status or "").lower()
        vehicle_status_default = "total loss" if is_total_loss else "repair"
        days_default = ""
        comments_default = self.current_claim.assignment_claim_notes or self.current_claim.office_additional_notes or ""
        prior_damage_default = "No"
        drivable_default = "No" if is_total_loss or "not drivable" in (self.current_claim.assignment_claim_notes or "").lower() else "Yes"

        route_box = QMessageBox(self)
        route_box.setWindowTitle("AIG Route")
        route_box.setText("Is it repairable?")
        route_box.setInformativeText("Yes = Repairable\nNo = Total Loss")
        route_box.setStandardButtons(QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel)
        route_choice = route_box.exec()
        if route_choice == QMessageBox.Cancel:
            return

        repairable = route_choice == QMessageBox.Yes
        route_label = ""
        questions: list[tuple] = []
        if repairable:
            original_box = QMessageBox(self)
            original_box.setWindowTitle("AIG Route")
            original_box.setText("Is this an original estimate?")
            original_box.setInformativeText("Yes = Original Estimate (E01)\nNo = Supplement (S01-S99)")
            original_box.setStandardButtons(QMessageBox.Yes | QMessageBox.No | QMessageBox.Cancel)
            original_choice = original_box.exec()
            if original_choice == QMessageBox.Cancel:
                return
            if original_choice == QMessageBox.Yes:
                route_label = "AIG Appraisal Notes - E01 - Repairable"
                questions = [
                    ("claim_number", "Claim #", False),
                    ("vehicle_status", "Vehicle status (repair / cash settlement option)", False, ["repair", "cash settlement option"]),
                    ("drivable", "Vehicle drivable?", False, ["Yes", "No"]),
                    ("owner_copy", "Copy of appraisal supplied to owner? If so, how?", False, ["Yes/Email", "No"]),
                    ("agreed_price", "Is this an agreed price with the shop of owner's choice? If so, with who?", False),
                    ("shop_copy", "Did you supply a copy of this appraisal to the shop? If so, how?", False, ["Yes/Email", "No"]),
                    ("shop_contact", "Shop email / fax #", False),
                    ("shop_tax", "Shop tax id", False),
                    ("prior_damage", "Any unrelated or prior damage?", False, ["Yes", "No"]),
                    ("upd_created", "If yes, was an UPD estimate created?", False, ["Yes", "No"]),
                    ("repair_days", "Number of days to repair (total labor hours divided by 4)", False),
                    ("comments", "Additional comments", True),
                ]
            else:
                route_label = "AIG Appraisal Notes - S01-S99 - Repairable"
                questions = [
                    ("claim_number", "Claim #", False),
                    ("vehicle_status", "Vehicle status (repair / cash settlement option)", False, ["repair", "cash settlement option"]),
                    ("owner_copy", "Copy of appraisal supplied to owner? If so, how?", False, ["Yes/Email", "No"]),
                    ("agreed_price", "Is this an agreed price with the shop of owner's choice? If so, with who?", False),
                    ("shop_copy", "Did you supply a copy of this appraisal to the shop? If so, how?", False, ["Yes/Email", "No"]),
                    ("shop_contact", "Shop email / fax #", False),
                    ("shop_tax", "Shop tax id", False),
                    ("repair_days", "Additional number of days to repair (total labor hours divided by 4)", False),
                    ("comments", "Additional comments", True),
                ]
        else:
            route_label = "AIG Appraisal Notes - Total Loss"
            questions = [
                ("claim_number", "Claim #", False),
                ("vehicle_status", "Vehicle status (total loss / constructive total loss)", False, ["total loss", "constructive total loss"]),
                ("drivable", "Vehicle drivable?", False, ["Yes", "No"]),
                ("prior_damage", "Any unrelated or prior damage?", False, ["Yes", "No"]),
                ("upd_created", "If yes, was an UPD estimate created?", False, ["Yes", "No"]),
                ("towing", "Towing?", False),
                ("storage", "Storage per day?", False),
                ("charges", "Additional charges?", False),
                ("repair_days", "Number of days to repair (total labor hours divided by 4)", False),
                ("comments", "Additional comments", True),
            ]

        defaults = {
            "claim_number": self.current_claim.claim_number or self.current_claim.claim_id or "",
            "vehicle_status": vehicle_status_default,
            "drivable": drivable_default,
            "owner_copy": "No",
            "agreed_price": "N/A",
            "shop_copy": "No",
            "shop_contact": self.current_claim.shop_email or self.current_claim.contact_phone or "",
            "shop_tax": "N/A",
            "prior_damage": prior_damage_default,
            "upd_created": "No",
            "repair_days": days_default,
            "comments": comments_default,
            "towing": "N/A",
            "storage": "N/A",
            "charges": "N/A",
        }
        dialog = WizardDialog(self, "AIG Claim Summary", questions, defaults)
        if dialog.exec() != QDialog.Accepted:
            return
        answers = dialog.values()
        output_lines = [route_label, ""]
        skip_next = False
        ordered = [(field[1], answers.get(field[0], "").strip()) for field in questions]
        for index, (label, value) in enumerate(ordered):
            if skip_next:
                skip_next = False
                continue
            if label == "Any unrelated or prior damage?" and index + 1 < len(ordered):
                next_label, next_value = ordered[index + 1]
                if next_label == "If yes, was an UPD estimate created?":
                    output_lines.append(f"{label} If so, was an UPD estimate created?: {value} / {next_value}")
                    skip_next = True
                    continue
            output_lines.append(f"{label}: {value}")
        QApplication.clipboard().setText("\n".join(output_lines).strip())
        QMessageBox.information(self, APP_NAME, "AIG claim summary copied to clipboard.")

    def _generate_default_claim_notes(self) -> None:
        if not self.current_claim:
            return
        route, ok = QInputDialog.getItem(
            self,
            "Default Claim Notes",
            "Which route do you want to use?",
            ["Repairable", "Total Loss", "Supplement"],
            0,
            False,
        )
        if not ok or not route:
            return

        today_default = datetime.now().strftime("%#m/%#d/%Y") if os.name == "nt" else datetime.now().strftime("%-m/%-d/%Y")
        is_total_loss = "total" in (self.current_claim.claim_type or "").lower() or "total" in (self.current_claim.status or "").lower()
        drivable_default = "No" if is_total_loss else "Yes"
        shop_default = self.current_claim.shop_name or "No Shop Chosen"
        agreed_price_default = "No shop chosen" if not self.current_claim.shop_name else "N/A"
        estimate_copy_default = "No shop chosen" if not self.current_claim.shop_name else "Yes/Email"
        comments_default = self.current_claim.assignment_claim_notes or self.current_claim.office_additional_notes or ""
        explain_default = self.current_claim.assignment_claim_notes or ""
        supplement_comments_default = self.current_claim.office_additional_notes or "Supplement for parts, labor rate and AP."

        if route == "Repairable":
            header = "*** REPAIRABLE VEHICLE PARTIAL LOSS SUMMARY ***"
            questions = [
                ("date_inspected", "1. Date Inspected", False),
                ("drivable", "2. Is vehicle driveable?", False, ["Yes", "No"]),
                ("shop", "3. Shop chosen? Name of shop/City?", False),
                ("agreed_price", "4. Agreed Price? AP with whom?", False),
                ("estimate_copy", "5. Was shop provided copy of estimate? How delivered?", False, ["Yes/Email", "No", "No shop chosen"]),
                ("cost_effective_parts", "6. Were the most cost effective parts utilized?", False, ["Yes", "No"]),
                ("quote", "7. Quote #/Location of available parts", False),
                ("direction_to_pay", "8. Direction to pay?", False, ["Yes", "No"]),
                ("contacted_owner", "9. Contacted Owner Post Insp. & explained process?", False, ["Yes", "No"]),
                ("days", "10. Days to repair?", False),
                ("damage_consistent", "11. Damage consistent w facts?", False, ["Yes", "No"]),
                ("damage_photos", "12. Damages shown in photos?", False, ["Yes", "No"]),
                ("damage_explain", "12a. Explain damages shown in photos", True),
                ("unrelated", "13. Unrelated damage?", False, ["Yes", "No"]),
                ("unrelated_list", "13a. List unrelated damage", False),
                ("comments", "14. Additional Repairable Comments", True),
            ]
            defaults = {
                "date_inspected": today_default,
                "drivable": drivable_default,
                "shop": shop_default,
                "agreed_price": agreed_price_default,
                "estimate_copy": estimate_copy_default,
                "cost_effective_parts": "Yes",
                "quote": "No LKQ or A/M available",
                "direction_to_pay": "No",
                "contacted_owner": "Yes",
                "days": "",
                "damage_consistent": "Yes",
                "damage_photos": "Yes",
                "damage_explain": explain_default,
                "unrelated": "No",
                "unrelated_list": "none",
                "comments": comments_default,
            }
        elif route == "Total Loss":
            header = "*** TOTAL LOSS SUMMARY ***"
            questions = [
                ("date_inspected", "1. Date Inspected", False),
                ("location", "2. Location of vehicle?", False),
                ("tow_bill", "3. Tow Bill Amount", False),
                ("storage_rate", "4. Storage Rate", False),
                ("drivable", "5. Is vehicle driveable?", False, ["Yes", "No"]),
                ("contacted_owner", "6. Contacted Owner Post Insp.? TL Process explained?", False, ["Yes", "No"]),
                ("valuation", "7. TL Valuation/Request#", False),
                ("acv", "8. ACV", False),
                ("salvage", "9. Salvage Value", False),
                ("move", "10. Does vehicle need to be moved?", False, ["Yes", "No"]),
                ("lot", "11. Lot#", False),
                ("unrelated", "12. Unrelated damage?", False, ["Yes", "No"]),
                ("unrelated_list", "12a. List unrelated damage", False),
                ("comments", "13. Additional TL Comments", True),
            ]
            defaults = {
                "date_inspected": today_default,
                "location": self.current_claim.location_of_vehicle or self.current_claim.shop_name or "Customer has vehicle",
                "tow_bill": "N/A",
                "storage_rate": "N/A",
                "drivable": drivable_default,
                "contacted_owner": "Yes",
                "valuation": "CCC",
                "acv": "N/A",
                "salvage": "Salvage value:",
                "move": "No",
                "lot": "N/A",
                "unrelated": "No",
                "unrelated_list": "none",
                "comments": comments_default,
            }
        else:
            header = "*** SUPPLEMENT PARTIAL LOSS SUMMARY ***"
            questions = [
                ("shop", "1. Name of shop/City", False),
                ("shop_tax", "2. Shop Tax ID", False),
                ("direction_to_pay", "3. Direction to pay attached?", False, ["Yes", "No"]),
                ("invoices", "4. Invoices and photos attached?", False, ["Yes", "No"]),
                ("comments", "5. Additional Supplemental Comments", True),
            ]
            defaults = {
                "shop": shop_default,
                "shop_tax": "N/A",
                "direction_to_pay": "Yes",
                "invoices": "Yes",
                "comments": supplement_comments_default,
            }

        dialog = WizardDialog(self, "Default Claim Notes", questions, defaults)
        if dialog.exec() != QDialog.Accepted:
            return
        answers = dialog.values()
        output_lines = [header, ""]
        ordered = [(field[1], answers.get(field[0], "").strip()) for field in questions]
        for label, value in ordered:
            if label == "12a. Explain damages shown in photos" or label == "13a. List unrelated damage" or label == "12a. List unrelated damage":
                continue
            if label == "12. Damages shown in photos?":
                explain_value = answers.get("damage_explain", "").strip()
                output_lines.append(f"{label} Explain.: {value} {explain_value}".strip())
                continue
            if label in {"13. Unrelated damage?", "12. Unrelated damage?"}:
                list_value = answers.get("unrelated_list", "").strip()
                output_lines.append(f"{label} List: {value} {list_value}".strip())
                continue
            output_lines.append(f"{label}: {value}")
        QApplication.clipboard().setText("\n".join(output_lines).strip())
        QMessageBox.information(self, APP_NAME, "Default Claim Notes copied to clipboard.")

    def _generate_mitchell_total_loss(self) -> None:
        if not self.current_claim:
            return
        if Document is None:
            QMessageBox.warning(self, APP_NAME, "python-docx is not available for the Mitchell form.")
            return

        existing_mitchell = self._load_existing_mitchell_answers(self.current_claim)
        claimant_name_default, insured_name_default = self._mitchell_party_defaults(self.current_claim)
        today_default = datetime.now().strftime("%#m/%#d/%y") if os.name == "nt" else datetime.now().strftime("%-m/%-d/%y")
        top_questions = [
            ("Claim-Suffix ID", "Claim-Suffix ID", False),
            ("Claimant Name", "Claimant Name", False),
            ("Claimant Phone", "Claimant Phone", False),
            ("Loss Date", "Loss Date", False),
            ("License Plate", "License Plate", False),
            ("Insured Name", "Insured Name", False),
            ("Insured Phone", "Insured Phone", False),
            ("Loss Type", "Loss Type", False),
            ("VIN", "VIN", False),
            ("Year", "Year", False),
            ("Make", "Make", False),
            ("Model", "Model", False),
            ("Sub-model", "Sub-model", False),
            ("Mileage", "Mileage", False),
            ("Body Style", "Body Style", False),
            ("Ext. Color", "Ext. Color", False),
            ("Engine", "Engine", False),
            ("Transmission", "Transmission", False),
            ("Drive Train", "Drive Train", False),
            ("Location of Vehicle", "Location of Vehicle", False),
            ("Zip Code", "Zip Code", False),
            ("Inspected By", "Inspected By", False),
            ("Date", "Date", False),
        ]
        top_defaults = {
            "Claim-Suffix ID": existing_mitchell.get("Claim-Suffix ID", self.current_claim.claim_number or self.current_claim.claim_id or ""),
            "Claimant Name": existing_mitchell.get("Claimant Name", claimant_name_default),
            "Claimant Phone": existing_mitchell.get("Claimant Phone", self.current_claim.contact_phone or ""),
            "Loss Date": existing_mitchell.get("Loss Date", self.current_claim.date_of_loss or ""),
            "License Plate": existing_mitchell.get("License Plate", ""),
            "Insured Name": existing_mitchell.get("Insured Name", insured_name_default),
            "Insured Phone": existing_mitchell.get("Insured Phone", ""),
            "Loss Type": existing_mitchell.get("Loss Type", "Collision"),
            "VIN": existing_mitchell.get("VIN", self.current_claim.vin or ""),
            "Year": existing_mitchell.get("Year", self._mitchell_year_from_vehicle(self.current_claim.vehicle)),
            "Make": existing_mitchell.get("Make", self._mitchell_make_from_vehicle(self.current_claim.vehicle)),
            "Model": existing_mitchell.get("Model", self._mitchell_model_from_vehicle(self.current_claim.vehicle)),
            "Sub-model": existing_mitchell.get("Sub-model", ""),
            "Mileage": existing_mitchell.get("Mileage", ""),
            "Body Style": existing_mitchell.get("Body Style", ""),
            "Ext. Color": existing_mitchell.get("Ext. Color", ""),
            "Engine": existing_mitchell.get("Engine", ""),
            "Transmission": existing_mitchell.get("Transmission", ""),
            "Drive Train": existing_mitchell.get("Drive Train", ""),
            "Location of Vehicle": existing_mitchell.get("Location of Vehicle", self.current_claim.location_of_vehicle or self.current_claim.shop_name or ""),
            "Zip Code": existing_mitchell.get("Zip Code", self._zip_from_address(self.current_claim.owner_address)),
            "Inspected By": existing_mitchell.get("Inspected By", "Fernando Marin"),
            "Date": existing_mitchell.get("Date", today_default),
        }
        top_dialog = WizardDialog(self, "Mitchell Total Loss - Vehicle Information", top_questions, top_defaults)
        if top_dialog.exec() != QDialog.Accepted:
            return
        top_answers = top_dialog.values()

        rating_options = [
            "5 - Excellent",
            "4 - Very good",
            "3 - Good",
            "2 - Fair",
            "1 - Poor",
            "U - Unknown",
        ]
        condition_items = [
            "Seats",
            "Carpet",
            "Headliner",
            "Doors / Interior Panels",
            "Dash/Console",
            "Glass",
            "Body",
            "Paint",
            "Trim",
            "Vinyl or Convertible Tops",
            "Engine",
            "Transmission",
            "Tires",
        ]
        condition_questions: list[tuple] = []
        condition_defaults: dict[str, str] = {}
        for item in condition_items:
            condition_questions.append((f"{item} rating", f"{item} rating", False, rating_options))
            condition_questions.append((f"{item} comments", f"{item} comments", True))
            condition_defaults[f"{item} rating"] = existing_mitchell.get(f"{item} rating", "3 - Good")
            condition_defaults[f"{item} comments"] = existing_mitchell.get(f"{item} comments", "")
        condition_dialog = WizardDialog(self, "Mitchell Total Loss - Vehicle Condition", condition_questions, condition_defaults)
        if condition_dialog.exec() != QDialog.Accepted:
            return
        condition_answers_map = condition_dialog.values()

        final_questions = [
            ("After-Market Installed Parts, Refurbishments and Prior Damage", "After-Market Installed Parts, Refurbishments and Prior Damage", True),
            ("Comments", "Comments", True),
        ]
        final_defaults = {
            "After-Market Installed Parts, Refurbishments and Prior Damage": existing_mitchell.get(
                "After-Market Installed Parts, Refurbishments And Prior Damage",
                existing_mitchell.get("After-Market Installed Parts, Refurbishments and Prior Damage", "N/A"),
            ),
            "Comments": existing_mitchell.get("Comments", self.current_claim.assignment_claim_notes or ""),
        }
        final_dialog = WizardDialog(self, "Mitchell Total Loss - Final Sections", final_questions, final_defaults)
        if final_dialog.exec() != QDialog.Accepted:
            return
        final_answers = final_dialog.values()

        condition_answers: list[tuple[str, str, str]] = []
        for item in condition_items:
            condition_answers.append(
                (
                    item,
                    condition_answers_map.get(f"{item} rating", "3 - Good"),
                    condition_answers_map.get(f"{item} comments", "").strip(),
                )
            )

        target_folder = self._record_folder_path(self.current_claim)
        claim_suffix = (self.current_claim.claim_id or "claim")[-4:]
        doc_path = self._unique_destination(target_folder / f"{claim_suffix}-MITCHELL-TOTAL-LOSS.doc")
        pdf_path = doc_path.with_suffix(".pdf")

        try:
            self._fill_mitchell_total_loss_template(
                doc_path,
                top_answers,
                condition_answers,
                final_answers.get("After-Market Installed Parts, Refurbishments and Prior Damage", ""),
                final_answers.get("Comments", "").strip() or self.current_claim.assignment_claim_notes or self.current_claim.office_additional_notes or "",
            )
            self._export_doc_to_pdf(doc_path, pdf_path)
            QDesktopServices.openUrl(QUrl.fromLocalFile(str(pdf_path)))
            QMessageBox.information(self, APP_NAME, f"Created:\n{pdf_path}")
        except Exception as exc:
            if doc_path.exists():
                QDesktopServices.openUrl(QUrl.fromLocalFile(str(doc_path)))
                QMessageBox.information(self, APP_NAME, f"Created:\n{doc_path}\n\nPDF export failed:\n{exc}")
            else:
                QMessageBox.warning(self, APP_NAME, f"Could not fill the Mitchell template:\n{exc}")

    def _mitchell_party_defaults(self, claim: ClaimView) -> tuple[str, str]:
        customer = (claim.customer_name or "").strip()
        return customer, customer

    def _load_existing_mitchell_answers(self, claim: ClaimView) -> dict[str, str]:
        target_folder = self._record_folder_path(claim)
        candidates = sorted(target_folder.glob("*MITCHELL-TOTAL-LOSS.docx"))
        if not candidates or Document is None:
            return {}
        try:
            doc = Document(candidates[-1])
        except Exception:
            return {}

        extracted: dict[str, str] = {}
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text or ":" not in text:
                continue
            label, value = text.split(":", 1)
            label = label.strip()
            value = value.strip()
            if label in {
                "Claim-Suffix ID", "Claimant Name", "Claimant Phone", "Loss Date", "License Plate",
                "Insured Name", "Insured Phone", "Loss Type", "VIN", "Year", "Make", "Model",
                "Sub-model", "Mileage", "Body Style", "Ext. Color", "Engine", "Transmission",
                "Drive Train", "Location of Vehicle", "Zip Code", "Inspected By", "Date",
                "After-Market Installed Parts, Refurbishments And Prior Damage", "Comments",
            }:
                extracted[label] = "" if value == "-" else value
        if doc.tables:
            for row in doc.tables[0].rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                if len(cells) >= 3 and cells[0]:
                    extracted[f"{cells[0]} rating"] = cells[1]
                    extracted[f"{cells[0]} comments"] = cells[2]
        return extracted

    def _mitchell_year_from_vehicle(self, vehicle_text: str) -> str:
        tokens = (vehicle_text or "").split()
        if not tokens:
            return ""
        first = tokens[0]
        if re.fullmatch(r"\d{4}", first):
            return first
        if re.fullmatch(r"\d{2}", first):
            return f"20{first}"
        return ""

    def _mitchell_make_from_vehicle(self, vehicle_text: str) -> str:
        tokens = (vehicle_text or "").split()
        return tokens[1].title() if len(tokens) >= 2 else ""

    def _mitchell_model_from_vehicle(self, vehicle_text: str) -> str:
        tokens = (vehicle_text or "").split()
        return " ".join(tokens[2:4]).title() if len(tokens) >= 3 else ""

    def _zip_from_address(self, address: str) -> str:
        match = re.search(r"\b(\d{5})(?:-\d{4})?\b", address or "")
        return match.group(1) if match else ""

    def _reselect_current_claim(self, target_key: str | None = None) -> None:
        if not self.current_claim and not target_key:
            return
        lookup_key = target_key or self.current_claim.key
        self.current_claim = next((record for record in self.claims if record.key == lookup_key), None)
        if self.current_claim:
            self._render_claim(self.current_claim)

    def _open_assignment_sheet(self) -> None:
        if not self.current_claim:
            return
        assignment_path = self._find_assignment_sheet_path(self.current_claim)
        if not assignment_path:
            QMessageBox.information(self, APP_NAME, "No assignment sheet was found for this claim yet.")
            return
        QDesktopServices.openUrl(assignment_path.as_uri())  # type: ignore[arg-type]

    def _open_claim_folder(self) -> None:
        if not self.current_claim or not self.current_claim.source_path:
            return
        folder = Path(self.current_claim.source_path)
        if not folder.exists():
            QMessageBox.information(self, APP_NAME, "The claim folder could not be found.")
            return
        if sys.platform.startswith("win"):
            os.startfile(str(folder))  # type: ignore[attr-defined]
        else:
            subprocess.Popen(["open", str(folder)])


def main() -> int:
    app = QApplication(sys.argv)
    app.setApplicationName(APP_NAME)
    app.setApplicationVersion(APP_VERSION)
    window = ClaimsDashboard()
    window.show()
    return app.exec()


if __name__ == "__main__":
    raise SystemExit(main())
