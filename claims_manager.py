import json
from dataclasses import dataclass, asdict, field
from datetime import datetime, timedelta
from email.message import EmailMessage
from html import unescape
from itertools import permutations
import math
import os
from pathlib import Path
import re
import shutil
import smtplib
import subprocess
import sys
import threading
import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk
from copy import copy
import ssl
import time
from urllib.parse import parse_qs, quote, unquote, urlparse
from urllib.request import Request, urlopen
import webbrowser


APP_DIR = Path(__file__).resolve().parent
DATA_FILE = APP_DIR / "claims_data.json"
SECRETS_FILE = APP_DIR / "claims_secrets.json"
BACKUP_DIR = APP_DIR / "Backups"
MAX_DATA_BACKUPS = 60
DATE_FORMAT = "%Y-%m-%d %H:%M"
DEFAULT_WATCHED_FOLDERS = [
    r"C:\Users\ferna\Desktop\PENDING CLAIMS",
    r"C:\Users\ferna\Desktop\Closed Claims",
]
DEFAULT_TOOLS_FOLDER = r"C:\Users\ferna\Desktop\Claim Tools"
PAYROLL_TEMPLATE = Path(DEFAULT_TOOLS_FOLDER) / "APPRAISER PAYROLL SPREADSHEET.xlsx"
PAYROLL_OUTPUT_DIR = APP_DIR / "Payroll"
CLIENT_INSTRUCTIONS_DIR = Path(DEFAULT_TOOLS_FOLDER) / "Client instruction sheets"
VENDOR_DIR = APP_DIR / "vendor"
WORKING_SHEETS_DIR = APP_DIR / "Working Sheets"
MITCHELL_TOTAL_LOSS_TEMPLATE = Path(DEFAULT_TOOLS_FOLDER) / "MITCHELL TOTAL LOSS.doc"
OFFICE_UPDATE_EMAIL_FROM = "fernandomarin27@gmail.com"
OFFICE_UPDATE_EMAIL_TO = "gferreira@duhamels.com"
OFFICE_UPDATE_EMAIL_CC = "joe@lasalallc.com"
OFFICE_UPDATE_EMAIL_SUBJECT = "Open sheet"
OFFICE_UPDATE_EMAIL_BODY = "thank you!"
ROUTE_HOME_ADDRESS = "5 Richlee Rd, Norwalk, CT 06851"
APPTRAK_DOCS_DIR = Path(r"C:\AMobile\docs")
APPTRAK_AUTOMATION_DIR = Path(r"C:\AMobile\automation")
APPTRAK_IMPORT_BAT = APPTRAK_AUTOMATION_DIR / "Run-AppTrakImport.bat"
APPTRAK_EXE = Path(r"C:\AMobile\app\apptrak.exe")
APPTRAK_CHECK_INTERVAL_MS = 5 * 60 * 1000
MITCHELL_TOTAL_LOSS_FACTORY_OPTIONS = """Exterior
Air Dam - Front
Air Dam - Rear
Automatic Headlights
Carriage Top
Fog Lights
Graphics - OEM
Ground Effects - OEM
Luggage Rack
Mirror - Heated and Power
Pickup - Dual Rear Wheels
Pickup - Exterior Rails
Pickup - Sliding Rear Window
Pickup - Snow Plow Prep. Package
Pickup - Truck Bed Liner
Pickup - Tonneau Cover
Power Mirrors
Privacy Glass
Rear Step Bumper - Chrome
Rear Step Bumper - Painted
Removable Top
Running Boards
Soft Top
Spoiler (Rear)
Tinted Glass
Trailer Tow Hitch / Receiver
Trailer Tow Package
Van - Dual Power Sliding Doors
Van - Power Sliding Side Door
Vinyl/Landau Top
Wheels - Alum/Alloy
Wheels - Chrome
Wheels - Wire

Interior
Adjustable Foot Pedals
Air Conditioning - Automatic
Air Conditioning - Manual
Air Conditioning - Rear
Audio - AM/FM
Audio - Cassette
Audio - CD Player
Audio - CD Remote Changer
Audio - Entertainment / DVD System
Audio - In-Dash CD Changer
Audio - MP3
Audio - Premium Sound System
Audio - Rear Entertainment / DVD System
Cruise Control
Moonroof - Power
Moonroof - Dual
Navigation System
OnStar
Power Locks
Power Windows
Rear Window Defogger
Seats - First Row Bench
Seats - First Row Buckets
Seats - Heated
Seats - Heated and Cooled
Seats - Leather
Seats - Power Driver
Seats - Power Passenger
Seats - Second Row Bench
Seats - Second Row Buckets
Seats - Third Row Bench
Seats - Third Row Buckets
Sunroof (Power)
Telescope Steering Wheel
Tilt Steering - Automatic
Tilt Steering - Manual

Mechanical
ABS Brake Systems
ABS - Four Wheel
ABS - Rear Only
Automatic Load-Leveling
Pickup - Dual Fuel Tanks
Pickup - Power Running Board
Power Brakes
Power Steering
Suspension - Rear Air
Suspension - Self-leveling
Traction Control

Safety
Air Bag - Driver Side
Air Bag - Front Side
Air Bag - Head
Air Bag - Passenger Side
Air Bag - Second Row Side
Air Bag - Side Air Bags - Curtain
Air Bag - Side Air Bags - Torso
Anti-Theft - System
Anti-Theft - Tracking/Notification
Garage Door Opener
Keyless Entry
Remote Ignition
Reverse Sensor
Safety Rollbar
Tire Inflation/Pressure Monitor"""

if VENDOR_DIR.exists():
    sys.path.insert(0, str(VENDOR_DIR))

try:
    from pypdf import PdfReader
except Exception:
    PdfReader = None

try:
    from openpyxl import load_workbook
except Exception:
    load_workbook = None

try:
    from docx import Document
except Exception:
    Document = None

try:
    from docx.enum.section import WD_ORIENT
    from docx.shared import Inches
except Exception:
    WD_ORIENT = None
    Inches = None

def now_stamp() -> str:
    return datetime.now().strftime(DATE_FORMAT)


def legacy_notes_to_history(text: str) -> list[dict[str, str]]:
    cleaned = text.strip()
    if not cleaned:
        return []
    return [{"timestamp": "Imported", "text": cleaned}]


def render_note_history(entries: list[dict[str, str]]) -> str:
    if not entries:
        return ""
    parts = []
    for entry in entries:
        stamp = entry.get("timestamp", "")
        text = entry.get("text", "").strip()
        if text:
            parts.append(f"[{stamp}] {text}" if stamp else text)
    return "\n\n".join(parts)


def extract_job_number_from_text(text: str) -> str:
    match = re.search(r"\b(\d{8})\b", text)
    if not match:
        match = re.search(r"\b(\d{5,7})\b", text)
    return match.group(1) if match else ""


@dataclass
class ClaimRecord:
    key: str
    claim_id: str
    title: str
    status: str
    source_path: str
    claim_type: str
    updated_at: str
    notes: str = ""
    note_history: list[dict[str, str]] = field(default_factory=list)
    customer_name: str = ""
    insurance_company: str = ""
    claim_number: str = ""
    policy_number: str = ""
    insured_name: str = ""
    claimant_name: str = ""
    date_of_loss: str = ""
    town: str = ""
    owner_address: str = ""
    location_of_vehicle: str = ""
    vehicle: str = ""
    vin: str = ""
    damage_description: str = ""
    facts_of_loss: str = ""
    total_loss: bool = False
    shop_name: str = ""
    contact_phone: str = ""
    contact_email: str = ""
    assign_pdf_path: str = ""
    assignment_claim_notes: str = ""
    route_address_override: str = ""
    office_contacted: str = "No"
    office_appt_when: str = ""
    office_progress_status: str = ""
    office_waiting_for_paperwork: str = "No"
    office_additional_notes: str = ""
    manual_overrides: dict[str, object] = field(default_factory=dict)


class ClaimsStore:
    def __init__(self, path: Path) -> None:
        self.path = path
        self.data = self._load()
        self.secrets = self._load_secrets()
        self._migrate_legacy_secrets()

    def _load(self) -> dict:
        default_data = {
            "watched_folders": DEFAULT_WATCHED_FOLDERS.copy(),
            "claim_tools_folder": DEFAULT_TOOLS_FOLDER,
            "claim_tools_contacts": [],
            "processed_apptrak_pdfs": [],
            "email_settings": {
                "office_update_from": OFFICE_UPDATE_EMAIL_FROM,
                "office_update_app_password": "",
            },
            "route_plan_keys": [],
            "route_geocode_cache": {},
            "claims": {},
        }
        if not self.path.exists():
            return default_data

        try:
            loaded = json.loads(self.path.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return default_data

        loaded.setdefault("watched_folders", default_data["watched_folders"])
        loaded.setdefault("claim_tools_folder", default_data["claim_tools_folder"])
        loaded.setdefault("claim_tools_contacts", default_data["claim_tools_contacts"])
        loaded.setdefault("processed_apptrak_pdfs", default_data["processed_apptrak_pdfs"])
        loaded.setdefault("email_settings", default_data["email_settings"])
        loaded.setdefault("route_plan_keys", default_data["route_plan_keys"])
        loaded.setdefault("route_geocode_cache", default_data["route_geocode_cache"])
        loaded.setdefault("claims", {})
        return loaded

    def _load_secrets(self) -> dict:
        default_secrets = {"office_update_app_password": ""}
        if not SECRETS_FILE.exists():
            return default_secrets
        try:
            loaded = json.loads(SECRETS_FILE.read_text(encoding="utf-8"))
        except json.JSONDecodeError:
            return default_secrets
        loaded.setdefault("office_update_app_password", "")
        return loaded

    def _save_secrets(self) -> None:
        temp_path = SECRETS_FILE.with_suffix(f"{SECRETS_FILE.suffix}.tmp")
        temp_path.write_text(json.dumps(self.secrets, indent=2), encoding="utf-8")
        temp_path.replace(SECRETS_FILE)

    def _migrate_legacy_secrets(self) -> None:
        legacy_settings = self.data.get("email_settings", {})
        legacy_password = (legacy_settings.get("office_update_app_password") or "").strip()
        current_secret = (self.secrets.get("office_update_app_password") or "").strip()
        if legacy_password and not current_secret:
            self.secrets["office_update_app_password"] = legacy_password
            self._save_secrets()
        if legacy_password:
            legacy_settings["office_update_app_password"] = ""
            self.save()

    def save(self) -> None:
        self._backup_existing_file()
        temp_path = self.path.with_suffix(f"{self.path.suffix}.tmp")
        temp_path.write_text(json.dumps(self.data, indent=2), encoding="utf-8")
        temp_path.replace(self.path)

    def _backup_existing_file(self) -> None:
        if not self.path.exists():
            return
        BACKUP_DIR.mkdir(parents=True, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        backup_path = BACKUP_DIR / f"{self.path.stem}-{timestamp}{self.path.suffix}"
        shutil.copy2(self.path, backup_path)
        backups = sorted(
            BACKUP_DIR.glob(f"{self.path.stem}-*{self.path.suffix}"),
            key=lambda item: item.stat().st_mtime,
            reverse=True,
        )
        for old_backup in backups[MAX_DATA_BACKUPS:]:
            try:
                old_backup.unlink()
            except OSError:
                pass

    @property
    def watched_folders(self) -> list[str]:
        return self.data.setdefault("watched_folders", [])

    @property
    def claims(self) -> dict:
        return self.data.setdefault("claims", {})

    @property
    def claim_tools_folder(self) -> str:
        return self.data.setdefault("claim_tools_folder", DEFAULT_TOOLS_FOLDER)

    @property
    def claim_tools_contacts(self) -> list[dict[str, str]]:
        return self.data.setdefault("claim_tools_contacts", [])

    @property
    def processed_apptrak_pdfs(self) -> list[str]:
        return self.data.setdefault("processed_apptrak_pdfs", [])

    @property
    def email_settings(self) -> dict:
        return self.data.setdefault(
            "email_settings",
            {
                "office_update_from": OFFICE_UPDATE_EMAIL_FROM,
                "office_update_app_password": "",
            },
        )

    @property
    def office_update_app_password(self) -> str:
        return (self.secrets.get("office_update_app_password") or "").strip()

    def set_office_update_app_password(self, value: str) -> None:
        self.secrets["office_update_app_password"] = (value or "").strip()
        self._save_secrets()

    @property
    def route_plan_keys(self) -> list[str]:
        return self.data.setdefault("route_plan_keys", [])

    @property
    def route_geocode_cache(self) -> dict:
        return self.data.setdefault("route_geocode_cache", {})

    def add_watched_folder(self, folder: str) -> bool:
        if folder in self.watched_folders:
            return False
        self.watched_folders.append(folder)
        self.save()
        return True

    def remove_watched_folder(self, folder: str) -> None:
        self.data["watched_folders"] = [item for item in self.watched_folders if item != folder]
        self.save()

    def upsert_claim(self, record: ClaimRecord) -> None:
        self.claims[record.key] = asdict(record)

    def get_claim(self, key: str) -> ClaimRecord | None:
        item = self.claims.get(key)
        if not item:
            return None
        return ClaimRecord(**self._normalize_claim_item(item))

    def normalize_all_claims(self) -> None:
        self.data["claims"] = {
            key: self._normalize_claim_item(value)
            for key, value in self.claims.items()
        }

    def _normalize_claim_item(self, item: dict) -> dict:
        normalized = dict(item)
        normalized.pop("assigned_to", None)
        normalized.pop("office_appt_scheduled", None)
        normalized.setdefault("note_history", legacy_notes_to_history(normalized.get("notes", "")))
        normalized_claim_id = extract_job_number_from_text(str(normalized.get("claim_id", "")))
        if normalized_claim_id:
            normalized["claim_id"] = normalized_claim_id
        normalized.setdefault("town", "")
        normalized.setdefault("owner_address", "")
        normalized.setdefault("location_of_vehicle", "")
        normalized.setdefault("damage_description", "")
        normalized.setdefault("facts_of_loss", "")
        normalized.setdefault("total_loss", False)
        normalized.setdefault("route_address_override", "")
        normalized.setdefault("office_contacted", "No")
        normalized.setdefault("office_appt_when", "")
        normalized.setdefault("office_progress_status", "")
        normalized.setdefault("office_waiting_for_paperwork", "No")
        normalized.setdefault("office_additional_notes", "")
        normalized.setdefault("assignment_claim_notes", "")
        normalized.setdefault("manual_overrides", {})
        return normalized


class ClaimsManagerApp:
    def __init__(self, root: tk.Tk) -> None:
        self.root = root
        self.root.title("Claims Manager")
        self.root.geometry("1280x760")
        self.store = ClaimsStore(DATA_FILE)
        self.store.normalize_all_claims()
        self.store.save()
        self.search_var = tk.StringVar()
        self.selected_key: str | None = None
        self.treeviews: dict[str, ttk.Treeview] = {}
        self.tree_sort_state: dict[str, tuple[str, bool]] = {}
        self.tree_context_menu: tk.Menu | None = None
        self.claim_notebook: ttk.Notebook | None = None
        self.office_tree: ttk.Treeview | None = None
        self.office_notes_text: tk.Text | None = None
        self.claim_tools_contacts_tree: ttk.Treeview | None = None
        self.route_available_tree: ttk.Treeview | None = None
        self.route_selected_tree: ttk.Treeview | None = None
        self.route_selected_stop_key: str | None = None
        self.route_map_button: ttk.Button | None = None
        self.scan_access_issues: list[str] = []
        self.last_dialog_geometry: tuple[int, int, int, int] | None = None
        self.apptrak_import_running = False
        self.apptrak_after_id: str | None = None
        self.apptrak_status_var = tk.StringVar(value="AppTrak import idle.")
        self.summary_vars = {
            "all": tk.StringVar(value="0"),
            "open": tk.StringVar(value="0"),
            "closed": tk.StringVar(value="0"),
        }
        self.closed_stats_vars = {
            "today": tk.StringVar(value="0"),
            "week": tk.StringVar(value="0"),
            "month": tk.StringVar(value="0"),
            "selected_date": tk.StringVar(value="-"),
            "selected_count": tk.StringVar(value="0"),
        }
        self.closed_date_var = tk.StringVar()
        self.range_start_var = tk.StringVar()
        self.range_end_var = tk.StringVar()
        self.tools_folder_var = tk.StringVar(value=self.store.claim_tools_folder)
        self.office_appt_when_var = tk.StringVar()
        self.office_progress_status_var = tk.StringVar()
        self.office_waiting_for_paperwork_var = tk.StringVar(value="No")
        self.route_home_var = tk.StringVar(value=ROUTE_HOME_ADDRESS)
        self.route_address_var = tk.StringVar()
        self.route_status_var = tk.StringVar(value="Build tomorrow's route from your house.")
        self.editable_detail_keys = {
            "claim_id",
            "customer_name",
            "insurance_company",
            "claim_number",
            "date_of_loss",
            "town",
            "owner_address",
            "location_of_vehicle",
            "vehicle",
            "vin",
            "damage_description",
            "facts_of_loss",
            "total_loss",
            "shop_name",
            "contact_phone",
            "contact_email",
            "assignment_claim_notes",
        }
        self.detail_vars = {
            "claim_id": tk.StringVar(),
            "title": tk.StringVar(),
            "status": tk.StringVar(),
            "source_path": tk.StringVar(),
            "claim_type": tk.StringVar(),
            "updated_at": tk.StringVar(),
            "customer_name": tk.StringVar(),
            "insurance_company": tk.StringVar(),
            "claim_number": tk.StringVar(),
            "policy_number": tk.StringVar(),
            "insured_name": tk.StringVar(),
            "claimant_name": tk.StringVar(),
            "date_of_loss": tk.StringVar(),
            "town": tk.StringVar(),
            "owner_address": tk.StringVar(),
            "location_of_vehicle": tk.StringVar(),
            "vehicle": tk.StringVar(),
            "vin": tk.StringVar(),
            "damage_description": tk.StringVar(),
            "facts_of_loss": tk.StringVar(),
            "total_loss": tk.StringVar(),
            "shop_name": tk.StringVar(),
            "contact_phone": tk.StringVar(),
            "contact_email": tk.StringVar(),
            "assign_pdf_path": tk.StringVar(),
            "assignment_claim_notes": tk.StringVar(),
        }

        self._build_ui()
        self.scan_folders()
        self._schedule_next_apptrak_check(initial_delay_ms=15000)

    def _build_ui(self) -> None:
        self.root.columnconfigure(0, weight=3)
        self.root.columnconfigure(1, weight=2)
        self.root.rowconfigure(1, weight=1)

        header = ttk.Frame(self.root, padding=16)
        header.grid(row=0, column=0, columnspan=2, sticky="ew")
        header.columnconfigure(0, weight=1)

        title_frame = ttk.Frame(header)
        title_frame.grid(row=0, column=0, sticky="w")
        ttk.Label(title_frame, text="Claims Dashboard", font=("Segoe UI", 20, "bold")).grid(row=0, column=0, sticky="w")
        ttk.Label(
            title_frame,
            text="Track open claims, closed claims, notes, and watched claim folders from one place.",
        ).grid(row=1, column=0, sticky="w", pady=(4, 0))

        controls = ttk.Frame(header)
        controls.grid(row=0, column=1, sticky="e")
        ttk.Entry(controls, textvariable=self.search_var, width=28).grid(row=0, column=0, padx=(0, 8))
        ttk.Button(controls, text="Search", command=self.refresh_views).grid(row=0, column=1, padx=(0, 8))
        ttk.Button(controls, text="Refresh Scan", command=self.scan_folders).grid(row=0, column=2, padx=(0, 8))
        ttk.Button(controls, text="Check AppTrak", command=self.run_apptrak_import_now).grid(row=0, column=3, padx=(0, 8))
        ttk.Button(controls, text="Add Folder", command=self.add_folder).grid(row=0, column=4, padx=(0, 8))
        ttk.Button(controls, text="New Claim", command=self.add_manual_claim).grid(row=0, column=5)
        ttk.Label(header, textvariable=self.apptrak_status_var).grid(row=1, column=1, sticky="e", pady=(8, 0))

        summary = ttk.Frame(self.root, padding=(16, 0, 16, 16))
        summary.grid(row=1, column=0, sticky="nsew")
        summary.columnconfigure((0, 1, 2), weight=1)
        summary.rowconfigure(1, weight=1)

        self._summary_card(summary, "All Claims", self.summary_vars["all"], 0)
        self._summary_card(summary, "Open Claims", self.summary_vars["open"], 1)
        self._summary_card(summary, "Closed Claims", self.summary_vars["closed"], 2)

        notebook = ttk.Notebook(summary)
        notebook.grid(row=1, column=0, columnspan=3, sticky="nsew", pady=(16, 0))
        self.claim_notebook = notebook
        for name, label in [("all", "All Claims"), ("open", "Open Claims"), ("closed", "Closed Claims")]:
            frame = ttk.Frame(notebook, padding=8)
            frame.columnconfigure(0, weight=1)
            frame.rowconfigure(0, weight=1)
            notebook.add(frame, text=label)
            tree = self._build_tree(frame)
            self.treeviews[name] = tree

        self.tree_context_menu = tk.Menu(self.root, tearoff=0)
        self.tree_context_menu.add_command(label="Open Working Sheet", command=self.open_working_sheet_for_selected_claim)
        self.tree_context_menu.add_command(label="Email Assignment Sheet", command=self.email_assignment_sheet_for_selected_claim)
        self.tree_context_menu.add_command(label="NADA Value PDF", command=self.generate_nada_pdf_for_selected_claim)
        self.claim_tools_menu = tk.Menu(self.tree_context_menu, tearoff=0)
        self.tree_context_menu.add_cascade(label="Claim Tools", menu=self.claim_tools_menu)
        self.tree_context_menu.add_command(label="Open Folder", command=self.open_selected_claim_folder_from_menu)

        reports_frame = ttk.Frame(notebook, padding=12)
        reports_frame.columnconfigure(0, weight=1)
        reports_frame.rowconfigure(1, weight=1)
        notebook.add(reports_frame, text="Reports")

        filters_frame = ttk.LabelFrame(reports_frame, text="Filters", padding=12)
        filters_frame.grid(row=0, column=0, sticky="ew", pady=(0, 12))
        filters_frame.columnconfigure(1, weight=1)
        filters_frame.columnconfigure(3, weight=1)
        ttk.Label(filters_frame, text="Specific Date").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=2)
        self.closed_date_combo = ttk.Combobox(filters_frame, textvariable=self.closed_date_var, state="readonly", height=10)
        self.closed_date_combo.grid(row=0, column=1, sticky="ew", pady=2)
        self.closed_date_combo.bind("<<ComboboxSelected>>", self.on_closed_date_selected)
        ttk.Label(filters_frame, text="Start Date").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=2)
        self.range_start_combo = ttk.Combobox(filters_frame, textvariable=self.range_start_var, height=10)
        self.range_start_combo.grid(row=1, column=1, sticky="ew", pady=2)
        self.range_start_combo.bind("<FocusOut>", lambda _event: self.normalize_report_date_field(self.range_start_var))
        self.range_start_combo.bind("<Return>", lambda _event: self.normalize_report_date_field(self.range_start_var))
        ttk.Label(filters_frame, text="End Date").grid(row=1, column=2, sticky="w", padx=(12, 8), pady=2)
        self.range_end_combo = ttk.Combobox(filters_frame, textvariable=self.range_end_var, height=10)
        self.range_end_combo.grid(row=1, column=3, sticky="ew", pady=2)
        self.range_end_combo.bind("<FocusOut>", lambda _event: self.normalize_report_date_field(self.range_end_var))
        self.range_end_combo.bind("<Return>", lambda _event: self.normalize_report_date_field(self.range_end_var))
        ttk.Button(filters_frame, text="Apply Date Range", command=self.apply_reports_range_filter).grid(row=1, column=4, sticky="ew", padx=(12, 0), pady=2)
        ttk.Button(filters_frame, text="Clear Filters", command=self.clear_reports_filters).grid(row=0, column=4, sticky="ew", padx=(12, 0), pady=2)

        body_frame = ttk.Frame(reports_frame)
        body_frame.grid(row=1, column=0, sticky="nsew")
        body_frame.columnconfigure(0, weight=1)
        body_frame.columnconfigure(1, weight=1)
        body_frame.rowconfigure(0, weight=1)

        stats_frame = ttk.LabelFrame(body_frame, text="Closed Volume", padding=12)
        stats_frame.grid(row=0, column=0, sticky="nsew", padx=(0, 8))
        stats_frame.columnconfigure(1, weight=1)
        ttk.Label(stats_frame, text="Closed Today").grid(row=0, column=0, sticky="w", padx=(0, 12), pady=2)
        ttk.Label(stats_frame, textvariable=self.closed_stats_vars["today"]).grid(row=0, column=1, sticky="w", pady=2)
        ttk.Label(stats_frame, text="Closed This Week").grid(row=1, column=0, sticky="w", padx=(0, 12), pady=2)
        ttk.Label(stats_frame, textvariable=self.closed_stats_vars["week"]).grid(row=1, column=1, sticky="w", pady=2)
        ttk.Label(stats_frame, text="Closed This Month").grid(row=2, column=0, sticky="w", padx=(0, 12), pady=2)
        ttk.Label(stats_frame, textvariable=self.closed_stats_vars["month"]).grid(row=2, column=1, sticky="w", pady=2)
        ttk.Label(stats_frame, text="Selected Date").grid(row=3, column=0, sticky="w", padx=(0, 12), pady=(10, 2))
        ttk.Label(stats_frame, textvariable=self.closed_stats_vars["selected_date"]).grid(row=3, column=1, sticky="w", pady=(10, 2))
        ttk.Label(stats_frame, text="Count For Date").grid(row=4, column=0, sticky="w", padx=(0, 12), pady=2)
        ttk.Label(stats_frame, textvariable=self.closed_stats_vars["selected_count"]).grid(row=4, column=1, sticky="w", pady=2)
        ttk.Button(stats_frame, text="Calculate Payroll", command=self.calculate_payroll).grid(row=5, column=0, columnspan=2, sticky="ew", pady=(16, 0))

        report_list_frame = ttk.LabelFrame(body_frame, text="Claims In Filter", padding=12)
        report_list_frame.grid(row=0, column=1, sticky="nsew", padx=(8, 0))
        report_list_frame.columnconfigure(0, weight=1)
        report_list_frame.rowconfigure(0, weight=1)
        self.reports_tree = ttk.Treeview(
            report_list_frame,
            columns=("closed_date", "claim_id", "customer", "insurance", "type"),
            show="headings",
            selectmode="browse",
        )
        for col, width in [
            ("closed_date", 90),
            ("claim_id", 90),
            ("customer", 170),
            ("insurance", 170),
            ("type", 90),
        ]:
            self.reports_tree.heading(col, text=col.replace("_", " ").title())
            self.reports_tree.column(col, width=width, anchor="w")
        self._configure_tree_sorting(self.reports_tree, ("closed_date", "claim_id", "customer", "insurance", "type"))
        self.reports_tree.grid(row=0, column=0, sticky="nsew")
        reports_scroll = ttk.Scrollbar(report_list_frame, orient="vertical", command=self.reports_tree.yview)
        reports_scroll.grid(row=0, column=1, sticky="ns")
        self.reports_tree.configure(yscrollcommand=reports_scroll.set)

        office_frame = ttk.Frame(notebook, padding=12)
        office_frame.columnconfigure(0, weight=1)
        office_frame.rowconfigure(0, weight=3)
        office_frame.rowconfigure(1, weight=2)
        notebook.add(office_frame, text="Office Update")

        office_list_frame = ttk.LabelFrame(office_frame, text="Open Claims", padding=12)
        office_list_frame.grid(row=0, column=0, sticky="nsew", pady=(0, 10))
        office_list_frame.columnconfigure(0, weight=1)
        office_list_frame.rowconfigure(0, weight=1)
        self.office_tree = ttk.Treeview(
            office_list_frame,
            columns=("claim_id", "customer", "status", "inspection_when"),
            show="headings",
            selectmode="browse",
        )
        for col, width in [
            ("claim_id", 90),
            ("customer", 260),
            ("status", 260),
            ("inspection_when", 140),
        ]:
            self.office_tree.heading(col, text=col.replace("_", " ").title())
            self.office_tree.column(col, width=width, anchor="w")
        self._configure_tree_sorting(self.office_tree, ("claim_id", "customer", "status", "inspection_when"))
        self.office_tree.grid(row=0, column=0, sticky="nsew")
        office_scroll = ttk.Scrollbar(office_list_frame, orient="vertical", command=self.office_tree.yview)
        office_scroll.grid(row=0, column=1, sticky="ns")
        self.office_tree.configure(yscrollcommand=office_scroll.set)
        self.office_tree.bind("<<TreeviewSelect>>", self.on_office_claim_selected)
        self.office_tree.bind("<Button-3>", self.show_tree_context_menu)

        office_bottom_frame = ttk.Frame(office_frame)
        office_bottom_frame.grid(row=1, column=0, sticky="nsew")
        office_bottom_frame.columnconfigure(0, weight=1)
        office_bottom_frame.columnconfigure(1, weight=1)
        office_bottom_frame.rowconfigure(0, weight=0)
        office_bottom_frame.rowconfigure(1, weight=1)

        office_form_frame = ttk.LabelFrame(office_bottom_frame, text="Update Details", padding=12)
        office_form_frame.grid(row=0, column=0, sticky="new", padx=(0, 8))
        office_form_frame.columnconfigure(1, weight=1)
        ttk.Label(office_form_frame, text="Progress").grid(row=0, column=0, sticky="w", padx=(0, 10), pady=4)
        ttk.Combobox(
            office_form_frame,
            textvariable=self.office_progress_status_var,
            values=[
                "No Contact Yet",
                "Contacted",
                "Appointment Scheduled",
                "Seen - Need To Write",
                "Written - Under Review",
                "Supplement - Waiting For Paperwork",
                "Waiting For Paperwork",
            ],
            state="readonly",
        ).grid(row=0, column=1, sticky="ew", pady=4)
        ttk.Label(office_form_frame, text="Date Of Inspection Set Up").grid(row=1, column=0, sticky="w", padx=(0, 10), pady=4)
        ttk.Entry(office_form_frame, textvariable=self.office_appt_when_var).grid(row=1, column=1, sticky="ew", pady=4)
        ttk.Label(office_form_frame, text="Waiting For Paperwork?").grid(row=2, column=0, sticky="w", padx=(0, 10), pady=4)
        ttk.Combobox(office_form_frame, textvariable=self.office_waiting_for_paperwork_var, values=["Yes", "No"], state="readonly").grid(row=2, column=1, sticky="ew", pady=4)
        ttk.Label(office_form_frame, text="Additional Notes").grid(row=3, column=0, sticky="nw", padx=(0, 10), pady=4)
        self.office_notes_text = tk.Text(office_form_frame, height=4, wrap="word")
        self.office_notes_text.grid(row=3, column=1, sticky="ew", pady=4)
        ttk.Button(office_form_frame, text="Save Update", command=self.save_office_update).grid(row=4, column=0, columnspan=2, sticky="ew", pady=(8, 0))

        office_summary_frame = ttk.LabelFrame(office_bottom_frame, text="Selected Claim Summary", padding=12)
        office_summary_frame.grid(row=0, column=1, rowspan=2, sticky="nsew", padx=(8, 0))
        office_summary_frame.columnconfigure(0, weight=1)
        office_summary_frame.rowconfigure(0, weight=1)
        self.office_summary_text = tk.Text(office_summary_frame, height=10, wrap="word")
        self.office_summary_text.grid(row=0, column=0, sticky="nsew", pady=(0, 8))
        office_actions = ttk.Frame(office_summary_frame)
        office_actions.grid(row=1, column=0, sticky="ew")
        office_actions.columnconfigure((0, 1, 2), weight=1)
        ttk.Button(office_actions, text="Copy Claim Summary", command=self.copy_office_summary).grid(row=0, column=0, sticky="ew", padx=(0, 6))
        ttk.Button(office_actions, text="Preview PDF", command=self.preview_office_summary_pdf).grid(row=0, column=1, sticky="ew", padx=6)
        ttk.Button(office_actions, text="Email Office Update", command=self.export_office_summary_pdf).grid(row=0, column=2, sticky="ew", padx=(6, 0))

        route_frame = ttk.Frame(notebook, padding=12)
        route_frame.columnconfigure(0, weight=1)
        route_frame.columnconfigure(1, weight=1)
        route_frame.rowconfigure(1, weight=1)
        notebook.add(route_frame, text="Route Planner")

        tools_tab = ttk.Frame(notebook, padding=12)
        tools_tab.columnconfigure(0, weight=1)
        tools_tab.rowconfigure(0, weight=1)
        notebook.add(tools_tab, text="Claim Tools")

        tools_contacts_frame = ttk.LabelFrame(tools_tab, text="Names And Numbers", padding=12)
        tools_contacts_frame.grid(row=0, column=0, sticky="nsew")
        tools_contacts_frame.columnconfigure(0, weight=1)
        tools_contacts_frame.rowconfigure(0, weight=1)

        self.claim_tools_contacts_tree = ttk.Treeview(
            tools_contacts_frame,
            columns=("name", "number", "prompt_guide", "notes"),
            show="headings",
            selectmode="browse",
        )
        self.claim_tools_contacts_tree.heading("name", text="Name")
        self.claim_tools_contacts_tree.heading("number", text="Number")
        self.claim_tools_contacts_tree.heading("prompt_guide", text="Prompt Guide")
        self.claim_tools_contacts_tree.heading("notes", text="Notes")
        self.claim_tools_contacts_tree.column("name", width=220, anchor="w")
        self.claim_tools_contacts_tree.column("number", width=180, anchor="w")
        self.claim_tools_contacts_tree.column("prompt_guide", width=280, anchor="w")
        self.claim_tools_contacts_tree.column("notes", width=320, anchor="w")
        self._configure_tree_sorting(self.claim_tools_contacts_tree, ("name", "number", "prompt_guide", "notes"))
        self.claim_tools_contacts_tree.grid(row=0, column=0, sticky="nsew")
        tools_contacts_scroll = ttk.Scrollbar(tools_contacts_frame, orient="vertical", command=self.claim_tools_contacts_tree.yview)
        tools_contacts_scroll.grid(row=0, column=1, sticky="ns")
        self.claim_tools_contacts_tree.configure(yscrollcommand=tools_contacts_scroll.set)
        self.claim_tools_contacts_tree.bind("<Double-1>", self.call_claim_tool_contact)

        tools_contacts_actions = ttk.Frame(tools_contacts_frame)
        tools_contacts_actions.grid(row=1, column=0, columnspan=2, sticky="ew", pady=(8, 0))
        tools_contacts_actions.columnconfigure((0, 1, 2), weight=1)
        ttk.Button(tools_contacts_actions, text="Add", command=self.add_claim_tool_contact).grid(row=0, column=0, sticky="ew", padx=(0, 6))
        ttk.Button(tools_contacts_actions, text="Edit", command=self.edit_claim_tool_contact).grid(row=0, column=1, sticky="ew", padx=6)
        ttk.Button(tools_contacts_actions, text="Remove", command=self.remove_claim_tool_contact).grid(row=0, column=2, sticky="ew", padx=(6, 0))

        route_top = ttk.LabelFrame(route_frame, text="Route Settings", padding=12)
        route_top.grid(row=0, column=0, columnspan=2, sticky="ew", pady=(0, 12))
        route_top.columnconfigure(1, weight=1)
        ttk.Label(route_top, text="Start From").grid(row=0, column=0, sticky="w", padx=(0, 8))
        ttk.Entry(route_top, textvariable=self.route_home_var).grid(row=0, column=1, sticky="ew")
        ttk.Label(route_top, textvariable=self.route_status_var).grid(row=1, column=0, columnspan=2, sticky="w", pady=(8, 0))

        route_available_frame = ttk.LabelFrame(route_frame, text="Open Claims", padding=12)
        route_available_frame.grid(row=1, column=0, sticky="nsew", padx=(0, 8))
        route_available_frame.columnconfigure(0, weight=1)
        route_available_frame.rowconfigure(0, weight=1)
        self.route_available_tree = ttk.Treeview(
            route_available_frame,
            columns=("claim_id", "customer", "address"),
            show="headings",
            selectmode="extended",
        )
        for col, width in [("claim_id", 90), ("customer", 170), ("address", 220)]:
            self.route_available_tree.heading(col, text=col.replace("_", " ").title())
            self.route_available_tree.column(col, width=width, anchor="w")
        self._configure_tree_sorting(self.route_available_tree, ("claim_id", "customer", "address"))
        self.route_available_tree.grid(row=0, column=0, sticky="nsew")
        route_available_scroll = ttk.Scrollbar(route_available_frame, orient="vertical", command=self.route_available_tree.yview)
        route_available_scroll.grid(row=0, column=1, sticky="ns")
        self.route_available_tree.configure(yscrollcommand=route_available_scroll.set)
        self.route_available_tree.bind("<Button-3>", self.show_tree_context_menu)
        ttk.Button(route_available_frame, text="Add Selected To Route", command=self.add_selected_claims_to_route).grid(row=1, column=0, columnspan=2, sticky="ew", pady=(8, 0))

        route_selected_frame = ttk.LabelFrame(route_frame, text="Tomorrow's Stops", padding=12)
        route_selected_frame.grid(row=1, column=1, sticky="nsew", padx=(8, 0))
        route_selected_frame.columnconfigure(0, weight=1)
        route_selected_frame.rowconfigure(0, weight=1)
        self.route_selected_tree = ttk.Treeview(
            route_selected_frame,
            columns=("order", "claim_id", "customer", "address"),
            show="headings",
            selectmode="browse",
        )
        for col, width in [("order", 50), ("claim_id", 90), ("customer", 150), ("address", 260)]:
            self.route_selected_tree.heading(col, text=col.replace("_", " ").title())
            self.route_selected_tree.column(col, width=width, anchor="w")
        self._configure_tree_sorting(self.route_selected_tree, ("order", "claim_id", "customer", "address"))
        self.route_selected_tree.grid(row=0, column=0, columnspan=3, sticky="nsew")
        route_selected_scroll = ttk.Scrollbar(route_selected_frame, orient="vertical", command=self.route_selected_tree.yview)
        route_selected_scroll.grid(row=0, column=3, sticky="ns")
        self.route_selected_tree.configure(yscrollcommand=route_selected_scroll.set)
        self.route_selected_tree.bind("<<TreeviewSelect>>", self.on_route_stop_selected)
        self.route_selected_tree.bind("<Button-3>", self.show_tree_context_menu)

        ttk.Label(route_selected_frame, text="Stop Address").grid(row=1, column=0, sticky="w", pady=(8, 0))
        ttk.Entry(route_selected_frame, textvariable=self.route_address_var).grid(row=2, column=0, sticky="ew", pady=(4, 8))
        ttk.Button(route_selected_frame, text="Save Address", command=self.save_route_stop_address).grid(row=2, column=1, sticky="ew", padx=(8, 0), pady=(4, 8))
        ttk.Button(route_selected_frame, text="Remove Stop", command=self.remove_route_stop).grid(row=2, column=2, sticky="ew", padx=(8, 0), pady=(4, 8))

        route_actions = ttk.Frame(route_selected_frame)
        route_actions.grid(row=3, column=0, columnspan=4, sticky="ew")
        route_actions.columnconfigure((0, 1, 2), weight=1)
        self.route_map_button = ttk.Button(route_actions, text="Open Optimized Route", command=self.open_route_in_google_maps)
        self.route_map_button.grid(row=0, column=0, sticky="ew", padx=(0, 6))
        ttk.Button(route_actions, text="Email Uninspected Only", command=self.email_uninspected_only_to_self).grid(row=0, column=1, sticky="ew", padx=6)
        ttk.Button(route_actions, text="Clear Route", command=self.clear_route_plan).grid(row=0, column=2, sticky="ew", padx=(6, 0))

        side = ttk.Frame(self.root, padding=(0, 0, 16, 16))
        side.grid(row=1, column=1, sticky="nsew")
        side.columnconfigure(0, weight=1)
        side.rowconfigure(1, weight=1)
        side.rowconfigure(2, weight=1)

        top_info = ttk.Frame(side)
        top_info.grid(row=0, column=0, sticky="ew", pady=(0, 12))
        top_info.columnconfigure(0, weight=1)
        top_info.columnconfigure(1, weight=1)

        watched = ttk.LabelFrame(top_info, text="Watched Folders", padding=8)
        watched.grid(row=0, column=0, sticky="ew")
        watched.columnconfigure(0, weight=1)
        self.folder_list = tk.Listbox(watched, height=3)
        self.folder_list.grid(row=0, column=0, sticky="ew")
        ttk.Button(watched, text="Remove Folder", command=self.remove_selected_folder).grid(row=1, column=0, sticky="ew", pady=(6, 0))

        tools = ttk.LabelFrame(top_info, text="Claim Tools Folder", padding=8)
        tools.grid(row=0, column=1, sticky="ew", padx=(12, 0))
        tools.columnconfigure(0, weight=1)
        ttk.Label(tools, textvariable=self.tools_folder_var, wraplength=360, justify="left").grid(row=0, column=0, sticky="w")

        details_wrap = ttk.LabelFrame(side, text="Claim Details", padding=8)
        details_wrap.grid(row=1, column=0, sticky="nsew", pady=(0, 12))
        details_wrap.columnconfigure(0, weight=1)
        details_wrap.rowconfigure(0, weight=1)

        details_canvas = tk.Canvas(details_wrap, highlightthickness=0)
        details_canvas.grid(row=0, column=0, sticky="nsew")
        details_scrollbar = ttk.Scrollbar(details_wrap, orient="vertical", command=details_canvas.yview)
        details_scrollbar.grid(row=0, column=1, sticky="ns")
        details_canvas.configure(yscrollcommand=details_scrollbar.set)

        details = ttk.Frame(details_canvas)
        details_window = details_canvas.create_window((0, 0), window=details, anchor="nw")
        details.columnconfigure(1, weight=1)

        def sync_details_scroll(_event: object) -> None:
            details_canvas.configure(scrollregion=details_canvas.bbox("all"))

        def resize_details(event: object) -> None:
            details_canvas.itemconfigure(details_window, width=getattr(event, "width", 0))

        details.bind("<Configure>", sync_details_scroll)
        details_canvas.bind("<Configure>", resize_details)

        detail_rows = [
            ("Claim ID", "claim_id"),
            ("Customer", "customer_name"),
            ("Insurance", "insurance_company"),
            ("Claim #", "claim_number"),
            ("Date of Loss", "date_of_loss"),
            ("Town", "town"),
            ("Owner Address", "owner_address"),
            ("Location", "location_of_vehicle"),
            ("Vehicle", "vehicle"),
            ("VIN", "vin"),
            ("Damage", "damage_description"),
            ("Facts of Loss", "facts_of_loss"),
            ("Total Loss", "total_loss"),
            ("Shop", "shop_name"),
            ("Contact Phone", "contact_phone"),
            ("Contact Email", "contact_email"),
            ("Claim Notes", "assignment_claim_notes"),
            ("Status", "status"),
            ("Updated", "updated_at"),
        ]
        for idx, (label, key) in enumerate(detail_rows):
            ttk.Label(details, text=label).grid(row=idx, column=0, sticky="nw", padx=(0, 10), pady=3)
            value_entry = ttk.Entry(details, textvariable=self.detail_vars[key], state="normal")
            value_entry.grid(row=idx, column=1, sticky="ew", pady=3)

        ttk.Label(details, text="Folder").grid(row=len(detail_rows), column=0, sticky="nw", padx=(0, 10), pady=(8, 3))
        folder_entry = ttk.Entry(details, textvariable=self.detail_vars["source_path"], state="readonly")
        folder_entry.grid(row=len(detail_rows), column=1, sticky="ew", pady=(8, 3))
        ttk.Button(details, text="Save Details", command=self.save_claim_details).grid(
            row=len(detail_rows) + 1,
            column=0,
            columnspan=2,
            sticky="ew",
            pady=(8, 0),
        )

        notes_frame = ttk.LabelFrame(side, text="Notes", padding=8)
        notes_frame.grid(row=2, column=0, sticky="nsew")
        notes_frame.columnconfigure(0, weight=1)
        notes_frame.rowconfigure(1, weight=1)
        notes_frame.rowconfigure(3, weight=1)

        ttk.Label(notes_frame, text="New Note").grid(row=0, column=0, sticky="w")
        self.new_note_text = tk.Text(notes_frame, height=4, wrap="word")
        self.new_note_text.grid(row=1, column=0, sticky="nsew", pady=(4, 8))

        ttk.Label(notes_frame, text="Saved Notes").grid(row=2, column=0, sticky="w")
        self.notes_history_text = tk.Text(notes_frame, height=6, wrap="word", state="disabled")
        self.notes_history_text.grid(row=3, column=0, sticky="nsew", pady=(4, 8))

        actions = ttk.Frame(notes_frame)
        actions.grid(row=4, column=0, sticky="ew")
        actions.columnconfigure((0, 1, 2, 3), weight=1)
        ttk.Button(actions, text="Add Note", command=self.save_selected_claim).grid(row=0, column=0, sticky="ew", padx=(0, 8))
        ttk.Button(actions, text="Reopen", command=self.reopen_selected_claim).grid(row=0, column=1, sticky="ew", padx=(0, 8))
        ttk.Button(actions, text="Reopen As Supp", command=self.reopen_selected_claim_as_supplement).grid(row=0, column=2, sticky="ew", padx=(0, 8))
        ttk.Button(actions, text="Mark Closed", command=lambda: self.set_selected_status("Closed")).grid(row=0, column=3, sticky="ew")

    def _summary_card(self, parent: ttk.Frame, label: str, value_var: tk.StringVar, column: int) -> None:
        card = ttk.LabelFrame(parent, text=label, padding=16)
        card.grid(row=0, column=column, sticky="ew", padx=(0, 10) if column < 2 else 0)
        ttk.Label(card, textvariable=value_var, font=("Segoe UI", 18, "bold")).grid(row=0, column=0, sticky="w")

    def _configure_tree_sorting(self, tree: ttk.Treeview, columns: tuple[str, ...] | list[str]) -> None:
        for col in columns:
            tree.heading(col, command=lambda c=col, t=tree: self._sort_treeview_column(t, c))

    def _parse_sort_value(self, column: str, value: object) -> tuple[int, object]:
        text = str(value or "").strip()
        if not text:
            return (2, "")

        if column in {"claim_id", "order"} and text.isdigit():
            return (0, int(text))

        normalized_date = text
        if column in {"date_of_loss", "closed_date", "inspection_when", "updated_at"}:
            for fmt in (DATE_FORMAT, "%Y-%m-%d", "%m/%d/%y", "%m/%d/%Y", "%m-%d-%y", "%m-%d-%Y"):
                try:
                    return (0, datetime.strptime(normalized_date, fmt))
                except ValueError:
                    continue

        numeric_text = text.replace(",", "")
        if re.fullmatch(r"-?\d+(?:\.\d+)?", numeric_text):
            try:
                return (0, float(numeric_text))
            except ValueError:
                pass

        return (1, text.lower())

    def _sort_treeview_column(self, tree: ttk.Treeview, column: str, descending: bool | None = None) -> None:
        tree_id = str(tree)
        current = self.tree_sort_state.get(tree_id)
        if descending is None:
            descending = current[0] == column and not current[1] if current else False
        self.tree_sort_state[tree_id] = (column, descending)

        items = [(self._parse_sort_value(column, tree.set(item, column)), item) for item in tree.get_children("")]
        items.sort(key=lambda pair: pair[0], reverse=descending)

        for index, (_, item) in enumerate(items):
            tree.move(item, "", index)

        for col in tree["columns"]:
            title = col.replace("_", " ").title()
            if col == column:
                title = f"{title} {'▼' if descending else '▲'}"
            tree.heading(col, text=title)
        self._configure_tree_sorting(tree, tree["columns"])

    def _apply_saved_tree_sort(self, tree: ttk.Treeview) -> None:
        state = self.tree_sort_state.get(str(tree))
        if state:
            self._sort_treeview_column(tree, state[0], state[1])
        else:
            self._configure_tree_sorting(tree, tree["columns"])

    def _build_tree(self, parent: ttk.Frame) -> ttk.Treeview:
        columns = (
            "claim_id",
            "customer_name",
            "insurance_company",
            "claim_number",
            "town",
            "date_of_loss",
            "status",
            "type",
            "updated_at",
        )
        tree = ttk.Treeview(parent, columns=columns, show="headings", selectmode="browse")
        for col, width in [
            ("claim_id", 120),
            ("customer_name", 180),
            ("insurance_company", 170),
            ("claim_number", 140),
            ("town", 110),
            ("date_of_loss", 100),
            ("status", 90),
            ("type", 90),
            ("updated_at", 140),
        ]:
            tree.heading(col, text=col.replace("_", " ").title())
            tree.column(col, width=width, anchor="w")
        self._configure_tree_sorting(tree, columns)
        tree.grid(row=0, column=0, sticky="nsew")
        scrollbar = ttk.Scrollbar(parent, orient="vertical", command=tree.yview)
        scrollbar.grid(row=0, column=1, sticky="ns")
        tree.configure(yscrollcommand=scrollbar.set)
        tree.bind("<<TreeviewSelect>>", self.on_claim_selected)
        tree.bind("<ButtonRelease-1>", self.on_claim_selected)
        tree.bind("<Double-1>", self.open_selected_claim_folder)
        tree.bind("<Button-3>", self.show_tree_context_menu)
        return tree

    def add_folder(self) -> None:
        folder = filedialog.askdirectory(title="Select a folder to watch")
        if not folder:
            return
        if self.store.add_watched_folder(folder):
            self.scan_folders()
        else:
            messagebox.showinfo("Already Added", "That folder is already being watched.")

    def remove_selected_folder(self) -> None:
        selection = self.folder_list.curselection()
        if not selection:
            return
        folder = self.folder_list.get(selection[0])
        self.store.remove_watched_folder(folder)
        self.scan_folders()

    def add_manual_claim(self) -> None:
        claim_id = simpledialog.askstring("New Claim", "Claim ID")
        if not claim_id:
            return
        normalized_claim_id = self._extract_job_number(claim_id.strip()) or claim_id.strip()
        title = simpledialog.askstring("New Claim", "Claim title (optional)") or normalized_claim_id
        pending_folder = next((folder for folder in self.store.watched_folders if "pending" in folder.lower()), "")
        folder_name = title.strip()
        if pending_folder:
            if normalized_claim_id not in folder_name:
                folder_name = f"{normalized_claim_id} - {folder_name}"
            claim_path = self._unique_destination(Path(pending_folder) / folder_name)
            claim_path.mkdir(parents=True, exist_ok=True)
            key = f"fs::{claim_path}"
            source_path = str(claim_path)
            record_title = claim_path.name
            claim_type = "Original"
        else:
            key = f"manual::{normalized_claim_id}"
            source_path = "Manual Entry"
            record_title = title.strip()
            claim_type = "Manual"

        existing = self.store.get_claim(key)
        record = ClaimRecord(
            key=key,
            claim_id=normalized_claim_id,
            title=record_title,
            status=existing.status if existing else "Open",
            source_path=source_path,
            claim_type=claim_type,
            updated_at=now_stamp(),
            notes=existing.notes if existing else "",
            note_history=existing.note_history if existing else [],
        )
        self.store.upsert_claim(record)
        self.store.save()
        self.scan_folders()
        self.refresh_views(select_key=key)

    def scan_folders(self, show_access_warnings: bool = True) -> None:
        existing_claims = self.store.claims.copy()
        discovered: dict[str, dict] = {}
        discovered_records: list[ClaimRecord] = []
        self.scan_access_issues = []

        for folder in self.store.watched_folders:
            folder_path = Path(folder)
            if not folder_path.exists():
                continue
            default_status = self._default_status_for_folder(folder_path)
            for item in self._claim_entries_for_folder(folder_path):
                key = f"fs::{item.resolve()}"
                previous = existing_claims.get(key, {})
                discovered[key] = asdict(
                    ClaimRecord(
                        key=key,
                        claim_id=self._claim_id_for_item(item),
                        title=item.name,
                        status=previous.get("status", default_status),
                        source_path=str(item.resolve()),
                        claim_type="Original",
                        updated_at=datetime.fromtimestamp(item.stat().st_mtime).strftime(DATE_FORMAT),
                        notes=previous.get("notes", ""),
                        note_history=previous.get("note_history", legacy_notes_to_history(previous.get("notes", ""))),
                        customer_name=previous.get("customer_name", ""),
                        insurance_company=previous.get("insurance_company", ""),
                        claim_number=previous.get("claim_number", ""),
                        policy_number=previous.get("policy_number", ""),
                        insured_name=previous.get("insured_name", ""),
                        claimant_name=previous.get("claimant_name", ""),
                        date_of_loss=previous.get("date_of_loss", ""),
                        town=previous.get("town", ""),
                        owner_address=previous.get("owner_address", ""),
                        location_of_vehicle=previous.get("location_of_vehicle", ""),
                        vehicle=previous.get("vehicle", ""),
                        vin=previous.get("vin", ""),
                        damage_description=previous.get("damage_description", ""),
                        facts_of_loss=previous.get("facts_of_loss", ""),
                        total_loss=previous.get("total_loss", False),
                        shop_name=previous.get("shop_name", ""),
                        contact_phone=previous.get("contact_phone", ""),
                        contact_email=previous.get("contact_email", ""),
                        assign_pdf_path=previous.get("assign_pdf_path", ""),
                        assignment_claim_notes=previous.get("assignment_claim_notes", ""),
                        office_contacted=previous.get("office_contacted", "No"),
                        office_appt_when=previous.get("office_appt_when", ""),
                        office_progress_status=previous.get("office_progress_status", ""),
                        office_waiting_for_paperwork=previous.get("office_waiting_for_paperwork", "No"),
                        office_additional_notes=previous.get("office_additional_notes", ""),
                        route_address_override=previous.get("route_address_override", ""),
                        manual_overrides=previous.get("manual_overrides", {}),
                    )
                )
                assign_data = self._extract_claim_details(item)
                if assign_data:
                    discovered[key].update(assign_data)
                    old_item = item
                    renamed_item = self._rename_plain_claim_folder_if_needed(item, discovered[key])
                    if renamed_item != item:
                        key = f"fs::{renamed_item.resolve()}"
                        discovered[key] = discovered.pop(f"fs::{item.resolve()}")
                        discovered[key]["key"] = key
                        discovered[key]["title"] = renamed_item.name
                        discovered[key]["source_path"] = str(renamed_item.resolve())
                        discovered[key]["updated_at"] = datetime.fromtimestamp(renamed_item.stat().st_mtime).strftime(DATE_FORMAT)
                        assign_path = str(discovered[key].get("assign_pdf_path", ""))
                        if assign_path:
                            try:
                                relative_assign = Path(assign_path).relative_to(old_item)
                                discovered[key]["assign_pdf_path"] = str((renamed_item / relative_assign).resolve())
                            except Exception:
                                if Path(assign_path).name.lower() == "assign.pdf":
                                    discovered[key]["assign_pdf_path"] = str((renamed_item / "assign.pdf").resolve())
                self._apply_manual_overrides(discovered[key], previous)
                discovered_records.append(ClaimRecord(**discovered[key]))

        for key, item in existing_claims.items():
            if key.startswith("manual::"):
                discovered[key] = item
                discovered_records.append(ClaimRecord(**item))

        enriched = self._enrich_records_from_matches(discovered_records)
        self._classify_records(enriched)
        self._normalize_record_fields(enriched)
        self.store.data["claims"] = {record.key: asdict(record) for record in enriched}

        self.store.save()
        self.refresh_claim_tools_contacts_view()
        self.refresh_views(select_key=self.selected_key)
        if show_access_warnings and self.scan_access_issues:
            issue_text = "\n".join(f"- {path}" for path in self.scan_access_issues[:10])
            if len(self.scan_access_issues) > 10:
                issue_text += f"\n- and {len(self.scan_access_issues) - 10} more"
            messagebox.showwarning(
                "Folder In Use",
                "Close this folder or window before opening the app again or running Refresh Scan:\n\n"
                f"{issue_text}",
            )

    def _schedule_next_apptrak_check(self, initial_delay_ms: int = APPTRAK_CHECK_INTERVAL_MS) -> None:
        if self.apptrak_after_id:
            try:
                self.root.after_cancel(self.apptrak_after_id)
            except Exception:
                pass
        self.apptrak_after_id = self.root.after(initial_delay_ms, self._run_scheduled_apptrak_import)

    def _run_scheduled_apptrak_import(self) -> None:
        self.apptrak_after_id = None
        self._run_apptrak_import_cycle(manual=False)

    def run_apptrak_import_now(self) -> None:
        self._run_apptrak_import_cycle(manual=True)

    def _run_apptrak_import_cycle(self, manual: bool) -> None:
        if self.apptrak_import_running:
            if manual:
                messagebox.showinfo("AppTrak Busy", "AppTrak import is already running.")
            return

        self.apptrak_import_running = True
        self.apptrak_status_var.set("Running AppTrak import..." if manual else "Checking AppTrak docs...")
        processed_snapshot = set(self.store.processed_apptrak_pdfs)

        def worker() -> None:
            result: dict[str, object] = {
                "manual": manual,
                "error": None,
                "staged": [],
                "import_ran": False,
            }
            run_started_at = datetime.now()
            try:
                if manual and APPTRAK_IMPORT_BAT.exists():
                    self._ensure_apptrak_running()
                    subprocess.run(
                        ["cmd", "/c", str(APPTRAK_IMPORT_BAT)],
                        cwd=str(APPTRAK_AUTOMATION_DIR),
                        check=True,
                        capture_output=True,
                        text=True,
                        timeout=300,
                    )
                    result["import_ran"] = True
                result["staged"] = self._stage_new_apptrak_pdfs(processed_snapshot, run_started_at if result["import_ran"] else None)
            except Exception as exc:
                result["error"] = str(exc)
            self.root.after(0, lambda: self._finish_apptrak_import_cycle(result))

        threading.Thread(target=worker, daemon=True).start()

    def _ensure_apptrak_running(self) -> None:
        if not APPTRAK_EXE.exists():
            raise RuntimeError(f"AppTrak executable was not found:\n{APPTRAK_EXE}")
        if self._is_apptrak_running():
            return
        try:
            subprocess.Popen([str(APPTRAK_EXE)], cwd=str(APPTRAK_EXE.parent))
        except Exception as exc:
            raise RuntimeError(f"Could not launch AppTrak:\n{exc}") from exc
        time.sleep(8)

    def _is_apptrak_running(self) -> bool:
        try:
            result = subprocess.run(
                ["tasklist", "/FI", "IMAGENAME eq apptrak.exe"],
                check=False,
                capture_output=True,
                text=True,
                timeout=15,
            )
        except Exception:
            return False
        return "apptrak.exe" in result.stdout.lower()

    def _finish_apptrak_import_cycle(self, result: dict[str, object]) -> None:
        self.apptrak_import_running = False
        self._schedule_next_apptrak_check()

        manual = bool(result.get("manual"))
        error = str(result.get("error") or "").strip()
        staged = list(result.get("staged") or [])

        if error:
            self.apptrak_status_var.set("AppTrak import failed.")
            if manual:
                messagebox.showerror("AppTrak Import Failed", error)
            return

        if staged:
            processed = set(self.store.processed_apptrak_pdfs)
            processed.update(item["pdf_name"] for item in staged if item.get("pdf_name"))
            self.store.data["processed_apptrak_pdfs"] = sorted(processed)
            self.store.save()
            self.scan_folders(show_access_warnings=False)
            folder_lines = "\n".join(f"- {item['folder_name']} ({item['pdf_name']})" for item in staged)
            self.apptrak_status_var.set(f"AppTrak staged {len(staged)} new import(s).")
            messagebox.showinfo(
                "AppTrak Imports Ready",
                f"Staged {len(staged)} new import(s):\n\n{folder_lines}",
            )
            return

        self.apptrak_status_var.set("AppTrak checked. No new numeric PDFs found.")
        if manual:
            messagebox.showinfo("AppTrak Imports", "No new numeric PDFs were found in AppTrak docs.")

    def _stage_new_apptrak_pdfs(self, processed_snapshot: set[str], imported_since: datetime | None = None) -> list[dict[str, str]]:
        pending_root = self._pending_claims_root()
        if not pending_root:
            raise RuntimeError("Pending Claims folder is not configured.")
        if not APPTRAK_DOCS_DIR.exists():
            raise RuntimeError(f"AppTrak docs folder was not found:\n{APPTRAK_DOCS_DIR}")

        pending_root.mkdir(parents=True, exist_ok=True)
        next_folder_number = self._next_pending_staging_number(pending_root)
        staged: list[dict[str, str]] = []
        pdf_paths = sorted(
            [path for path in APPTRAK_DOCS_DIR.glob("*.pdf") if re.fullmatch(r"\d+", path.stem)],
            key=lambda path: path.name.lower(),
        )
        for pdf_path in pdf_paths:
            if pdf_path.name in processed_snapshot:
                continue
            if imported_since is not None:
                try:
                    modified_at = datetime.fromtimestamp(pdf_path.stat().st_mtime)
                except OSError:
                    continue
                if modified_at < (imported_since - timedelta(seconds=10)):
                    continue
            target_folder = pending_root / str(next_folder_number)
            while target_folder.exists():
                next_folder_number += 1
                target_folder = pending_root / str(next_folder_number)
            target_folder.mkdir(parents=True, exist_ok=False)
            shutil.copy2(pdf_path, target_folder / pdf_path.name)
            staged.append(
                {
                    "pdf_name": pdf_path.name,
                    "folder_name": target_folder.name,
                    "folder_path": str(target_folder),
                }
            )
            next_folder_number += 1
        return staged

    def _pending_claims_root(self) -> Path | None:
        for folder in self.store.watched_folders:
            if "pending" in folder.lower():
                return Path(folder)
        return None

    def _next_pending_staging_number(self, pending_root: Path) -> int:
        numbers = []
        try:
            for item in pending_root.iterdir():
                if item.is_dir() and re.fullmatch(r"\d+", item.name):
                    numbers.append(int(item.name))
        except OSError:
            pass
        return (max(numbers) + 1) if numbers else 1

    def _rename_plain_claim_folder_if_needed(self, item: Path, record_data: dict[str, object]) -> Path:
        if not item.is_dir():
            return item
        folder_token = item.name.strip()
        if not re.fullmatch(r"\d+", folder_token):
            return item

        owner_name = str(record_data.get("customer_name") or record_data.get("claimant_name") or record_data.get("insured_name") or "").strip()
        if not owner_name or owner_name == "-" or owner_name.lower() == "unknown":
            return item

        cleaned_name = re.sub(r'[<>:"/\\|?*]', "", owner_name).strip()
        cleaned_name = " ".join(cleaned_name.split())
        if not cleaned_name:
            return item

        claim_identifier = str(record_data.get("claim_id") or "").strip()
        prefix = claim_identifier if re.fullmatch(r"\d{8}", claim_identifier) else folder_token
        target = item.with_name(f"{prefix} - {cleaned_name}")
        if target == item:
            return item
        target = self._unique_destination(target)
        try:
            item.rename(target)
        except Exception:
            return item
        return target

    def _claim_entries_for_folder(self, folder_path: Path) -> list[Path]:
        name = folder_path.name.strip().lower()
        if "closed" in name:
            return self._find_closed_claim_folders(folder_path)
        try:
            entries = [
                entry
                for entry in folder_path.iterdir()
                if not (entry.is_file() and entry.name.lower() in {"desktop.ini", "thumbs.db"})
            ]
        except (PermissionError, OSError):
            self.scan_access_issues.append(str(folder_path))
            return []
        return sorted(entries, key=lambda entry: entry.name.lower())

    def _find_closed_claim_folders(self, root: Path) -> list[Path]:
        claim_folders: list[Path] = []
        try:
            month_folders = [item for item in root.iterdir() if item.is_dir()]
        except (PermissionError, OSError):
            self.scan_access_issues.append(str(root))
            return claim_folders

        for month_folder in sorted(month_folders, key=lambda entry: entry.name.lower()):
            try:
                day_folders = [item for item in month_folder.iterdir() if item.is_dir() and item.name.lower().startswith("claims ")]
            except (PermissionError, OSError):
                self.scan_access_issues.append(str(month_folder))
                continue

            for day_folder in sorted(day_folders, key=lambda entry: entry.name.lower()):
                try:
                    entries = [item for item in day_folder.iterdir() if item.is_dir()]
                except (PermissionError, OSError):
                    self.scan_access_issues.append(str(day_folder))
                    continue
                claim_folders.extend(sorted(entries, key=lambda entry: entry.name.lower()))
        return claim_folders

    def _claim_id_for_item(self, item: Path) -> str:
        source = item.stem if item.is_file() else item.name
        return self._extract_job_number(source) or source

    def _extract_job_number(self, text: str) -> str:
        return extract_job_number_from_text(text)

    def _apply_manual_overrides(self, target: dict, previous: dict | None) -> None:
        if not previous:
            return
        overrides = previous.get("manual_overrides", {}) or {}
        if not isinstance(overrides, dict):
            return
        for field_name, value in overrides.items():
            target[field_name] = value

    def _extract_claim_details(self, claim_item: Path) -> dict[str, str]:
        if not claim_item.is_dir() or PdfReader is None:
            return {}

        pdf_paths = self._candidate_pdf_paths(claim_item)
        for pdf_path in pdf_paths:
            details = self._extract_details_from_pdf(pdf_path)
            if details:
                return details
        return {}

    def _candidate_pdf_paths(self, claim_item: Path) -> list[Path]:
        pdfs = sorted(claim_item.glob("*.pdf"), key=lambda path: path.name.lower())
        if not pdfs:
            return []
        is_closed_claim = "closed claims" in str(claim_item).lower()

        def rank(path: Path) -> tuple[int, str]:
            name = path.name.lower()
            if is_closed_claim and "claim summary" in name:
                return (0, name)
            if name == "assign.pdf":
                return (1, name)
            if "assign" in name:
                return (2, name)
            if "claim summary" in name:
                return (3, name)
            if "estimate" in name:
                return (4, name)
            return (5, name)

        return sorted(pdfs, key=rank)

    def _extract_details_from_pdf(self, pdf_path: Path) -> dict[str, str]:
        if PdfReader is None:
            return {}

        try:
            reader = PdfReader(str(pdf_path))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:5])
        except Exception:
            return {"assign_pdf_path": str(pdf_path)}

        if "Claim Summary" in text and "Owner:" in text:
            return self._parse_claim_summary_pdf_text(text, pdf_path)
        if "Assignment Sheet - Duhamel & Duhamel, LLC" in text:
            return self._parse_assignment_pdf_text(text, pdf_path)
        if all(fragment in text.lower() for fragment in ("instructions to estimator", "facts of loss", "vehicle information")):
            return self._parse_email_thread_assignment_pdf_text(text, pdf_path)
        return self._parse_generic_pdf_text(text, pdf_path)

    def _parse_email_thread_assignment_pdf_text(self, text: str, pdf_path: Path) -> dict[str, str]:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        normalized_text = "\n".join(lines)
        details = {"assign_pdf_path": str(pdf_path)}

        file_number_match = re.search(r"\b(\d{5,8})\.pdf\b", pdf_path.name, re.IGNORECASE)
        if file_number_match:
            details["claim_id"] = file_number_match.group(1)

        claim_number_match = re.search(r"\bClaim\s*:\s*([A-Z0-9 \-]+)", normalized_text, re.IGNORECASE)
        if claim_number_match:
            details["claim_number"] = " ".join(claim_number_match.group(1).split())

        date_of_loss_match = re.search(r"Date Of Loss\s*:\s*([0-9/:-]+\s+[AP]M|[0-9/:-]+)", normalized_text, re.IGNORECASE)
        if date_of_loss_match:
            raw_dol = " ".join(date_of_loss_match.group(1).split())
            date_only_match = re.match(r"(\d{2}/\d{2}/\d{4})", raw_dol)
            if date_only_match:
                details["date_of_loss"] = date_only_match.group(1)
            else:
                details["date_of_loss"] = raw_dol

        from_matches = re.findall(r"\bFrom\s*:\s*([^\n]+)", normalized_text, re.IGNORECASE)
        for insurer in from_matches:
            cleaned_insurer = insurer.strip(" ,")
            upper_insurer = cleaned_insurer.upper()
            if any(blocked in upper_insurer for blocked in ("DO_NOT_REPLY", "ASSIGNMENT")):
                continue
            details["insurance_company"] = cleaned_insurer
            break

        instructions_match = re.search(
            r"Instructions\s+to\s+Estimator\s*:\s*(.*?)Facts\s+of\s+Loss\s*:",
            text,
            re.IGNORECASE | re.DOTALL,
        )
        instructions_block = " ".join(instructions_match.group(1).split()) if instructions_match else ""
        if instructions_block:
            details["assignment_claim_notes"] = instructions_block
            damage_inline_match = re.search(r"Damage Description\s*:\s*(.+)", instructions_block, re.IGNORECASE)
            if damage_inline_match:
                details["damage_description"] = " ".join(damage_inline_match.group(1).split())

        facts_match = re.search(
            r"Facts\s+of\s+Loss\s*:\s*(.*?)Vehicle\s+Location",
            text,
            re.IGNORECASE | re.DOTALL,
        )
        facts_block = " ".join(facts_match.group(1).split()) if facts_match else ""
        if facts_block:
            details["facts_of_loss"] = facts_block

        vehicle_location_name_match = re.search(r"Vehicle Location Name\s*:\s*([^\n]+)", normalized_text, re.IGNORECASE)
        vehicle_location_match = re.search(r"Vehicle Location\s*:\s*([^\n]+)", normalized_text, re.IGNORECASE)
        vehicle_location_phone_match = re.search(r"Vehicle Location[\s\S]{0,200}?Phone\s*:\s*\(?(\d{3})\)?[-.\s]*(\d{3})[-.\s]*(\d{4})", normalized_text, re.IGNORECASE)
        if vehicle_location_name_match:
            details["shop_name"] = vehicle_location_name_match.group(1).strip()
        if vehicle_location_match:
            details["location_of_vehicle"] = " ".join(vehicle_location_match.group(1).split())
            if not details.get("shop_name"):
                details["shop_name"] = details["location_of_vehicle"]
        if vehicle_location_phone_match and not details.get("contact_phone"):
            details["contact_phone"] = f"{vehicle_location_phone_match.group(1)}-{vehicle_location_phone_match.group(2)}-{vehicle_location_phone_match.group(3)}"

        owner_block = self._extract_labeled_section(
            normalized_text,
            ["Owner", "Claimant"],
            ["Insured", "Vehicle Information"],
        )
        owner_info = self._parse_contact_section(owner_block)
        if owner_info["name"]:
            details["customer_name"] = owner_info["name"]
            details["claimant_name"] = owner_info["name"]
        if owner_info["address"]:
            details["owner_address"] = owner_info["address"]
        if owner_info["town"]:
            details["town"] = owner_info["town"]
        if owner_info["phone"]:
            details["contact_phone"] = owner_info["phone"]
        if owner_info["email"]:
            details["contact_email"] = owner_info["email"]

        insured_block = self._extract_labeled_section(
            normalized_text,
            ["Insured"],
            ["Vehicle Information"],
        )
        insured_info = self._parse_contact_section(insured_block)
        if insured_info["name"]:
            details["insured_name"] = insured_info["name"]
            if not details.get("customer_name"):
                details["customer_name"] = insured_info["name"]
        if not details.get("owner_address") and insured_info["address"]:
            details["owner_address"] = insured_info["address"]
        if not details.get("town") and insured_info["town"]:
            details["town"] = insured_info["town"]
        if not details.get("contact_phone") and insured_info["phone"]:
            details["contact_phone"] = insured_info["phone"]
        if not details.get("contact_email") and insured_info["email"]:
            details["contact_email"] = insured_info["email"]

        vehicle_block = self._extract_labeled_section(
            normalized_text,
            ["Vehicle Information"],
            ["Notice", "$"],
        )
        if vehicle_block:
            vin_match = re.search(r"VIN\s*:\s*([A-HJ-NPR-Z0-9]{17})", vehicle_block, re.IGNORECASE)
            if vin_match:
                details["vin"] = vin_match.group(1).strip()

            vehicle_match = re.search(r"Vehicle\s*:\s*([^\n]+)", vehicle_block, re.IGNORECASE)
            if vehicle_match:
                details["vehicle"] = " ".join(vehicle_match.group(1).split())

            plate_match = re.search(r"License Plate\s*:\s*([A-Z0-9-]+)", vehicle_block, re.IGNORECASE)
            if plate_match:
                details["assignment_claim_notes"] = details.get("assignment_claim_notes", "")

            if not details.get("damage_description"):
                poi_match = re.search(r"Primary Point of Impact\s*:\s*([^\n]+)", vehicle_block, re.IGNORECASE)
                if poi_match:
                    details["damage_description"] = poi_match.group(1).strip()

            drivable_match = re.search(r"Drivable\s*:\s*([YN])", vehicle_block, re.IGNORECASE)
            if drivable_match:
                prefix = details.get("assignment_claim_notes", "")
                drivable_text = "Drivable: Yes" if drivable_match.group(1).upper() == "Y" else "Drivable: No"
                if drivable_text.lower() not in prefix.lower():
                    details["assignment_claim_notes"] = (prefix + "\n" + drivable_text).strip() if prefix else drivable_text

            total_loss_match = re.search(r"Total Loss Status\s*:\s*([^\n]+)", vehicle_block, re.IGNORECASE)
            if total_loss_match and total_loss_match.group(1).strip():
                details["total_loss"] = "total" in total_loss_match.group(1).lower()

        if re.search(r"\bpossible total\b|\btotal loss\b", normalized_text, re.IGNORECASE):
            details["total_loss"] = True

        return details

    def _extract_labeled_section(self, text: str, headings: list[str], next_headings: list[str]) -> str:
        start_pattern = "|".join(re.escape(item) for item in headings)
        if "$" in next_headings:
            end_pattern = "$"
        else:
            end_pattern = "|".join(re.escape(item) for item in next_headings)
        pattern = rf"(?:{start_pattern})\s*:?\s*(.*?)(?=(?:{end_pattern})\s*:|$)"
        match = re.search(pattern, text, re.IGNORECASE | re.DOTALL)
        return match.group(1).strip() if match else ""

    def _parse_contact_section(self, section_text: str) -> dict[str, str]:
        info = {"name": "", "address": "", "town": "", "phone": "", "email": ""}
        if not section_text:
            return info

        lines = [line.strip() for line in section_text.splitlines() if line.strip()]
        if not lines:
            return info

        name_match = re.search(r"\b(?:Owner|Claimant|Insured)\s*:\s*([^\n]+)", section_text, re.IGNORECASE)
        if name_match:
            name_candidate = name_match.group(1).strip()
        else:
            name_candidate = re.sub(r"^\((Claimant|Insured|Owner)\)\s*", "", lines[0], flags=re.IGNORECASE).strip()
        if not re.search(r"[@:]", name_candidate) and not re.search(r"\b(phone|cell|email|address)\b", name_candidate, re.IGNORECASE):
            info["name"] = name_candidate.title() if name_candidate.isupper() else name_candidate

        email_match = re.search(r"\bEmail\s*:\s*([A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,})\b", section_text, re.IGNORECASE)
        if not email_match:
            email_match = re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", section_text, re.IGNORECASE)
        if email_match:
            info["email"] = email_match.group(1) if email_match.lastindex else email_match.group(0)

        phone_match = re.search(r"\bPhone\s*:\s*\(?(\d{3})\)?[-.\s]*(\d{3})[-.\s]*(\d{4})", section_text, re.IGNORECASE)
        if not phone_match:
            phone_match = re.search(r"\(?(\d{3})\)?[-.\s]*(\d{3})[-.\s]*(\d{4})", section_text)
        if phone_match:
            info["phone"] = f"{phone_match.group(1)}-{phone_match.group(2)}-{phone_match.group(3)}"

        address_match = re.search(r"\bAddress\s*:\s*([^\n]+)", section_text, re.IGNORECASE)
        if address_match:
            info["address"] = " ".join(address_match.group(1).split())
        else:
            address_lines: list[str] = []
            for line in lines[1:] if info["name"] else lines:
                if info["email"] and info["email"] in line:
                    continue
                if re.search(r"\b(phone|cell|work|home|email|owner:|claimant:|insured:|address:)\b", line, re.IGNORECASE):
                    continue
                address_lines.append(line.title() if line.isupper() else line)
            if address_lines:
                info["address"] = " ".join(address_lines).strip()

        if info["address"]:
            town_match = re.search(r"\b([A-Z][A-Za-z .'-]+),\s*([A-Z]{2}),?\s*(\d{5}(?:-\d{4})?)\b", info["address"])
            if not town_match:
                town_match = re.search(r"\b([A-Z][A-Za-z .'-]+),\s*([A-Z]{2})\s+(\d{5}(?:-\d{4})?)\b", info["address"])
            if town_match:
                info["town"] = town_match.group(1).strip().title()

        return info

    def _parse_claim_summary_pdf_text(self, text: str, pdf_path: Path) -> dict[str, str]:
        details = {"assign_pdf_path": str(pdf_path)}

        owner_match = re.search(r"Owner:\s*([^\n\r]+)", text, re.IGNORECASE)
        if owner_match:
            details["customer_name"] = owner_match.group(1).strip()

        insured_match = re.search(r"Insured:\s*([^\n\r]+)", text, re.IGNORECASE)
        if insured_match:
            details["insured_name"] = insured_match.group(1).strip()

        insurer_match = re.search(r"Company:\s*([^\n\r]+?)\s+Claim #:", text, re.IGNORECASE)
        if insurer_match:
            details["insurance_company"] = insurer_match.group(1).strip()

        claim_match = re.search(r"Claim #:\s*([A-Z0-9 \-_]+)", text, re.IGNORECASE)
        if claim_match:
            details["claim_number"] = claim_match.group(1).strip()

        vin_match = re.search(r"VIN:\s*([A-HJ-NPR-Z0-9]{17})", text, re.IGNORECASE)
        if vin_match:
            details["vin"] = vin_match.group(1)

        vehicle_match = re.search(r"Owner:\s*[^\n\r]+\s*\n([^\n\r]+)", text, re.IGNORECASE)
        if vehicle_match:
            details["vehicle"] = vehicle_match.group(1).strip()

        town_match = re.search(r"\n([A-Z][A-Za-z .'-]+),\s*([A-Z]{2})\s+(\d{5})", text)
        if town_match:
            details["town"] = town_match.group(1).strip()

        phone_match = re.search(r"Owner:\s*[^\n\r]+\s+Cell:\s*\((\d{3})\)\s*(\d{3})-(\d{4})", text, re.IGNORECASE)
        if phone_match:
            details["contact_phone"] = f"{phone_match.group(1)}-{phone_match.group(2)}-{phone_match.group(3)}"

        shop_match = re.search(r"Place of Inspection:\s*([^\n\r]+)", text, re.IGNORECASE)
        if shop_match:
            details["shop_name"] = shop_match.group(1).strip()

        if re.search(r"Settlement Type:\s*Total Loss", text, re.IGNORECASE) or re.search(r"Total loss threshold reached", text, re.IGNORECASE):
            details["total_loss"] = True

        return details

    def _parse_assignment_pdf_text(self, text: str, pdf_path: Path) -> dict[str, str]:
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        insurance_company = ""
        claim_number = ""
        policy_number = ""
        customer_name = ""
        insured_name = ""
        claimant_name = ""
        date_of_loss = ""
        town = ""
        owner_address = ""
        location_of_vehicle = ""
        vehicle = ""
        vin = ""
        damage_description = ""
        facts_of_loss = ""
        shop_name = ""
        contact_phone = ""
        contact_email = ""
        assignment_claim_notes = ""

        def clean_line(value: str) -> str:
            value = value.strip()
            if not value or re.fullmatch(r"[-\s/@]+", value):
                return ""
            if value.upper() in {"UNKNOWN", "NOT AVAILABLE", "NOT READABLE"}:
                return ""
            return value

        def extract_name_from_owner_line(value: str) -> str:
            cleaned = clean_line(re.sub(r"^\((Claimant|Insured)\)\s*", "", value, flags=re.IGNORECASE))
            if not cleaned:
                return ""
            if re.search(r"\b(vehicle owner|vehicle information|home ph|work ph|pager|cell|ext)\b", cleaned, re.IGNORECASE):
                return ""
            if re.search(r"\b(not readable|manual|unknown|type|color|engine|trans|plate|style|mileage|damage|make|model|vin)\b", cleaned, re.IGNORECASE):
                return ""
            if re.search(r"\d", cleaned):
                return ""
            if not re.fullmatch(r"[A-Za-z][A-Za-z0-9&.,'() /-]*", cleaned):
                return ""
            return cleaned

        try:
            insurance_idx = lines.index("Insurance Information")
            if insurance_idx >= 4:
                insurance_company = clean_line(lines[insurance_idx - 7]) if insurance_idx >= 7 else ""
                insured_name = clean_line(lines[insurance_idx - 6]) if insurance_idx >= 6 else ""
                claim_number = clean_line(lines[insurance_idx - 5]) if insurance_idx >= 5 else ""
                policy_number = clean_line(lines[insurance_idx - 4]) if insurance_idx >= 4 and lines[insurance_idx - 4] != "Not Available" else ""
        except ValueError:
            pass

        date_match = re.search(r"(\d{2}/\d{2}/\d{2})\s+Date of Loss:", text)
        if date_match:
            date_of_loss = date_match.group(1)

        location_match = re.search(r"Location:\s*([^\n\r]+)", text, re.IGNORECASE)
        if location_match:
            location_of_vehicle = clean_line(location_match.group(1))

        try:
            owner_idx = lines.index("Vehicle Owner")
            owner_block = lines[max(0, owner_idx - 8):owner_idx]
            filtered = [line for line in owner_block if clean_line(line)]
            preferred_owner_names = [extract_name_from_owner_line(line) for line in filtered]
            preferred_owner_names = [line for line in preferred_owner_names if line]
            if preferred_owner_names:
                customer_name = preferred_owner_names[0]
            address_lines: list[str] = []
            for line in filtered:
                town_match = re.search(r"\b([A-Z][A-Z ]+)\s+([A-Z]{2})\s+(\d{5})\b", line)
                if town_match:
                    town = town_match.group(1).title()
                    owner_address = " ".join(part for part in [*address_lines, f"{town_match.group(1).title()}, {town_match.group(2)} {town_match.group(3)}"] if part).strip()
                    break
                if not preferred_owner_names or line != preferred_owner_names[0]:
                    cleaned_line = clean_line(line)
                    if not cleaned_line:
                        continue
                    if "not readable" in cleaned_line.lower():
                        continue
                    if re.search(r"\b(home ph|work ph|pager|cell|ext|vehicle owner|vehicle information)\b", cleaned_line, re.IGNORECASE):
                        continue
                    address_lines.append(cleaned_line.title() if cleaned_line.isupper() else cleaned_line)
            if not vehicle:
                for line in filtered:
                    vehicle_candidate = clean_line(line)
                    if re.match(r"^\d{2}\s+[A-Z]", vehicle_candidate, re.IGNORECASE):
                        vehicle = vehicle_candidate.title() if vehicle_candidate.isupper() else vehicle_candidate
                        break
        except ValueError:
            pass

        owner_role_match = re.search(r"\((Claimant|Insured)\)\s+([A-Z][A-Z .'\-]+)$", text, re.IGNORECASE | re.MULTILINE)
        if owner_role_match:
            owner_role = owner_role_match.group(1).lower()
            owner_name = clean_line(owner_role_match.group(2).title() if owner_role_match.group(2).isupper() else owner_role_match.group(2))
            customer_name = customer_name or owner_name
            if owner_role == "claimant":
                claimant_name = owner_name
            else:
                insured_name = owner_name
        elif "(Claimant)" in text:
            claimant_name = customer_name
        elif "(Insured)" in text:
            insured_name = insured_name or customer_name

        vin_match = re.search(r"\b[A-HJ-NPR-Z0-9]{17}\b", text)
        if vin_match:
            vin = vin_match.group(0)

        if vin:
            for index, line in enumerate(lines):
                if line == vin and index > 0:
                    vehicle_line = lines[index - 1]
                    if not vehicle_line.endswith(":"):
                        vehicle = vehicle_line
                    break

        damage_match = re.search(r"Damage:\s*([^\n\r]+)", text, re.IGNORECASE)
        if damage_match:
            damage_description = clean_line(damage_match.group(1))

        facts_match = re.search(r"Facts of Loss:\s*([^\n\r]+)", text, re.IGNORECASE)
        if facts_match:
            facts_of_loss = clean_line(facts_match.group(1))

        shop_match = re.search(r"\(\s*[A-Z0-9]+\s*\)\s*\n([^\n]+)\n([^\n]+)\n([^\n]+)\n([\d-]+)", text)
        if shop_match:
            shop_name = clean_line(shop_match.group(1))
        if location_of_vehicle and not shop_name:
            shop_name = location_of_vehicle

        appraisal_idx = next((idx for idx, value in enumerate(lines) if value.lower().startswith("appraisal notes:")), -1)
        claimant_context = lines[max(0, appraisal_idx - 14):min(len(lines), appraisal_idx + 10)] if appraisal_idx >= 0 else []
        if claimant_context:
            claimant_phones: list[str] = []
            for line in claimant_context:
                for match in re.findall(r"\b\d{3}-\d{3}-\d{4}\b", line):
                    if match not in {"203-792-2150", "203-791-8066", "888-345-7797", "888-884-0885"}:
                        claimant_phones.append(match)
                for groups in re.findall(r"\(?(\d{3})\)?\s*(\d{3})-(\d{4})", line):
                    phone = f"{groups[0]}-{groups[1]}-{groups[2]}"
                    if phone not in {"203-792-2150", "203-791-8066", "888-345-7797", "888-884-0885"}:
                        claimant_phones.append(phone)
            if claimant_phones:
                contact_phone = claimant_phones[-1]
            for line in claimant_context:
                email_match = re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", line, re.IGNORECASE)
                if email_match:
                    email_value = email_match.group(0)
                    if all(blocked not in email_value.lower() for blocked in ("claimsolution", "duhamel", "apptrak", "autobid", "claims@innovativeclaims")):
                        contact_email = email_value
                        break

            if not vehicle:
                make_match = re.search(r"Make:\s*([A-Z0-9]+)", text, re.IGNORECASE)
                model_match = re.search(r"Model:\s*([A-Z0-9][A-Z0-9 \-]+)", text, re.IGNORECASE)
                if make_match and model_match:
                    vehicle = f"{make_match.group(1)} {model_match.group(1)}".strip()

        note_lines: list[str] = []
        note_prefixes = ("instructions to estimator:", "appraisal notes:")
        stop_prefixes = (
            "ext:",
            "location:",
            "fax:",
            "pager:",
            "cell:",
            "registration expiration:",
            "inspection number:",
            "estimate by:",
            "amount:",
        )
        for idx, raw_line in enumerate(lines):
            lowered = raw_line.lower()
            if not lowered.startswith(note_prefixes):
                continue
            for follow_line in lines[idx:]:
                lowered_follow = follow_line.lower()
                if note_lines and lowered_follow.startswith(stop_prefixes):
                    break
                if "app trak" in lowered_follow or re.search(r"page \d+ of \d+", lowered_follow):
                    break
                if lowered_follow.startswith(note_prefixes):
                    cleaned_follow = follow_line.strip()
                else:
                    cleaned_follow = clean_line(follow_line)
                if not cleaned_follow:
                    continue
                note_lines.append(cleaned_follow)
            if note_lines:
                assignment_claim_notes = " ".join(note_lines).strip()
                break
        if appraisal_idx >= 0:
            pre_notes: list[str] = []
            for line in reversed(lines[max(0, appraisal_idx - 6):appraisal_idx]):
                lowered = line.lower()
                if lowered.startswith(("type:", "color:", "engine:", "trans:", "estimate by:", "amount:", "(claimant)", "phone:", "vehicle information", "vehicle owner", "facts of loss:")):
                    break
                cleaned_pre = clean_line(line)
                if not cleaned_pre:
                    continue
                if re.search(r"\b\d{3}-\d{3}-\d{4}\b", cleaned_pre):
                    continue
                pre_notes.append(cleaned_pre)
            pre_notes.reverse()
            if pre_notes:
                combined = " ".join(pre_notes).strip()
                assignment_claim_notes = f"{combined} {assignment_claim_notes}".strip() if assignment_claim_notes else combined

        return {
            "assign_pdf_path": str(pdf_path),
            "customer_name": customer_name,
            "insurance_company": insurance_company,
            "claim_number": claim_number,
            "policy_number": policy_number,
            "insured_name": insured_name,
            "claimant_name": claimant_name,
            "date_of_loss": date_of_loss,
            "town": town,
            "owner_address": owner_address,
            "location_of_vehicle": location_of_vehicle,
            "vehicle": vehicle,
            "vin": vin,
            "damage_description": damage_description,
            "facts_of_loss": facts_of_loss,
            "shop_name": shop_name,
            "contact_phone": contact_phone,
            "contact_email": contact_email,
            "assignment_claim_notes": assignment_claim_notes,
        }

    def _parse_generic_pdf_text(self, text: str, pdf_path: Path) -> dict[str, str]:
        details = {"assign_pdf_path": str(pdf_path)}

        claim_match = re.search(r"Claim Number\s+([A-Z0-9-]+)", text, re.IGNORECASE)
        if claim_match:
            details["claim_number"] = claim_match.group(1)

        owner_match = re.search(r"Owner\s+([^\n]+)", text, re.IGNORECASE)
        if owner_match:
            details["customer_name"] = owner_match.group(1).strip()

        town_match = re.search(r"\n[A-Z0-9 .'-]+\n([A-Z][A-Z ]+)\s+[A-Z]{2}\s+\d{5}\b", text)
        if town_match:
            details["town"] = town_match.group(1).title()

        insurer_match = re.search(r"\n([A-Z][A-Za-z &]+Insurance[^\n]*)\n", text)
        if not insurer_match:
            insurer_match = re.search(r"\n([A-Z][A-Za-z &]+Mutual[^\n]*)\n", text)
        if insurer_match:
            details["insurance_company"] = insurer_match.group(1).strip()

        vin_match = re.search(r"\b[A-HJ-NPR-Z0-9]{17}\b", text)
        if vin_match:
            details["vin"] = vin_match.group(0)

        vehicle_match = re.search(r"\b(20\d{2}\s+[A-Z][A-Za-z0-9]+\s+[A-Z0-9][A-Za-z0-9 ]+)\b", text)
        if vehicle_match:
            details["vehicle"] = vehicle_match.group(1).strip()

        phone_match = re.search(r"Owner\s+[^\n]+\s+\((\d{3})\)\s*(\d{3})-(\d{4})", text, re.IGNORECASE)
        if phone_match:
            details["contact_phone"] = f"{phone_match.group(1)}-{phone_match.group(2)}-{phone_match.group(3)}"

        email_match = re.search(r"\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b", text, re.IGNORECASE)
        if email_match:
            details["contact_email"] = email_match.group(0)

        estimate_shop_match = re.search(
            r"(?:MY WAY AUTO BODY|[A-Z][A-Za-z0-9&'.,\- ]+(?:AUTO BODY|BODY SHOP|COLLISION|COLLISION CENTER|MOTORS|SERVICE|REPAIR SERVICE|REPAIR))",
            text,
            re.IGNORECASE,
        )
        if estimate_shop_match:
            details["shop_name"] = estimate_shop_match.group(0).strip()

        return details

    def _enrich_records_from_matches(self, records: list[ClaimRecord]) -> list[ClaimRecord]:
        by_claim_id: dict[str, list[ClaimRecord]] = {}
        by_claim_number: dict[str, list[ClaimRecord]] = {}
        by_name: dict[str, list[ClaimRecord]] = {}

        for record in records:
            by_claim_id.setdefault(record.claim_id.strip().lower(), []).append(record)
            if record.claim_number:
                by_claim_number.setdefault(record.claim_number.strip().lower(), []).append(record)
            for name in {self._normalized_name(record.customer_name), self._normalized_name(record.title)}:
                if name:
                    by_name.setdefault(name, []).append(record)

        enriched: list[ClaimRecord] = []
        for record in records:
            donors = []
            donors.extend(by_claim_id.get(record.claim_id.strip().lower(), []))
            if record.claim_number:
                donors.extend(by_claim_number.get(record.claim_number.strip().lower(), []))
            donors.extend(by_name.get(self._normalized_name(record.customer_name), []))
            donors.extend(by_name.get(self._normalized_name(record.title), []))

            best = self._best_donor(record, donors)
            if best:
                for field in [
                    "customer_name",
                    "insurance_company",
                    "claim_number",
                    "policy_number",
                    "insured_name",
                    "claimant_name",
                    "date_of_loss",
                    "town",
                    "vehicle",
                    "vin",
                    "shop_name",
                    "contact_phone",
                    "contact_email",
                ]:
                    if not getattr(record, field) and getattr(best, field):
                        setattr(record, field, getattr(best, field))
                if (not record.assign_pdf_path or Path(record.assign_pdf_path).name.lower() != "assign.pdf") and best.assign_pdf_path:
                    record.assign_pdf_path = best.assign_pdf_path
            enriched.append(record)
        return enriched

    def _best_donor(self, record: ClaimRecord, donors: list[ClaimRecord]) -> ClaimRecord | None:
        def score(candidate: ClaimRecord) -> tuple[int, int, int]:
            if candidate.key == record.key:
                return (-1, -1, -1)
            completeness = sum(
                1 for field in [
                    candidate.customer_name,
                    candidate.insurance_company,
                    candidate.claim_number,
                    candidate.policy_number,
                    candidate.insured_name,
                    candidate.date_of_loss,
                    candidate.town,
                    candidate.vehicle,
                    candidate.vin,
                ] if field
            )
            same_claim_id = int(candidate.claim_id.strip().lower() == record.claim_id.strip().lower())
            same_claim_number = int(bool(record.claim_number) and candidate.claim_number.strip().lower() == record.claim_number.strip().lower())
            return (same_claim_id + same_claim_number, completeness, int(candidate.status == "Closed"))

        valid = [candidate for candidate in donors if candidate.key != record.key]
        if not valid:
            return None
        best = max(valid, key=score)
        return best if score(best)[0] > 0 or self._normalized_name(best.customer_name) == self._normalized_name(record.customer_name) or self._normalized_name(best.title) == self._normalized_name(record.title) else None

    def _normalized_name(self, value: str) -> str:
        cleaned = re.sub(r"[^a-z0-9]+", " ", value.lower()).strip()
        tokens = [token for token in cleaned.split() if token not in {"supp", "supplement", "inreview", "ready4review", "waiting4review", "closed", "claims"}]
        return " ".join(tokens)

    def _classify_records(self, records: list[ClaimRecord]) -> None:
        groups: dict[str, list[ClaimRecord]] = {}
        for record in records:
            if record.claim_id:
                groups.setdefault(record.claim_id, []).append(record)

        for record in records:
            if record.source_path == "Manual Entry":
                record.claim_type = "Manual"
                continue

            siblings = groups.get(record.claim_id, [])
            title_lower = record.title.lower()
            if len(siblings) <= 1:
                record.claim_type = "Original"
                continue

            explicit_originals = [item for item in siblings if "supp" not in item.title.lower()]
            explicit_supplements = [item for item in siblings if "supp" in item.title.lower()]

            if "supp" in title_lower:
                record.claim_type = "Supplement"
                continue

            if explicit_originals:
                first_original = min(explicit_originals, key=self._record_sort_key)
                record.claim_type = "Original" if record.key == first_original.key else "Supplement"
                continue

            first_record = min(siblings, key=self._record_sort_key)
            record.claim_type = "Original" if record.key == first_record.key else "Supplement"

    def _normalize_record_fields(self, records: list[ClaimRecord]) -> None:
        for record in records:
            preferred_name = self._preferred_customer_name(record)
            if preferred_name:
                record.customer_name = preferred_name
            if self._is_bad_name(record.claimant_name):
                record.claimant_name = ""

    def _preferred_customer_name(self, record: ClaimRecord) -> str:
        for candidate in (
            record.customer_name,
            record.claimant_name,
            record.insured_name,
            self._customer_name_from_title(record.title),
        ):
            if candidate and not self._is_bad_name(candidate):
                return candidate.strip()
        return ""

    def _is_bad_name(self, value: str) -> bool:
        cleaned = (value or "").strip()
        if not cleaned:
            return True
        lowered = cleaned.lower()
        return any(
            fragment in lowered
            for fragment in (
                "unknown",
                "not available",
                "not readable",
                "manual",
                "vehicle location and damage",
                "app trak",
                "all rights reserved",
                "page 1 of",
                "page 2 of",
            )
        )

    def _customer_name_from_title(self, title: str) -> str:
        cleaned = re.sub(r"^\d{8}\s*-\s*", "", title).strip()
        cleaned = re.sub(r"\s*\([^)]*\)\s*$", "", cleaned).strip()
        if " - " in cleaned:
            cleaned = cleaned.split(" - ", 1)[0].strip()
        return cleaned

    def _normalize_insurer_key(self, insurer_name: str) -> str:
        cleaned = re.sub(r"[^A-Z0-9 ]+", " ", (insurer_name or "").upper())
        tokens = [
            token for token in cleaned.split()
            if token not in {"INSURANCE", "ASSURANCE", "COMPANY", "CO", "GROUP", "INC", "INC.", "CORP", "CORPORATION"}
        ]
        return " ".join(tokens).strip()

    def _display_insurer_name(self, insurer_name: str) -> str:
        cleaned = " ".join((insurer_name or "").replace("_", " ").split()).strip()
        if not cleaned:
            return "Unknown Insurer"
        return cleaned.title() if cleaned.isupper() else cleaned

    def _existing_instruction_keys(self) -> dict[str, str]:
        existing: dict[str, str] = {}
        if not CLIENT_INSTRUCTIONS_DIR.exists():
            return existing
        for path in CLIENT_INSTRUCTIONS_DIR.iterdir():
            if path.suffix.lower() not in {".docx", ".pdf"}:
                continue
            insurer_name = path.stem.replace(" - Client Instructions", "")
            key = self._normalize_insurer_key(insurer_name)
            if key:
                existing[key] = insurer_name
        return existing

    def _ensure_client_instruction_sheets(self, records: list[ClaimRecord]) -> list[str]:
        if Document is None:
            return []
        CLIENT_INSTRUCTIONS_DIR.mkdir(parents=True, exist_ok=True)
        existing_keys = self._existing_instruction_keys()
        insurer_records: dict[str, list[ClaimRecord]] = {}
        display_names: dict[str, str] = {}

        for record in records:
            insurer = (record.insurance_company or "").strip()
            if not insurer:
                continue
            key = self._normalize_insurer_key(insurer)
            if not key or key in existing_keys:
                continue
            insurer_records.setdefault(key, []).append(record)
            display_names.setdefault(key, self._display_insurer_name(insurer))

        generated: list[str] = []
        for key, insurer_group in insurer_records.items():
            sections = self._build_client_instruction_sections(insurer_group)
            if not sections:
                continue
            display_name = display_names[key]
            base_name = f"{display_name} - Client Instructions"
            docx_path = CLIENT_INSTRUCTIONS_DIR / f"{base_name}.docx"
            pdf_path = CLIENT_INSTRUCTIONS_DIR / f"{base_name}.pdf"
            self._write_instruction_docx(display_name, sections, docx_path)
            try:
                self._export_docx_to_pdf(docx_path, pdf_path)
            except Exception:
                pass
            generated.append(display_name)
        return generated

    def _build_client_instruction_sections(self, insurer_records: list[ClaimRecord]) -> dict[str, list[str]]:
        aggregated: dict[str, list[str]] = {}
        for record in insurer_records:
            claim_path = Path(record.source_path)
            if not claim_path.exists() or not claim_path.is_dir():
                continue
            text = self._extract_assignment_sheet_text(claim_path)
            if not text:
                continue
            parsed = self._parse_client_instruction_sections(text)
            for heading, bullets in parsed.items():
                aggregated.setdefault(heading, [])
                for bullet in bullets:
                    if bullet not in aggregated[heading]:
                        aggregated[heading].append(bullet)
        return aggregated

    def _extract_assignment_sheet_text(self, claim_path: Path) -> str:
        if PdfReader is None:
            return ""
        preferred = []
        for path in sorted(claim_path.glob("*.pdf"), key=lambda item: item.name.lower()):
            name = path.name.lower()
            if name == "assign.pdf" or "assign" in name:
                preferred.append(path)
        for pdf_path in preferred:
            try:
                text = "\n".join((page.extract_text() or "") for page in PdfReader(str(pdf_path)).pages[:5])
            except Exception:
                continue
            if "Assignment Sheet - Duhamel & Duhamel, LLC" in text:
                return text
        return ""

    def _parse_client_instruction_sections(self, text: str) -> dict[str, list[str]]:
        section_specs = [
            ("Required Photos", ["PHOTOS"]),
            ("General Rules", ["GENERAL RULES"]),
            ("Labor Rates", ["LABOR RATES"]),
            ("Parts Usage", ["PARTS USAGE", "PART USAGE"]),
            ("Agreed Price", ["AGREED PRICE"]),
            ("Betterment", ["BETTERMENT"]),
            ("Repair Time", ["DAYS TO REPAR", "DAYS TO REPAIR"]),
            ("Total Loss Procedures", ["TOTAL LOSSES", "TOTAL LOSS PROCEDURE", "TOTAL LOSSES:"]),
            ("Specific Client Notes", ["SPECIFIC CLIENT NOTES"]),
            ("Supervisor", ["SUPERVISOR", "SUPERVISORS"]),
        ]
        stop_markers = {
            "VEHICLE LOCATION AND DAMAGE",
            "ESTIMATE BY:",
            "AMOUNT:",
            "FACTS OF LOSS:",
            "APPRAISAL NOTES:",
            "INSPECTION NUMBER:",
            "CLIENT INSTRUCTIONS:",
            "COMPANY:",
            "ATTENTION:",
            "PHONE: FAX:",
            "CONTACT PHONE: E-MAIL:",
            "CONTACT PHONE:",
            "E-MAIL:",
        }

        lines = [line.strip() for line in text.splitlines()]
        sections: dict[str, list[str]] = {}
        current_heading: str | None = None

        def is_heading(line: str) -> tuple[str, str] | None:
            upper = line.upper().rstrip(":")
            for heading, keywords in section_specs:
                for keyword in keywords:
                    if upper == keyword.rstrip(":"):
                        return heading, keyword
            return None

        for raw_line in lines:
            line = raw_line.strip()
            if not line:
                continue
            upper = line.upper()
            if (
                upper in stop_markers
                or upper.startswith("(CLAIMANT)")
                or upper.startswith("(INSURED)")
                or upper.startswith("(")
                or re.fullmatch(r"[A-Z]+,\s*[A-Z][A-Z ]+", upper)
            ):
                current_heading = None
                continue

            heading_match = is_heading(line)
            if heading_match:
                current_heading = heading_match[0]
                sections.setdefault(current_heading, [])
                continue

            if not current_heading:
                continue

            cleaned = self._clean_instruction_line(line)
            if not cleaned:
                continue
            if cleaned not in sections[current_heading]:
                sections[current_heading].append(cleaned)

        return {heading: bullets for heading, bullets in sections.items() if bullets}

    def _clean_instruction_line(self, line: str) -> str:
        cleaned = re.sub(r"\s+", " ", line).strip(" -;")
        if not cleaned or len(cleaned) < 6:
            return ""
        upper = cleaned.upper()
        bad_fragments = [
            "APP TRAK",
            "ASSIGNMENT SHEET",
            "FILE #:",
            "APPRAISER:",
            "INSURED:",
            "OWNER:",
            "VIN:",
            "PLATE:",
            "DAMAGE:",
            "MAKE:",
            "MODEL:",
            "CONFIRMED AT SHOP",
            "INSTRUCTIONS TO ESTIMATOR",
            "FACTS OF LOSS",
            "VEHICLE HAS BEEN",
            "PLEASE REACH OUT TO",
            "WORKFILE",
            "ESTIMATE OF RECORD",
            "CLAIM SUMMARY FILE CREATED",
            "IMAGE WORKFILE CREATED",
            "ESTIMATE REPORT FILE CREATED",
            "DATE INSPECTED:",
            "LOCATION OF VEHICLE?",
            "TOW BILL AMOUNT:",
            "STORAGE RATE:",
            "CONTACTED OWNER POST INSP",
            "LOT#:",
            "UNRELATED DAMAGE?",
            "ADDITIONAL TL COMMENTS:",
            "PRIMARY IMPACT POINT:",
            "SECONDARY IMPACT POINT:",
            "PLACE OF INSPECTION:",
            "ADDRESS:",
        ]
        if any(fragment in upper for fragment in bad_fragments):
            return ""
        if re.fullmatch(r"[A-Z]+,\s*[A-Z][A-Z ]+", upper):
            return ""
        if re.fullmatch(r"\(?\s*[A-Z0-9]{3,}\s*\)?", cleaned):
            return ""
        if re.fullmatch(r"[A-Z0-9\-() ]{2,}", cleaned) and len(cleaned.split()) <= 3:
            return ""
        if re.search(r"\b\d{1,4}\s+\w+", cleaned) and any(state in upper for state in (" CT ", " NY ", " MA ", " OH ", " VA ")):
            return ""
        return cleaned

    def _write_instruction_docx(self, insurer_name: str, sections: dict[str, list[str]], path: Path) -> None:
        document = Document()
        document.add_heading(insurer_name.upper(), level=0)
        document.add_paragraph("Client Instructions")
        document.add_paragraph("Organized reference sheet for field/client instructions")
        document.add_paragraph("Cleaned and organized from the assignment sheet's client instructions only.")
        for heading, bullets in sections.items():
            document.add_heading(heading, level=1)
            for bullet in bullets:
                document.add_paragraph(bullet, style="List Bullet")
        document.save(path)

    def _export_docx_to_pdf(self, docx_path: Path, pdf_path: Path) -> None:
        script = (
            "$word = New-Object -ComObject Word.Application; "
            "$word.Visible = $false; "
            "$word.DisplayAlerts = 0; "
            f"$doc = $word.Documents.Open('{str(docx_path)}'); "
            f"$doc.SaveAs([ref]'{str(pdf_path)}', [ref]17); "
            "$doc.Close(); "
            "$word.Quit();"
        )
        subprocess.run(
            ["powershell", "-NoProfile", "-Command", script],
            check=True,
            capture_output=True,
            text=True,
        )

    def _record_sort_key(self, record: ClaimRecord) -> tuple[str, str]:
        closed_date = self._closed_date_for_record(record)
        date_key = closed_date.strftime("%Y-%m-%d") if closed_date else "9999-12-31"
        return (date_key, record.updated_at or "9999-12-31 23:59", record.title.lower())

    def _default_status_for_folder(self, folder_path: Path) -> str:
        name = folder_path.name.strip().lower()
        if "closed" in name:
            return "Closed"
        if "pending" in name or "open" in name:
            return "Open"
        return "Open"

    def _closed_date_for_record(self, record: ClaimRecord) -> datetime | None:
        if record.status != "Closed":
            return None
        path = record.source_path.replace("/", "\\")
        match = re.search(r"\\Claims (\d{1,2})-(\d{1,2})-(\d{2})\\", path)
        if not match:
            return None
        month, day, year = match.groups()
        try:
            return datetime.strptime(f"{month}-{day}-{year}", "%m-%d-%y")
        except ValueError:
            return None

    def _update_closed_stats(self, records: list[ClaimRecord]) -> None:
        closed_dates = [value for value in (self._closed_date_for_record(record) for record in records) if value]
        now = datetime.now()
        today = now.date()
        iso_year, iso_week, _ = now.isocalendar()

        today_count = sum(1 for value in closed_dates if value.date() == today)
        week_count = sum(1 for value in closed_dates if value.isocalendar()[:2] == (iso_year, iso_week))
        month_count = sum(1 for value in closed_dates if value.year == now.year and value.month == now.month)

        self.closed_stats_vars["today"].set(str(today_count))
        self.closed_stats_vars["week"].set(str(week_count))
        self.closed_stats_vars["month"].set(str(month_count))

        date_counts: dict[str, int] = {}
        for value in closed_dates:
            label = value.strftime("%Y-%m-%d")
            date_counts[label] = date_counts.get(label, 0) + 1

        options = sorted(date_counts.keys(), reverse=True)
        self.closed_date_combo["values"] = options
        self.range_start_combo["values"] = options
        self.range_end_combo["values"] = options
        current = self.closed_date_var.get()
        if options:
            if current not in date_counts:
                current = options[0]
                self.closed_date_var.set(current)
            self.closed_stats_vars["selected_date"].set(current)
            self.closed_stats_vars["selected_count"].set(str(date_counts[current]))
        else:
            self.closed_date_var.set("")
            self.closed_stats_vars["selected_date"].set("-")
            self.closed_stats_vars["selected_count"].set("0")

        self._refresh_reports_claims(records)

    def on_closed_date_selected(self, _event: object) -> None:
        selected = self.closed_date_var.get()
        if not selected:
            self.closed_stats_vars["selected_count"].set("0")
            self._refresh_reports_claims([ClaimRecord(**self.store._normalize_claim_item(item)) for item in self.store.claims.values()])
            return
        self.range_start_var.set("")
        self.range_end_var.set("")
        records = [ClaimRecord(**self.store._normalize_claim_item(item)) for item in self.store.claims.values()]
        count = sum(
            1
            for record in records
            if (closed_date := self._closed_date_for_record(record)) and closed_date.strftime("%Y-%m-%d") == selected
        )
        self.closed_stats_vars["selected_date"].set(selected)
        self.closed_stats_vars["selected_count"].set(str(count))
        self._refresh_reports_claims(records)

    def _refresh_reports_claims(self, records: list[ClaimRecord]) -> None:
        filtered = self._filtered_report_records(records)
        self.reports_tree.delete(*self.reports_tree.get_children())
        for record in filtered:
            closed_date = self._closed_date_for_record(record)
            self.reports_tree.insert(
                "",
                tk.END,
                iid=record.key,
                values=(
                    closed_date.strftime("%Y-%m-%d") if closed_date else "",
                    record.claim_id,
                    record.customer_name or record.title,
                    record.insurance_company,
                    record.claim_type,
                ),
            )
        self._apply_saved_tree_sort(self.reports_tree)

    def _filtered_report_records(self, records: list[ClaimRecord]) -> list[ClaimRecord]:
        filtered = [record for record in records if record.status == "Closed" and self._closed_date_for_record(record)]
        start = self._parse_report_date(self.range_start_var.get().strip())
        end = self._parse_report_date(self.range_end_var.get().strip())
        if start or end:
            if start:
                filtered = [record for record in filtered if self._closed_date_for_record(record).date() >= start.date()]
            if end:
                filtered = [record for record in filtered if self._closed_date_for_record(record).date() <= end.date()]
        else:
            selected_date = self.closed_date_var.get().strip()
            if selected_date:
                filtered = [
                    record for record in filtered
                    if self._closed_date_for_record(record).strftime("%Y-%m-%d") == selected_date
                ]

        return sorted(filtered, key=lambda record: (self._closed_date_for_record(record), record.claim_id))

    def _parse_report_date(self, value: str) -> datetime | None:
        if not value:
            return None
        normalized = value.strip()
        for fmt in ("%Y-%m-%d", "%m/%d/%y", "%m/%d/%Y", "%m-%d-%y", "%m-%d-%Y"):
            try:
                return datetime.strptime(normalized, fmt)
            except ValueError:
                continue
        return None

    def normalize_report_date_field(self, variable: tk.StringVar) -> None:
        parsed = self._parse_report_date(variable.get().strip())
        if parsed:
            variable.set(parsed.strftime("%Y-%m-%d"))

    def apply_reports_range_filter(self) -> None:
        start = self.range_start_var.get().strip()
        end = self.range_end_var.get().strip()
        self.normalize_report_date_field(self.range_start_var)
        self.normalize_report_date_field(self.range_end_var)
        start = self.range_start_var.get().strip()
        end = self.range_end_var.get().strip()
        if start and not self._parse_report_date(start):
            messagebox.showerror("Invalid Start Date", "Use YYYY-MM-DD or MM/DD/YY for the start date.")
            return
        if end and not self._parse_report_date(end):
            messagebox.showerror("Invalid End Date", "Use YYYY-MM-DD or MM/DD/YY for the end date.")
            return
        if start or end:
            self.closed_date_var.set("")
            self.closed_stats_vars["selected_date"].set("-")
            self.closed_stats_vars["selected_count"].set("0")
        records = [ClaimRecord(**self.store._normalize_claim_item(item)) for item in self.store.claims.values()]
        self._refresh_reports_claims(records)

    def clear_reports_filters(self) -> None:
        self.closed_date_var.set("")
        self.range_start_var.set("")
        self.range_end_var.set("")
        records = [ClaimRecord(**self.store._normalize_claim_item(item)) for item in self.store.claims.values()]
        self._refresh_reports_claims(records)

    def calculate_payroll(self) -> None:
        records = [ClaimRecord(**self.store._normalize_claim_item(item)) for item in self.store.claims.values()]
        if load_workbook is None:
            messagebox.showerror("Excel Support Missing", "openpyxl is not available for payroll export.")
            return
        if not PAYROLL_TEMPLATE.exists():
            messagebox.showerror("Template Missing", f"Payroll template not found:\n{PAYROLL_TEMPLATE}")
            return

        start = self._parse_report_date(self.range_start_var.get().strip())
        end = self._parse_report_date(self.range_end_var.get().strip())
        if not start or not end:
            messagebox.showerror(
                "Missing Date Range",
                "Enter both Start Date and End Date in Reports before calculating payroll.",
            )
            return
        if start.date() > end.date():
            messagebox.showerror("Invalid Date Range", "Start Date must be on or before End Date.")
            return

        filtered = self._filtered_report_records(records)
        if not filtered:
            messagebox.showinfo("No Claims", "No closed claims match the current report filters.")
            return

        wb = load_workbook(PAYROLL_TEMPLATE)
        ws = wb[wb.sheetnames[0]]
        ws["A2"] = "Bill Date:"
        ws["B2"] = f"{start.month}/{start.day}/{start.year} - {end.month}/{end.day}/{end.year}"
        ws["B2"].number_format = "@"
        start_row = 5
        end_row = 46
        max_rows = end_row - start_row + 1

        if len(filtered) > max_rows:
            messagebox.showwarning(
                "Too Many Claims",
                f"The template fits {max_rows} claims. Only the first {max_rows} matching claims will be exported.",
            )

        for row in range(start_row, end_row + 1):
            for col in ("A", "B", "C", "D", "E", "F", "G", "H", "I", "J", "K", "L", "M"):
                ws[f"{col}{row}"] = None
            ws[f"F{row}"] = f"=D{row}+E{row}"
            ws[f"D{row}"].number_format = "0.00"
            ws[f"E{row}"].number_format = "0.00"
            ws[f"F{row}"].number_format = "0.00"

        exported = filtered[:max_rows]
        for idx, record in enumerate(exported, start=start_row):
            ws[f"A{idx}"] = self._payroll_file_number(record)
            ws[f"B{idx}"] = self._payroll_insurance_name(record)
            ws[f"C{idx}"] = self._payroll_owner_name(record)
            ws[f"D{idx}"] = self._payroll_base_amount(record)
            ws[f"E{idx}"] = self._payroll_total_loss_amount(record)
            ws[f"F{idx}"] = f"=D{idx}+E{idx}"
            self._normalize_payroll_row_style(ws, idx)
            ws[f"D{idx}"].number_format = "0.00"
            ws[f"E{idx}"].number_format = "0.00"
            ws[f"F{idx}"].number_format = "0.00"

        total_row = start_row + len(exported)
        if total_row <= end_row:
            ws[f"F{total_row}"] = f"=SUM(F{start_row}:F{total_row - 1})" if exported else "0.00"
            ws[f"F{total_row}"].number_format = "0.00"
            for row in range(total_row + 1, end_row + 1):
                for col in ("A", "B", "C", "D", "E", "F"):
                    ws[f"{col}{row}"] = None

        last_print_row = max(total_row, 4)
        ws.print_area = f"$A$1:$F${last_print_row}"
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 1
        ws.page_margins.left = 0.25
        ws.page_margins.right = 0.25
        ws.page_margins.top = 0.35
        ws.page_margins.bottom = 0.35
        ws.page_margins.header = 0.1
        ws.page_margins.footer = 0.1
        ws.row_breaks.brk = []

        PAYROLL_OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
        base_name = f"Fernando pay {end.month}-{end.day}-{end.strftime('%y')}"
        output_path = PAYROLL_OUTPUT_DIR / f"{base_name}.xlsx"
        pdf_path = PAYROLL_OUTPUT_DIR / f"{base_name}.pdf"
        wb.save(output_path)

        try:
            self._export_workbook_to_pdf(output_path, pdf_path)
            os.startfile(str(pdf_path))
            messagebox.showinfo(
                "Payroll Created",
                f"Saved payroll files:\n{output_path}\n{pdf_path}\n\nThe PDF preview should open now.",
            )
        except Exception as exc:
            messagebox.showerror(
                "PDF Export Failed",
                f"Saved workbook:\n{output_path}\n\nCould not create payroll PDF preview:\n{exc}",
            )

    def _payroll_file_number(self, record: ClaimRecord) -> str:
        return (record.claim_id or record.claim_number or "").strip()

    def _payroll_insurance_name(self, record: ClaimRecord) -> str:
        return (record.insurance_company or "").strip()

    def _payroll_owner_name(self, record: ClaimRecord) -> str:
        name = (record.customer_name or record.title or "").strip()
        if not name:
            return ""
        if "," in name:
            return name
        if name.isupper():
            return name.title()
        return name

    def _payroll_base_amount(self, record: ClaimRecord) -> float:
        return 25.0 if record.claim_type == "Supplement" else 57.0

    def _payroll_total_loss_amount(self, record: ClaimRecord) -> float | None:
        return 5.0 if record.total_loss else None

    def _normalize_payroll_row_style(self, ws, row: int) -> None:
        source_row = 5
        for col in ("A", "B", "C", "D", "E", "F"):
            target = ws[f"{col}{row}"]
            source = ws[f"{col}{source_row}"]
            target.font = copy(source.font)
            target.fill = copy(source.fill)
            target.border = copy(source.border)
            target.alignment = copy(source.alignment)
            target.protection = copy(source.protection)

    def _export_workbook_to_pdf(self, workbook_path: Path, pdf_path: Path) -> None:
        script = (
            "$excel = New-Object -ComObject Excel.Application; "
            "$excel.Visible = $false; "
            "$excel.DisplayAlerts = $false; "
            f"$workbook = $excel.Workbooks.Open('{str(workbook_path)}'); "
            f"$workbook.ExportAsFixedFormat(0, '{str(pdf_path)}'); "
            "$workbook.Close($false); "
            "$excel.Quit();"
        )
        subprocess.run(
            ["powershell", "-NoProfile", "-Command", script],
            check=True,
            capture_output=True,
            text=True,
        )

    def refresh_views(self, select_key: str | None = None) -> None:
        watched = self.store.watched_folders
        self.folder_list.delete(0, tk.END)
        for folder in watched:
            self.folder_list.insert(tk.END, folder)

        records = [ClaimRecord(**self.store._normalize_claim_item(item)) for item in self.store.claims.values()]
        search = self.search_var.get().strip().lower()
        if search:
            records = [
                record
                for record in records
                if search in record.claim_id.lower()
                or search in record.title.lower()
                or search in record.notes.lower()
                or search in record.source_path.lower()
                or search in record.customer_name.lower()
                or search in record.insurance_company.lower()
                or search in record.claim_number.lower()
                or search in record.policy_number.lower()
                or search in record.insured_name.lower()
                or search in record.claimant_name.lower()
                or search in record.date_of_loss.lower()
                or search in record.town.lower()
                or search in record.vehicle.lower()
                or search in record.vin.lower()
                or search in record.shop_name.lower()
                or search in record.contact_phone.lower()
                or search in record.contact_email.lower()
            ]

        all_records = sorted(records, key=lambda record: (record.status != "Open", record.title.lower()))
        filtered = {
            "all": all_records,
            "open": [record for record in all_records if record.status == "Open"],
            "closed": [record for record in all_records if record.status == "Closed"],
        }

        self.summary_vars["all"].set(str(len(all_records)))
        self.summary_vars["open"].set(str(len(filtered["open"])))
        self.summary_vars["closed"].set(str(len(filtered["closed"])))
        self._update_closed_stats(records)
        self._refresh_office_updates(filtered["open"])
        self._refresh_route_planner(filtered["open"])

        for name, tree in self.treeviews.items():
            tree.delete(*tree.get_children())
            for record in filtered[name]:
                tree.insert(
                    "",
                    tk.END,
                    iid=record.key,
                    values=(
                        record.claim_id,
                        record.customer_name or record.title,
                        record.insurance_company,
                        record.claim_number,
                        record.town,
                        record.date_of_loss,
                        record.status,
                        record.claim_type,
                        record.updated_at,
                    ),
                )
            self._apply_saved_tree_sort(tree)

        if select_key and self.store.claims.get(select_key):
            self.selected_key = select_key
            self.populate_details(select_key)
            for tree in self.treeviews.values():
                if tree.exists(select_key):
                    tree.selection_set(select_key)
                    tree.focus(select_key)
                    tree.see(select_key)
        elif self.selected_key and self.store.claims.get(self.selected_key):
            self.populate_details(self.selected_key)
        else:
            self.clear_details()

    def _refresh_office_updates(self, open_records: list[ClaimRecord]) -> None:
        if not self.office_tree:
            return
        self.office_tree.delete(*self.office_tree.get_children())
        for record in open_records:
            self.office_tree.insert(
                "",
                tk.END,
                iid=record.key,
                values=(
                    record.claim_id,
                    record.customer_name or record.title,
                    record.office_progress_status or "No update",
                    record.office_appt_when,
                ),
            )
        self._apply_saved_tree_sort(self.office_tree)
        if self.selected_key and any(record.key == self.selected_key for record in open_records):
            self.office_tree.selection_set(self.selected_key)
            self.office_tree.focus(self.selected_key)
            self.office_tree.see(self.selected_key)
            self.populate_office_update(self.selected_key)
            return
        if open_records:
            self.selected_key = open_records[0].key
            self.office_tree.selection_set(self.selected_key)
            self.office_tree.focus(self.selected_key)
            self.office_tree.see(self.selected_key)
            self.populate_office_update(self.selected_key)
            return
        self.selected_key = None
        self.office_appt_when_var.set("")
        self.office_progress_status_var.set("")
        self.office_waiting_for_paperwork_var.set("No")
        if self.office_notes_text:
            self.office_notes_text.delete("1.0", tk.END)
        self._set_office_summary("")

    def on_office_claim_selected(self, _event: object) -> None:
        if not self.office_tree:
            return
        selection = self.office_tree.selection()
        if not selection:
            return
        self.selected_key = selection[0]
        self.populate_details(self.selected_key)
        self.populate_office_update(self.selected_key)

    def populate_office_update(self, key: str) -> None:
        record = self.store.get_claim(key)
        if not record:
            return
        self.office_appt_when_var.set(record.office_appt_when or "")
        self.office_progress_status_var.set(record.office_progress_status or "")
        self.office_waiting_for_paperwork_var.set(record.office_waiting_for_paperwork or "No")
        if self.office_notes_text:
            self.office_notes_text.delete("1.0", tk.END)
            if record.office_additional_notes:
                self.office_notes_text.insert("1.0", record.office_additional_notes)
        self._set_office_summary(self._build_office_summary(record))

    def save_office_update(self) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select an open claim in Office Update first.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return
        previous_appt_when = record.office_appt_when or ""
        record.office_appt_when = self.office_appt_when_var.get().strip()
        record.office_progress_status = self.office_progress_status_var.get().strip()
        record.office_waiting_for_paperwork = self.office_waiting_for_paperwork_var.get() or "No"
        record.office_additional_notes = self.office_notes_text.get("1.0", tk.END).strip() if self.office_notes_text else ""
        manual_overrides = dict(record.manual_overrides or {})
        manual_overrides["office_appt_when"] = record.office_appt_when
        manual_overrides["office_progress_status"] = record.office_progress_status
        manual_overrides["office_waiting_for_paperwork"] = record.office_waiting_for_paperwork
        manual_overrides["office_additional_notes"] = record.office_additional_notes
        record.manual_overrides = manual_overrides
        record.updated_at = now_stamp()
        self.store.upsert_claim(record)
        self.store.save()
        if record.office_appt_when and record.office_appt_when != previous_appt_when:
            try:
                self._send_office_appointment_calendar_invite(record)
            except Exception as exc:
                messagebox.showwarning("Calendar Reminder Failed", f"Saved the office update, but could not send the calendar reminder:\n{exc}")
        self.populate_office_update(record.key)
        self.refresh_views(select_key=record.key)

    def _build_office_summary(self, record: ClaimRecord) -> str:
        lines = [
            f"{record.claim_id} - {record.customer_name or record.title}",
            f"Insurance: {record.insurance_company or '-'}",
        ]
        if record.office_appt_when:
            lines.append(f"Inspection Date Set Up: {record.office_appt_when}")
        if record.office_progress_status:
            lines.append(f"Progress: {record.office_progress_status}")
        if record.office_waiting_for_paperwork == "Yes":
            lines.append("Waiting for paperwork: Yes")
        if record.claim_type == "Supplement":
            lines.append("Claim Type: Supplement")
        if record.shop_name:
            lines.append(f"Shop: {record.shop_name}")
        if record.office_additional_notes:
            lines.append(f"Additional Notes: {record.office_additional_notes}")
        return "\n".join(lines)

    def _build_combined_office_summary(self, open_records: list[ClaimRecord]) -> str:
        if not open_records:
            return ""
        parts = ["Open Claims Office Update", ""]
        for record in open_records:
            parts.append(self._build_office_summary(record))
            parts.append("")
        return "\n".join(parts).strip()

    def _set_office_summary(self, value: str) -> None:
        self.office_summary_text.delete("1.0", tk.END)
        if value:
            self.office_summary_text.insert("1.0", value)

    def copy_office_summary(self) -> None:
        text = self.office_summary_text.get("1.0", tk.END).strip()
        if not text:
            return
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        messagebox.showinfo("Copied", "Office summary copied to clipboard.")

    def export_office_summary_pdf(self) -> None:
        open_records = self._get_open_records_for_office_update()
        if not open_records:
            messagebox.showinfo("No Summary", "There is no office summary to export.")
            return
        try:
            pdf_path = self._generate_office_update_pdf(open_records)
            self._send_office_update_email(pdf_path)
        except Exception as exc:
            messagebox.showerror("Export Failed", f"Could not create office update PDF:\n{exc}")
            return
        messagebox.showinfo("Office Update Sent", f"Exported and emailed:\n{pdf_path}")

    def preview_office_summary_pdf(self) -> None:
        open_records = self._get_open_records_for_office_update()
        if not open_records:
            messagebox.showinfo("No Summary", "There is no office summary to export.")
            return
        try:
            pdf_path = self._generate_office_update_pdf(open_records)
            os.startfile(str(pdf_path))
        except Exception as exc:
            messagebox.showerror("Preview Failed", f"Could not create office update PDF:\n{exc}")
            return

    def email_uninspected_only_to_self(self) -> None:
        records = self._get_uninspected_records_for_self_copy()
        if not records:
            messagebox.showinfo("No Uninspected Claims", "There are no uninspected open claims to email.")
            return
        try:
            pdf_path = self._generate_detailed_office_update_pdf(records, title="Uninspected Open Claims", base_name_prefix="Uninspected Open Claims")
            self._send_office_update_email_to_self(pdf_path, subject="Uninspected Open Claims")
        except Exception as exc:
            messagebox.showerror("Export Failed", f"Could not create office update PDF:\n{exc}")
            return
        messagebox.showinfo("Office Update Sent", f"Uninspected-only copy emailed to yourself:\n{pdf_path}")

    def _generate_office_update_pdf(self, open_records: list[ClaimRecord]) -> Path:
        if Document is None:
            raise RuntimeError("python-docx is not available.")
        output_dir = APP_DIR / "Office Updates"
        output_dir.mkdir(parents=True, exist_ok=True)
        base_name = f"Office Update {datetime.now().strftime('%Y-%m-%d %H-%M')}"
        docx_path = self._unique_destination(output_dir / f"{base_name}.docx")
        pdf_path = docx_path.with_suffix(".pdf")
        document = Document()
        document.add_heading("Office Update", level=0)
        document.add_paragraph(f"Generated: {now_stamp()}")
        if WD_ORIENT is not None:
            section = document.sections[0]
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
            if Inches is not None:
                section.left_margin = Inches(0.4)
                section.right_margin = Inches(0.4)
                section.top_margin = Inches(0.4)
                section.bottom_margin = Inches(0.4)

        columns = ["Job #", "Customer", "Progress", "Inspection", "Paperwork", "Shop", "Additional Notes"]
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        header_cells = table.rows[0].cells
        for idx, label in enumerate(columns):
            header_cells[idx].text = label

        for record in open_records:
            row = table.add_row().cells
            row[0].text = record.claim_id or "-"
            row[1].text = record.customer_name or record.title or "-"
            row[2].text = record.office_progress_status or "-"
            row[3].text = record.office_appt_when or "-"
            row[4].text = record.office_waiting_for_paperwork or "-"
            row[5].text = record.shop_name or "-"
            row[6].text = record.office_additional_notes or ""
        document.save(docx_path)
        self._export_docx_to_pdf(docx_path, pdf_path)
        return pdf_path

    def _generate_detailed_office_update_pdf(
        self,
        open_records: list[ClaimRecord],
        title: str = "Office Update - Detailed",
        base_name_prefix: str = "Office Update Detailed",
    ) -> Path:
        if Document is None:
            raise RuntimeError("python-docx is not available.")
        output_dir = APP_DIR / "Office Updates"
        output_dir.mkdir(parents=True, exist_ok=True)
        base_name = f"{base_name_prefix} {datetime.now().strftime('%Y-%m-%d %H-%M')}"
        docx_path = self._unique_destination(output_dir / f"{base_name}.docx")
        pdf_path = docx_path.with_suffix(".pdf")
        document = Document()
        document.add_heading(title, level=0)
        document.add_paragraph(f"Generated: {now_stamp()}")
        if WD_ORIENT is not None:
            section = document.sections[0]
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width
            if Inches is not None:
                section.left_margin = Inches(0.3)
                section.right_margin = Inches(0.3)
                section.top_margin = Inches(0.3)
                section.bottom_margin = Inches(0.3)

        columns = [
            "Job #",
            "Customer",
            "Insurance",
            "Address",
            "Phone",
            "Progress",
            "Inspection",
            "Paperwork",
            "Appraisal Notes",
            "Claim Notes",
            "Additional Notes",
        ]
        table = document.add_table(rows=1, cols=len(columns))
        table.style = "Table Grid"
        header_cells = table.rows[0].cells
        for idx, label in enumerate(columns):
            header_cells[idx].text = label

        for record in open_records:
            row = table.add_row().cells
            row[0].text = record.claim_id or "-"
            row[1].text = record.customer_name or record.title or "-"
            row[2].text = record.insurance_company or "-"
            row[3].text = record.owner_address or record.location_of_vehicle or record.town or "-"
            row[4].text = record.contact_phone or "-"
            row[5].text = record.office_progress_status or "-"
            row[6].text = record.office_appt_when or "-"
            row[7].text = record.office_waiting_for_paperwork or "-"
            row[8].text = record.assignment_claim_notes or "-"
            row[9].text = record.notes or "-"
            row[10].text = record.office_additional_notes or ""

        document.save(docx_path)
        self._export_docx_to_pdf(docx_path, pdf_path)
        return pdf_path

    def _send_office_update_email(self, attachment_path: Path) -> None:
        sender_email, app_password = self._get_office_update_email_credentials()
        if not sender_email or not app_password:
            raise RuntimeError("A Gmail address and app password are required for office update email.")

        message = EmailMessage()
        message["From"] = sender_email
        message["To"] = OFFICE_UPDATE_EMAIL_TO
        message["Cc"] = OFFICE_UPDATE_EMAIL_CC
        message["Subject"] = OFFICE_UPDATE_EMAIL_SUBJECT
        message.set_content(OFFICE_UPDATE_EMAIL_BODY)
        message.add_attachment(
            attachment_path.read_bytes(),
            maintype="application",
            subtype="pdf",
            filename=attachment_path.name,
        )

        context = ssl.create_default_context()
        with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as smtp:
            smtp.ehlo()
            smtp.starttls(context=context)
            smtp.ehlo()
            smtp.login(sender_email, app_password)
            smtp.send_message(message)

    def _send_office_update_email_to_self(self, attachment_path: Path, subject: str | None = None) -> None:
        sender_email, app_password = self._get_office_update_email_credentials()
        if not sender_email or not app_password:
            raise RuntimeError("A Gmail address and app password are required for office update email.")

        message = EmailMessage()
        message["From"] = sender_email
        message["To"] = sender_email
        message["Subject"] = subject or f"{OFFICE_UPDATE_EMAIL_SUBJECT} - Copy"
        message.set_content(OFFICE_UPDATE_EMAIL_BODY)
        message.add_attachment(
            attachment_path.read_bytes(),
            maintype="application",
            subtype="pdf",
            filename=attachment_path.name,
        )

        context = ssl.create_default_context()
        with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as smtp:
            smtp.ehlo()
            smtp.starttls(context=context)
            smtp.ehlo()
            smtp.login(sender_email, app_password)
            smtp.send_message(message)

    def _get_uninspected_records_for_self_copy(self) -> list[ClaimRecord]:
        records: list[ClaimRecord] = []
        for record in self._get_open_records_for_office_update():
            progress = (record.office_progress_status or "").strip().lower()
            waiting = (record.office_waiting_for_paperwork or "").strip().lower()
            if waiting == "yes":
                continue
            if progress in {"seen - need to write", "written - under review", "supplement - waiting for paperwork", "waiting for paperwork"}:
                continue
            records.append(record)
        return records

    def _parse_office_appointment_datetime(self, value: str) -> datetime | None:
        normalized = " ".join((value or "").split()).strip()
        if not normalized:
            return None
        for fmt in (
            "%Y-%m-%d %H:%M",
            "%Y-%m-%d %I:%M %p",
            "%m/%d/%y %H:%M",
            "%m/%d/%Y %H:%M",
            "%m/%d/%y %I:%M %p",
            "%m/%d/%Y %I:%M %p",
            "%m-%d-%y %H:%M",
            "%m-%d-%Y %H:%M",
            "%m-%d-%y %I:%M %p",
            "%m-%d-%Y %I:%M %p",
        ):
            try:
                return datetime.strptime(normalized, fmt)
            except ValueError:
                continue
        date_only = self._parse_report_date(normalized)
        if date_only:
            return date_only.replace(hour=10, minute=0)
        return None

    def _build_office_appointment_ics(self, record: ClaimRecord, start_at: datetime) -> str:
        end_at = start_at + timedelta(hours=1)
        uid = f"{record.claim_id}-{start_at.strftime('%Y%m%dT%H%M%S')}@claimsmanager"
        summary = f"Inspection - {record.claim_id} {record.customer_name or record.title}"
        description_lines = [
            f"Claim ID: {record.claim_id or '-'}",
            f"Customer: {record.customer_name or record.title or '-'}",
            f"Insurance: {record.insurance_company or '-'}",
            f"Shop: {record.shop_name or '-'}",
            f"Notes: {record.office_additional_notes or '-'}",
        ]
        location = record.owner_address or record.shop_name or record.town or ""
        summary = summary.replace(",", "\\,").replace(";", "\\;")
        description = "\\n".join(line.replace(",", "\\,").replace(";", "\\;") for line in description_lines)
        location = location.replace(",", "\\,").replace(";", "\\;")
        return "\r\n".join(
            [
                "BEGIN:VCALENDAR",
                "VERSION:2.0",
                "PRODID:-//Claims Manager//Office Update//EN",
                "CALSCALE:GREGORIAN",
                "BEGIN:VEVENT",
                f"UID:{uid}",
                f"DTSTAMP:{datetime.utcnow().strftime('%Y%m%dT%H%M%SZ')}",
                f"DTSTART;TZID=America/New_York:{start_at.strftime('%Y%m%dT%H%M%S')}",
                f"DTEND;TZID=America/New_York:{end_at.strftime('%Y%m%dT%H%M%S')}",
                f"SUMMARY:{summary}",
                f"DESCRIPTION:{description}",
                f"LOCATION:{location}",
                "BEGIN:VALARM",
                "ACTION:DISPLAY",
                "DESCRIPTION:Inspection tomorrow night reminder",
                "TRIGGER:-PT15H",
                "END:VALARM",
                "BEGIN:VALARM",
                "ACTION:DISPLAY",
                "DESCRIPTION:Inspection 9 AM reminder",
                "TRIGGER:-PT1H",
                "END:VALARM",
                "END:VEVENT",
                "END:VCALENDAR",
                "",
            ]
        )

    def _send_office_appointment_calendar_invite(self, record: ClaimRecord) -> None:
        sender_email, app_password = self._get_office_update_email_credentials()
        if not sender_email or not app_password:
            raise RuntimeError("A Gmail address and app password are required for calendar reminders.")
        start_at = self._parse_office_appointment_datetime(record.office_appt_when)
        if not start_at:
            raise RuntimeError("Use a date like YYYY-MM-DD or MM/DD/YY for the inspection date.")

        ics_content = self._build_office_appointment_ics(record, start_at)
        message = EmailMessage()
        message["From"] = sender_email
        message["To"] = sender_email
        message["Subject"] = f"Calendar Reminder - {record.claim_id or 'Inspection'}"
        message.set_content("Calendar invite attached for your inspection reminder.")
        message.add_attachment(
            ics_content.encode("utf-8"),
            maintype="text",
            subtype="calendar",
            filename=f"{record.claim_id or 'inspection'}-appointment.ics",
            params={"method": "REQUEST", "charset": "utf-8"},
        )

        context = ssl.create_default_context()
        with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as smtp:
            smtp.ehlo()
            smtp.starttls(context=context)
            smtp.ehlo()
            smtp.login(sender_email, app_password)
            smtp.send_message(message)

    def _send_route_email(self, route_url: str) -> None:
        sender_email, app_password = self._get_office_update_email_credentials()
        if not sender_email or not app_password:
            raise RuntimeError("A Gmail address and app password are required for route email.")

        message = EmailMessage()
        message["From"] = sender_email
        message["To"] = sender_email
        message["Subject"] = "Tomorrow Route"
        message.set_content(f"Google Maps route link:\n\n{route_url}")

        context = ssl.create_default_context()
        with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as smtp:
            smtp.ehlo()
            smtp.starttls(context=context)
            smtp.ehlo()
            smtp.login(sender_email, app_password)
            smtp.send_message(message)

    def _get_office_update_email_credentials(self) -> tuple[str, str]:
        settings = self.store.email_settings
        sender_email = (settings.get("office_update_from") or OFFICE_UPDATE_EMAIL_FROM).strip()
        app_password = self.store.office_update_app_password

        if not sender_email:
            sender_email = simpledialog.askstring(
                "Gmail Address",
                "Enter the Gmail address to send office updates from:",
                initialvalue=OFFICE_UPDATE_EMAIL_FROM,
                parent=self.root,
            ) or ""
            sender_email = sender_email.strip()
            settings["office_update_from"] = sender_email
            self.store.save()

        if not app_password:
            app_password = simpledialog.askstring(
                "Gmail App Password",
                "Enter your Gmail app password for office update email:",
                parent=self.root,
                show="*",
            ) or ""
            app_password = app_password.strip().replace(" ", "")
            self.store.set_office_update_app_password(app_password)

        return sender_email, app_password

    def _get_open_records_for_office_update(self) -> list[ClaimRecord]:
        records: list[ClaimRecord] = []
        if not self.office_tree:
            return records
        for key in self.office_tree.get_children():
            record = self.store.get_claim(key)
            if record and record.status == "Open":
                records.append(record)
        return records

    def _route_default_address(self, record: ClaimRecord) -> str:
        if (record.route_address_override or "").strip():
            return self._normalize_route_address(record.route_address_override)

        candidates = []
        if record.location_of_vehicle and record.town and record.town.lower() not in record.location_of_vehicle.lower():
            candidates.append(f"{record.location_of_vehicle}, {record.town}, CT")
        if record.shop_name and record.town and record.town.lower() not in record.shop_name.lower():
            candidates.append(f"{record.shop_name}, {record.town}, CT")
        candidates.extend(
            [
                record.location_of_vehicle,
                record.owner_address,
                record.shop_name,
                record.town,
            ]
        )

        for value in candidates:
            cleaned = (value or "").strip()
            if cleaned and cleaned != "-":
                return self._normalize_route_address(cleaned)
        return ""

    def _normalize_route_address(self, address: str) -> str:
        cleaned = " ".join((address or "").replace("\n", " ").split()).strip(" ,")
        if not cleaned:
            return ""
        cleaned = re.sub(r"\s*,\s*", ", ", cleaned)
        city_state_zip = re.search(r"\b([A-Za-z][A-Za-z .'-]+)\s+([A-Z]{2})\s+(\d{5})\b", cleaned)
        if city_state_zip and "," not in cleaned[:city_state_zip.start(1)]:
            prefix = cleaned[:city_state_zip.start(1)].strip(" ,")
            cleaned = f"{prefix}, {city_state_zip.group(1).strip()}, {city_state_zip.group(2)} {city_state_zip.group(3)}" if prefix else f"{city_state_zip.group(1).strip()}, {city_state_zip.group(2)} {city_state_zip.group(3)}"
        return cleaned

    def _route_address_variants(self, address: str) -> list[str]:
        normalized = self._normalize_route_address(address)
        if not normalized:
            return []

        variants: list[str] = []

        def add_variant(value: str) -> None:
            cleaned = self._normalize_route_address(value)
            if cleaned and cleaned not in variants:
                variants.append(cleaned)

        add_variant(normalized)
        add_variant(re.sub(r"\s+#\w+\b", "", normalized, flags=re.IGNORECASE))
        add_variant(
            re.sub(
                r"\b(?:APT|APT\.|UNIT|STE|SUITE|FL|FLOOR)\s+[A-Z0-9-]+\b",
                "",
                normalized,
                flags=re.IGNORECASE,
            )
        )

        number_match = re.search(r"\b\d{1,6}\s+[A-Za-z0-9].*", normalized)
        if number_match:
            add_variant(number_match.group(0))

        if "," in normalized:
            parts = [part.strip() for part in normalized.split(",") if part.strip()]
            if len(parts) >= 2:
                add_variant(", ".join(parts[-3:]))
                add_variant(", ".join(parts[-2:]))

        return variants

    def _route_cache_key(self, address: str) -> str:
        return self._normalize_route_address(address).lower()

    def _geocode_route_address(self, address: str) -> tuple[float, float] | None:
        variants = self._route_address_variants(address)
        if not variants:
            return None

        cache = self.store.route_geocode_cache
        for variant in variants:
            cache_key = self._route_cache_key(variant)
            cached = cache.get(cache_key)
            if isinstance(cached, list) and len(cached) == 2:
                try:
                    return float(cached[0]), float(cached[1])
                except (TypeError, ValueError):
                    pass

        for variant in variants:
            query = quote(variant)
            url = (
                "https://nominatim.openstreetmap.org/search"
                f"?q={query}&format=jsonv2&limit=1&countrycodes=us"
            )
            request = Request(
                url,
                headers={
                    "User-Agent": "ClaimsManagerRoutePlanner/1.0 (route optimization)",
                    "Accept": "application/json",
                },
            )
            try:
                with urlopen(request, timeout=6) as response:
                    payload = json.loads(response.read().decode("utf-8"))
            except Exception:
                continue

            if not payload:
                continue

            first = payload[0]
            try:
                lat = float(first["lat"])
                lon = float(first["lon"])
            except (KeyError, TypeError, ValueError):
                continue

            cache[self._route_cache_key(variant)] = [lat, lon]
            self.store.save()
            return lat, lon

        return None

    def _geocode_route_record(self, record: ClaimRecord) -> tuple[float, float] | None:
        address = self._route_default_address(record)
        if not address:
            return None

        coords = self._geocode_route_address(address)
        if coords:
            return coords

        fallback_variants: list[str] = []
        town = (record.town or "").strip()
        owner_address = self._normalize_route_address(record.owner_address)

        city_state_zip = re.search(r"\b([A-Za-z][A-Za-z .'-]+),?\s+([A-Z]{2})\s+(\d{5})\b", owner_address or address)
        if city_state_zip:
            fallback_variants.append(
                f"{city_state_zip.group(1).strip()}, {city_state_zip.group(2)} {city_state_zip.group(3)}"
            )
            fallback_variants.append(city_state_zip.group(3))

        if town:
            fallback_variants.append(f"{town}, CT")

        for variant in fallback_variants:
            coords = self._geocode_route_address(variant)
            if coords:
                return coords

        return None

    @staticmethod
    def _route_distance(point_a: tuple[float, float], point_b: tuple[float, float]) -> float:
        lat1, lon1 = point_a
        lat2, lon2 = point_b
        radius_miles = 3958.8
        dlat = math.radians(lat2 - lat1)
        dlon = math.radians(lon2 - lon1)
        a = (
            math.sin(dlat / 2) ** 2
            + math.cos(math.radians(lat1))
            * math.cos(math.radians(lat2))
            * math.sin(dlon / 2) ** 2
        )
        return 2 * radius_miles * math.asin(math.sqrt(a))

    def _route_path_distance(
        self,
        ordered_keys: list[str],
        coords_by_key: dict[str, tuple[float, float]],
        start_coord: tuple[float, float],
    ) -> float:
        total = 0.0
        previous = start_coord
        for key in ordered_keys:
            current = coords_by_key[key]
            total += self._route_distance(previous, current)
            previous = current
        return total

    def _nearest_neighbor_route_order(
        self,
        route_keys: list[str],
        coords_by_key: dict[str, tuple[float, float]],
        start_coord: tuple[float, float],
    ) -> list[str]:
        remaining = route_keys[:]
        ordered: list[str] = []
        current = start_coord
        while remaining:
            next_key = min(remaining, key=lambda key: self._route_distance(current, coords_by_key[key]))
            ordered.append(next_key)
            remaining.remove(next_key)
            current = coords_by_key[next_key]
        return ordered

    def _two_opt_route_order(
        self,
        route_keys: list[str],
        coords_by_key: dict[str, tuple[float, float]],
        start_coord: tuple[float, float],
    ) -> list[str]:
        best = route_keys[:]
        best_distance = self._route_path_distance(best, coords_by_key, start_coord)
        improved = True
        while improved:
            improved = False
            for i in range(len(best) - 1):
                for j in range(i + 1, len(best)):
                    candidate = best[:i] + list(reversed(best[i : j + 1])) + best[j + 1 :]
                    candidate_distance = self._route_path_distance(candidate, coords_by_key, start_coord)
                    if candidate_distance + 0.01 < best_distance:
                        best = candidate
                        best_distance = candidate_distance
                        improved = True
                        break
                if improved:
                    break
        return best

    def _osrm_route_order(
        self,
        route_keys: list[str],
        coords_by_key: dict[str, tuple[float, float]],
        start_coord: tuple[float, float],
    ) -> list[str] | None:
        coordinates = [start_coord] + [coords_by_key[key] for key in route_keys]
        coord_value = ";".join(f"{lon:.6f},{lat:.6f}" for lat, lon in coordinates)
        url = (
            "https://router.project-osrm.org/trip/v1/driving/"
            f"{coord_value}?source=first&roundtrip=false&steps=false&overview=false"
        )
        request = Request(
            url,
            headers={
                "User-Agent": "ClaimsManagerRoutePlanner/1.0 (route optimization)",
                "Accept": "application/json",
            },
        )
        try:
            with urlopen(request, timeout=8) as response:
                payload = json.loads(response.read().decode("utf-8"))
        except Exception:
            return None

        if payload.get("code") != "Ok":
            return None

        optimized_pairs: list[tuple[int, str]] = []
        for waypoint in payload.get("waypoints", []):
            input_index = waypoint.get("trips_index")
            optimized_index = waypoint.get("waypoint_index")
            if input_index in (None, 0) or optimized_index is None:
                continue
            route_list_index = int(input_index) - 1
            if 0 <= route_list_index < len(route_keys):
                optimized_pairs.append((int(optimized_index), route_keys[route_list_index]))

        if len(optimized_pairs) != len(route_keys):
            return None

        optimized_pairs.sort(key=lambda item: item[0])
        return [key for _, key in optimized_pairs]

    def _refresh_route_planner(self, open_records: list[ClaimRecord]) -> None:
        if not self.route_available_tree or not self.route_selected_tree:
            return
        self.route_available_tree.delete(*self.route_available_tree.get_children())
        self.route_selected_tree.delete(*self.route_selected_tree.get_children())

        route_keys = [key for key in self.store.route_plan_keys if any(record.key == key for record in open_records)]
        if route_keys != self.store.route_plan_keys:
            self.store.data["route_plan_keys"] = route_keys
            self.store.save()

        for record in open_records:
            self.route_available_tree.insert(
                "",
                tk.END,
                iid=record.key,
                values=(record.claim_id, record.customer_name or record.title, self._route_default_address(record)),
            )

        for index, key in enumerate(route_keys, start=1):
            record = self.store.get_claim(key)
            if not record:
                continue
            self.route_selected_tree.insert(
                "",
                tk.END,
                iid=record.key,
                values=(index, record.claim_id, record.customer_name or record.title, self._route_default_address(record)),
            )

        self._apply_saved_tree_sort(self.route_available_tree)
        self._apply_saved_tree_sort(self.route_selected_tree)
        self.route_selected_stop_key = None
        self.route_address_var.set("")
        if self.route_map_button:
            state = "normal" if route_keys else "disabled"
            self.route_map_button.configure(state=state)

    def add_selected_claims_to_route(self) -> None:
        if not self.route_available_tree:
            return
        selected = self.route_available_tree.selection()
        if not selected:
            return
        route_keys = self.store.route_plan_keys
        for key in selected:
            if key not in route_keys:
                route_keys.append(key)
        self.store.save()
        self.refresh_views()

    def on_route_stop_selected(self, _event: object) -> None:
        if not self.route_selected_tree:
            return
        selection = self.route_selected_tree.selection()
        if not selection:
            self.route_selected_stop_key = None
            self.route_address_var.set("")
            return
        self.route_selected_stop_key = selection[0]
        record = self.store.get_claim(self.route_selected_stop_key)
        if record:
            self.route_address_var.set(self._route_default_address(record))

    def save_route_stop_address(self) -> None:
        if not self.route_selected_stop_key:
            messagebox.showinfo("No Stop Selected", "Select a stop in Tomorrow's Stops first.")
            return
        record = self.store.get_claim(self.route_selected_stop_key)
        if not record:
            return
        record.route_address_override = self.route_address_var.get().strip()
        self.store.upsert_claim(record)
        self.store.save()
        self.refresh_views(select_key=record.key)

    def remove_route_stop(self) -> None:
        if not self.route_selected_stop_key:
            return
        self.store.data["route_plan_keys"] = [key for key in self.store.route_plan_keys if key != self.route_selected_stop_key]
        self.store.save()
        self.refresh_views()

    def clear_route_plan(self) -> None:
        self.store.data["route_plan_keys"] = []
        self.store.save()
        self.refresh_views()

    def _build_google_maps_route_url(self, route_keys: list[str]) -> str:
        addresses = []
        for key in route_keys:
            record = self.store.get_claim(key)
            if not record:
                continue
            address = self._route_default_address(record)
            if not address:
                raise RuntimeError(f"{record.claim_id} is missing a stop address.")
            addresses.append(address)
        if not addresses:
            raise RuntimeError("No stop addresses are available.")

        origin = quote(self.route_home_var.get().strip() or ROUTE_HOME_ADDRESS)
        destination = quote(addresses[-1])
        waypoint_addresses = addresses[:-1]
        waypoint_value = "|".join(quote(address) for address in waypoint_addresses)
        url = f"https://www.google.com/maps/dir/?api=1&origin={origin}&destination={destination}&travelmode=driving"
        if waypoint_value:
            url += f"&waypoints={waypoint_value}"
        return url

    def _optimize_route_keys(self, show_messages: bool = True) -> list[str] | None:
        route_keys = list(self.store.route_plan_keys)
        if len(route_keys) < 2:
            if show_messages:
                messagebox.showinfo("Not Enough Stops", "Add at least two stops to optimize the route.")
            return None

        home_address = self.route_home_var.get().strip() or ROUTE_HOME_ADDRESS
        start_coord = self._geocode_route_address(home_address)
        if not start_coord:
            if show_messages:
                messagebox.showerror("Home Address Error", "Could not locate the starting address for the route.")
            return None

        coords_by_key: dict[str, tuple[float, float]] = {}
        failed_records: list[ClaimRecord] = []

        for key in route_keys:
            record = self.store.get_claim(key)
            if not record:
                continue
            address = self._route_default_address(record)
            if not address:
                if show_messages:
                    messagebox.showerror("Missing Address", f"{record.claim_id} is missing a stop address.")
                return None
            coords = self._geocode_route_record(record)
            if not coords:
                failed_records.append(record)
                continue
            coords_by_key[key] = coords

        if failed_records:
            if show_messages:
                items = "\n".join(f"{record.claim_id} - {record.customer_name or record.title}" for record in failed_records)
                messagebox.showerror(
                    "Route Address Error",
                    "Could not locate these route stops. Edit their stop addresses first:\n"
                    f"{items}",
                )
            return None

        optimized = self._osrm_route_order(route_keys, coords_by_key, start_coord)
        used_fallback = False
        if not optimized:
            used_fallback = True
            if len(route_keys) <= 8:
                optimized = min(
                    (list(order) for order in permutations(route_keys)),
                    key=lambda order: self._route_path_distance(order, coords_by_key, start_coord),
                )
            else:
                optimized = self._nearest_neighbor_route_order(route_keys, coords_by_key, start_coord)
                optimized = self._two_opt_route_order(optimized, coords_by_key, start_coord)

        if optimized != route_keys:
            self.store.data["route_plan_keys"] = optimized
            self.store.save()
            self.refresh_views()

        total_miles = self._route_path_distance(optimized, coords_by_key, start_coord)
        if used_fallback:
            self.route_status_var.set(
                f"Optimized {len(optimized)} stops with fallback routing. Approx route length: {total_miles:.1f} miles."
            )
        else:
            self.route_status_var.set(
                f"Optimized {len(optimized)} stops using road routing. Approx route length: {total_miles:.1f} miles."
            )
        return optimized

    def optimize_route_plan(self) -> None:
        if self._optimize_route_keys(show_messages=True):
            messagebox.showinfo("Route Ready", "Route is ready.")

    def open_route_in_google_maps(self) -> None:
        route_keys = list(self.store.route_plan_keys)
        if not route_keys:
            messagebox.showinfo("No Route", "Add stops to Tomorrow's Stops first.")
            return
        optimized_keys = self._optimize_route_keys(show_messages=True)
        if not optimized_keys:
            return
        try:
            url = self._build_google_maps_route_url(optimized_keys)
        except Exception as exc:
            messagebox.showerror("Route Failed", str(exc))
            return
        try:
            self._send_route_email(url)
        except Exception as exc:
            messagebox.showerror("Route Email Failed", f"Could not email the route link:\n{exc}")
            return
        webbrowser.open(url)
        self.route_status_var.set("Opened the optimized route in Google Maps and emailed the route link to your Gmail.")

    def on_claim_selected(self, event: object) -> None:
        tree = getattr(event, "widget", None)
        if not isinstance(tree, ttk.Treeview):
            return
        selection = tree.selection()
        if not selection:
            return
        self.selected_key = selection[0]
        self.populate_details(self.selected_key)

    def open_selected_claim_folder(self, event: object) -> None:
        tree = getattr(event, "widget", None)
        if not isinstance(tree, ttk.Treeview):
            return
        selection = tree.selection()
        if not selection:
            return

        record = self.store.get_claim(selection[0])
        if not record:
            return
        if record.source_path == "Manual Entry":
            messagebox.showinfo("No Folder", "This claim does not have a folder yet.")
            return

        path = Path(record.source_path)
        target = path if path.exists() else path.parent
        if not target.exists():
            messagebox.showerror("Folder Missing", f"Could not find:\n{record.source_path}")
            return
        os.startfile(str(target))

    def open_selected_claim_folder_from_menu(self) -> None:
        if not self.selected_key:
            return
        dummy_event = type("Dummy", (), {"widget": None})()
        record = self.store.get_claim(self.selected_key)
        if not record:
            return
        if record.source_path == "Manual Entry":
            messagebox.showinfo("No Folder", "This claim does not have a folder yet.")
            return
        path = Path(record.source_path)
        target = path if path.exists() else path.parent
        if not target.exists():
            messagebox.showerror("Folder Missing", f"Could not find:\n{record.source_path}")
            return
        os.startfile(str(target))

    def show_tree_context_menu(self, event: object) -> None:
        tree = getattr(event, "widget", None)
        if not isinstance(tree, ttk.Treeview):
            return
        row_id = tree.identify_row(getattr(event, "y", 0))
        if row_id:
            tree.selection_set(row_id)
            self.selected_key = row_id
            self.populate_details(row_id)
        if self.tree_context_menu:
            self._refresh_claim_tools_context_menu()
            self.tree_context_menu.tk_popup(getattr(event, "x_root", 0), getattr(event, "y_root", 0))

    def _refresh_claim_tools_context_menu(self) -> None:
        if not self.claim_tools_menu:
            return
        self.claim_tools_menu.delete(0, tk.END)
        self.claim_tools_menu.add_command(label="AIG Claim Summary", command=self.generate_aig_claim_summary_for_selected_claim)
        self.claim_tools_menu.add_command(label="Default Claim Notes", command=self.generate_default_claim_notes_for_selected_claim)
        self.claim_tools_menu.add_command(label="Mitchell Total Loss", command=self.generate_mitchell_total_loss_for_selected_claim)
        self.claim_tools_menu.add_separator()

        tools_root = Path(self.store.claim_tools_folder)
        if not tools_root.exists():
            self.claim_tools_menu.add_command(label="Claim Tools Folder Missing", state="disabled")
            return

        files = [path for path in tools_root.iterdir() if path.is_file()]
        files.sort(key=lambda path: path.name.lower())
        if not files:
            self.claim_tools_menu.add_command(label="No Files Found", state="disabled")
            return

        for path in files:
            label = path.name
            self.claim_tools_menu.add_command(
                label=label,
                command=lambda file_path=path: self._open_claim_tool_file(file_path),
            )

    def _open_claim_tool_file(self, file_path: Path) -> None:
        if not file_path.exists():
            messagebox.showerror("Claim Tool Missing", f"Could not find:\n{file_path}")
            return
        os.startfile(str(file_path))

    def refresh_claim_tools_contacts_view(self) -> None:
        if not self.claim_tools_contacts_tree:
            return
        self.claim_tools_contacts_tree.delete(*self.claim_tools_contacts_tree.get_children())
        for index, entry in enumerate(self.store.claim_tools_contacts):
            name = (entry.get("name") or "").strip()
            number = (entry.get("number") or "").strip()
            prompt_guide = " ".join((entry.get("prompt_guide") or "").split())
            notes = " ".join((entry.get("notes") or "").split())
            self.claim_tools_contacts_tree.insert("", "end", iid=str(index), values=(name, number, prompt_guide, notes))
        self._apply_saved_tree_sort(self.claim_tools_contacts_tree)

    def _prompt_claim_tool_contact(
        self,
        title: str,
        initial_name: str = "",
        initial_number: str = "",
        initial_prompt_guide: str = "",
        initial_notes: str = "",
    ) -> dict[str, str] | None:
        dialog = tk.Toplevel(self.root)
        dialog.title(title)
        dialog.transient(self.root)
        dialog.resizable(True, True)
        dialog.geometry("720x360")
        dialog.grab_set()

        result: dict[str, dict[str, str] | None] = {"value": None}
        name_var = tk.StringVar(value=initial_name)
        number_var = tk.StringVar(value=initial_number)

        frame = ttk.Frame(dialog, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(1, weight=1)
        frame.rowconfigure(2, weight=1)
        frame.rowconfigure(3, weight=1)

        ttk.Label(frame, text="Name").grid(row=0, column=0, sticky="w", padx=(0, 8), pady=(0, 8))
        name_entry = ttk.Entry(frame, textvariable=name_var, width=42)
        name_entry.grid(row=0, column=1, sticky="ew", pady=(0, 8))

        ttk.Label(frame, text="Number").grid(row=1, column=0, sticky="w", padx=(0, 8), pady=(0, 8))
        ttk.Entry(frame, textvariable=number_var, width=42).grid(row=1, column=1, sticky="ew", pady=(0, 8))

        ttk.Label(frame, text="Prompt Guide").grid(row=2, column=0, sticky="nw", padx=(0, 8), pady=(0, 8))
        prompt_text = tk.Text(frame, height=5, wrap="word")
        prompt_text.grid(row=2, column=1, sticky="nsew", pady=(0, 8))
        prompt_text.insert("1.0", initial_prompt_guide or "")

        ttk.Label(frame, text="Notes").grid(row=3, column=0, sticky="nw", padx=(0, 8))
        notes_text = tk.Text(frame, height=5, wrap="word")
        notes_text.grid(row=3, column=1, sticky="nsew")
        notes_text.insert("1.0", initial_notes or "")

        button_row = ttk.Frame(frame)
        button_row.grid(row=4, column=0, columnspan=2, sticky="e", pady=(12, 0))

        def submit() -> None:
            name = name_var.get().strip()
            number = number_var.get().strip()
            prompt_guide = prompt_text.get("1.0", tk.END).strip()
            notes = notes_text.get("1.0", tk.END).strip()
            if not name:
                messagebox.showinfo("Name Required", "Enter a name first.", parent=dialog)
                return
            result["value"] = {
                "name": name,
                "number": number,
                "prompt_guide": prompt_guide,
                "notes": notes,
            }
            self._remember_dialog_geometry(dialog)
            dialog.destroy()

        def cancel() -> None:
            result["value"] = None
            self._remember_dialog_geometry(dialog)
            dialog.destroy()

        ttk.Button(button_row, text="Save", command=submit).pack(side="left", padx=(0, 8))
        ttk.Button(button_row, text="Cancel", command=cancel).pack(side="left")

        dialog.protocol("WM_DELETE_WINDOW", cancel)
        dialog.update_idletasks()
        self._place_dialog(dialog)
        name_entry.focus_set()
        self.root.wait_window(dialog)
        return result["value"]

    def add_claim_tool_contact(self) -> None:
        result = self._prompt_claim_tool_contact("Add Claim Tool Contact")
        if result is None:
            return
        self.store.claim_tools_contacts.append(result)
        self.store.save()
        self.refresh_claim_tools_contacts_view()

    def edit_claim_tool_contact(self) -> None:
        if not self.claim_tools_contacts_tree:
            return
        selection = self.claim_tools_contacts_tree.selection()
        if not selection:
            messagebox.showinfo("No Selection", "Select a contact first.")
            return
        index = int(selection[0])
        entry = self.store.claim_tools_contacts[index]
        result = self._prompt_claim_tool_contact(
            "Edit Claim Tool Contact",
            entry.get("name", ""),
            entry.get("number", ""),
            entry.get("prompt_guide", ""),
            entry.get("notes", ""),
        )
        if result is None:
            return
        self.store.claim_tools_contacts[index] = result
        self.store.save()
        self.refresh_claim_tools_contacts_view()
        self.claim_tools_contacts_tree.selection_set(str(index))

    def remove_claim_tool_contact(self) -> None:
        if not self.claim_tools_contacts_tree:
            return
        selection = self.claim_tools_contacts_tree.selection()
        if not selection:
            messagebox.showinfo("No Selection", "Select a contact first.")
            return
        index = int(selection[0])
        del self.store.claim_tools_contacts[index]
        self.store.save()
        self.refresh_claim_tools_contacts_view()

    def call_claim_tool_contact(self, _event: object | None = None) -> None:
        if not self.claim_tools_contacts_tree:
            return
        selection = self.claim_tools_contacts_tree.selection()
        if not selection:
            return
        index = int(selection[0])
        entry = self.store.claim_tools_contacts[index]
        number = (entry.get("number") or "").strip()
        if not number:
            messagebox.showinfo("No Number", "This entry does not have a phone number yet.")
            return

        tel_number = re.sub(r"[^0-9+#*,;]", "", number)
        if not tel_number:
            messagebox.showinfo("Invalid Number", "This number cannot be used for a call action.")
            return

        try:
            os.startfile(f"tel:{tel_number}")
        except Exception:
            messagebox.showinfo(
                "Call Not Available",
                f"Windows could not start a call for:\n{number}\n\nIf a calling app is installed later, this should work.",
            )

    def _find_assignment_sheet_path(self, record: ClaimRecord) -> Path | None:
        candidates: list[Path] = []
        if record.assign_pdf_path:
            candidates.append(Path(record.assign_pdf_path))

        if record.source_path and record.source_path != "Manual Entry":
            folder_path = Path(record.source_path)
            if folder_path.exists() and folder_path.is_dir():
                candidates.append(folder_path / "assign.pdf")
                for child in folder_path.iterdir():
                    if not child.is_file() or child.suffix.lower() != ".pdf":
                        continue
                    lower_name = child.name.lower()
                    if "assign" in lower_name or re.fullmatch(r"\d+\.pdf", lower_name):
                        candidates.append(child)

        for candidate in candidates:
            try:
                if not candidate.exists() or not candidate.is_file() or candidate.suffix.lower() != ".pdf":
                    continue
                lower_name = candidate.name.lower()
                if "assign" in lower_name or re.fullmatch(r"\d+\.pdf", lower_name):
                    return candidate
            except OSError:
                continue
        return None

    def open_working_sheet_for_selected_claim(self) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select a claim first.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return
        assign_path = self._find_assignment_sheet_path(record)
        if assign_path:
            os.startfile(str(assign_path))
            return

        if record.source_path == "Manual Entry":
            messagebox.showinfo("No Assignment Sheet", "This claim does not have an assignment sheet.")
            return

        folder_path = Path(record.source_path)
        if folder_path.exists():
            messagebox.showinfo("Assignment Sheet Missing", "Could not find the assignment sheet. Opening the claim folder instead.")
            os.startfile(str(folder_path))
            return

        messagebox.showerror("Assignment Sheet Missing", "Could not find the original assignment sheet or claim folder.")

    def email_assignment_sheet_for_selected_claim(self) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select a claim first.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return

        assign_path = self._find_assignment_sheet_path(record)
        if not assign_path:
            messagebox.showerror("Assignment Sheet Missing", "Could not find the original assignment sheet for this claim.")
            return

        try:
            sender_email, app_password = self._get_office_update_email_credentials()
            if not sender_email or not app_password:
                raise RuntimeError("A Gmail address and app password are required to email the assignment sheet.")

            subject_parts = [part for part in (record.claim_id, record.customer_name or record.title) if part]
            subject = "Assignment Sheet"
            if subject_parts:
                subject = f"Assignment Sheet - {' - '.join(subject_parts)}"

            message = EmailMessage()
            message["From"] = sender_email
            message["To"] = sender_email
            message["Subject"] = subject
            message.set_content("Attached is the original assignment sheet.")
            message.add_attachment(
                assign_path.read_bytes(),
                maintype="application",
                subtype="pdf",
                filename=assign_path.name,
            )

            context = ssl.create_default_context()
            with smtplib.SMTP("smtp.gmail.com", 587, timeout=30) as smtp:
                smtp.ehlo()
                smtp.starttls(context=context)
                smtp.ehlo()
                smtp.login(sender_email, app_password)
                smtp.send_message(message)

            messagebox.showinfo("Assignment Sheet Emailed", f"Emailed to {sender_email}.")
        except Exception as exc:
            messagebox.showerror("Email Failed", f"Could not email the assignment sheet:\n{exc}")

    def generate_nada_pdf_for_selected_claim(self) -> None:
        webbrowser.open("https://www.jdpower.com/cars/manufacturers")

    def _prompt_text(self, title: str, prompt: str, initialvalue: str = "") -> str | None:
        dialog = tk.Toplevel(self.root)
        dialog.title(title)
        dialog.transient(self.root)
        dialog.resizable(True, True)
        dialog.geometry("560x260")
        dialog.grab_set()

        result: dict[str, str | None] = {"value": None}

        frame = ttk.Frame(dialog, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(1, weight=1)

        ttk.Label(frame, text=prompt, wraplength=520, justify="left").grid(row=0, column=0, sticky="w", pady=(0, 8))
        text = tk.Text(frame, wrap="word", height=8)
        text.grid(row=1, column=0, sticky="nsew")
        text.insert("1.0", initialvalue or "")
        text.focus_set()

        button_row = ttk.Frame(frame)
        button_row.grid(row=2, column=0, sticky="e", pady=(10, 0))

        def submit() -> None:
            result["value"] = text.get("1.0", tk.END).strip()
            self._remember_dialog_geometry(dialog)
            dialog.destroy()

        def cancel() -> None:
            result["value"] = None
            self._remember_dialog_geometry(dialog)
            dialog.destroy()

        def submit_event(_event: object | None = None) -> str:
            submit()
            return "break"

        def insert_newline(_event: object | None = None) -> str:
            text.insert(tk.INSERT, "\n")
            return "break"

        ttk.Button(button_row, text="OK", command=submit).pack(side="left", padx=(0, 8))
        ttk.Button(button_row, text="Cancel", command=cancel).pack(side="left")

        dialog.protocol("WM_DELETE_WINDOW", cancel)
        dialog.bind("<Escape>", lambda _event: cancel())
        dialog.bind("<Return>", submit_event)
        text.bind("<Return>", submit_event)
        text.bind("<Shift-Return>", insert_newline)
        dialog.update_idletasks()
        self._place_dialog(dialog)
        self.root.wait_window(dialog)
        return result["value"]

    def _prompt_choice(self, title: str, prompt: str, values: list[str], initialvalue: str = "") -> str | None:
        dialog = tk.Toplevel(self.root)
        dialog.title(title)
        dialog.transient(self.root)
        dialog.resizable(False, False)
        dialog.geometry("420x150")
        dialog.grab_set()

        result: dict[str, str | None] = {"value": None}
        selected = tk.StringVar(value=initialvalue or (values[0] if values else ""))

        frame = ttk.Frame(dialog, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(0, weight=1)

        ttk.Label(frame, text=prompt, wraplength=380, justify="left").grid(row=0, column=0, sticky="w", pady=(0, 10))
        combo = ttk.Combobox(frame, textvariable=selected, values=values, state="readonly")
        combo.grid(row=1, column=0, sticky="ew")
        combo.focus_set()

        button_row = ttk.Frame(frame)
        button_row.grid(row=2, column=0, sticky="e", pady=(12, 0))

        def submit() -> None:
            result["value"] = selected.get().strip()
            self._remember_dialog_geometry(dialog)
            dialog.destroy()

        def cancel() -> None:
            result["value"] = None
            self._remember_dialog_geometry(dialog)
            dialog.destroy()

        def submit_event(_event: object | None = None) -> str:
            submit()
            return "break"

        ttk.Button(button_row, text="OK", command=submit).pack(side="left", padx=(0, 8))
        ttk.Button(button_row, text="Cancel", command=cancel).pack(side="left")

        dialog.protocol("WM_DELETE_WINDOW", cancel)
        dialog.bind("<Escape>", lambda _event: cancel())
        dialog.bind("<Return>", submit_event)
        combo.bind("<Return>", submit_event)
        dialog.update_idletasks()
        self._place_dialog(dialog)
        self.root.wait_window(dialog)
        return result["value"]

    def _center_dialog(self, dialog: tk.Toplevel) -> None:
        dialog.update_idletasks()
        width = dialog.winfo_width()
        height = dialog.winfo_height()
        root_x = self.root.winfo_rootx()
        root_y = self.root.winfo_rooty()
        root_width = self.root.winfo_width()
        root_height = self.root.winfo_height()
        x = root_x + max((root_width - width) // 2, 0)
        y = root_y + max((root_height - height) // 2, 0)
        dialog.geometry(f"{width}x{height}+{x}+{y}")

    def _place_dialog(self, dialog: tk.Toplevel) -> None:
        dialog.update_idletasks()
        width = dialog.winfo_width()
        height = dialog.winfo_height()
        if self.last_dialog_geometry:
            _last_width, _last_height, x, y = self.last_dialog_geometry
            dialog.geometry(f"{width}x{height}+{x}+{y}")
            return
        self._center_dialog(dialog)

    def _remember_dialog_geometry(self, dialog: tk.Toplevel) -> None:
        dialog.update_idletasks()
        self.last_dialog_geometry = (
            dialog.winfo_width(),
            dialog.winfo_height(),
            dialog.winfo_x(),
            dialog.winfo_y(),
        )

    def _prompt_yes_no_text(self, title: str, prompt: str) -> str | None:
        result = messagebox.askyesnocancel(title, prompt, parent=self.root)
        if result is None:
            return None
        return "Yes" if result else "No"

    def _prompt_wizard(self, title: str, questions: list[dict[str, object]]) -> dict[str, str] | None:
        if not questions:
            return {}

        dialog = tk.Toplevel(self.root)
        dialog.title(title)
        dialog.transient(self.root)
        dialog.resizable(True, True)
        dialog.geometry("720x360")
        dialog.grab_set()

        result: dict[str, dict[str, str] | None] = {"value": None}
        answers = {str(question["key"]): str(question.get("default", "")) for question in questions}
        index_var = tk.IntVar(value=0)

        frame = ttk.Frame(dialog, padding=12)
        frame.pack(fill="both", expand=True)
        frame.columnconfigure(0, weight=1)
        frame.rowconfigure(2, weight=1)

        step_label = ttk.Label(frame, text="")
        step_label.grid(row=0, column=0, sticky="w")
        prompt_label = ttk.Label(frame, text="", wraplength=660, justify="left")
        prompt_label.grid(row=1, column=0, sticky="w", pady=(6, 10))

        input_frame = ttk.Frame(frame)
        input_frame.grid(row=2, column=0, sticky="nsew")
        input_frame.columnconfigure(0, weight=1)
        input_frame.rowconfigure(0, weight=1)

        text_widget = tk.Text(input_frame, wrap="word", height=8)
        combo_var = tk.StringVar()
        combo_widget = ttk.Combobox(input_frame, textvariable=combo_var, state="readonly")

        button_row = ttk.Frame(frame)
        button_row.grid(row=3, column=0, sticky="ew", pady=(12, 0))
        button_row.columnconfigure(1, weight=1)

        back_button = ttk.Button(button_row, text="< Back")
        next_button = ttk.Button(button_row, text="Next >")
        cancel_button = ttk.Button(button_row, text="Cancel")
        back_button.grid(row=0, column=0, sticky="w")
        next_button.grid(row=0, column=2, sticky="e", padx=(8, 0))
        cancel_button.grid(row=0, column=3, sticky="e", padx=(8, 0))

        def store_current_value() -> None:
            question = questions[index_var.get()]
            key = str(question["key"])
            input_type = str(question.get("input_type", "text"))
            if input_type == "choice":
                answers[key] = combo_var.get().strip()
            else:
                answers[key] = text_widget.get("1.0", tk.END).strip()

        def render_current_question() -> None:
            question = questions[index_var.get()]
            key = str(question["key"])
            prompt = str(question.get("label", key))
            input_type = str(question.get("input_type", "text"))
            step_label.configure(text=f"Question {index_var.get() + 1} of {len(questions)}")
            prompt_label.configure(text=prompt)

            text_widget.grid_forget()
            combo_widget.grid_forget()

            if input_type == "choice":
                combo_widget.configure(values=list(question.get("choices", [])))
                combo_var.set(answers.get(key, str(question.get("default", ""))))
                combo_widget.grid(row=0, column=0, sticky="ew")
                combo_widget.focus_set()
            else:
                text_widget.grid(row=0, column=0, sticky="nsew")
                text_widget.delete("1.0", tk.END)
                text_widget.insert("1.0", answers.get(key, str(question.get("default", ""))))
                text_widget.focus_set()

            back_button.configure(state="normal" if index_var.get() > 0 else "disabled")
            next_button.configure(text="Finish" if index_var.get() == len(questions) - 1 else "Next >")

        def go_back() -> None:
            store_current_value()
            if index_var.get() > 0:
                index_var.set(index_var.get() - 1)
                render_current_question()

        def go_next() -> None:
            store_current_value()
            if index_var.get() >= len(questions) - 1:
                result["value"] = dict(answers)
                self._remember_dialog_geometry(dialog)
                dialog.destroy()
                return
            index_var.set(index_var.get() + 1)
            render_current_question()

        def cancel() -> None:
            result["value"] = None
            self._remember_dialog_geometry(dialog)
            dialog.destroy()

        def next_event(_event: object | None = None) -> str:
            go_next()
            return "break"

        def back_event(_event: object | None = None) -> str:
            go_back()
            return "break"

        def insert_newline(_event: object | None = None) -> str:
            text_widget.insert(tk.INSERT, "\n")
            return "break"

        back_button.configure(command=go_back)
        next_button.configure(command=go_next)
        cancel_button.configure(command=cancel)
        dialog.protocol("WM_DELETE_WINDOW", cancel)
        dialog.bind("<Escape>", lambda _event: cancel())
        dialog.bind("<Return>", next_event)
        dialog.bind("<Shift-Left>", back_event)
        text_widget.bind("<Return>", next_event)
        text_widget.bind("<Shift-Return>", insert_newline)
        combo_widget.bind("<Return>", next_event)
        dialog.update_idletasks()
        self._place_dialog(dialog)
        render_current_question()
        self.root.wait_window(dialog)
        return result["value"]

    def generate_aig_claim_summary_for_selected_claim(self) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select a claim first.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return

        vehicle_status_default = "repair" if not record.total_loss else "total loss"
        days_default = ""
        shop_contact_default = record.contact_email or record.contact_phone or ""
        comments_default = record.assignment_claim_notes or record.office_additional_notes or ""
        prior_damage_default = "No"
        drivable_default = "Yes"
        if "not drivable" in (record.assignment_claim_notes or "").lower():
            drivable_default = "No"
        if record.total_loss:
            drivable_default = "No"

        repairable = messagebox.askyesnocancel(
            "AIG Route",
            "Is it repairable?\n\nYes = Repairable\nNo = Total Loss",
            parent=self.root,
        )
        if repairable is None:
            return

        route_label = ""
        questions: list[tuple[str, str]] = []
        if repairable:
            original = messagebox.askyesnocancel(
                "AIG Route",
                "Is this an original estimate?\n\nYes = Original Estimate (E01)\nNo = Supplement (S01-S99)",
                parent=self.root,
            )
            if original is None:
                return
            if original:
                route_label = "AIG Appraisal Notes - E01 - Repairable"
                questions = [
                    ("Claim #", record.claim_number or record.claim_id or "", "text"),
                    ("Vehicle status (repair / cash settlement option)", vehicle_status_default, "choice", ["repair", "cash settlement option"]),
                    ("Vehicle drivable?", drivable_default, "choice", ["Yes", "No"]),
                    ("Copy of appraisal supplied to owner? If so, how?", "No", "choice", ["Yes/Email", "No"]),
                    ("Is this an agreed price with the shop of owner's choice? If so, with who?", "N/A", "text"),
                    ("Did you supply a copy of this appraisal to the shop? If so, how?", "No", "choice", ["Yes/Email", "No"]),
                    ("Shop email / fax #", "N/A", "text"),
                    ("Shop tax id", "N/A", "text"),
                    ("Any unrelated or prior damage?", prior_damage_default, "choice", ["Yes", "No"]),
                    ("If yes, was an UPD estimate created?", "No", "choice", ["Yes", "No"]),
                    ("Number of days to repair (total labor hours divided by 4)", days_default, "text"),
                    ("Additional comments", comments_default, "text"),
                ]
            else:
                route_label = "AIG Appraisal Notes - S01-S99 - Repairable"
                questions = [
                    ("Claim #", record.claim_number or record.claim_id or "", "text"),
                    ("Vehicle status (repair / cash settlement option)", vehicle_status_default, "choice", ["repair", "cash settlement option"]),
                    ("Copy of appraisal supplied to owner? If so, how?", "No", "choice", ["Yes/Email", "No"]),
                    ("Is this an agreed price with the shop of owner's choice? If so, with who?", "N/A", "text"),
                    ("Did you supply a copy of this appraisal to the shop? If so, how?", "No", "choice", ["Yes/Email", "No"]),
                    ("Shop email / fax #", "N/A", "text"),
                    ("Shop tax id", "N/A", "text"),
                    ("Additional number of days to repair (total labor hours divided by 4)", days_default, "text"),
                    ("Additional comments", comments_default, "text"),
                ]
        else:
            route_label = "AIG Appraisal Notes - Total Loss"
            questions = [
                ("Claim #", record.claim_number or record.claim_id or "", "text"),
                ("Vehicle status (total loss / constructive total loss)", "total loss", "choice", ["total loss", "constructive total loss"]),
                ("Vehicle drivable?", drivable_default, "choice", ["Yes", "No"]),
                ("Any unrelated or prior damage?", prior_damage_default, "choice", ["Yes", "No"]),
                ("If yes, was an UPD estimate created?", "No", "choice", ["Yes", "No"]),
                ("Towing?", "N/A", "text"),
                ("Storage per day?", "N/A", "text"),
                ("Additional charges?", "N/A", "text"),
                ("Number of days to repair (total labor hours divided by 4)", days_default, "text"),
                ("Additional comments", comments_default, "text"),
            ]

        wizard_questions: list[dict[str, object]] = []
        for index, question in enumerate(questions):
            label = str(question[0])
            default = str(question[1])
            input_type = str(question[2]) if len(question) > 2 else "text"
            entry: dict[str, object] = {
                "key": f"q{index}",
                "label": label,
                "default": default,
                "input_type": input_type,
            }
            if input_type == "choice":
                entry["choices"] = list(question[3]) if len(question) > 3 else ["Yes", "No"]
            wizard_questions.append(entry)

        answers_map = self._prompt_wizard("AIG Claim Summary", wizard_questions)
        if answers_map is None:
            return
        answers = [(str(question[0]), str(answers_map.get(f"q{index}", "")).strip()) for index, question in enumerate(questions)]

        merged_answers: list[tuple[str, str]] = []
        skip_next = False
        for index, (label, value) in enumerate(answers):
            if skip_next:
                skip_next = False
                continue
            if label == "Any unrelated or prior damage?" and index + 1 < len(answers):
                next_label, next_value = answers[index + 1]
                if next_label == "If yes, was an UPD estimate created?":
                    merged_answers.append((f"{label} If so, was an UPD estimate created?", f"{value} / {next_value}"))
                    skip_next = True
                    continue
            merged_answers.append((label, value))

        lines = [route_label, ""]
        for label, value in merged_answers:
            lines.append(f"{label}: {value}")
        output = "\n".join(lines).strip()

        self.root.clipboard_clear()
        self.root.clipboard_append(output)
        messagebox.showinfo("Copied", "AIG claim summary copied to clipboard. Paste it into the estimate.")

    def generate_default_claim_notes_for_selected_claim(self) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select a claim first.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return

        route_choice = self._prompt_wizard(
            "Default Claim Notes",
            [
                {
                    "key": "route",
                    "label": "Which route do you want to use?",
                    "default": "Repairable",
                    "input_type": "choice",
                    "choices": ["Repairable", "Total Loss", "Supplement"],
                }
            ],
        )
        if route_choice is None:
            return
        route = str(route_choice.get("route", "Repairable"))

        today_default = datetime.now().strftime("%-m/%-d/%Y") if os.name != "nt" else datetime.now().strftime("%#m/%#d/%Y")
        drivable_default = "No" if record.total_loss else "Yes"
        shop_default = record.shop_name or "No"
        agreed_price_default = "No shop chosen" if not record.shop_name else "N/A"
        estimate_copy_default = "No shop chosen" if not record.shop_name else "Yes/Email"
        parts_default = "Yes"
        quote_default = "No LKQ or A/M available"
        direction_to_pay_default = "No"
        contacted_owner_default = "Yes"
        days_default = ""
        damage_consistent_default = "Yes"
        damage_photos_default = "Yes"
        damage_explain_default = record.damage_description or ""
        unrelated_default = "No"
        unrelated_list_default = "none"
        comments_default = record.assignment_claim_notes or record.office_additional_notes or ""
        tl_location_default = record.location_of_vehicle or record.shop_name or "Customer has vehicle"
        tl_valuation_default = "CCC"
        salvage_default = "Salvage value:"
        move_default = "No"
        supplement_comments_default = record.office_additional_notes or "Supplement for parts, labor rate and AP."

        if route == "Repairable":
            header = "*** REPAIRABLE VEHICLE PARTIAL LOSS SUMMARY ***"
            questions = [
                ("1. Date Inspected", today_default, "text"),
                ("2. Is vehicle driveable?", drivable_default, "choice", ["Yes", "No"]),
                ("3. Shop chosen? Name of shop/City?", shop_default, "text"),
                ("4. Agreed Price? AP with whom?", agreed_price_default, "text"),
                ("5. Was shop provided copy of estimate? How delivered?", estimate_copy_default, "choice", ["Yes/Email", "No", "No shop chosen"]),
                ("6. Were the most cost effective parts utilized?", parts_default, "choice", ["Yes", "No"]),
                ("7. Quote #/Location of available parts", quote_default, "text"),
                ("8. Direction to pay?", direction_to_pay_default, "choice", ["Yes", "No"]),
                ("9. Contacted Owner Post Insp. & explained process?", contacted_owner_default, "choice", ["Yes", "No"]),
                ("10. Days to repair?", days_default, "text"),
                ("11. Damage consistent w facts?", damage_consistent_default, "choice", ["Yes", "No"]),
                ("12. Damages shown in photos?", damage_photos_default, "choice", ["Yes", "No"]),
                ("12a. Explain damages shown in photos", damage_explain_default, "text"),
                ("13. Unrelated damage?", unrelated_default, "choice", ["Yes", "No"]),
                ("13a. List unrelated damage", unrelated_list_default, "text"),
                ("14. Additional Repairable Comments", comments_default, "text"),
            ]
        elif route == "Total Loss":
            header = "*** TOTAL LOSS SUMMARY ***"
            questions = [
                ("1. Date Inspected", today_default, "text"),
                ("2. Location of vehicle?", tl_location_default, "text"),
                ("3. Tow Bill Amount", "N/A", "text"),
                ("4. Storage Rate", "N/A", "text"),
                ("5. Is vehicle driveable?", drivable_default, "choice", ["Yes", "No"]),
                ("6. Contacted Owner Post Insp.? TL Process explained?", contacted_owner_default, "choice", ["Yes", "No"]),
                ("7. TL Valuation/Request#", tl_valuation_default, "text"),
                ("8. ACV", "N/A", "text"),
                ("9. Salvage Value", salvage_default, "text"),
                ("10. Does vehicle need to be moved?", move_default, "choice", ["Yes", "No"]),
                ("11. Lot#", "N/A", "text"),
                ("12. Unrelated damage?", unrelated_default, "choice", ["Yes", "No"]),
                ("12a. List unrelated damage", unrelated_list_default, "text"),
                ("13. Additional TL Comments", comments_default, "text"),
            ]
        else:
            header = "*** SUPPLEMENT PARTIAL LOSS SUMMARY ***"
            questions = [
                ("1. Name of shop/City", shop_default, "text"),
                ("2. Shop Tax ID", "N/A", "text"),
                ("3. Direction to pay attached?", "Yes", "choice", ["Yes", "No"]),
                ("4. Invoices and photos attached?", "Yes", "choice", ["Yes", "No"]),
                ("5. Additional Supplemental Comments", supplement_comments_default, "text"),
            ]

        wizard_questions = []
        for index, question in enumerate(questions):
            label = str(question[0])
            default = str(question[1])
            input_type = str(question[2])
            entry: dict[str, object] = {
                "key": f"q{index}",
                "label": label,
                "default": default,
                "input_type": input_type,
            }
            if input_type == "choice":
                entry["choices"] = list(question[3]) if len(question) > 3 else ["Yes", "No"]
            wizard_questions.append(entry)

        answers_map = self._prompt_wizard("Default Claim Notes", wizard_questions)
        if answers_map is None:
            return
        answers = [(str(question[0]), str(answers_map.get(f"q{index}", "")).strip()) for index, question in enumerate(questions)]

        output_lines = [header, ""]
        for label, value in answers:
            if label == "12a. Explain damages shown in photos":
                continue
            if label == "13a. List unrelated damage":
                continue
            if label == "12. Damages shown in photos?":
                explain_value = next((item[1] for item in answers if item[0] == "12a. Explain damages shown in photos"), "")
                output_lines.append(f"{label} Explain.: {value} {explain_value}".strip())
                continue
            if label in {"13. Unrelated damage?", "12. Unrelated damage?"}:
                list_value = next((item[1] for item in answers if item[0] == "13a. List unrelated damage"), "")
                output_lines.append(f"{label} List: {value} {list_value}".strip())
                continue
            output_lines.append(f"{label}: {value}")

        output = "\n".join(output_lines).strip()
        self.root.clipboard_clear()
        self.root.clipboard_append(output)
        messagebox.showinfo("Copied", "Default Claim Notes copied to clipboard. Paste it into the estimate.")

    def generate_mitchell_total_loss_for_selected_claim(self) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select a claim first.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return
        if Document is None:
            messagebox.showerror("Word Support Missing", "python-docx is not available for the Mitchell form.")
            return
        existing_mitchell = self._load_existing_mitchell_answers(record)

        claimant_name_default, insured_name_default = self._mitchell_party_defaults(record)
        top_questions = [
            {"key": "Claim-Suffix ID", "label": "Claim-Suffix ID", "default": existing_mitchell.get("Claim-Suffix ID", record.claim_number or record.claim_id or ""), "input_type": "text"},
            {"key": "Claimant Name", "label": "Claimant Name", "default": existing_mitchell.get("Claimant Name", claimant_name_default), "input_type": "text"},
            {"key": "Claimant Phone", "label": "Claimant Phone", "default": existing_mitchell.get("Claimant Phone", record.contact_phone or ""), "input_type": "text"},
            {"key": "Loss Date", "label": "Loss Date", "default": existing_mitchell.get("Loss Date", record.date_of_loss or ""), "input_type": "text"},
            {"key": "License Plate", "label": "License Plate", "default": existing_mitchell.get("License Plate", ""), "input_type": "text"},
            {"key": "Insured Name", "label": "Insured Name", "default": existing_mitchell.get("Insured Name", insured_name_default), "input_type": "text"},
            {"key": "Insured Phone", "label": "Insured Phone", "default": existing_mitchell.get("Insured Phone", ""), "input_type": "text"},
            {"key": "Loss Type", "label": "Loss Type", "default": existing_mitchell.get("Loss Type", "Collision"), "input_type": "text"},
            {"key": "VIN", "label": "VIN", "default": existing_mitchell.get("VIN", record.vin or ""), "input_type": "text"},
            {"key": "Year", "label": "Year", "default": existing_mitchell.get("Year", self._mitchell_year_from_vehicle(record.vehicle)), "input_type": "text"},
            {"key": "Make", "label": "Make", "default": existing_mitchell.get("Make", self._mitchell_make_from_vehicle(record.vehicle)), "input_type": "text"},
            {"key": "Model", "label": "Model", "default": existing_mitchell.get("Model", self._mitchell_model_from_vehicle(record.vehicle)), "input_type": "text"},
            {"key": "Sub-model", "label": "Sub-model", "default": existing_mitchell.get("Sub-model", ""), "input_type": "text"},
            {"key": "Mileage", "label": "Mileage", "default": existing_mitchell.get("Mileage", ""), "input_type": "text"},
            {"key": "Body Style", "label": "Body Style", "default": existing_mitchell.get("Body Style", ""), "input_type": "text"},
            {"key": "Ext. Color", "label": "Ext. Color", "default": existing_mitchell.get("Ext. Color", ""), "input_type": "text"},
            {"key": "Engine", "label": "Engine", "default": existing_mitchell.get("Engine", ""), "input_type": "text"},
            {"key": "Transmission", "label": "Transmission", "default": existing_mitchell.get("Transmission", ""), "input_type": "text"},
            {"key": "Drive Train", "label": "Drive Train", "default": existing_mitchell.get("Drive Train", ""), "input_type": "text"},
            {"key": "Location of Vehicle", "label": "Location of Vehicle", "default": existing_mitchell.get("Location of Vehicle", record.location_of_vehicle or record.shop_name or ""), "input_type": "text"},
            {"key": "Zip Code", "label": "Zip Code", "default": existing_mitchell.get("Zip Code", self._zip_from_address(record.owner_address)), "input_type": "text"},
            {"key": "Inspected By", "label": "Inspected By", "default": existing_mitchell.get("Inspected By", "Fernando Marin"), "input_type": "text"},
            {"key": "Date", "label": "Date", "default": existing_mitchell.get("Date", datetime.now().strftime("%#m/%#d/%y") if os.name == "nt" else datetime.now().strftime("%-m/%-d/%y")), "input_type": "text"},
        ]
        top_answers = self._prompt_wizard("Mitchell Total Loss - Vehicle Information", top_questions)
        if top_answers is None:
            return

        rating_options = [
            "5 - Excellent",
            "4 - Very good",
            "3 - Good",
            "2 - Fair",
            "1 - Poor",
            "U - Unknown",
        ]
        condition_items = [
            "Seats",
            "Carpet",
            "Headliner",
            "Doors / Interior Panels",
            "Dash/Console",
            "Glass",
            "Body",
            "Paint",
            "Trim",
            "Vinyl or Convertible Tops",
            "Engine",
            "Transmission",
            "Tires",
        ]
        condition_questions: list[dict[str, object]] = []
        for item in condition_items:
            condition_questions.append(
                {"key": f"{item} rating", "label": f"{item} rating", "default": existing_mitchell.get(f"{item} rating", "3 - Good"), "input_type": "choice", "choices": rating_options}
            )
            condition_questions.append(
                {"key": f"{item} comments", "label": f"{item} comments", "default": existing_mitchell.get(f"{item} comments", ""), "input_type": "text"}
            )
        final_questions = [
            {
                "key": "After-Market Installed Parts, Refurbishments and Prior Damage",
                "label": "After-Market Installed Parts, Refurbishments and Prior Damage",
                "default": existing_mitchell.get("After-Market Installed Parts, Refurbishments and Prior Damage", "N/A"),
                "input_type": "text",
            },
            {
                "key": "Comments",
                "label": "Comments",
                "default": existing_mitchell.get("Comments", record.assignment_claim_notes or ""),
                "input_type": "text",
            },
        ]
        condition_answers_map = self._prompt_wizard("Mitchell Total Loss - Vehicle Condition", condition_questions)
        if condition_answers_map is None:
            return
        final_answers = self._prompt_wizard("Mitchell Total Loss - Final Sections", final_questions)
        if final_answers is None:
            return

        condition_answers: list[tuple[str, str, str]] = []
        for item in condition_items:
            condition_answers.append(
                (
                    item,
                    condition_answers_map.get(f"{item} rating", "3 - Good"),
                    condition_answers_map.get(f"{item} comments", "").strip(),
                )
            )

        aftermarket = final_answers.get("After-Market Installed Parts, Refurbishments and Prior Damage", "")
        final_comments = final_answers.get("Comments", "").strip() or record.assignment_claim_notes or record.office_additional_notes or ""

        target_folder = self._record_folder_path(record)
        claim_suffix = (record.claim_id or "claim")[-4:]
        doc_path = self._unique_destination(target_folder / f"{claim_suffix}-MITCHELL-TOTAL-LOSS.doc")
        pdf_path = doc_path.with_suffix(".pdf")

        try:
            self._fill_mitchell_total_loss_template(
                doc_path,
                top_answers,
                condition_answers,
                aftermarket,
                final_comments,
            )
        except Exception as exc:
            messagebox.showerror("Mitchell Total Loss Failed", f"Could not fill the Mitchell template:\n{exc}")
            return
        try:
            self._export_doc_to_pdf(doc_path, pdf_path)
            os.startfile(str(pdf_path))
            messagebox.showinfo("Mitchell Total Loss Created", f"Created:\n{pdf_path}")
        except Exception:
            os.startfile(str(doc_path))
            messagebox.showinfo("Mitchell Total Loss Created", f"Created:\n{doc_path}")

    def _mitchell_party_defaults(self, record: ClaimRecord) -> tuple[str, str]:
        claimant = (record.claimant_name or "").strip()
        insured = (record.insured_name or "").strip()
        customer = (record.customer_name or "").strip()

        if not claimant and customer and customer != insured:
            claimant = customer
        if not insured and customer and customer != claimant:
            insured = customer
        if not claimant:
            claimant = customer or insured
        if not insured:
            insured = customer or claimant
        return claimant, insured

    def _load_existing_mitchell_answers(self, record: ClaimRecord) -> dict[str, str]:
        target_folder = self._record_folder_path(record)
        candidates = sorted(target_folder.glob("*MITCHELL-TOTAL-LOSS.docx"))
        if not candidates or Document is None:
            return {}
        path = candidates[-1]
        try:
            doc = Document(path)
        except Exception:
            return {}

        extracted: dict[str, str] = {}
        for paragraph in doc.paragraphs:
            text = paragraph.text.strip()
            if not text or ":" not in text:
                continue
            label, value = text.split(":", 1)
            label = label.strip()
            value = value.strip()
            if label in {
                "Claim-Suffix ID", "Claimant Name", "Claimant Phone", "Loss Date", "License Plate",
                "Insured Name", "Insured Phone", "Loss Type", "VIN", "Year", "Make", "Model",
                "Sub-model", "Mileage", "Body Style", "Ext. Color", "Engine", "Transmission",
                "Drive Train", "Location of Vehicle", "Zip Code", "Inspected By", "Date",
                "After-Market Installed Parts, Refurbishments And Prior Damage", "Comments",
            }:
                extracted[label] = "" if value == "-" else value

        if doc.tables:
            for row in doc.tables[0].rows[1:]:
                cells = [cell.text.strip() for cell in row.cells]
                if len(cells) >= 3 and cells[0]:
                    extracted[f"{cells[0]} rating"] = cells[1]
                    extracted[f"{cells[0]} comments"] = cells[2]

        return extracted

    def _mitchell_year_from_vehicle(self, vehicle_text: str) -> str:
        tokens = (vehicle_text or "").split()
        if not tokens:
            return ""
        first = tokens[0]
        if re.fullmatch(r"\d{4}", first):
            return first
        if re.fullmatch(r"\d{2}", first):
            return f"20{first}"
        return ""

    def _mitchell_make_from_vehicle(self, vehicle_text: str) -> str:
        tokens = (vehicle_text or "").split()
        if len(tokens) < 2:
            return ""
        return tokens[1].title()

    def _mitchell_model_from_vehicle(self, vehicle_text: str) -> str:
        tokens = (vehicle_text or "").split()
        if len(tokens) < 3:
            return ""
        return " ".join(tokens[2:4]).title()

    def _zip_from_address(self, address: str) -> str:
        match = re.search(r"\b(\d{5})(?:-\d{4})?\b", address or "")
        return match.group(1) if match else ""

    def _suggest_nada_vehicle_inputs(self, record: ClaimRecord) -> dict[str, str]:
        result = {"year": "", "make": "", "model": "", "trim": "", "mileage": ""}
        vehicle_text = " ".join((record.vehicle or "").split()).strip()
        if not vehicle_text:
            return result

        tokens = vehicle_text.split()
        if tokens and re.fullmatch(r"\d{2}|\d{4}", tokens[0]):
            year_token = tokens[0]
            result["year"] = f"20{year_token}" if len(year_token) == 2 else year_token
            tokens = tokens[1:]

        if tokens:
            result["make"] = tokens[0].title()
        if len(tokens) >= 2:
            if len(tokens) == 2:
                result["model"] = tokens[1].title()
            else:
                result["model"] = " ".join(token.title() for token in tokens[1:-1])
                result["trim"] = tokens[-1].upper()

        return result

    def _fetch_text_url(self, url: str) -> str:
        request = Request(
            url,
            headers={
                "User-Agent": "ClaimsManager/1.0",
                "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8",
            },
        )
        with urlopen(request, timeout=15) as response:
            return response.read().decode("utf-8", errors="ignore")

    def _extract_jdpower_urls_from_search_html(self, html: str) -> list[str]:
        candidates: list[str] = []

        def add_candidate(value: str) -> None:
            cleaned = unescape(value).strip()
            if not cleaned:
                return
            if cleaned.startswith("//"):
                cleaned = f"https:{cleaned}"
            if "jdpower.com/cars/" not in cleaned.lower():
                return
            if "/specs" in cleaned.lower():
                return
            if cleaned not in candidates:
                candidates.append(cleaned)

        for match in re.findall(r'href="([^"]+)"', html, flags=re.IGNORECASE):
            if "duckduckgo.com/l/?" in match:
                try:
                    parsed = urlparse(unescape(match))
                    uddg = parse_qs(parsed.query).get("uddg", [])
                    if uddg:
                        add_candidate(unquote(uddg[0]))
                except Exception:
                    pass
            else:
                add_candidate(match)

        for match in re.findall(r'https?://www\.jdpower\.com/cars/[^\s"<>]+', html, flags=re.IGNORECASE):
            add_candidate(match)

        return candidates

    def _nada_candidate_score(self, page_text: str, year: str, make: str, model: str, trim: str) -> int:
        score = 0
        lowered = page_text.lower()
        if "pricing & values" in lowered:
            score += 10
        if "used car values" in lowered:
            score += 8
        if "estimated trade in value" in lowered or "estimated trade-in value" in lowered:
            score += 4
        if year and year in page_text:
            score += 8
        if make and make.lower() in lowered:
            score += 6
        if model and model.lower() in lowered:
            score += 6
        if trim and trim.lower() in lowered:
            score += 4
        if "selected trim" in lowered:
            score += 2
        return score

    def _slugify_vehicle_value_part(self, value: str) -> str:
        slug = re.sub(r"[^a-z0-9]+", "-", (value or "").lower()).strip("-")
        return re.sub(r"-+", "-", slug)

    def _normalize_trim_value(self, value: str) -> str:
        normalized = re.sub(r"[^a-z0-9]+", " ", (value or "").lower())
        return " ".join(normalized.split())

    def _find_nada_value_candidates(self, year: str, make: str, model: str, trim: str) -> list[dict[str, str | int]]:
        search_queries = [
            f'site:jdpower.com/cars "{year} {make} {model}" "Pricing & Values"',
            f'site:jdpower.com/cars "{year} {make} {model}" "Used Car Values"',
            f'site:jdpower.com/cars "{year} {make} {model}" values',
            f'site:jdpower.com/cars "{year} {make}" "{model}" values',
            f'site:jdpower.com/cars "{make}" "{year}" "{model}" values',
        ]

        direct_candidates = []
        make_slug = self._slugify_vehicle_value_part(make)
        model_slug = self._slugify_vehicle_value_part(model)
        if year and make_slug and model_slug:
            direct_candidates.append(f"https://www.jdpower.com/cars/{year}/{make_slug}/{model_slug}")

        candidate_urls: list[str] = []
        for candidate in direct_candidates:
            if candidate not in candidate_urls:
                candidate_urls.append(candidate)

        for query in search_queries:
            search_url = f"https://duckduckgo.com/html/?q={quote(query)}"
            try:
                html = self._fetch_text_url(search_url)
            except Exception:
                continue
            for candidate in self._extract_jdpower_urls_from_search_html(html):
                if candidate not in candidate_urls:
                    candidate_urls.append(candidate)

        candidates: list[dict[str, str | int]] = []
        for candidate in candidate_urls:
            try:
                page_html = self._fetch_text_url(candidate)
            except Exception:
                continue
            page_text = self._html_to_text(page_html)
            score = self._nada_candidate_score(page_text, year, make, model, trim)
            selected_trim = ""
            trim_match = re.search(r"Selected Trim:\s*(.*?)\s*Pricing & Values", page_text, flags=re.IGNORECASE)
            if trim_match:
                selected_trim = trim_match.group(1).strip()
            normalized_trim = self._normalize_trim_value(selected_trim)
            if not normalized_trim:
                continue
            if any(
                bad in normalized_trim
                for bad in ["exterior photos", "interior photos", "photos", "specs", "reviews", "overview"]
            ):
                continue
            page_title_match = re.search(
                r"(\d{4} .*? (?:Pricing\s*&\s*Values|Used Car Values|Estimated Trade-In Value))",
                page_text,
                flags=re.IGNORECASE,
            )
            page_title = page_title_match.group(1).strip() if page_title_match else candidate
            candidates.append(
                {
                    "url": candidate,
                    "page_html": page_html,
                    "score": score,
                    "selected_trim": selected_trim,
                    "page_title": page_title,
                }
            )

        candidates.sort(key=lambda item: int(item["score"]), reverse=True)
        return candidates

    def _find_nada_value_page(self, year: str, make: str, model: str, trim: str) -> tuple[str, str]:
        candidates = self._find_nada_value_candidates(year, make, model, trim)
        if not candidates:
            raise RuntimeError("Could not find a matching JD Power / NADA value page for that vehicle.")

        normalized_input_trim = self._normalize_trim_value(trim)
        trim_words = [word for word in normalized_input_trim.split() if len(word) > 1]
        matched_candidates: list[dict[str, str | int]] = []
        for candidate in candidates:
            candidate_trim = self._normalize_trim_value(str(candidate.get("selected_trim") or ""))
            if not candidate_trim:
                continue
            if not trim_words:
                matched_candidates.append(candidate)
                continue
            if all(word in candidate_trim for word in trim_words):
                matched_candidates.append(candidate)

        if matched_candidates:
            chosen = matched_candidates[0]
        else:
            available = [
                str(candidate.get("selected_trim") or "").strip()
                for candidate in candidates[:12]
                if str(candidate.get("selected_trim") or "").strip()
            ]
            if available:
                raise RuntimeError(
                    "Could not match that trim.\n\nAvailable options:\n" + "\n".join(f"- {item}" for item in available)
                )
            raise RuntimeError("Could not match that trim.")

        url = str(chosen["url"])
        page_html = str(chosen["page_html"])
        if not url or not page_html:
            raise RuntimeError("Could not find a matching JD Power / NADA value page for that vehicle.")
        return url, page_html

    def _html_to_text(self, html: str) -> str:
        text = re.sub(r"<script.*?</script>", " ", html, flags=re.IGNORECASE | re.DOTALL)
        text = re.sub(r"<style.*?</style>", " ", text, flags=re.IGNORECASE | re.DOTALL)
        text = re.sub(r"<[^>]+>", " ", text)
        text = unescape(text)
        return " ".join(text.split())

    def _parse_nada_values(self, page_html: str) -> dict[str, str]:
        page_text = self._html_to_text(page_html)

        def grab(pattern: str) -> str:
            match = re.search(pattern, page_text, flags=re.IGNORECASE)
            return match.group(1).strip() if match else ""

        return {
            "page_title": grab(r"(\d{4} .*? (?:Pricing\s*&\s*Values|Used Car Values|Estimated Trade-In Value))"),
            "selected_trim": grab(r"Selected Trim:\s*(.*?)\s*Pricing & Values"),
            "mileage_context": grab(r"with\s+(.+?)\s+miles\s+are what people paid"),
            "average_price_paid": grab(r"Average Price Paid\s*\$([0-9,]+)"),
            "buy_range": grab(r"80% of People Paid\s*\$([0-9,]+\s*-\s*[0-9,]+)"),
            "trade_low": grab(r"Low\s*\$([0-9,]+)"),
            "trade_average": grab(r"Average\s*\$([0-9,]+)"),
            "trade_high": grab(r"High\s*\$([0-9,]+)"),
            "private_party": grab(r"Private (?:Party|Sale)\s*\$([0-9,]+)"),
            "clean_retail": grab(r"Clean Retail\s*\$([0-9,]+)"),
            "retail": grab(r"Retail(?: Value)?\s*\$([0-9,]+)"),
        }

    def _generate_nada_value_pdf(
        self,
        record: ClaimRecord,
        year: str,
        make: str,
        model: str,
        trim: str,
        mileage: str,
        target_folder: Path,
    ) -> Path:
        value_url, page_html = self._find_nada_value_page(year, make, model, trim)
        values = self._parse_nada_values(page_html)

        if not any(
            [
                values["clean_retail"],
                values["retail"],
                values["private_party"],
                values["trade_average"],
                values["average_price_paid"],
            ]
        ):
            if not any(
                [
                    values["clean_retail"],
                    values["retail"],
                    values["private_party"],
                    values["trade_average"],
                    values["average_price_paid"],
                ]
            ):
                raise RuntimeError("Found the trim page, but could not pull usable values from it.")

        claim_suffix = (record.claim_id or "claim")[-4:]
        base_name = f"{claim_suffix}-NADA"
        docx_path = self._unique_destination(target_folder / f"{base_name}.docx")
        pdf_path = docx_path.with_suffix(".pdf")

        document = Document()
        document.add_heading("NADA / JD Power Vehicle Value", level=0)
        document.add_paragraph(f"Generated: {now_stamp()}")

        summary_fields = [
            ("Claim ID", record.claim_id or "-"),
            ("Customer", record.customer_name or record.title or "-"),
            ("Insurance", record.insurance_company or "-"),
            ("Vehicle", f"{year} {make} {model} {trim}".strip()),
            ("Mileage Entered", mileage or "-"),
            ("JD Power Page", value_url),
            ("Page Title", values["page_title"] or "-"),
            ("Selected Trim", values["selected_trim"] or trim or "-"),
            ("Public Page Mileage Context", values["mileage_context"] or "-"),
        ]
        for label, value in summary_fields:
            p = document.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(value)

        document.add_heading("Values", level=1)
        value_rows = [
            ("Retail Value", values["clean_retail"] or values["retail"]),
            ("Private Sale", values["private_party"]),
            ("Average Price Paid", values["average_price_paid"]),
            ("80% Buy Range", values["buy_range"]),
        ]
        if not any(row_value for _, row_value in value_rows):
            value_rows.extend(
                [
                    ("Trade-In Low", values["trade_low"]),
                    ("Trade-In Average", values["trade_average"]),
                    ("Trade-In High", values["trade_high"]),
                ]
            )

        for label, value in value_rows:
            if not value:
                continue
            p = document.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(f"${value}" if value and value != "-" and not value.startswith("$") else value)

        note = document.add_paragraph()
        note.add_run("Note: ").bold = True
        note.add_run(
            "This PDF was generated from the public JD Power / NADA value page. "
            "Mileage-adjusted public values may vary depending on the page and option flow exposed by JD Power."
        )

        document.save(docx_path)
        self._export_docx_to_pdf(docx_path, pdf_path)
        return pdf_path

    def _generate_working_sheet(self, record: ClaimRecord) -> Path:
        if Document is None:
            raise RuntimeError("python-docx is not available.")
        WORKING_SHEETS_DIR.mkdir(parents=True, exist_ok=True)
        base_name = f"{record.claim_id or 'claim'} - Working Sheet"
        docx_path = self._unique_destination(WORKING_SHEETS_DIR / f"{base_name}.docx")
        pdf_path = docx_path.with_suffix(".pdf")
        document = Document()
        document.add_heading(f"{record.claim_id} Working Sheet", level=0)
        intro = document.add_paragraph()
        intro.add_run("Customer: ").bold = True
        intro.add_run(record.customer_name or "-")
        for label, value in [
            ("Insurance", record.insurance_company or "-"),
            ("Claim #", record.claim_number or "-"),
            ("Date of Loss", record.date_of_loss or "-"),
            ("Town", record.town or "-"),
            ("Owner Address", record.owner_address or "-"),
            ("Location", record.location_of_vehicle or "-"),
            ("Claimant", record.claimant_name or "-"),
            ("Insured", record.insured_name or "-"),
            ("Contact Phone", record.contact_phone or "-"),
            ("Contact Email", record.contact_email or "-"),
            ("Shop", record.shop_name or "-"),
            ("Vehicle", record.vehicle or "-"),
            ("VIN", record.vin or "-"),
            ("Damage", record.damage_description or "-"),
            ("Facts of Loss", record.facts_of_loss or "-"),
            ("Type", record.claim_type or "-"),
            ("Total Loss", "Yes" if record.total_loss else "No"),
            ("Folder", record.source_path or "-"),
        ]:
            p = document.add_paragraph()
            p.add_run(f"{label}: ").bold = True
            p.add_run(value)

        document.add_heading("Assignment Claim Notes", level=1)
        document.add_paragraph(record.assignment_claim_notes or "No assignment claim notes found.")

        document.add_heading("Claim Notes", level=1)
        note_text = render_note_history(record.note_history).strip()
        document.add_paragraph(note_text if note_text else "No saved notes.")

        insurer_instructions = self._load_client_instruction_text(record.insurance_company)
        document.add_heading("Client Instructions", level=1)
        if insurer_instructions:
            for line in insurer_instructions:
                if line.startswith("## "):
                    document.add_heading(line[3:], level=2)
                else:
                    document.add_paragraph(line, style="List Bullet")
        else:
            document.add_paragraph("No generated client instruction sheet found for this insurer yet.")

        document.save(docx_path)
        self._export_docx_to_pdf(docx_path, pdf_path)
        return pdf_path

    def _load_client_instruction_text(self, insurer_name: str) -> list[str]:
        key = self._normalize_insurer_key(insurer_name)
        if not key:
            return []
        if not CLIENT_INSTRUCTIONS_DIR.exists():
            return []
        candidates = []
        for path in CLIENT_INSTRUCTIONS_DIR.iterdir():
            if path.suffix.lower() not in {".docx", ".pdf"}:
                continue
            insurer_part = path.stem.replace(" - Client Instructions", "")
            if self._normalize_insurer_key(insurer_part) == key:
                candidates.append(path)
        candidates.sort(key=lambda p: (p.suffix.lower() != ".docx", p.name.lower()))
        if not candidates:
            return []
        path = candidates[0]
        if path.suffix.lower() == ".docx" and Document is not None:
            doc = Document(path)
            lines: list[str] = []
            for para in doc.paragraphs:
                text = para.text.strip()
                if not text:
                    continue
                style = para.style.name.lower() if para.style and para.style.name else ""
                if "heading 1" in style and "client instructions" not in text.lower():
                    lines.append(f"## {text}")
                elif text.lower() not in {
                    insurer_name.lower(),
                    "client instructions",
                    "organized reference sheet for field/client instructions",
                    "cleaned and organized from the assignment sheet's client instructions only.",
                }:
                    lines.append(text.lstrip("â€¢ ").strip())
            return lines
        if path.suffix.lower() == ".pdf" and PdfReader is not None:
            text = "\n".join((page.extract_text() or "") for page in PdfReader(str(path)).pages[:5])
            lines = [line.strip() for line in text.splitlines() if line.strip()]
            output: list[str] = []
            for line in lines:
                if line.lower() in {
                    insurer_name.lower(),
                    "client instructions",
                    "organized reference sheet for field/client instructions",
                    "cleaned and organized from the assignment sheet's client instructions only.",
                }:
                    continue
                if re.fullmatch(r"[A-Z][A-Za-z /&'-]+", line) and line.lower() not in {"client instructions"}:
                    output.append(f"## {line.title()}")
                else:
                    output.append(line.lstrip("â€¢ ").strip())
            return output
        return []

    def populate_details(self, key: str) -> None:
        record = self.store.get_claim(key)
        if not record:
            self.clear_details()
            return

        self.detail_vars["claim_id"].set(record.claim_id)
        self.detail_vars["title"].set(record.title)
        self.detail_vars["status"].set(record.status)
        self.detail_vars["source_path"].set(record.source_path)
        self.detail_vars["claim_type"].set(record.claim_type)
        self.detail_vars["updated_at"].set(record.updated_at)
        self.detail_vars["customer_name"].set(record.customer_name)
        self.detail_vars["insurance_company"].set(record.insurance_company)
        self.detail_vars["claim_number"].set(record.claim_number)
        self.detail_vars["policy_number"].set(record.policy_number)
        self.detail_vars["insured_name"].set(record.insured_name)
        self.detail_vars["claimant_name"].set(record.claimant_name)
        self.detail_vars["date_of_loss"].set(record.date_of_loss)
        self.detail_vars["town"].set(record.town)
        self.detail_vars["owner_address"].set(record.owner_address)
        self.detail_vars["location_of_vehicle"].set(record.location_of_vehicle)
        self.detail_vars["vehicle"].set(record.vehicle)
        self.detail_vars["vin"].set(record.vin)
        self.detail_vars["damage_description"].set(record.damage_description)
        self.detail_vars["facts_of_loss"].set(record.facts_of_loss)
        self.detail_vars["total_loss"].set("Yes" if record.total_loss else "No")
        self.detail_vars["shop_name"].set(record.shop_name)
        self.detail_vars["contact_phone"].set(record.contact_phone)
        self.detail_vars["contact_email"].set(record.contact_email)
        self.detail_vars["assignment_claim_notes"].set(record.assignment_claim_notes)
        self.detail_vars["assign_pdf_path"].set(record.assign_pdf_path)
        self.new_note_text.delete("1.0", tk.END)
        self._set_notes_history_text(render_note_history(record.note_history))

    def clear_details(self) -> None:
        for value in self.detail_vars.values():
            value.set("")
        self.new_note_text.delete("1.0", tk.END)
        self._set_notes_history_text("")

    def _set_notes_history_text(self, value: str) -> None:
        self.notes_history_text.configure(state="normal")
        self.notes_history_text.delete("1.0", tk.END)
        if value:
            self.notes_history_text.insert("1.0", value)
        self.notes_history_text.configure(state="disabled")

    def _clean_detail_value(self, key: str) -> str:
        value = self.detail_vars[key].get().strip()
        if value == "-":
            return ""
        return value

    def save_claim_details(self) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select a claim before saving details.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return

        manual_overrides = dict(record.manual_overrides or {})
        for key in self.editable_detail_keys:
            cleaned = self._clean_detail_value(key)
            if key == "total_loss":
                normalized_bool = cleaned.lower() in {"yes", "y", "true", "1", "total", "total loss"}
                record.total_loss = normalized_bool
                manual_overrides[key] = normalized_bool
            else:
                setattr(record, key, cleaned)
                manual_overrides[key] = cleaned

        record.manual_overrides = manual_overrides
        record.updated_at = now_stamp()
        self.store.upsert_claim(record)
        self.store.save()
        self.refresh_views(select_key=record.key)

    def save_selected_claim(self) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select a claim before saving.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return

        note_text = self.new_note_text.get("1.0", tk.END).strip()
        if not note_text:
            messagebox.showinfo("No Note", "Type a note before saving.")
            return
        record.note_history.append({"timestamp": now_stamp(), "text": note_text})
        record.notes = render_note_history(record.note_history)
        record.updated_at = now_stamp()
        self.store.upsert_claim(record)
        self.store.save()
        self.refresh_views(select_key=record.key)

    def set_selected_status(self, status: str) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select a claim first.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return

        if status == "Closed":
            moved_path = self._move_claim_to_closed(record)
            if moved_path is None:
                return
            record.key = f"fs::{moved_path}"
            record.source_path = moved_path
            record.title = Path(moved_path).name
            record.claim_id = self._extract_job_number(record.title) or record.claim_id
        record.status = status
        note_text = self.new_note_text.get("1.0", tk.END).strip()
        if note_text:
            record.note_history.append({"timestamp": now_stamp(), "text": note_text})
            record.notes = render_note_history(record.note_history)
        record.updated_at = now_stamp()
        self.store.upsert_claim(record)
        self.store.save()
        self.scan_folders()
        self.refresh_views(select_key=record.key)

    def reopen_selected_claim(self) -> None:
        self._reopen_selected_claim(copy_only=False)

    def reopen_selected_claim_as_supplement(self) -> None:
        self._reopen_selected_claim(copy_only=True)

    def _reopen_selected_claim(self, copy_only: bool) -> None:
        if not self.selected_key:
            messagebox.showinfo("No Claim Selected", "Select a claim first.")
            return
        record = self.store.get_claim(self.selected_key)
        if not record:
            return

        new_path = self._move_or_copy_claim_to_pending(record, copy_only=copy_only)
        if new_path is None:
            return

        note_text = self.new_note_text.get("1.0", tk.END).strip()
        if note_text:
            record.note_history.append({"timestamp": now_stamp(), "text": note_text})
            record.notes = render_note_history(record.note_history)

        if copy_only:
            new_key = f"fs::{new_path}"
            new_record = ClaimRecord(
                key=new_key,
                claim_id=self._extract_job_number(Path(new_path).name) or record.claim_id,
                title=Path(new_path).name,
                status="Open",
                source_path=new_path,
                claim_type="Supplement",
                updated_at=now_stamp(),
                notes=record.notes,
                note_history=list(record.note_history),
                customer_name=record.customer_name,
                insurance_company=record.insurance_company,
                claim_number=record.claim_number,
                policy_number=record.policy_number,
                insured_name=record.insured_name,
                claimant_name=record.claimant_name,
                date_of_loss=record.date_of_loss,
                town=record.town,
                vehicle=record.vehicle,
                vin=record.vin,
                shop_name=record.shop_name,
                contact_phone=record.contact_phone,
                contact_email=record.contact_email,
                assign_pdf_path=record.assign_pdf_path,
            )
            self.store.upsert_claim(new_record)
            self.store.save()
            self.scan_folders()
            self.refresh_views(select_key=new_key)
            return

        record.source_path = new_path
        record.key = f"fs::{new_path}"
        record.title = Path(new_path).name
        record.claim_id = self._extract_job_number(record.title) or record.claim_id
        record.status = "Open"
        record.updated_at = now_stamp()
        self.store.upsert_claim(record)
        self.store.save()
        self.scan_folders()
        self.refresh_views(select_key=record.key)

    def _move_claim_to_closed(self, record: ClaimRecord) -> str | None:
        source = Path(record.source_path)
        if record.source_path == "Manual Entry" or not source.exists() or not source.is_dir():
            return record.source_path

        closed_folder = next((folder for folder in self.store.watched_folders if "closed" in folder.lower()), "")
        if not closed_folder:
            messagebox.showerror("Missing Closed Folder", "Closed Claims folder is not configured.")
            return None

        today = datetime.now()
        closed_root = Path(closed_folder)
        month_folder = closed_root / f"{today.month}-{today.year}"
        day_folder = month_folder / f"Claims {today.month}-{today.day}-{today.strftime('%y')}"
        day_folder.mkdir(parents=True, exist_ok=True)
        destination = self._unique_destination(day_folder / source.name)
        if source.resolve() != destination.resolve():
            shutil.move(str(source), str(destination))
        return str(destination)

    def _move_or_copy_claim_to_pending(self, record: ClaimRecord, copy_only: bool) -> str | None:
        source = Path(record.source_path)
        if record.source_path == "Manual Entry" or not source.exists() or not source.is_dir():
            return record.source_path

        pending_folder = next((folder for folder in self.store.watched_folders if "pending" in folder.lower()), "")
        if not pending_folder:
            messagebox.showerror("Missing Pending Folder", "Pending Claims folder is not configured.")
            return None

        target_name = self._supplement_folder_name(source.name) if copy_only else source.name
        destination = self._unique_destination(Path(pending_folder) / target_name)
        if copy_only:
            shutil.copytree(str(source), str(destination))
        else:
            shutil.move(str(source), str(destination))
        return str(destination)

    def _supplement_folder_name(self, folder_name: str) -> str:
        if "supp" in folder_name.lower():
            return folder_name
        return f"{folder_name} - supp"

    def _record_folder_path(self, record: ClaimRecord) -> Path:
        if record.source_path != "Manual Entry":
            source = Path(record.source_path)
            if source.exists() and source.is_dir():
                return source
        WORKING_SHEETS_DIR.mkdir(parents=True, exist_ok=True)
        return WORKING_SHEETS_DIR

    def _fill_mitchell_total_loss_template(
        self,
        output_path: Path,
        top_answers: dict[str, str],
        condition_answers: list[tuple[str, str, str]],
        aftermarket: str,
        final_comments: str,
    ) -> None:
        if not MITCHELL_TOTAL_LOSS_TEMPLATE.exists():
            raise RuntimeError("Mitchell total loss template is missing.")

        text_field_map = {
            "Text1": top_answers.get("Claim-Suffix ID", ""),
            "Text2": top_answers.get("Claimant Name", ""),
            "Text3": top_answers.get("Claimant Phone", ""),
            "Text4": top_answers.get("Loss Date", ""),
            "Text5": top_answers.get("License Plate", ""),
            "Text6": top_answers.get("Insured Name", ""),
            "Text7": top_answers.get("Insured Phone", ""),
            "Text8": top_answers.get("Loss Type", ""),
            "Text9": top_answers.get("VIN", ""),
            "Text10": top_answers.get("Year", ""),
            "Text11": top_answers.get("Make", ""),
            "Text12": top_answers.get("Model", ""),
            "Text13": top_answers.get("Sub-model", ""),
            "Text14": top_answers.get("Mileage", ""),
            "Text15": top_answers.get("Body Style", ""),
            "Text16": top_answers.get("Ext. Color", ""),
            "Text17": top_answers.get("Engine", ""),
            "Text18": top_answers.get("Location of Vehicle", ""),
            "Text19": top_answers.get("Zip Code", ""),
            "Text20": top_answers.get("Inspected By", ""),
            "Text21": top_answers.get("Date", ""),
            "Text22": condition_answers[0][2] if len(condition_answers) > 0 else "",
            "Text23": condition_answers[1][2] if len(condition_answers) > 1 else "",
            "Text24": condition_answers[2][2] if len(condition_answers) > 2 else "",
            "Text25": condition_answers[3][2] if len(condition_answers) > 3 else "",
            "Text26": condition_answers[4][2] if len(condition_answers) > 4 else "",
            "Text27": condition_answers[5][2] if len(condition_answers) > 5 else "",
            "Text28": condition_answers[6][2] if len(condition_answers) > 6 else "",
            "Text29": condition_answers[7][2] if len(condition_answers) > 7 else "",
            "Text30": condition_answers[8][2] if len(condition_answers) > 8 else "",
            "Text31": condition_answers[9][2] if len(condition_answers) > 9 else "",
            "Text32": condition_answers[10][2] if len(condition_answers) > 10 else "",
            "Text33": condition_answers[11][2] if len(condition_answers) > 11 else "",
            "Text34": condition_answers[12][2] if len(condition_answers) > 12 else "",
            "Text35": aftermarket,
            "Text36": final_comments,
        }

        checkbox_values: list[bool] = [False] * 216

        # Factory-installed options that should be on by default in the Mitchell form.
        for field_number in [31, 58, 63, 66, 79, 80, 105, 108, 109]:
            checkbox_values[field_number - 1] = True

        transmission_indices = [18, 19, 20, 21]
        transmission_labels = ["Adaptive", "Automatic", "Interactive", "Manual"]
        self._apply_choice_to_checkbox_indices(checkbox_values, transmission_indices, transmission_labels, top_answers.get("Transmission", ""))

        drive_train_indices = [22, 23, 24]
        drive_train_labels = ["2WD", "4WD", "AWD"]
        self._apply_choice_to_checkbox_indices(checkbox_values, drive_train_indices, drive_train_labels, top_answers.get("Drive Train", ""))

        rating_groups = [
            [124, 125, 126, 127, 128, 129],
            [131, 132, 133, 134, 135, 136],
            [138, 139, 140, 141, 142, 143],
            [145, 146, 147, 148, 149, 150],
            [152, 153, 154, 155, 156, 157],
            [159, 160, 161, 162, 163, 164],
            [166, 167, 168, 169, 170, 171],
            [173, 174, 175, 176, 177, 178],
            [180, 181, 182, 183, 184, 185],
            [187, 188, 189, 190, 191, 192],
            [194, 195, 196, 197, 198, 199],
            [201, 202, 203, 204, 205, 206],
            [208, 209, 210, 211, 212, 213],
        ]
        rating_labels = ["5 - Excellent", "4 - Very good", "3 - Good", "2 - Fair", "1 - Poor", "U - Unknown"]
        for idx, (_, rating, _) in enumerate(condition_answers):
            if idx < len(rating_groups):
                self._apply_choice_to_checkbox_indices(checkbox_values, rating_groups[idx], rating_labels, rating)

        payload = {
            "template_path": str(MITCHELL_TOTAL_LOSS_TEMPLATE),
            "output_path": str(output_path),
            "text_fields": text_field_map,
            "checkbox_values": checkbox_values,
        }
        payload_path = APP_DIR / "mitchell_fill_payload.json"
        payload_path.write_text(json.dumps(payload), encoding="utf-8")

        script = r"""
$payload = Get-Content -Raw '__PAYLOAD__' | ConvertFrom-Json
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
$word.ScreenUpdating = $false
$word.Options.SaveNormalPrompt = $false
$word.Options.ConfirmConversions = $false
$word.Options.WarnBeforeSavingPrintingSendingMarkup = $false
$word.AutomationSecurity = 3
$doc = $word.Documents.Open($payload.template_path)
try {
    foreach ($field in $doc.FormFields) {
        $name = [string]$field.Name
        if ($payload.text_fields.PSObject.Properties.Name -contains $name) {
            $field.Result = [string]$payload.text_fields.$name
        }
    }
    for ($i = 0; $i -lt $doc.FormFields.Count; $i++) {
        $field = $doc.FormFields.Item($i + 1)
        if ($field.Type -eq 71) {
            $field.CheckBox.Value = [bool]$payload.checkbox_values[$i]
        }
    }
    $doc.SaveAs2($payload.output_path, 0)
}
finally {
    if ($null -ne $doc) { $doc.Close($false) }
    if ($null -ne $word) { $word.Quit() }
}
"""
        script = script.replace("__PAYLOAD__", str(payload_path))
        try:
            subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
            check=True,
            capture_output=True,
            text=True,
            timeout=90,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("Word timed out while filling the Mitchell template. Close any open Word windows and try again.") from exc

    def _apply_choice_to_checkbox_indices(
        self,
        checkbox_values: list[bool],
        indices: list[int],
        labels: list[str],
        value: str,
    ) -> None:
        normalized = (value or "").strip().lower()
        selected_idx = None
        for idx, label in enumerate(labels):
            if normalized == label.lower() or normalized in label.lower():
                selected_idx = idx
                break
        if selected_idx is None:
            return
        for idx, field_number in enumerate(indices):
            checkbox_values[field_number - 1] = idx == selected_idx

    def _export_doc_to_pdf(self, doc_path: Path, pdf_path: Path) -> None:
        script = rf"""
$word = New-Object -ComObject Word.Application
$word.Visible = $false
$word.DisplayAlerts = 0
$word.ScreenUpdating = $false
$word.Options.SaveNormalPrompt = $false
$word.Options.ConfirmConversions = $false
$word.Options.WarnBeforeSavingPrintingSendingMarkup = $false
$word.AutomationSecurity = 3
$doc = $word.Documents.Open('{doc_path}')
try {{
    $doc.ExportAsFixedFormat('{pdf_path}', 17)
}}
finally {{
    if ($null -ne $doc) {{ $doc.Close($false) }}
    if ($null -ne $word) {{ $word.Quit() }}
}}
"""
        try:
            subprocess.run(
                ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                check=True,
                capture_output=True,
                text=True,
                timeout=90,
            )
        except subprocess.TimeoutExpired as exc:
            raise RuntimeError("Word timed out while exporting the Mitchell PDF. Close any open Word windows and try again.") from exc

    def _unique_destination(self, destination: Path) -> Path:
        if not destination.exists():
            return destination
        counter = 2
        while True:
            candidate = destination.with_name(f"{destination.name} ({counter})")
            if not candidate.exists():
                return candidate
            counter += 1


def main() -> None:
    root = tk.Tk()
    ttk.Style().theme_use("clam")
    ClaimsManagerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()

